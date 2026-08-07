"""Fix BANANAS31 in CASE_INDEX: clear row 3, write correct row 14."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font
from pathlib import Path

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
idx_xlsx = CASE_DIR / "CASE_INDEX.xlsx"

wb = openpyxl.load_workbook(idx_xlsx)
ws = wb["INDEX"]

# 1. Clear wrong row 3
for ci in range(1, ws.max_column+2):
    ws.cell(3, ci).value = None

# 2. Write correct entry to row 14
# Columns: file, token, status, loai, action_class, ngay_mo + extras
entry = [
    "CASE_BANANAS31.xlsx",
    "BANANAS31",
    "COLD_START — PENDING",
    "Alpha fullyDelisted, dump -80% ATH, pos_r~6.7 (BTW/phan phoi?)",
    "COLD_START",
    "2026-08-07",
    "",  # col 7
    "",  # col 8
    "",  # col 9
    "",  # col 10
    "ALPHA_21 | contract 0x3d4f...a760 BSC | cold-start 07/08/2026",
]
for ci, val in enumerate(entry, 1):
    c = ws.cell(14, ci)
    c.value = val
    c.font  = Font(size=9)

wb.save(idx_xlsx)
print(f"Fixed: row 3 cleared, row 14 written")

# Verify
wb2 = openpyxl.load_workbook(idx_xlsx)
ws2 = wb2["INDEX"]
for ri in [3, 13, 14]:
    vals = [ws2.cell(ri, ci).value for ci in range(1, 8)]
    print(f"  Row {ri}: {vals}")
