#!/usr/bin/env python3
"""T5: Export data for cooldown tokens PARTI/EVAA/CLO/BLUAI Jul 10-29."""
import csv
from pathlib import Path

f = Path('/root/Data_back_test/whale_data/data/export/freshwatch_daily_2026-07.csv')
rows = list(csv.DictReader(f.open(encoding='utf-8')))

for sym in ['PARTIUSDT', 'EVAAUSDT', 'CLOUSDT', 'BLUAIUSDT']:
    data = [(r['date'], r['lsr_pos_r'], r['lsr_acc_r'], r['lsr_global_r'],
             r['oi_usd'], r['funding_avg'], r['funding_n'], r['data_ok'])
            for r in rows if r['symbol'] == sym and r['date'] >= '2026-07-10']
    print(f"\n{sym} ({len(data)} rows Jul10-29):")
    print("  date       | pos_r  | acc_r  | glsr   | oi_usd        | f_avg    | fn | ok")
    for d in data:
        oi_v = float(d[4]) if d[4] else 0
        print(f"  {d[0]} | {d[1]:>6} | {d[2]:>6} | {d[3]:>6} | {oi_v:>13,.0f} | {d[5]:>8} | {d[6]} | {d[7]}")

# CAMPAIGNS sheet
import openpyxl
wb = openpyxl.load_workbook('/root/Data_back_test/whale_data/data/ledger/ledger_v2.xlsx', read_only=True, data_only=True)
ws = wb['CAMPAIGNS']
header = [str(c) if c is not None else '' for c in next(ws.iter_rows(max_row=1, values_only=True))]
print('\nCAMPAIGNS cols:', header[:20])
for row in ws.iter_rows(min_row=2, values_only=True):
    sym = str(row[0]).strip() if row else ''
    if sym in ('PARTI', 'EVAA', 'CLO', 'BLUAI'):
        d = {header[i]: row[i] for i in range(min(len(header), len(row))) if row[i] is not None}
        print(f"\n  {sym}: {d}")
