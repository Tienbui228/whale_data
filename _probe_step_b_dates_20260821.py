"""STEP B — Date verify + neo B2 mốc kiểm. Read-only."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

CASES = Path(r"D:\Data back test\whale_data\docs\case_files")

def load_daily(sym):
    fp = CASES / f"CASE_{sym}.xlsx"
    if not fp.exists(): return None, None, None
    wb = openpyxl.load_workbook(fp, data_only=True)
    if "DAILY" not in wb.sheetnames: return wb, None, None
    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
    return wb, ws, ci

def get_row(ws, ci, target_date):
    for r in range(3, ws.max_row+1):
        d = ws.cell(r, 1).value
        if d and str(d)[:10] == target_date:
            return r
    return None

def get_val(ws, ci, r, key_candidates):
    for k in key_candidates:
        if k in ci:
            v = ws.cell(r, ci[k]).value
            if v is not None:
                return v
    return None

FIELDS = {
    "close": ["c"],
    "vol_perp": ["vol_perp_usd"],
    "OI": ["oi_usd"],
    "pos_r": ["pos_r", "lsr_pos_r"],
    "acc_r": ["acc_r", "lsr_acc_r"],
    "gLSR": ["gLSR", "lsr_global_r"],
    "taker_buy": ["taker_buy_pct (k[10]/k[7])", "taker_buy_pct"],
    "f_avg": ["f_avg"],
}

MOCS = {
    "BANK": [("2026-07-02", "vol ≈$16.4M probe"),
             ("2026-07-10", "pos_r ≈3.7"),
             ("2026-07-19", "pos_r sập →≈1.17"),
             ("2026-07-20", "pos_r sập →≈1.17")],
    "TUT":  [("2026-07-11", "vol ≈$14M + OI ≈$3.3M"),
             ("2026-07-14", "pos_r ≈5.68 acc ≈1.65"),
             ("2026-08-05", "harvest pos_r −≥25%")],
    "AKE":  [("2026-07-14", "taker ≈57.9% + OI −15%"),
             ("2026-07-18", "vol ≈$480k")],
    "BTW":  [("2026-06-20", "close ≈0.20 vol ≈$638M funding ≈+0.29%"),
             ("2026-06-27", "close 0.098"), ("2026-06-29", "close 0.052"),
             ("2026-08-10", "GOM fire")],
    "CYS":  [("2026-08-04", "funding ≈−1.3e-3 squeeze"),
             ("2026-08-07", "GOM-C fire"),
             ("2026-08-13", "mg14 ≈−18.3%")],
}

D_1 = "2026-08-20"
print("="*70)
print("B1 — max_date phải = 2026-08-20 (D-1). Không tương lai, không trùng.")
print("="*70)
for sym in ["BANK","TUT","AKE","BTW","CYS"]:
    wb, ws, ci = load_daily(sym)
    if ws is None:
        print(f"  {sym:<6}: FILE/DAILY KHÔNG TỒN TẠI — CHƯA CHẠY")
        continue
    dates = sorted({str(ws.cell(r,1).value)[:10] for r in range(3, ws.max_row+1) if ws.cell(r,1).value})
    max_d = dates[-1] if dates else None
    future = [d for d in dates if d > D_1]
    ok_max = max_d == D_1
    print(f"  {sym:<6}: max_date={max_d} · target={D_1} · match={ok_max} · future dates={len(future)}")

print("\n" + "="*70)
print("B2 — Neo đối chiếu 3 mốc/token. Paste DAILY values.")
print("="*70)
for sym, mocs in MOCS.items():
    wb, ws, ci = load_daily(sym)
    if ws is None:
        print(f"\n[{sym}] CHƯA CHẠY (file/DAILY not found)")
        continue
    print(f"\n[{sym}]")
    for d_str, note in mocs:
        r = get_row(ws, ci, d_str)
        if r is None:
            print(f"  {d_str}: KHÔNG TỒN TẠI DAILY row · flag DATE-SHIFT-CANDIDATE · note={note}")
            continue
        vals = {label: get_val(ws, ci, r, cand) for label, cand in FIELDS.items()}
        vs = ", ".join(f"{k}={v}" for k, v in vals.items() if v is not None)
        print(f"  {d_str}: {vs}")
        print(f"    (expected note: {note})")

print("\n" + "="*70)
print("B3 — BTW min_date ≥ 60d TRƯỚC 20/06 (target min ≤ 2026-04-21).")
print("="*70)
wb, ws, ci = load_daily("BTW")
if ws:
    dates = sorted({str(ws.cell(r,1).value)[:10] for r in range(3, ws.max_row+1) if ws.cell(r,1).value})
    min_d = dates[0] if dates else None
    ok = min_d and min_d <= "2026-04-21"
    print(f"  BTW min_date={min_d} · target ≤2026-04-21 · pass={ok}" + ("" if ok else " → THIẾU-NỀN 60d, KHÔNG tự collect thêm"))

print("\n" + "="*70)
print("B4 — TUT close 08/08 (T+30) + 20/08 [DATA]")
print("="*70)
wb, ws, ci = load_daily("TUT")
if ws:
    for d_str in ["2026-08-08", "2026-08-20"]:
        r = get_row(ws, ci, d_str)
        if r is None:
            print(f"  {d_str}: KHÔNG TỒN TẠI DAILY row (TUT max={sorted({str(ws.cell(r,1).value)[:10] for r in range(3, ws.max_row+1) if ws.cell(r,1).value})[-1]})")
        else:
            c = get_val(ws, ci, r, ["c"])
            print(f"  {d_str}: close={c}")

print("\n[DONE STEP B]")
