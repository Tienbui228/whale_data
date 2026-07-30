#!/usr/bin/env python3
"""Append L-38 to CASE_INDEX.xlsx LESSONS sheet và regenerate lessons_learned.md."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from pathlib import Path
from datetime import datetime

IDX   = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
MDOUT = Path(r"D:\Data back test\whale_data\docs\lessons_learned.md")

L38 = [
    "L-38",
    "2026-07-31",
    "BULLA 07-20 (B4 probe — 4 nguon 4 so)",
    ("Mot metric mot ngay ra 4 con so khac nhau (1.75/1.68/1.981/1.99) tu 4 nguon: "
     "(1) u50 loose = 1.75 (nguon khong ro, retired); "
     "(2) top_pos_1h keep-last 23:00 UTC = 1.6811; "
     "(3) derived_1h = 00:00 UTC dau ngay = 1.981; "
     "(4) native_1d repo goc = 1.99. "
     "Nguyen nhan: quy uoc canonical CHUA TON TAI khi calib duoc viet. "
     "He qua: verdict buoc phai chon nguon theo ket luan muon co = vi pham nguyen tac. "
     "Fix: quy uoc phai duoc ky TRUOC khi bat ky so nao duoc trich dan vao calib/DAILY. "
     "Xem proposals/ratio_daily_canonical_20260731.md cho de xuat keep-last canonical."),
    "rule",
    "GHI NHAN — proposals/ratio_daily_canonical_20260731.md cho operator ky",
    "ky proposals/ratio_daily_canonical_20260731.md truoc khi chay backfill tu ngay 31/07",
    None,
]

wb = openpyxl.load_workbook(str(IDX), data_only=False)
ws = wb["LESSONS"]

# Find last L-xx row
last_lesson_row = 4
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v and str(v).startswith("L-"):
        last_lesson_row = r

last_id = ws.cell(last_lesson_row, 1).value
print(f"Last lesson: {last_id} at row {last_lesson_row}")
assert last_id == "L-37", f"Expected L-37, got {last_id!r}"

row_n = last_lesson_row + 1
for c, v in enumerate(L38, start=1):
    ws.cell(row_n, c).value = v
print(f"Wrote L-38 at row {row_n}")

wb.save(str(IDX))
print(f"Saved: {IDX}")

# Regenerate md
wb2 = openpyxl.load_workbook(str(IDX), data_only=True)
ws2 = wb2["LESSONS"]

lessons_count = sum(1 for r in range(5, ws2.max_row + 1)
                    if ws2.cell(r, 1).value and str(ws2.cell(r, 1).value).startswith("L-"))
print(f"Total lessons: {lessons_count}")

lines = [
    "# Lessons Learned — FreshWatch",
    "",
    "> Auto-generated tu CASE_INDEX.xlsx sheet LESSONS. Khong edit file nay truc tiep.",
    f"> Lan sinh: {datetime.now().strftime('%Y-%m-%d')} (rev.4) · Tong: {lessons_count} bai hoc (L-01 → L-38)",
    "",
]

for r in range(5, ws2.max_row + 1):
    lid        = ws2.cell(r, 1).value
    ngay       = ws2.cell(r, 2).value
    case_src   = ws2.cell(r, 3).value
    bai_hoc    = ws2.cell(r, 4).value
    loai       = ws2.cell(r, 5).value
    trang_thai = ws2.cell(r, 6).value
    hanh_dong  = ws2.cell(r, 7).value

    if not lid or not str(lid).startswith("L-"):
        continue

    lines.append(f"## {lid} [{ngay}] {case_src} — `{loai}` · {trang_thai}")
    lines.append(str(bai_hoc or ""))
    if hanh_dong:
        lines.append(f"*Hanh dong tiep: {hanh_dong}*")
    lines.append("")

wb2.close()
MDOUT.write_text("\n".join(lines), encoding="utf-8")
print(f"Written: {MDOUT} ({lessons_count} lessons)")
