"""_task_insert_lessons_20260821.py
Insert L-52, L-53, L-54 vào LESSONS sheet CASE_INDEX.xlsx.
Nguồn nội dung: proposals/lessons_pending_batch_20260818.md (§D.1, §E.1, §E.2).
Backup L-37 đã tạo trước (scratch/CASE_INDEX_bak_pre_lesson_insert_20260821.xlsx).
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
BAK = Path(r"D:\Data back test\whale_data\scratch\CASE_INDEX_bak_pre_lesson_insert_20260821.xlsx")
assert BAK.exists(), "Backup missing — refuse write"
print(f"Backup verified: {BAK}")

wb = openpyxl.load_workbook(XLSX)
ws = wb["LESSONS"]
print(f"LESSONS before: max_row={ws.max_row}")

# Find first empty row after L-BUG
next_r = ws.max_row + 1
for r in range(ws.max_row, 4, -1):
    v = ws.cell(r, 1).value
    if v and str(v).strip():
        next_r = r + 1
        break
print(f"Insert starting at row {next_r}")

# ── L-52: Venue coverage ──
lessons_new = [
    {
        "id": "L-52",
        "ngay": "2026-08-19",
        "case_nguon": "TRIA + recon 4 case (BTW/GUA/APR/CLO)",
        "bai_hoc": ("Venue coverage — 'không có spot Binance' ≠ 'không có spot'. "
                    "Băng thông xả và chân hàng phải tính trên venue spot LỚN NHẤT THỰC TẾ (không phải venue mặc định đầu tiên tra). "
                    "Trước khi kết luận 'Alpha-only' hay 'perp-only chain': tra spot venues lớn (Bybit · OKX · Gate · KuCoin · MEXC · Bitget) trước, kết luận sau. "
                    "LAYER COVERAGE mọi report case phải khai `spot-ext: <venue>/N/A/CHƯA-TRA` — quy ước CHƯA-TRA ≠ N/A. "
                    "Cột canonical CASE schema v2.1: vol_spot_ext_usd (col 51) · spot_ext_venue (col 52) — CẤM ghi Bybit/OKX vào vol_spot_usd (col 31 = Binance api/v3 canonical). "
                    "Exhibit TRIA 19/08: Bybit turnover Aug18 = $2,304,652 (3× Alpha $770K) — nếu không tra Bybit → mù nguyên một van xả 3×. "
                    "Recon: GUA có Gate spot đã bỏ sót."),
        "loai": "VENUE",
        "trang_thai": "DA VAO — spec §D4-ext + amendment 19/08",
        "hanh_dong_tiep": ("Case-open probe: tra ≥3 venue spot (Bybit + OKX + Gate) tối thiểu trước khi ghi verdict Alpha/perp-only. "
                           "Identity-check bit-level ≥3 mẫu chống mirror trước khi dùng field mới."),
        "mm": "V1-20260819",
    },
    # ── L-53: Captive-kho vùng-giữa ──
    {
        "id": "L-53",
        "ngay": "2026-08-20",
        "case_nguon": "CASE_HEMI (autopsy 20/08) — evidence #2 G-committed-long (BTW-July = #1)",
        "bai_hoc": ("Loài captive-kho vùng-giữa — máy perp mù theo CẤU TRÚC, không phải theo ngưỡng. "
                    "Evidence HEMI 07-08/2026: 4 điểm mù cùng lúc — "
                    "(1) div_avg3 vùng giữa 0.9-1.1 (pos cao 3.5 bị acc cao 3.9 che → mọi rule so-tương-đối mù) · "
                    "(2) un-hedge làm 5 tuần trước ignite (episode B1 chết 09/07, ignite 15/08 — ngoài mọi cửa sổ T−2 khuôn ON và ngoài recent_hedge_days=2) · "
                    "(3) vị thế long ngồi TĨNH từ trước window (không delta → mọi máy đo chuyển động câm) · "
                    "(4) đòn chính = kho 490M nằm sẵn trên sàn. "
                    "CẤM chỉnh A1/B1/GOM-E theo HEMI (p-hack). "
                    "Hai biến thể G-committed-long: BTW div 1.3-1.5, HEMI div 0.9-1.1 + ghế tĩnh. "
                    "Sensor duy trì duy nhất của loài này = on-chain kho (operator). "
                    "Máy đúng tầng: PHASE-TAG bắt S3-HARVEST + TEST-BAIT real-time 18/08 — hệ nhiều tầng đúng thiết kế: tầng nền chịu mù, tầng campaign kêu."),
        "loai": "SPECIES",
        "trang_thai": "ACTIVE — vá bằng §E-R Un-hedge Registry v1.0",
        "hanh_dong_tiep": ("Watch-item candidate STATIC-SEAT (PENDING n=1 HEMI, chờ đối chiếu BTW-July + fw-139 screen). "
                           "Bắt bằng MỨC thay vì DELTA: pos_r ≥2.5-3 bền ≥14d ∧ OI ≤$5M ∧ funding pin ∧ kho lớn."),
        "mm": "H1-20260820",
    },
    # ── L-54: Intel vs entry ──
    {
        "id": "L-54",
        "ngay": "2026-08-20",
        "case_nguon": "CASE_HEMI — luật entry L-54 củng cố",
        "bai_hoc": ("Phát hiện được ≠ giữ được lệnh. Tín hiệu sớm không có chuỗi xác nhận kế tiếp = tín hiệu không giao dịch được. "
                    "Evidence HEMI: giả sử hệ HOÀN HẢO bắt cú un-hedge 05-09/07, lệnh mở tại đó phải ngồi qua 5 tuần giá đè phẳng, "
                    "không tín hiệu trung gian nào (không pos_rerate, không unpin, không vol, không on-chain rút) trước ignite 15/08. "
                    "Không cấu trúc xác nhận = không cấu trúc giữ = thoát trước màn hoặc gồng mù. "
                    "Nguyên tắc: giá trị tín hiệu sớm = tín hiệu × chuỗi xác nhận nối được tới ignite. "
                    "Tín hiệu sớm mồ côi có giá trị INTEL (registry/watchlist) chứ KHÔNG có giá trị ENTRY. "
                    "Hệ quả: luật hiện hành 'entry duy nhất loài staged = S1-CONFIRMED churn ≥8×' được HEMI CỦNG CỐ chứ không bác. "
                    "'Con này không phải của mình' = kết luận đúng, có lời (tiền-không-mất), không phải thất bại hệ."),
        "loai": "ENTRY",
        "trang_thai": "ACTIVE — reinforces existing entry rule",
        "hanh_dong_tiep": ("(1) Tín hiệu sớm không kèm chuỗi xác nhận → xếp INTEL (registry/watchlist), KHÔNG mở entry. "
                           "(2) Entry chỉ mở khi chuỗi xác nhận nối tới ignite proximate (≤3d default). "
                           "(3) Verdict 'không phải của mình' = kết luận hợp lệ, cần thời gian ANALYSIS như positive verdict."),
        "mm": "H1-20260820",
    },
]

for i, L in enumerate(lessons_new):
    row = next_r + i
    ws.cell(row, 1).value = L["id"]
    ws.cell(row, 2).value = L["ngay"]
    ws.cell(row, 3).value = L["case_nguon"]
    ws.cell(row, 4).value = L["bai_hoc"]
    ws.cell(row, 5).value = L["loai"]
    ws.cell(row, 6).value = L["trang_thai"]
    ws.cell(row, 7).value = L["hanh_dong_tiep"]
    ws.cell(row, 8).value = L["mm"]
    print(f"  row {row}: {L['id']} · {L['loai']} · {L['trang_thai'][:40]}")

wb.save(XLSX)
print(f"\nSAVED CASE_INDEX.xlsx — LESSONS sheet now has {ws.max_row} rows total")

# Verify
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2["LESSONS"]
print(f"\n=== Verify — 3 last rows ===")
for r in range(ws2.max_row-2, ws2.max_row+1):
    print(f"row {r}: id={ws2.cell(r,1).value} · loai={ws2.cell(r,5).value} · mm={ws2.cell(r,8).value}")

print("\n[DONE insert]")
