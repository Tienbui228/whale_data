"""
_probe_cookie_c6_crossref_20260809.py
COOKIE C-6: Cross-reference — VPS grep + sổ chéo + peer comparison
"""
import sys, io, json, csv, statistics
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone

ROOT    = Path(r"D:\Data back test\whale_data")
SCRATCH = ROOT / "scratch"
DATE_TAG = "20260809"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"COOKIE C-6 CROSSREF  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

def read_csv(path):
    if not path.exists(): return []
    with open(path, encoding="utf-8", errors="replace") as f:
        return list(csv.DictReader(f))

def gf(d, *keys):
    for k in keys:
        v = d.get(k)
        if v is not None and v != "" and v != "None":
            try: return float(v)
            except: pass
    return None

# ── C-6a: agent_log.md grep for COOKIE mentions ───────────────────────────────
log("\n=== C-6a: AGENT_LOG grep COOKIE ===")
agent_log = ROOT / "agent_log.md"
if agent_log.exists():
    lines = agent_log.read_text(encoding="utf-8", errors="replace").splitlines()
    cookie_mentions = [(i+1, l) for i,l in enumerate(lines)
                       if "COOKIE" in l.upper() or "cookie" in l.lower()]
    log(f"  Total lines: {len(lines)}, COOKIE mentions: {len(cookie_mentions)}")
    for ln, txt in cookie_mentions[-20:]:
        log(f"    L{ln:5d}: {txt[:120]}")
else:
    log("  agent_log.md not found")

# ── C-6b: gom_state_v8 COOKIE chain full detail ───────────────────────────────
log("\n=== C-6b: GOM STATE v8 COOKIE FULL CHAIN ===")
gom_v8 = ROOT / "gom_state_v8.json"
if gom_v8.exists():
    gom_data = json.loads(gom_v8.read_text(encoding="utf-8"))
    cookie_chain = gom_data.get("COOKIEUSDT", {})
    log(f"  chain keys: {list(cookie_chain.keys())}")
    for k,v in cookie_chain.items():
        log(f"  {k}: {v}")
else:
    log("  gom_state_v8.json not found")

# Also check production gom_state.json
gom_prod = ROOT / "gom_state.json"
if gom_prod.exists():
    gom_prod_data = json.loads(gom_prod.read_text(encoding="utf-8", errors="replace"))
    c_prod = gom_prod_data.get("COOKIEUSDT", {})
    log(f"\n  gom_state.json (prod) COOKIEUSDT: {c_prod}")
else:
    log("\n  gom_state.json not found")

# ── C-6c: gom_scan.py GOM-C exact C1 condition ───────────────────────────────
log("\n=== C-6c: GOM-C C1 stealth_grind exact spec ===")
gom_scan = ROOT / "gom_scan.py"
if gom_scan.exists():
    txt = gom_scan.read_text(encoding="utf-8", errors="replace")
    # Find C1 / stealth_grind definition
    lines = txt.splitlines()
    c1_lines = []
    in_gomc = False
    for i, l in enumerate(lines):
        if "GOM-C" in l or "GOM_C" in l or "gomc" in l.lower():
            in_gomc = True
        if in_gomc and ("stealth" in l.lower() or "C1" in l or "c1_" in l.lower()):
            c1_lines.append((i+1, l))
        if in_gomc and len(c1_lines) > 30:
            break
    if c1_lines:
        log(f"  GOM-C C1 stealth_grind lines:")
        for ln, l in c1_lines[:30]:
            log(f"    L{ln}: {l}")
    else:
        log("  C1 / stealth_grind not found in gom_scan.py (may be named differently)")

    # Also grep for C2_quiet, C3_lax
    c23_lines = [(i+1, l) for i,l in enumerate(lines)
                 if "C2_quiet" in l or "C3_lax" in l or "c2_quiet" in l or "c3_lax" in l]
    if c23_lines:
        log(f"\n  C2_quiet / C3_lax lines:")
        for ln, l in c23_lines[:15]:
            log(f"    L{ln}: {l}")
else:
    log("  gom_scan.py not found")

# ── C-6d: Peer comparison (other GOM-C fires from v8) ────────────────────────
log("\n=== C-6d: PEER COMPARISON (GOM-C fires from gom_state_v8) ===")
if gom_v8.exists():
    gom_all = json.loads(gom_v8.read_text(encoding="utf-8"))
    peers = []
    for sym, chain in gom_all.items():
        if isinstance(chain, dict) and chain.get("track") == "gom_c":
            peers.append((sym, chain))
    log(f"  Total GOM-C chains in v8: {len(peers)}")
    log(f"  {'Symbol':20s} {'start':12s} {'close':12s} {'n':>3s} {'Q1_cum':>12s} {'VWAP':>12s} {'open?':>6s}")
    for sym, ch in sorted(peers, key=lambda x: x[1].get("start","") or ""):
        log(f"  {sym:20s} {str(ch.get('start','?')):12s} {str(ch.get('close','—')):12s} "
            f"{ch.get('n',0):>3d} ${ch.get('Q1_cum_usd',0):>11,.0f} "
            f"${ch.get('vwap',0):>11.5f} {'YES' if ch.get('chain_open') else 'no':>6s}")

# ── C-6e: FRESHWATCH consistency check vs probe C-3 numbers ───────────────────
log("\n=== C-6e: FRESHWATCH vs PROBE CONSISTENCY ===")
perp = read_csv(SCRATCH / f"cookie_perp_{DATE_TAG}.csv")
fw_csv = SCRATCH / "freshwatch_daily_2026-08.csv"
perp_by_d = {r["date"]: r for r in perp}

if fw_csv.exists():
    fw_rows = [r for r in read_csv(fw_csv)
               if r.get("symbol","").upper().replace("USDT","") == "COOKIE"]
    log(f"  FW rows: {len(fw_rows)}")
    log(f"  {'date':12s} {'fw_px_avg':>12s} {'perp_close':>12s} {'Δ%':>8s} {'fw_vol':>12s} {'perp_vol':>12s}")
    for r in fw_rows:
        dt = r.get("date","")
        fw_px = gf(r, "px_avg")
        fw_vol = gf(r, "vol_quote")
        p = perp_by_d.get(dt, {})
        px = gf(p, "close")
        vol = gf(p, "vol_quote_usd")
        delta = (fw_px/px - 1)*100 if (fw_px and px) else None
        log(f"  {dt:12s} {fw_px:>12.8f} {px:>12.8f} {f'{delta:+.2f}%' if delta is not None else 'N/A':>8s} "
            f"${fw_vol:>11,.0f} ${vol:>11,.0f}"
            if fw_px and px and fw_vol and vol else f"  {dt:12s} (partial data)")
else:
    log("  FW csv not found in scratch")

# ── C-6f: CASE_INDEX cross-check ─────────────────────────────────────────────
log("\n=== C-6f: CASE_INDEX check ===")
try:
    import openpyxl
    idx_path = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"
    if idx_path.exists():
        wb = openpyxl.load_workbook(str(idx_path), read_only=True)
        log(f"  CASE_INDEX sheets: {wb.sheetnames}")
        ws = wb["INDEX"]
        headers = None
        cookie_row = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = row
                continue
            if row and row[0] and "COOKIE" in str(row[0]).upper():
                cookie_row = row
                break
        if cookie_row:
            log(f"  COOKIE row in INDEX:")
            for h,v in zip(headers, cookie_row):
                if v is not None:
                    log(f"    {h}: {v}")
        else:
            log("  COOKIE not yet in CASE_INDEX (will be added in C-8)")
        wb.close()
    else:
        log("  CASE_INDEX.xlsx not found")
except ImportError:
    log("  openpyxl not available")

log(f"\n[C-6 DONE]")
