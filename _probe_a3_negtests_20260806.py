#!/usr/bin/env python3
"""A3: Negative tests for case_daily.py — 4 conditions.
   (i)  CSV rename       → ERROR export_csv_missing + exit≠0
   (ii) INDEX placeholder→ INFO case_registered_not_open
   (iii) Fake xlsx row   → ERROR case_daily_missing_xlsx + exit 1
   (iv) Orphan file      → WARNING orphan_case_file
   Each test: setup → run → teardown → verdict.
"""
import sys, subprocess, shutil, os
from pathlib import Path
from datetime import datetime, timezone

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

import openpyxl

WHALE      = Path(r"D:\Data back test\whale_data")
CASE_DIR   = WHALE / "docs" / "case_files"
EXPORT_DIR = Path(r"D:\Check onchain fresh wallets\data\export")
SCRIPT     = WHALE / "scripts" / "case_daily.py"

# Use a real date that has export CSV; 2026-08-05 (yesterday)
TEST_DATE  = "2026-08-05"
YM         = "2026-08"
CSV_PATH   = EXPORT_DIR / f"freshwatch_daily_{YM}.csv"
INDEX_PATH = CASE_DIR / "CASE_INDEX.xlsx"
BACKUP_IDX = WHALE / "scratch" / "backup_a3" / f"CASE_INDEX_a3_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.xlsx"

PYTHON = sys.executable

results = []

def run_cd(args: list[str]) -> tuple[int, str]:
    """Run case_daily.py with args, capture combined stdout+stderr, return (exit_code, output)."""
    cmd = [PYTHON, str(SCRIPT)] + args
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                       errors="replace", cwd=str(WHALE))
    return r.returncode, r.stdout + r.stderr

def verdict(label: str, ok: bool, detail: str) -> str:
    s = "PASS" if ok else "FAIL"
    results.append((label, s, detail))
    return f"  [{s}] {label}: {detail}"

print(f"A3 Negative tests — {datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')}")
print(f"TEST_DATE={TEST_DATE}  SCRIPT={SCRIPT.name}")
print("=" * 70)

# ── Backup INDEX before any modifications ────────────────────────────────────
BACKUP_IDX.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(INDEX_PATH, BACKUP_IDX)
print(f"\nINDEX backup → {BACKUP_IDX.name}")

# ════════════════════════════════════════════════════════════════════════════
# Test (i): CSV rename → ERROR export_csv_missing + exit≠0
# ════════════════════════════════════════════════════════════════════════════
print(f"\n{'─'*70}")
print(f"(i) CSV rename → ERROR+exit≠0")
CSV_BAK = CSV_PATH.with_suffix(".csv.bak_a3")
try:
    if not CSV_PATH.exists():
        print(verdict("(i) csv_missing_exit", False, f"SKIP: {CSV_PATH.name} does not exist"))
    else:
        CSV_PATH.rename(CSV_BAK)
        code, out = run_cd(["--date", TEST_DATE, "--dry-run"])
        CSV_BAK.rename(CSV_PATH)  # restore immediately

        has_err  = "export_csv_missing" in out
        nonzero  = (code != 0)
        ok = has_err and nonzero
        print(f"  exit_code={code}  has_err={has_err}")
        if not ok:
            print(f"  OUTPUT (last 20 lines):")
            for ln in out.splitlines()[-20:]:
                print(f"    {ln}")
        print(verdict("(i) csv_missing_exit", ok,
                      f"exit={code} error_in_output={has_err}"))
except Exception as ex:
    if CSV_BAK.exists() and not CSV_PATH.exists():
        CSV_BAK.rename(CSV_PATH)
    print(verdict("(i) csv_missing_exit", False, f"EXCEPTION: {ex}"))

# ════════════════════════════════════════════════════════════════════════════
# Test (ii): INDEX placeholder row → INFO case_registered_not_open
# ════════════════════════════════════════════════════════════════════════════
print(f"\n{'─'*70}")
print(f"(ii) INDEX placeholder row → INFO case_registered_not_open")
try:
    wb = openpyxl.load_workbook(str(INDEX_PATH), data_only=False)
    ws = wb["INDEX"]
    # Find first empty row at/after row 5
    insert_row = None
    for r in range(5, ws.max_row + 2):
        if ws.cell(r, 1).value is None and ws.cell(r, 2).value is None:
            insert_row = r
            break
    if insert_row is None:
        insert_row = ws.max_row + 1

    ws.cell(insert_row, 1).value = "(hang cho)"
    ws.cell(insert_row, 2).value = "TESTTOKEN_A3"
    ws.cell(insert_row, 3).value = "OPEN"
    wb.save(str(INDEX_PATH))
    wb.close()

    code, out = run_cd(["--date", TEST_DATE, "--dry-run"])

    # Restore INDEX
    shutil.copy2(BACKUP_IDX, INDEX_PATH)

    has_info = "case_registered_not_open" in out
    ok = has_info
    print(f"  exit_code={code}  has_info={has_info}")
    if not ok:
        print(f"  OUTPUT (last 20 lines):")
        for ln in out.splitlines()[-20:]:
            print(f"    {ln}")
    print(verdict("(ii) placeholder_info", ok,
                  f"exit={code} info_in_output={has_info}"))
except Exception as ex:
    shutil.copy2(BACKUP_IDX, INDEX_PATH)
    print(verdict("(ii) placeholder_info", False, f"EXCEPTION: {ex}"))

# ════════════════════════════════════════════════════════════════════════════
# Test (iii): Fake .xlsx row → ERROR case_daily_missing_xlsx + exit 1
# ════════════════════════════════════════════════════════════════════════════
print(f"\n{'─'*70}")
print(f"(iii) Fake xlsx row → ERROR+ping+exit 1")
try:
    wb = openpyxl.load_workbook(str(INDEX_PATH), data_only=False)
    ws = wb["INDEX"]
    insert_row = None
    for r in range(5, ws.max_row + 2):
        if ws.cell(r, 1).value is None and ws.cell(r, 2).value is None:
            insert_row = r
            break
    if insert_row is None:
        insert_row = ws.max_row + 1

    ws.cell(insert_row, 1).value = "CASE_FAKE_A3TEST.xlsx"
    ws.cell(insert_row, 2).value = "FAKETOKEN"
    ws.cell(insert_row, 3).value = "OPEN"
    wb.save(str(INDEX_PATH))
    wb.close()

    code, out = run_cd(["--date", TEST_DATE, "--dry-run"])

    # Restore INDEX
    shutil.copy2(BACKUP_IDX, INDEX_PATH)

    has_err  = "case_daily_missing_xlsx" in out
    exit1    = (code == 1)
    ok = has_err and exit1
    print(f"  exit_code={code}  has_err={has_err}")
    if not ok:
        print(f"  OUTPUT (last 20 lines):")
        for ln in out.splitlines()[-20:]:
            print(f"    {ln}")
    print(verdict("(iii) fake_xlsx_exit1", ok,
                  f"exit={code} error_in_output={has_err}"))
except Exception as ex:
    shutil.copy2(BACKUP_IDX, INDEX_PATH)
    print(verdict("(iii) fake_xlsx_exit1", False, f"EXCEPTION: {ex}"))

# ════════════════════════════════════════════════════════════════════════════
# Test (iv): Orphan xlsx in CASE_DIR → WARNING orphan_case_file
# ════════════════════════════════════════════════════════════════════════════
print(f"\n{'─'*70}")
print(f"(iv) Orphan file → WARNING orphan_case_file")
ORPHAN = CASE_DIR / "CASE_ORPHAN_A3TEST.xlsx"
try:
    # Create minimal xlsx
    wb_o = openpyxl.Workbook()
    wb_o.save(str(ORPHAN))
    wb_o.close()

    code, out = run_cd(["--date", TEST_DATE, "--dry-run"])

    ORPHAN.unlink(missing_ok=True)

    has_warn = "orphan_case_file" in out
    ok = has_warn
    print(f"  exit_code={code}  has_warn={has_warn}")
    if not ok:
        print(f"  OUTPUT (last 20 lines):")
        for ln in out.splitlines()[-20:]:
            print(f"    {ln}")
    print(verdict("(iv) orphan_warning", ok,
                  f"exit={code} warning_in_output={has_warn}"))
except Exception as ex:
    ORPHAN.unlink(missing_ok=True)
    print(verdict("(iv) orphan_warning", False, f"EXCEPTION: {ex}"))

# ════════════════════════════════════════════════════════════════════════════
print(f"\n{'='*70}")
print(f"SUMMARY ({datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')})")
all_pass = all(s == "PASS" for _, s, _ in results)
for label, s, detail in results:
    print(f"  [{s}] {label}: {detail}")
print(f"\nOVERALL: {'ALL PASS' if all_pass else 'CÓ FAIL'}")
