"""
_task_register_tst_20260809.py
Register CASE_TST in CASE_INDEX.xlsx + INDEX.md.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font

ROOT     = Path(r"D:\Data back test\whale_data")
CASE_DIR = ROOT / "docs" / "case_files"
DOCS     = ROOT / "docs"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"REGISTER TST  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

# ── 1. CASE_INDEX.xlsx ────────────────────────────────────────────────────────
idx_xlsx = CASE_DIR / "CASE_INDEX.xlsx"
wb = openpyxl.load_workbook(str(idx_xlsx))
ws = wb["INDEX"]

# Find next empty row
max_r = ws.max_row
new_row = max_r + 1
for ri in range(2, max_r + 2):
    v = ws.cell(ri, 1).value
    if v in (None, "") or (str(v).strip() == ""):
        new_row = ri; break

# Check if TST already registered
already = False
for ri in range(1, max_r + 1):
    if "TST" in str(ws.cell(ri, 1).value or ""):
        log(f"  TST already at row {ri}: {[ws.cell(ri,c).value for c in range(1,5)]}")
        already = True; break

if not already:
    entry = [
        "CASE_TST.xlsx",
        "TST",
        "COLD_START — PENDING",
        "Alpha fullyDelisted, ATH=listing day, -94% drawdown. Massive vol spike 07/08 (+20%, 27x, funding negative). Cau hoi: co luc gom hang khong?",
        "COLD_START",
        "2026-08-08",
        "",  # col 7
        "",  # col 8
        "",  # col 9
        "",  # col 10
        "ALPHA_87 | BSC 0x86bb94...666429 | cold-start 08/08/2026",
    ]
    for ci, val in enumerate(entry, 1):
        c = ws.cell(new_row, ci)
        c.value = val
        c.font  = Font(size=9)
    log(f"  CASE_INDEX: row {new_row} written for TST")
    wb.save(str(idx_xlsx))
    log(f"  CASE_INDEX.xlsx saved")
else:
    log(f"  TST already registered — skip")

# ── 2. INDEX.md ───────────────────────────────────────────────────────────────
idx_md = DOCS / "INDEX.md"
if idx_md.exists():
    content = idx_md.read_text(encoding="utf-8")
    new_line = "| `docs/case_files/CASE_TST.xlsx` | TST | COLD_START — PENDING (Alpha fullyDelisted, vol spike 07/08, -94% ATH) |"
    if "CASE_TST" in content:
        log(f"  INDEX.md: TST already present")
    else:
        # Insert after BANANAS31 row
        insert_after = "CASE_BANANAS31.xlsx"
        if insert_after in content:
            lines = content.splitlines()
            new_lines = []
            for ln in lines:
                new_lines.append(ln)
                if insert_after in ln and "BANANAS31" in ln:
                    new_lines.append(new_line)
            content = "\n".join(new_lines)
        else:
            # Append to table section
            content += f"\n{new_line}\n"
        idx_md.write_text(content, encoding="utf-8")
        log(f"  INDEX.md: TST row added")
else:
    log(f"  INDEX.md not found at {idx_md}")

log(f"\n[REGISTER DONE]")
