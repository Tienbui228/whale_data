"""
_task_koma_8b_register_20260810.py
K-8 fix: register KOMA in CASE_INDEX using known row-4 header structure
(same as BROCCOLIF3B case: Row1=title, Row2=desc, Row3=TST entry, Row4=headers, Row5+=data)
"""
import sys, io, json, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT     = Path(r"D:\Data back test\whale_data")
DATE_TAG = "20260810"
SYM_BASE = "KOMA"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"K-8b REGISTER FIX  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

IDX_PATH  = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"
IMDX_PATH = ROOT / "docs" / "case_files" / "INDEX.md"
SCRATCH   = ROOT / "scratch"

# ── Inspect all rows ──────────────────────────────────────────────────────
log("\n=== Inspect all rows ===")
wb = openpyxl.load_workbook(str(IDX_PATH))
ws = wb["INDEX"]
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    log(f"  Row {i}: {[str(c)[:40] if c else None for c in row[:6]]}")

# ── Find header row: look for short 'file'/'token' exact match ────────────
log("\n=== Find header row ===")
header_row_idx = None; headers = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    # Check if first cell is literally 'file' (short, not a sentence)
    first = str(row[0] or "").lower().strip()
    if first == "file":
        header_row_idx = i; headers = list(row); break

log(f"  Header row: {header_row_idx} → {headers}")

if not headers:
    # fallback: scan for row where col values include short keywords
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(c).lower().strip() for c in row if c]
        if "file" in vals and "token" in vals:
            header_row_idx = i; headers = list(row); break
    log(f"  Fallback header row: {header_row_idx} → {headers}")

# Build column map from header
col_map = {}
if headers:
    for ci, h in enumerate(headers, start=1):
        if h:
            col_map[str(h).lower().strip()] = ci
log(f"  col_map: {col_map}")

# ── Check if KOMA already registered ──────────────────────────────────────
log("\n=== Check existing KOMA ===")
existing_row = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i <= (header_row_idx or 0): continue
    row_str = " ".join(str(c) for c in row if c).upper()
    if "KOMA" in row_str:
        existing_row = i
        log(f"  Found KOMA at row {i}: {[str(c)[:30] if c else None for c in row]}")
        break
if not existing_row:
    log(f"  KOMA not registered — appending")

# ── Write row ─────────────────────────────────────────────────────────────
new_row = {
    "file":         f"CASE_{SYM_BASE}.xlsx",
    "token":        SYM_BASE,
    "status":       "OPEN",
    "loai":         "GOM-C PROVISIONAL",
    "action_class": "COLD_START",
    "ngay_mo":      "2026-08-10",
    "t0":           "",
    "moc_ke_tiep":  "depth snap#2 >=05:51 UTC; GOM-C gate 15/08",
    "han":          "",
    "mm_id":        "ALPHA_1",
    "ghi_chu":      "K audit-collect 10/08; pump 07-31 +100%/$604M; OI UNCALIB; GOM-C PROV 07-23+08-02",
}

target_row = existing_row if existing_row else (ws.max_row + 1)
log(f"\n=== Writing to row {target_row} ===")

for col_key, value in new_row.items():
    if col_key in col_map:
        ci = col_map[col_key]
        ws.cell(row=target_row, column=ci, value=value)
        log(f"  col {ci} ({col_key}): {str(value)[:60]}")
    else:
        # Try partial match
        matched = next((hk for hk in col_map if col_key in hk or hk in col_key), None)
        if matched:
            ci = col_map[matched]
            ws.cell(row=target_row, column=ci, value=value)
            log(f"  col {ci} ({matched}≈{col_key}): {str(value)[:60]}")
        else:
            log(f"  SKIP {col_key}")

wb.save(str(IDX_PATH)); wb.close()
log(f"  Saved CASE_INDEX.xlsx")

# ── Verify ────────────────────────────────────────────────────────────────
log("\n=== Verify ===")
wb2 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws2 = wb2["INDEX"]
for i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
    row_str = " ".join(str(c) for c in row if c).upper()
    if "KOMA" in row_str:
        log(f"  Row {i}: {[str(c)[:40] if c else None for c in row]}")
wb2.close()

# ── Regenerate INDEX.md ────────────────────────────────────────────────────
log("\n=== Regenerate INDEX.md ===")
wb3 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws3 = wb3["INDEX"]
all_rows = list(ws3.iter_rows(values_only=True))
wb3.close()

# Find header again in all_rows
hdr_i = None; hdr = None
for i, row in enumerate(all_rows):
    first = str(row[0] or "").lower().strip()
    if first == "file":
        hdr_i = i; hdr = row; break
if hdr_i is None:
    for i, row in enumerate(all_rows):
        vals = [str(c).lower().strip() for c in row if c]
        if "file" in vals and "token" in vals:
            hdr_i = i; hdr = row; break

lines = [f"# FRESHWATCH CASE INDEX", "", f"Generated: {NOW.strftime('%Y-%m-%d %H:%M UTC')}", ""]
if hdr_i is not None:
    data_rows = [r for r in all_rows[hdr_i+1:] if any(c for c in r)]
    hdr_c = [str(c) if c else "" for c in hdr]
    lines.append("| " + " | ".join(hdr_c) + " |")
    lines.append("| " + " | ".join("---" for _ in hdr_c) + " |")
    for row in data_rows:
        cells = [str(c)[:60] if c is not None else "" for c in row]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")

IMDX_PATH.write_text("\n".join(lines), encoding="utf-8")
log(f"  INDEX.md: {IMDX_PATH.stat().st_size:,} bytes")

# ── L-37 backup ───────────────────────────────────────────────────────────
log("\n=== L-37 backup ===")
case_path = ROOT / "docs" / "case_files" / f"CASE_{SYM_BASE}.xlsx"
if case_path.exists():
    bkp = SCRATCH / f"CASE_{SYM_BASE}_bak_reg2_{NOW.strftime('%Y%m%d_%H%M%S')}.xlsx"
    shutil.copy2(case_path, bkp)
    log(f"  Backup: {bkp.name}")

log(f"\n[K-8b DONE]")
