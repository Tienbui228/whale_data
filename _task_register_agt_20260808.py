"""
_task_register_agt_20260808.py
Register CASE_AGT in CASE_INDEX.xlsx + INDEX.md.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font

ROOT     = Path(r"D:\Data back test\whale_data")
CASE_DIR = ROOT / "docs" / "case_files"
DOCS     = ROOT / "docs"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"REGISTER AGT  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

# ── 1. CASE_INDEX.xlsx ────────────────────────────────────────────────────────
idx_xlsx = CASE_DIR / "CASE_INDEX.xlsx"
wb = openpyxl.load_workbook(str(idx_xlsx))
ws = wb["INDEX"]

max_r = ws.max_row
new_row = max_r + 1
for ri in range(2, max_r + 2):
    v = ws.cell(ri, 1).value
    if v in (None, "") or (str(v).strip() == ""):
        new_row = ri; break

already = False
for ri in range(1, max_r + 1):
    if "AGT" in str(ws.cell(ri, 1).value or "") and "CASE_AGT" in str(ws.cell(ri, 1).value or ""):
        log(f"  AGT already at row {ri}")
        already = True; break

if not already:
    entry = [
        "CASE_AGT.xlsx",
        "AGT",
        "COLD_START — PENDING",
        "Alaya Governance Token. Perp onboard 2025-05-20, ATH=2025-05-27 (0.034016). "
        "Drawdown -56%. 19 vol spikes >=8x. gLSR luon >1.5 (retail crowded long). "
        "GOM-A 0 fires. Cau hoi: co luc gom, dang o pha nao?",
        "COLD_START",
        "2026-08-08",
        "",  # col 7
        "",  # col 8
        "",  # col 9
        "",  # col 10
        "ALPHA_191 | BSC 0x5dbde81...e274 | cold-start 08/08/2026",
    ]
    for ci, val in enumerate(entry, 1):
        c = ws.cell(new_row, ci)
        c.value = val
        c.font  = Font(size=9)
    log(f"  CASE_INDEX: row {new_row} written for AGT")
    wb.save(str(idx_xlsx))
    log(f"  CASE_INDEX.xlsx saved")
else:
    log(f"  AGT already registered — skip")

# ── 2. INDEX.md ───────────────────────────────────────────────────────────────
idx_md = DOCS / "INDEX.md"
if idx_md.exists():
    content = idx_md.read_text(encoding="utf-8")
    new_line = "| `docs/case_files/CASE_AGT.xlsx` | AGT | COLD_START — PENDING (Alaya Governance Token, perp 2025-05-20, ATH -56%, 19 vol spikes, gLSR >1.5) |"
    if "CASE_AGT" in content:
        log(f"  INDEX.md: AGT already present")
    else:
        insert_after = "CASE_TST.xlsx"
        if insert_after in content:
            lines = content.splitlines()
            new_lines = []
            for ln in lines:
                new_lines.append(ln)
                if insert_after in ln and "TST" in ln and "|" in ln:
                    new_lines.append(new_line)
            content = "\n".join(new_lines)
        else:
            content += f"\n{new_line}\n"
        idx_md.write_text(content, encoding="utf-8")
        log(f"  INDEX.md: AGT row added")
else:
    log(f"  INDEX.md not found at {idx_md}")

log(f"\n[REGISTER DONE]")
