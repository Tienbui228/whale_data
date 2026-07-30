"""
Append 2 bài học mới vào CASE_INDEX.xlsx sheet LESSONS, sau đó regenerate lessons_learned.md.
DRY-RUN mặc định; truyền --apply để ghi thật.
"""
import sys
import openpyxl
from pathlib import Path

APPLY = "--apply" in sys.argv
IDX = Path("D:/Data back test/whale_data/docs/case_files/CASE_INDEX.xlsx")
MD  = Path("D:/Data back test/whale_data/docs/lessons_learned.md")

wb = openpyxl.load_workbook(IDX, data_only=False)
ws = wb["LESSONS"]

# Tìm row cuối thực tế (có data ở cột A)
last_used = 4  # header là row 4
last_id_num = 0
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v and str(v).startswith("L-"):
        last_used = r
        try:
            last_id_num = int(str(v).split("-")[1])
        except (IndexError, ValueError):
            pass

print(f"Last lesson row: {last_used}, last id: L-{last_id_num:02d}")

new_lessons = [
    [
        f"L-{last_id_num+1:02d}",
        "2026-07-30",
        "BANK",
        "Backfill save nhầm file (_NEW suffix) do file gốc bị lock trong Excel — INDEX vẫn trỏ tên cũ → case_daily ghi vào file cũ (4 dòng) thay vì canonical (35 dòng). Fix: sau bất kỳ backfill nào, verify INDEX trỏ đúng canonical; canonical phải có đúng tên gốc.",
        "process/rule",
        "GHI NHAN",
        "thêm bước verify file name + row count vào checklist sau backfill",
        None,
    ],
    [
        f"L-{last_id_num+2:02d}",
        "2026-07-30",
        "BANK",
        "fn=1 fallback (Zone-A-agg) chỉ dành cho cột taker_buy_pct — KHÔNG mượn sang gLSR/pos_r/acc_r. Metric ratio phải đến từ export CSV hoặc để TRỐNG + comment, không ghép chéo metric. Guard bằng assert EXPORT_COL_MAP tại import time trong case_daily.py.",
        "rule",
        "GHI NHAN",
        "assert đã vào case_daily.py dòng sau FORMULA_COLS",
        None,
    ],
]

print(f"\nSẽ append tại rows {last_used+1} và {last_used+2}:")
for lesson in new_lessons:
    print(f"  {lesson[0]}: {lesson[3][:80]}...")

if APPLY:
    for i, row_data in enumerate(new_lessons):
        r = last_used + 1 + i
        for c, val in enumerate(row_data, start=1):
            ws.cell(r, c).value = val
    wb.save(IDX)
    print(f"\n[APPLY] Saved {IDX}")

    # Regenerate lessons_learned.md
    wb2 = openpyxl.load_workbook(IDX, data_only=True)
    ws2 = wb2["LESSONS"]
    lines = [
        "# Lessons Learned — FreshWatch",
        "",
        "> Auto-generated từ CASE_INDEX.xlsx sheet LESSONS. Không edit file này trực tiếp.",
        "",
    ]
    for r in range(5, ws2.max_row + 1):
        lid = ws2.cell(r, 1).value
        if not lid or not str(lid).startswith("L-"):
            continue
        ngay     = ws2.cell(r, 2).value or "—"
        case_    = ws2.cell(r, 3).value or "—"
        bai_hoc  = ws2.cell(r, 4).value or "—"
        loai     = ws2.cell(r, 5).value or "—"
        status   = ws2.cell(r, 6).value or "—"
        action   = ws2.cell(r, 7).value or "—"
        lines.append(f"## {lid} [{ngay}] {case_} — `{loai}` · {status}")
        lines.append(f"{bai_hoc}")
        lines.append(f"*Hành động tiếp: {action}*")
        lines.append("")
    wb2.close()

    MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"[APPLY] Regenerated {MD} ({len(lines)} lines)")
else:
    wb.close()
    print("\n[DRY-RUN] Không ghi. Chạy với --apply để thực hiện.")
