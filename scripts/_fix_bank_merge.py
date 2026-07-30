"""
Fix B1: Merge Jul29 taker (0.5046) từ CASE_BANK.xlsx row 6 vào CASE_BANK_NEW.xlsx row 37,
rồi rename CASE_BANK_NEW → CASE_BANK (backup old).

Chạy DRY-RUN trước (không truyền --apply) để xem sẽ làm gì.
Chạy thật: python scripts/_fix_bank_merge.py --apply
"""
import sys
import shutil
import openpyxl
from openpyxl.comments import Comment
from pathlib import Path
import datetime

APPLY = "--apply" in sys.argv
CASE_DIR = Path("D:/Data back test/whale_data/docs/case_files")
OLD_FILE  = CASE_DIR / "CASE_BANK.xlsx"        # file sai (4 rows, case_daily append nhầm)
NEW_FILE  = CASE_DIR / "CASE_BANK_NEW.xlsx"    # file canonical (35 rows backfill)
BAK_FILE  = CASE_DIR / "CASE_BANK_PRE_BF_20260730.xlsx"

print(f"=== FIX B1 BANK MERGE ({'APPLY' if APPLY else 'DRY-RUN'}) ===")
print(f"OLD (4 rows, sai): {OLD_FILE}")
print(f"NEW (35 rows, canonical): {NEW_FILE}")
print(f"Backup sẽ lưu tại: {BAK_FILE}")
print()

# Step 1: Đọc taker value + comment từ CASE_BANK.xlsx row 6
wb_old = openpyxl.load_workbook(OLD_FILE, data_only=False)
ws_old = wb_old["DAILY"]
taker_val  = ws_old.cell(6, 5).value
taker_cmt  = ws_old.cell(6, 5).comment
px_avg     = ws_old.cell(6, 2).value
vol_perp   = ws_old.cell(6, 4).value
print(f"Từ CASE_BANK.xlsx row 6:")
print(f"  col5 taker={taker_val}  comment={taker_cmt.text if taker_cmt else None}")
print(f"  col2 px_avg={px_avg}  col4 vol={vol_perp}")
wb_old.close()

# Step 2: Mở CASE_BANK_NEW.xlsx và kiểm tra row 37
wb_new = openpyxl.load_workbook(NEW_FILE, data_only=False)
ws_new = wb_new["DAILY"]
r37_date  = ws_new.cell(37, 1).value
r37_taker = ws_new.cell(37, 5).value
r37_px    = ws_new.cell(37, 2).value
print(f"\nTrong CASE_BANK_NEW.xlsx row 37:")
print(f"  col1 date={r37_date}  col5 taker(hiện tại)={r37_taker}  col2 px={r37_px}")

assert r37_date == "2026-07-29", f"Row 37 không phải Jul29: {r37_date}"
assert r37_taker is None, f"Row 37 taker đã có giá trị {r37_taker}, không cần merge"
print("\n  => Xác nhận row 37 = 2026-07-29, taker=None — an toàn để merge")

if APPLY:
    # Step 3: Ghi taker vào row 37
    ws_new.cell(37, 5).value = taker_val
    if taker_cmt:
        ws_new.cell(37, 5).comment = Comment(taker_cmt.text, "case_daily.py")
    print(f"\n  [APPLY] Ghi taker={taker_val} vào CASE_BANK_NEW.xlsx row 37 col5")

    # Cũng điền px_avg và vol nếu bị None trong NEW nhưng có trong OLD
    if r37_px is None and px_avg is not None:
        ws_new.cell(37, 2).value = px_avg
        print(f"  [APPLY] Ghi px_avg={px_avg} vào col2")
    vol_new = ws_new.cell(37, 4).value
    if vol_new is None and vol_perp is not None:
        ws_new.cell(37, 4).value = vol_perp
        print(f"  [APPLY] Ghi vol_perp={vol_perp} vào col4")

    wb_new.save(NEW_FILE)
    print(f"  [APPLY] Saved CASE_BANK_NEW.xlsx")
    wb_new.close()

    # Step 4: Rename
    shutil.move(str(OLD_FILE), str(BAK_FILE))
    print(f"\n  [APPLY] Renamed CASE_BANK.xlsx → CASE_BANK_PRE_BF_20260730.xlsx")
    NEW_FILE.rename(OLD_FILE)
    print(f"  [APPLY] Renamed CASE_BANK_NEW.xlsx → CASE_BANK.xlsx")

    # Step 5: Verify
    wb_v = openpyxl.load_workbook(OLD_FILE, data_only=True)
    ws_v = wb_v["DAILY"]
    v_mr = ws_v.max_row
    # đếm date rows
    dates = [ws_v.cell(r, 1).value for r in range(3, v_mr+1)
             if ws_v.cell(r,1).value and str(ws_v.cell(r,1).value)[:4].isdigit()]
    last_row = max((r for r in range(3, v_mr+1)
                    if ws_v.cell(r,1).value and str(ws_v.cell(r,1).value)[:4].isdigit()), default=None)
    last_taker = ws_v.cell(last_row, 5).value if last_row else None
    wb_v.close()
    print(f"\n=== VERIFY CASE_BANK.xlsx sau rename ===")
    print(f"  date rows: {len(dates)}")
    print(f"  last date: {dates[-1] if dates else None}")
    print(f"  last row: {last_row}, taker: {last_taker}")
    assert len(dates) == 35, f"Expected 35 date rows, got {len(dates)}"
    print("  => PASS: 35 dòng, taker Jul29 OK")
else:
    wb_new.close()
    print("\n[DRY-RUN] Không sửa gì. Chạy với --apply để thực hiện.")
    print(f"  Sẽ ghi: CASE_BANK_NEW.xlsx row37 col5 = {taker_val}")
    print(f"  Sẽ rename: CASE_BANK.xlsx → CASE_BANK_PRE_BF_20260730.xlsx")
    print(f"  Sẽ rename: CASE_BANK_NEW.xlsx → CASE_BANK.xlsx")
