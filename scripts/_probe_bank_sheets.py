"""
Probe: kiểm tra toàn bộ sheet trong CASE_BANK.xlsx — phát hiện sheet thừa,
row count, first/last date mỗi sheet, và đường dẫn tuyệt đối file đang mở.
"""
import os
import sys
from pathlib import Path
import openpyxl

TARGET = Path("D:/Data back test/whale_data/docs/case_files/CASE_BANK.xlsx")
print(f"File tuyệt đối: {TARGET.resolve()}")
print(f"Exists: {TARGET.exists()}")
print(f"mtime: {os.path.getmtime(TARGET):.0f}")
import datetime
print(f"mtime human: {datetime.datetime.fromtimestamp(os.path.getmtime(TARGET))}")
print()

wb = openpyxl.load_workbook(TARGET, data_only=True)
print(f"wb.sheetnames = {wb.sheetnames}")
print()

for sname in wb.sheetnames:
    ws = wb[sname]
    mr = ws.max_row
    mc = ws.max_column
    print(f"--- Sheet: '{sname}' ---")
    print(f"  max_row={mr}  max_col={mc}")
    # Tìm first/last date ở cột A (bỏ qua header)
    dates = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        val = row[0]
        if val is not None:
            dates.append(val)
    if dates:
        print(f"  cột A first value: {dates[0]}")
        print(f"  cột A last value : {dates[-1]}")
        print(f"  số dòng data (cột A not None): {len(dates)}")
    else:
        print(f"  cột A: không có giá trị (sheet rỗng hoặc không có ngày)")
    # In header row
    headers = [ws.cell(row=1, column=c).value for c in range(1, mc+1)]
    print(f"  row 1 headers: {headers}")
    # In 3 dòng đầu data
    for r in range(2, min(5, mr+1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(mc+1, 8))]
        print(f"  row {r}: {row_vals}")
    # In dòng cuối
    if mr >= 5:
        row_vals = [ws.cell(row=mr, column=c).value for c in range(1, min(mc+1, 8))]
        print(f"  row {mr} (last): {row_vals}")
    print()

wb.close()
