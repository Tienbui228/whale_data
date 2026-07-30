"""Probe: kiểm tra CASE_INDEX.xlsx xem BANK file name là gì."""
import openpyxl
from pathlib import Path

IDX = Path("D:/Data back test/whale_data/docs/case_files/CASE_INDEX.xlsx")
wb = openpyxl.load_workbook(IDX, data_only=True)
ws = wb["INDEX"]
print(f"INDEX max_row={ws.max_row}")
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 6)]
    if any(v is not None for v in vals):
        if any("BANK" in str(v).upper() for v in vals if v):
            print(f"  row {r}: {vals}  *** BANK ***")
        else:
            print(f"  row {r}: {vals}")
wb.close()
