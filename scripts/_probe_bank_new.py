"""
Probe: kiểm tra CASE_BANK_NEW.xlsx — xem có phải file backfill 35 dòng không.
"""
import os
from pathlib import Path
import openpyxl
import datetime

TARGET = Path("D:/Data back test/whale_data/docs/case_files/CASE_BANK_NEW.xlsx")
print(f"File: {TARGET.resolve()}")
print(f"Exists: {TARGET.exists()}")
print(f"Size: {TARGET.stat().st_size} bytes")
print(f"mtime: {datetime.datetime.fromtimestamp(os.path.getmtime(TARGET))}")
print()

wb = openpyxl.load_workbook(TARGET, data_only=True)
print(f"sheetnames = {wb.sheetnames}")
print()

for sname in wb.sheetnames:
    ws = wb[sname]
    mr = ws.max_row
    mc = ws.max_column
    print(f"--- Sheet: '{sname}' max_row={mr} max_col={mc} ---")
    if sname == "DAILY":
        # in header row 1 và 2
        h1 = [ws.cell(row=1, column=c).value for c in range(1, min(mc+1, 6))]
        h2 = [ws.cell(row=2, column=c).value for c in range(1, min(mc+1, 6))]
        print(f"  row1: {h1}")
        print(f"  row2 (col headers): {h2}")
        # đếm dòng date
        dates = []
        for row in ws.iter_rows(min_row=3, max_col=1, values_only=True):
            val = row[0]
            if val is not None:
                dates.append(str(val))
        print(f"  data rows (from row3): {len(dates)}")
        if dates:
            print(f"  first date: {dates[0]}")
            print(f"  last date : {dates[-1]}")
        # in row cuối thực tế
        for r in range(mr, max(mr-3, 2), -1):
            rv = [ws.cell(row=r, column=c).value for c in range(1, min(mc+1, 8))]
            if any(v is not None for v in rv):
                print(f"  row {r}: {rv}")
        # check col O (col 15) = pos_r at Jul20
        print(f"  Checking col O (pos_r) around row 37...")
        for r in range(2, mr+1):
            date_val = ws.cell(row=r, column=1).value
            if str(date_val) == "2026-07-20":
                col_o = ws.cell(row=r, column=15).value
                print(f"  row {r} date=2026-07-20 col_O(pos_r)={col_o}")
                break
wb.close()
