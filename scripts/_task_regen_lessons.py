#!/usr/bin/env python3
"""Regenerate docs/lessons_learned.md from CASE_INDEX.xlsx LESSONS sheet."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime
import openpyxl

IDX   = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
MDOUT = Path(r"D:\Data back test\whale_data\docs\lessons_learned.md")

wb = openpyxl.load_workbook(str(IDX), data_only=True)
ws = wb["LESSONS"]

lessons = []
for r in range(5, ws.max_row + 1):
    lid       = ws.cell(r, 1).value
    ngay      = ws.cell(r, 2).value
    case_src  = ws.cell(r, 3).value
    bai_hoc   = ws.cell(r, 4).value
    loai      = ws.cell(r, 5).value
    trang_thai = ws.cell(r, 6).value
    hanh_dong = ws.cell(r, 7).value

    if not lid or not str(lid).startswith("L-"):
        continue
    lessons.append((lid, ngay, case_src, bai_hoc, loai, trang_thai, hanh_dong))

print(f"Found {len(lessons)} lessons (L-01 → {lessons[-1][0] if lessons else '?'})")

lines = [
    "# Lessons Learned — FreshWatch",
    "",
    "> Auto-generated tu CASE_INDEX.xlsx sheet LESSONS. Khong edit file nay truc tiep.",
    f"> Lan sinh: {datetime.now().strftime('%Y-%m-%d')} · Tong: {len(lessons)} bai hoc ({lessons[0][0] if lessons else '—'} → {lessons[-1][0] if lessons else '—'})",
    "",
]

for lid, ngay, case_src, bai_hoc, loai, trang_thai, hanh_dong in lessons:
    date_s    = str(ngay or "")
    case_s    = str(case_src or "")
    loai_s    = str(loai or "")
    status_s  = str(trang_thai or "")
    bai_s     = str(bai_hoc or "")
    action_s  = str(hanh_dong or "")

    lines.append(f"## {lid} [{date_s}] {case_s} — `{loai_s}` · {status_s}")
    lines.append(bai_s)
    if action_s:
        lines.append(f"*Hanh dong tiep: {action_s}*")
    lines.append("")

wb.close()
MDOUT.write_text("\n".join(lines), encoding="utf-8")
print(f"Written: {MDOUT} ({len(lines)} lines)")
