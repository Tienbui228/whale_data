"""Probe: xem cấu trúc sheet LESSONS trong CASE_INDEX.xlsx."""
import openpyxl
from pathlib import Path

IDX = Path("D:/Data back test/whale_data/docs/case_files/CASE_INDEX.xlsx")
wb = openpyxl.load_workbook(IDX, data_only=True)
print(f"Sheets: {wb.sheetnames}")
if "LESSONS" in wb.sheetnames:
    ws = wb["LESSONS"]
    print(f"LESSONS max_row={ws.max_row} max_col={ws.max_column}")
    for r in range(1, min(ws.max_row+1, 10)):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
        if any(v is not None for v in vals):
            print(f"  row {r}: {vals}")
wb.close()
