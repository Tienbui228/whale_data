# -*- coding: utf-8 -*-
"""Check CASE_IDOL column structure + last data row."""
import openpyxl
wb = openpyxl.load_workbook("docs/case_files/CASE_IDOL.xlsx")
ws = wb["DAILY"]
print("Row 2 (col names):")
for c in ws[2]:
    print(f"  col {c.column:2d}: {c.value!r}")
print("\nRow 3 (first data - raw value, number_format):")
for c in ws[3]:
    if c.value is not None:
        print(f"  col {c.column:2d}: value={c.value!r:20} fmt={c.number_format!r}")
print("\nLast row values:")
last = ws.max_row
for c in ws[last]:
    if c.value is not None:
        print(f"  col {c.column:2d}: {c.value!r}")
# EVENTS last row
ws_e = wb["EVENTS"]
print(f"\nEVENTS max_row: {ws_e.max_row}")
for row in ws_e.iter_rows(min_row=1, max_row=ws_e.max_row, values_only=True):
    if any(v is not None for v in row):
        print("  ROW:", [str(v)[:40] if v else "" for v in row])
