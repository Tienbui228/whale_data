"""
Probe CASE_TUT.xlsx structure — 2026-08-06
"""
import openpyxl

PATH = r"D:\Data back test\whale_data\docs\case_files\CASE_TUT.xlsx"

wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)

print("=" * 60)
print("1. SHEET NAMES")
print("=" * 60)
for name in wb.sheetnames:
    print(f"  {name}")

# ── DAILY ──────────────────────────────────────────────────────
print("\n" + "=" * 60)
print("2. SHEET: DAILY")
print("=" * 60)
if "DAILY" in wb.sheetnames:
    ws = wb["DAILY"]
    rows = list(ws.iter_rows(values_only=True))
    if rows:
        headers = rows[0]
        print(f"Headers (row 1): {list(headers)}")
        data_rows = [r for r in rows[1:] if any(c is not None for c in r)]
        print(f"Data rows (non-empty): {len(data_rows)}")
    else:
        print("  (empty sheet)")
else:
    print("  Sheet DAILY not found")

# ── CASE ───────────────────────────────────────────────────────
print("\n" + "=" * 60)
print("3. SHEET: CASE")
print("=" * 60)
if "CASE" in wb.sheetnames:
    ws = wb["CASE"]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows:
        if any(c is not None for c in r):
            label = r[0] if len(r) > 0 else None
            value = r[1] if len(r) > 1 else None
            print(f"  {label!r:40s} | {value!r}")
else:
    print("  Sheet CASE not found")

# ── TRIGGERS ───────────────────────────────────────────────────
print("\n" + "=" * 60)
print("4. SHEET: TRIGGERS")
print("=" * 60)
if "TRIGGERS" in wb.sheetnames:
    ws = wb["TRIGGERS"]
    rows = list(ws.iter_rows(values_only=True))
    if rows:
        print(f"Headers (row 1): {list(rows[0])}")
        data_rows = [r for r in rows[1:] if any(c is not None for c in r)]
        print(f"Data rows (non-empty): {len(data_rows)}")
        if data_rows:
            print("First 3 data rows:")
            for r in data_rows[:3]:
                print(f"  {list(r)}")
    else:
        print("  (empty sheet)")
else:
    print("  Sheet TRIGGERS not found")

# ── OUTCOME ────────────────────────────────────────────────────
print("\n" + "=" * 60)
print("5. SHEET: OUTCOME")
print("=" * 60)
if "OUTCOME" in wb.sheetnames:
    ws = wb["OUTCOME"]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows:
        if any(c is not None for c in r):
            label = r[0] if len(r) > 0 else None
            value = r[1] if len(r) > 1 else None
            print(f"  {label!r:40s} | {value!r}")
else:
    print("  Sheet OUTCOME not found")

# ── MM_PROFILE ─────────────────────────────────────────────────
print("\n" + "=" * 60)
print("6. SHEET: MM_PROFILE")
print("=" * 60)
if "MM_PROFILE" in wb.sheetnames:
    ws = wb["MM_PROFILE"]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(c is not None for c in r)]
    print(f"Total non-empty rows: {len(non_empty)}")
    print("First 10 rows:")
    for r in non_empty[:10]:
        print(f"  {list(r)}")
else:
    print("  Sheet MM_PROFILE not found")

# ── UNLOCK ─────────────────────────────────────────────────────
print("\n" + "=" * 60)
print("7. SHEET: UNLOCK")
print("=" * 60)
if "UNLOCK" in wb.sheetnames:
    ws = wb["UNLOCK"]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(c is not None for c in r)]
    print(f"Total non-empty rows: {len(non_empty)}")
    print("First 10 rows:")
    for r in non_empty[:10]:
        print(f"  {list(r)}")
else:
    print("  Sheet UNLOCK not found")

wb.close()
print("\n" + "=" * 60)
print("PROBE DONE")
print("=" * 60)
