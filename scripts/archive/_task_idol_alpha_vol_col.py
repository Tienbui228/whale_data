#!/usr/bin/env python3
"""Patch: thêm col vol_alpha_usd (c42) + n_trades_alpha (c43) vào DAILY sheet."""
import sys, io, json, urllib.request
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pathlib import Path
from datetime import datetime, timezone as tz
import openpyxl
from openpyxl.cell.cell import MergedCell

WHALE    = Path(r'D:\Data back test\whale_data')
IDX_PATH = WHALE / 'docs' / 'case_files' / 'CASE_IDOL.xlsx'

# Fetch lại data
SYMBOL = 'ALPHA_223USDT'
start_ms = int(datetime(2026, 6, 1, tzinfo=tz.utc).timestamp() * 1000)
url = (f'https://www.binance.com/bapi/defi/v1/public/alpha-trade/klines'
       f'?symbol={SYMBOL}&interval=1d&startTime={start_ms}&limit=90')
req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0', 'Accept': '*/*'})
with urllib.request.urlopen(req, timeout=15) as r:
    rows_raw = json.loads(r.read())
if isinstance(rows_raw, dict): rows_raw = rows_raw.get('data', [])

alpha_daily = {}
for k in rows_raw:
    if not isinstance(k, (list, tuple)) or len(k) < 8: continue
    dt = datetime.fromtimestamp(int(k[0]) / 1000, tz.utc).strftime('%Y-%m-%d')
    alpha_daily[dt] = {'quote_vol': float(k[7] or 0), 'n_trades': int(k[8]) if len(k) > 8 else 0}

wb = openpyxl.load_workbook(str(IDX_PATH))
ws = wb['DAILY']

# Add headers at c42, c43
VCOL = 42; NCOL = 43
hdr_row = 2
h42 = ws.cell(hdr_row, VCOL)
h43 = ws.cell(hdr_row, NCOL)
if not isinstance(h42, MergedCell): h42.value = 'vol_alpha_usd'
if not isinstance(h43, MergedCell): h43.value = 'n_trades_alpha'
print(f'Headers: c{VCOL}=vol_alpha_usd  c{NCOL}=n_trades_alpha')

written = 0
for r in range(3, ws.max_row + 1):
    date_cell = ws.cell(r, 1)
    if isinstance(date_cell, MergedCell) or not date_cell.value: continue
    d_raw = date_cell.value
    d_str = d_raw.strftime('%Y-%m-%d') if isinstance(d_raw, datetime) else str(d_raw).strip()[:10]
    row = alpha_daily.get(d_str)
    if not row: continue
    c42 = ws.cell(r, VCOL)
    c43 = ws.cell(r, NCOL)
    if not isinstance(c42, MergedCell): c42.value = round(row['quote_vol'], 0)
    if not isinstance(c43, MergedCell): c43.value = row['n_trades']
    written += 1
    print(f'  {d_str}: vol=${row["quote_vol"]/1e6:.3f}M  n={row["n_trades"]:,}')

wb.save(str(IDX_PATH))
print(f'\nWritten {written} rows → SAVED')
