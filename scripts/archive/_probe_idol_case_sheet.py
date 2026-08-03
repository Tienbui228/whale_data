#!/usr/bin/env python3
"""Quick probe: CASE sheet structure in CASE_IDOL.xlsx."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pathlib import Path
import openpyxl
from openpyxl.cell.cell import MergedCell

IDX_PATH = Path(r'D:\Data back test\whale_data\docs\case_files\CASE_IDOL.xlsx')
wb = openpyxl.load_workbook(str(IDX_PATH), data_only=False)
wc = wb['CASE']

print(f'CASE sheet: max_row={wc.max_row} max_col={wc.max_column}')
for r in range(1, min(wc.max_row + 1, 60)):
    row_data = []
    for c in range(1, min((wc.max_column or 10) + 1, 11)):
        cell = wc.cell(r, c)
        if isinstance(cell, MergedCell): continue
        v = cell.value
        if v is not None:
            row_data.append(f'c{c}={str(v)[:50]!r}')
    if row_data:
        print(f'  row{r}: ' + '  '.join(row_data))
