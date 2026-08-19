#!/usr/bin/env python3
"""[1][2] CASE_GUA finalize: EVENTS row 49 + col52 rename + CASE note.

L-37 backup done pre-run (scratch/CASE_GUA_bak_pre_finalize_20260819.xlsx).
L-50 write-after-guard: max_col unchanged, max_row EVENTS +1, DAILY unchanged,
CASE +2.
"""
import sys, io
from pathlib import Path
import openpyxl

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_GUA.xlsx")

wb = openpyxl.load_workbook(str(XLSX), data_only=False)

# ─── snapshot before ─────────────────────────────────────────────────────
snap_before = {}
for s in ["CASE", "DAILY", "EVENTS"]:
    ws = wb[s]
    snap_before[s] = (ws.max_row, ws.max_column)
print(f"[before] {snap_before}")

# ═══ [1] EVENTS append row 49 ════════════════════════════════════════════
ws_e = wb["EVENTS"]
row_e = ws_e.max_row + 1
ev_row = (
    "2026-08-19",
    "GATE_ISOLATED_DEPOSIT_DISABLED",
    "deposit_disabled=TRUE → pool trapped, arb-bridge cắt, premium 4.5× cấu trúc. "
    "Depth $215 ±2%. Gate KHÔNG phải van xả — loại khỏi băng-thông vĩnh viễn.",
    "gua_identity_forensics_20260819_F7",
    "closed",
)
for i, v in enumerate(ev_row, 1):
    ws_e.cell(row_e, i).value = v
print(f"[1] EVENTS row {row_e}: appended GATE_ISOLATED_DEPOSIT_DISABLED")

# ═══ [2] DAILY col52 rename all 'gate' → 'gate[ISOLATED-DEPOSIT-DISABLED]' ═
ws_d = wb["DAILY"]
n_renamed = 0
for r in range(3, ws_d.max_row + 1):
    v = ws_d.cell(r, 52).value
    if v == "gate":
        ws_d.cell(r, 52).value = "gate[ISOLATED-DEPOSIT-DISABLED]"
        n_renamed += 1
print(f"[2a] DAILY col52 renamed 'gate' → 'gate[ISOLATED-DEPOSIT-DISABLED]': {n_renamed} rows")

# ═══ [2b] CASE sheet append note rows 22-23 ══════════════════════════════
ws_c = wb["CASE"]
note_rows = [
    ("Bandwidth EXCLUDE flag",
     "gate[ISOLATED-DEPOSIT-DISABLED] loại khỏi mọi phép băng-thông-xả"),
    ("L-52 vế đường-hàng-thông",
     "Luật: pair tồn tại ∧ deposit đóng = van vẽ trên giấy. "
     "EXCLUDE cờ ISOLATED/SUSPECT khỏi bandwidth. Ref: EVENTS r49 + forensics F7."),
]
for i, (c1, c2) in enumerate(note_rows):
    r = ws_c.max_row + 1
    ws_c.cell(r, 1).value = c1
    ws_c.cell(r, 2).value = c2
    print(f"[2b] CASE row {r}: c1={c1!r} c2={c2[:50]!r}...")

# ─── L-50 write-after-guard ──────────────────────────────────────────────
snap_after = {}
for s in ["CASE", "DAILY", "EVENTS"]:
    ws = wb[s]
    snap_after[s] = (ws.max_row, ws.max_column)
print(f"\n[L-50 guard] before={snap_before} after={snap_after}")

assert snap_after["EVENTS"] == (snap_before["EVENTS"][0] + 1, snap_before["EVENTS"][1]), \
    f"L-50 EVENTS: expected +1 row unchanged col, got {snap_after['EVENTS']}"
assert snap_after["DAILY"] == snap_before["DAILY"], \
    f"L-50 DAILY: shape must be unchanged, got {snap_after['DAILY']}"
assert snap_after["CASE"] == (snap_before["CASE"][0] + 2, snap_before["CASE"][1]), \
    f"L-50 CASE: expected +2 rows unchanged col, got {snap_after['CASE']}"
assert snap_after["DAILY"][1] >= 52, f"L-50 col52 guard: got {snap_after['DAILY'][1]}"
print(f"[L-50 guard] PASS")

wb.save(str(XLSX))
print(f"[save] {XLSX.name}")

# ─── verify: reopen data_only ────────────────────────────────────────────
wb2 = openpyxl.load_workbook(str(XLSX), data_only=True, read_only=True)
ws2_e = wb2["EVENTS"]
print(f"\n[verify] EVENTS row {ws2_e.max_row}:")
for c in range(1, 6):
    v = ws2_e.cell(ws2_e.max_row, c).value
    print(f"  col {c}: {str(v)[:70]!r}")

ws2_d = wb2["DAILY"]
n_flag = 0; n_bare = 0
for r in range(3, ws2_d.max_row + 1):
    v = ws2_d.cell(r, 52).value
    if v == "gate":
        n_bare += 1
    elif v == "gate[ISOLATED-DEPOSIT-DISABLED]":
        n_flag += 1
print(f"\n[verify] DAILY col52 counts: flagged={n_flag}, bare 'gate' remaining={n_bare}")
assert n_bare == 0, f"L-50 col52 verify fail: {n_bare} bare 'gate' remain"

ws2_c = wb2["CASE"]
print(f"\n[verify] CASE last 3 rows (r{ws2_c.max_row-2}..{ws2_c.max_row}):")
for r in range(ws2_c.max_row - 2, ws2_c.max_row + 1):
    v1 = ws2_c.cell(r, 1).value
    v2 = ws2_c.cell(r, 2).value
    print(f"  r{r}: c1={str(v1)[:30]!r}  c2={str(v2)[:60]!r}")
wb2.close()
print("\n[DONE]")
