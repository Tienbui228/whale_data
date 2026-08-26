#!/usr/bin/env python3
"""[BATCH V3 SYNC 27/08] Append L-62/63/64/65 vào CASE_INDEX LESSONS + regen md.

Operator "ok batch lessons" 27/08 → sync 4 new rows 67-70.
"""
from __future__ import annotations
import sys, io, hashlib, shutil, subprocess
from pathlib import Path
from datetime import datetime
import openpyxl
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

REPO = Path(r"D:/Data back test/whale_data")
IDX = REPO / "docs/case_files/CASE_INDEX.xlsx"
BAK = REPO / "docs/case_files/CASE_INDEX.xlsx.bak_pre_batchV3_20260827"

def _md5(p):
    h = hashlib.md5()
    with open(p, "rb") as f:
        for c in iter(lambda: f.read(1<<16), b""): h.update(c)
    return h.hexdigest()

print("═" * 90)
print(" BATCH V3 SYNC — CASE_INDEX LESSONS append L-62/63/64/65 rows 67-70")
print("═" * 90)

# Pre-sync
print(f"\n[PRE-SYNC]")
print(f"  file: {IDX}")
print(f"  mtime: {datetime.fromtimestamp(IDX.stat().st_mtime)}")
print(f"  md5:   {_md5(IDX)}")

# Backup
if not BAK.exists():
    shutil.copy2(IDX, BAK)
    print(f"  backup: {BAK.name}  md5={_md5(BAK)}")
else:
    print(f"  backup already exists: {BAK.name}")

# Load workbook (writable)
wb = openpyxl.load_workbook(str(IDX))
ws = wb["LESSONS"]
print(f"\n  Current LESSONS rows: {ws.max_row}")

# Prepare 4 new rows
# Schema (row 4 header): id | ngay | case_nguon | bai_hoc | loai | trang_thai | hanh_dong_tiep | MM
new_rows = [
    {
        "id": "L-62",
        "ngay": "2026-08-27",
        "case_nguon": "ERRATA-BAND-20260827 (⑬-CLOSE probe 5b + Scope-A hardcode band mismatch)",
        "bai_hoc": ("Đường-đọc-thật: số/nhãn/verdict do probe/gate xuất ra CHỈ hợp lệ khi declare "
                    "đường-đọc-thật của mọi ngưỡng dùng trong tính toán. Không declare = thước tách "
                    "khỏi số = bảng vô hiệu. Rule bắt buộc header cfg-provenance: Nguồn thresholds "
                    "(yaml path | hardcode | free-layer) + Giá trị THỰC từng ngưỡng dùng (không dựa "
                    "comment) + Timestamp yaml mtime nếu source = yaml (staleness detect). "
                    "Evidence: ERRATA-BAND-20260827 d353dc4 probe hardcode 0.505 vs prod 53.0."),
        "loai": "RULE-DISCIPLINE",
        "trang_thai": "SIGNED batch V3 27/08 → sheet 67",
        "hanh_dong_tiep": ("Rà lại hardcode probe/task; thay yaml load + provenance header. "
                          "Pre-commit hook 'no hardcoded thresholds' pending scope ký."),
        "MM": None,
    },
    {
        "id": "L-63",
        "ngay": "2026-08-27",
        "case_nguon": "⑩-AMENDMENT chuỗi 24-26/08 (⑩ stub → 3 vế mở rộng, R4-v2 avg_px scope-miss)",
        "bai_hoc": ("SCOPE-BY-MANIFEST: scope regen/canonicalize/migrate định nghĩa bằng DANH SÁCH "
                    "BIẾN CANONICAL BỊ ẢNH HƯỞNG (fields cụ thể: taker_buy_pct, avg_px, lsr_pos, ...) "
                    "và block chứa nó (KLINES/RATIO/OI), KHÔNG bằng tên slice mờ ('taker slice', "
                    "'LSR slice'). Manifest bảng per (block, col, source, consumer, canonical, verdict) "
                    "làm nguồn duy nhất; xung đột → CONFLICT column, DỪNG hỏi. "
                    "Evidence: ⑩ initial stub → R4-v1 chạm RATIO taker → B3 fail hội tụ do avg_px "
                    "(KLINES) chưa regen → R4-v2 mở rộng scope. Manifest v1 SIGNED 4cd6ec3."),
        "loai": "SCOPE-DISCIPLINE",
        "trang_thai": "SIGNED batch V3 27/08 → sheet 68",
        "hanh_dong_tiep": ("Mọi ⑥ regen/⑩ migrate/canonicalize ⇒ manifest bảng trước; verdict per col: "
                          "REGEN / VERIFY-THEN-KEEP / DEPRECATE-RENAME / N/A-BLOCK. Missing col = CONFLICT DỪNG."),
        "MM": None,
    },
    {
        "id": "L-64",
        "ngay": "2026-08-27",
        "case_nguon": "ERRATA-BAND-20260827 fix 2 probe hardcode",
        "bai_hoc": ("cfg-provenance header: mọi probe/gate BẮT BUỘC in header cfg-provenance tại đầu "
                    "run: [CFG-PROVENANCE <artifact>] Source: <yaml path | hardcode | free-layer> · "
                    "Method: <load method> · Values ACTUALLY USED: <threshold_name> = <actual_value> "
                    "[comment: source line]. Kèm assertions fail-loud khi mismatch. "
                    "Không header = bảng vô hiệu (L-62 companion). "
                    "Evidence: ERRATA-BAND d353dc4 fix _probe_5b_consumer_diff + _task_backfill_apply."),
        "loai": "PROBE-DISCIPLINE",
        "trang_thai": "SIGNED batch V3 27/08 → sheet 69",
        "hanh_dong_tiep": ("Template header in gom_scan.load_thresholds() wrapper; probe/task print auto. "
                          "Lint rule pending scope ký."),
        "MM": None,
    },
    {
        "id": "L-65",
        "ngay": "2026-08-27",
        "case_nguon": "ERRATA-BAND-20260827 subset check FAIL 8+27 extras + ARIA/STAR spot",
        "bai_hoc": ("Flip-set XOR hai chiều (band-change không đơn điệu): rule cũ 'band nới chỉ bớt "
                    "flip' = GIẢ ĐỊNH SAI (L-28 DEV-instance). Rule ĐÚNG: flip-set dưới band-change "
                    "là XOR HAI CHIỀU — mọi thay đổi band ⇒ RERUN toàn bộ bằng chứng flip, subset "
                    "không được giả định. Mechanism: wider band = more 7d-window all-pass cases → "
                    "flip transition NEW cases không tồn tại narrow band. "
                    "Evidence: ERRATA subset FAIL (5b 8 extras, Scope-A 27 extras); ARIA 07-30 "
                    "canonical crosses 0.53 boundary; STAR 06-28 canonical 06-24=0.5304 > 0.53 → "
                    "PROD flip leg=T can=F. Documentation phải note 'subset không được giả định' "
                    "cạnh mọi flip-count report."),
        "loai": "RULE-CORRECTION",
        "trang_thai": "SIGNED batch V3 27/08 → sheet 70",
        "hanh_dong_tiep": ("R6 skeleton vocabulary + E2-backlog note updated. Mọi flip-report post-⑬ "
                          "dùng rule mới."),
        "MM": None,
    },
]

start_row = ws.max_row + 1
print(f"\n[SYNC] Append 4 rows starting at row {start_row}")
for i, r in enumerate(new_rows):
    row_num = start_row + i
    ws.cell(row_num, 1).value = r["id"]
    ws.cell(row_num, 2).value = r["ngay"]
    ws.cell(row_num, 3).value = r["case_nguon"]
    ws.cell(row_num, 4).value = r["bai_hoc"]
    ws.cell(row_num, 5).value = r["loai"]
    ws.cell(row_num, 6).value = r["trang_thai"]
    ws.cell(row_num, 7).value = r["hanh_dong_tiep"]
    ws.cell(row_num, 8).value = r["MM"]
    print(f"  row {row_num}: {r['id']}  {r['loai']}  {r['trang_thai']}")

# Save
wb.save(str(IDX))
print(f"\n[POST-SYNC]")
print(f"  mtime: {datetime.fromtimestamp(IDX.stat().st_mtime)}")
print(f"  md5:   {_md5(IDX)}")

# Verify — reload and read rows 67-70
wb2 = openpyxl.load_workbook(str(IDX), data_only=True)
ws2 = wb2["LESSONS"]
print(f"\n[VERIFY] Reload rows 67-70:")
for row_num in range(start_row, start_row + 4):
    vals = [ws2.cell(row_num, c).value for c in range(1, 9)]
    print(f"  row {row_num}: {vals[0]} · {vals[1]} · loai={vals[4]} · status={vals[5]}")

# Regen lessons_learned.md
print(f"\n[REGEN lessons_learned.md]")
lessons_md = REPO / "docs/lessons_learned.md"
md_pre_mtime = datetime.fromtimestamp(lessons_md.stat().st_mtime) if lessons_md.exists() else None
md_pre_md5 = _md5(lessons_md) if lessons_md.exists() else None
print(f"  pre-regen mtime: {md_pre_mtime}  md5: {md_pre_md5}")

# Run regen tool
r = subprocess.run(["python", str(REPO / "scripts/_task_regen_lessons.py")],
                   capture_output=True, text=True)
print(f"  regen rc: {r.returncode}")
if r.stdout: print(f"  stdout: {r.stdout[:200]}")
if r.stderr: print(f"  stderr: {r.stderr[:200]}")

md_post_mtime = datetime.fromtimestamp(lessons_md.stat().st_mtime)
md_post_md5 = _md5(lessons_md)
print(f"  post-regen mtime: {md_post_mtime}  md5: {md_post_md5}")

# Verify L-62/63/64/65 in md
md_content = lessons_md.read_text(encoding="utf-8", errors="replace")
for lid in ["L-62", "L-63", "L-64", "L-65"]:
    found = f"## {lid} —" in md_content
    print(f"  {lid} in md: {'✓' if found else '✗'}")
