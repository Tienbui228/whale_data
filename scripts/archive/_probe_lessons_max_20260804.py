#!/usr/bin/env python3
"""Probe: find max lesson number in CASE_INDEX.xlsx LESSONS sheet."""
import io, sys
import openpyxl
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
wb = openpyxl.load_workbook(CASE_DIR / "CASE_INDEX.xlsx")
print(f"Sheets: {wb.sheetnames}")

if "LESSONS" in wb.sheetnames:
    ws = wb["LESSONS"]
    print(f"Max row: {ws.max_row}")
    print(f"Max col: {ws.max_column}")
    # Print header
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
    print(f"Headers: {headers}")
    # Print last 5 data rows
    print("\nLast 5 rows:")
    for r in range(max(2, ws.max_row-4), ws.max_row+1):
        row_vals = [ws.cell(r, c).value for c in range(1, min(6, ws.max_column+1))]
        print(f"  row {r}: {row_vals}")
    # Find max lesson number
    max_num = 0
    for r in range(2, ws.max_row+1):
        val = ws.cell(r, 1).value  # Assuming col A = lesson_id
        if val and str(val).startswith("L-"):
            try:
                n = int(str(val).split("-")[1])
                if n > max_num:
                    max_num = n
            except:
                pass
    print(f"\nMax lesson number found: L-{max_num}")
    print(f"Next lesson will be: L-{max_num+1}")
else:
    print("LESSONS sheet not found")
    print(f"Available sheets: {wb.sheetnames}")
