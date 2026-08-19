#!/usr/bin/env python3
"""Inspect CASE_GUA sheets + find col52 cells to rename + note placement."""
import sys, io
from pathlib import Path
import openpyxl

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_GUA.xlsx")
wb = openpyxl.load_workbook(str(XLSX), data_only=True, read_only=True)
print(f"sheets: {wb.sheetnames}")
for s in wb.sheetnames:
    ws = wb[s]
    print(f"\n== {s}: max_row={ws.max_row} max_col={ws.max_column}")

# DAILY col52 values
print(f"\n== DAILY col52 (spot_ext_venue) values ==")
ws = wb["DAILY"]
n_gate = 0
n_other = 0
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 52).value
    if v is not None and v != "":
        if v == "gate":
            n_gate += 1
        else:
            n_other += 1
            print(f"  row {r} col52 = {v!r} (non-'gate' non-null)")
print(f"  → n_gate={n_gate}, n_other={n_other}")

# CASE sheet look at r14-25 to find NOTES area
print(f"\n== CASE sheet (r1-r30) ==")
ws = wb["CASE"]
for r in range(1, min(ws.max_row + 1, 31)):
    v1 = ws.cell(r, 1).value
    v2 = ws.cell(r, 2).value
    print(f"  r{r:>2}: c1={str(v1)[:40]!r:>42}  c2={str(v2)[:50]!r}")
wb.close()
