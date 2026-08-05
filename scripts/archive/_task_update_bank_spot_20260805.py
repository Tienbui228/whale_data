# -*- coding: utf-8 -*-
"""Update CASE_BANK.xlsx DAILY sheet with Binance spot vol + taker from W2 probe."""
import openpyxl
from openpyxl.styles import PatternFill
from datetime import datetime, timezone
import urllib.request, json, ssl, time

CASE_PATH = r"docs/case_files/CASE_BANK.xlsx"

# W2 data from probe run 2026-08-05 (hardcoded from scratch/bank_spot_20260805.txt)
# Format: date -> (vol_spot_usd, taker_spot, perp_vol_usd, perp_px)
SPOT_DATA = {
    "2026-07-01": (286000, 0.670, 1300000, 0.0370),
    "2026-07-02": (3750000, 0.512, 16400000, 0.0376),
    "2026-07-03": (397000, 0.363, 2500000, 0.0374),
    "2026-07-04": (203000, 0.376, 1700000, 0.0375),
    "2026-07-05": (243000, 0.527, 1300000, 0.0374),
    "2026-07-06": (230000, 0.507, 1500000, 0.0358),
    "2026-07-07": (491000, 0.455, 2000000, 0.0360),
    "2026-07-08": (530000, 0.543, 2000000, 0.0347),
    "2026-07-09": (131000, 0.426, 900000, 0.0355),
    "2026-07-10": (675000, 0.433, 4100000, 0.0365),
    "2026-07-11": (2440000, 0.477, 19900000, 0.0414),
    "2026-07-12": (1610000, 0.437, 14900000, 0.0410),
    "2026-07-13": (573000, 0.491, 3700000, 0.0423),
    "2026-07-14": (459000, 0.605, 2800000, 0.0441),
    "2026-07-15": (2460000, 0.530, 37000000, 0.0511),
    "2026-07-16": (6400000, 0.518, 147700000, 0.0612),
    "2026-07-17": (21450000, 0.497, 432000000, 0.0705),  # UNLOCK
    "2026-07-18": (29800000, 0.528, 480100000, 0.1111),  # DEP_80M
    "2026-07-19": (114270000, 0.522, 1618300000, 0.2259),
    "2026-07-20": (77060000, 0.528, 1348700000, 0.2672),
    "2026-07-21": (97030000, 0.499, 1397400000, 0.1824),
    "2026-07-22": (97100000, 0.509, 1226600000, 0.2304),
    "2026-07-23": (55960000, 0.505, 889500000, 0.2746),
    "2026-07-24": (65500000, 0.507, 911400000, 0.3014),
    "2026-07-25": (45180000, 0.511, 723400000, 0.3466),
    "2026-07-26": (35560000, 0.514, 721200000, 0.3793),
    "2026-07-27": (124750000, 0.517, 1529100000, 0.3398),  # SPIKE
    "2026-07-28": (116370000, 0.511, 1152200000, 0.2171),
    "2026-07-29": (66910000, 0.507, 738400000, 0.1722),
    "2026-07-30": (85580000, 0.465, 1039600000, 0.0610),
    "2026-07-31": (35330000, 0.466, 562700000, 0.0575),
    "2026-08-01": (19910000, 0.462, 253300000, 0.0508),
    "2026-08-02": (12180000, 0.493, 172600000, 0.0447),
    "2026-08-03": (12350000, 0.481, 143000000, 0.0442),
    "2026-08-04": (46830000, 0.483, 471600000, 0.0524),
    "2026-08-05": (6860000, 0.473, 61800000, 0.0498),
}

def run():
    wb = openpyxl.load_workbook(CASE_PATH)
    ws = wb["DAILY"]

    # Find header row (row 2 = col names)
    # Find date column index and spot columns
    header_row = 2
    col_map = {}
    for cell in ws[header_row]:
        if cell.value:
            col_map[str(cell.value).strip()] = cell.column

    print(f"Columns found: {list(col_map.keys())[:10]}...")
    print(f"Total cols in header: {len(col_map)}")

    # Find date col
    date_col = None
    for k, v in col_map.items():
        if k.lower() in ("date", "ngay", "d"):
            date_col = v
            break
    if not date_col:
        # Try col 1
        date_col = 1
    print(f"Date col: {date_col}")

    # Find spot cols by name (actual CASE_BANK col names)
    venue_col = col_map.get("spot_venue") or 27
    vol_spot_col = col_map.get("vol_spot_usd") or 28
    taker_spot_col = col_map.get("taker_spot_pct") or 29
    print(f"Found: venue={venue_col}, vol={vol_spot_col}, taker={taker_spot_col}")

    updated = 0
    skipped = 0
    for row_idx in range(3, ws.max_row + 1):
        cell_date = ws.cell(row=row_idx, column=date_col)
        if not cell_date.value:
            continue
        try:
            if isinstance(cell_date.value, str):
                d_str = cell_date.value[:10]
            elif hasattr(cell_date.value, 'strftime'):
                d_str = cell_date.value.strftime("%Y-%m-%d")
            else:
                continue
        except:
            continue

        if d_str not in SPOT_DATA:
            skipped += 1
            continue

        vol_spot, taker_spot, perp_vol, perp_px = SPOT_DATA[d_str]

        # Write venue (col 27)
        if venue_col:
            ws.cell(row=row_idx, column=venue_col).value = "Binance_spot"

        # Write vol_spot_usd (col 28)
        c_vol = ws.cell(row=row_idx, column=vol_spot_col)
        c_vol.value = vol_spot
        c_vol.number_format = '#,##0'

        # Write taker_spot (col 29)
        c_tak = ws.cell(row=row_idx, column=taker_spot_col)
        c_tak.value = taker_spot
        c_tak.number_format = '0.0%'

        updated += 1

    wb.save(CASE_PATH)
    print(f"DONE: updated {updated} rows, skipped {skipped}")
    print(f"Saved: {CASE_PATH}")

run()
