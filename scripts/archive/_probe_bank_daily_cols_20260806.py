import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(r"D:\Data back test\whale_data\docs\case_files\CASE_BANK.xlsx", data_only=True)
ws = wb["DAILY"]

print("=== DAILY sheet: row1/row2 headers (A through AQ) ===")
for col_idx in range(1, 44):  # A=1 to AQ=43
    col_letter = get_column_letter(col_idx)
    r1 = ws.cell(row=1, column=col_idx).value
    r2 = ws.cell(row=2, column=col_idx).value
    print(f"{col_letter:>3} | {str(r1):<30} | {str(r2)}")

print()
# Count data rows: starting from row 3 until first empty row in col A
data_rows = 0
for row in ws.iter_rows(min_row=3, values_only=True):
    if row[0] is None:
        break
    data_rows += 1

print(f"Total data rows (from row 3, col A non-empty): {data_rows}")
print(f"Sheet max_row reported: {ws.max_row}")
