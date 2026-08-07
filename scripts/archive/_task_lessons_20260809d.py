"""Step 4 only: update CASE_BANANAS31.xlsx CASE/TRIGGERS/OUTCOME. Run after closing the file in Excel."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl
from openpyxl.styles import Font

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
bn_path  = CASE_DIR / "CASE_BANANAS31.xlsx"

def log(s=""): print(s)
log("=== STEP 4: CASE_BANANAS31.xlsx CASE/TRIGGERS/OUTCOME ===")

wb_bn = openpyxl.load_workbook(str(bn_path))

# CASE sheet
ws_case = wb_bn["CASE"]
verdict_text = ("NO_SETUP — TRACKING-passive. Build 06/26 het han (L-47). "
                "Ho MM-BAIT-B-shape serial, dong ho cum #4 Aug-Oct/26 [n=3 INFERENCE]. "
                "Operator signed 09/08.")
updated = False
for r in range(1, ws_case.max_row + 1):
    for c in range(1, ws_case.max_column + 1):
        val = ws_case.cell(r, c).value
        if val and "COLD_START" in str(val) and "PENDING" in str(val):
            ws_case.cell(r, c).value = verdict_text
            ws_case.cell(r, c).font  = Font(size=9)
            log(f"  CASE: verdict at row {r} col {c} updated")
            updated = True; break
    if updated: break

if not updated:
    new_r = ws_case.max_row + 2
    ws_case.cell(new_r, 1).value = "verdict"
    ws_case.cell(new_r, 2).value = verdict_text
    for c in [1, 2]: ws_case.cell(new_r, c).font = Font(size=9)
    log(f"  CASE: verdict appended at row {new_r}")

for r in range(1, ws_case.max_row + 1):
    if "archetype" in str(ws_case.cell(r, 1).value or "").lower():
        ws_case.cell(r, 2).value = "MM-BAIT-B-shape serial (INFERENCE, n=3)"
        ws_case.cell(r, 2).font  = Font(size=9)
        log(f"  CASE: archetype row {r} updated")

# TRIGGERS sheet
ws_trg = wb_bn["TRIGGERS"]
triggers = [
    ["ID",  "Trigger",                              "Threshold",                   "Status", "Ghi chu"],
    ["T1",  "gLSR drop + pos rebuild",              "gLSR <= 0.7 + pos_r tang 3d", "WAIT",   "Fuel nap tu reset crowd"],
    ["T2",  "Volume bung ben vung gia",              "vol >= 3x med14 ben >= 2d",   "WAIT",   "Confirm build momentum"],
    ["T3",  "GOM-A fire + yaml screen pass",        "GOM-A=TRUE + screen pass",    "WAIT",   "Xem screen_thresholds.yaml"],
    ["INV", "gLSR >= 1.2 hien tai (block entry)",   "gLSR >= 1.2",                 "BLOCK",  "Crowd da long het — mat fuel"],
]
for ri, row_data in enumerate(triggers, 1):
    for ci, val in enumerate(row_data, 1):
        ws_trg.cell(ri, ci).value = val
        ws_trg.cell(ri, ci).font  = Font(size=9, bold=(ri == 1))
log(f"  TRIGGERS: {len(triggers)} rows written")

# OUTCOME sheet
ws_out = wb_bn["OUTCOME"]
outcome_rows = [
    ["Field",            "Value"],
    ["verdict_operator", "NO_SETUP — TRACKING-passive"],
    ["ngay_verdict",     "2026-08-09"],
    ["archetype",        "MM-BAIT-B-shape serial (INFERENCE n=3)"],
    ["ghi_chu",          "Case giao khoa L-46/47/48 — giu lam doi chung, khong dong file"],
    ["clock_cum_4",      "Du kien Aug-Oct 2026 — theo doi gLSR drop ve <= 0.7"],
    ["operator_signed",  "09/08/2026"],
]
for ri, row_data in enumerate(outcome_rows, 1):
    for ci, val in enumerate(row_data, 1):
        ws_out.cell(ri, ci).value = val
        ws_out.cell(ri, ci).font  = Font(size=9, bold=(ri == 1))
log(f"  OUTCOME: {len(outcome_rows)} rows written")

wb_bn.save(str(bn_path))
log(f"  CASE_BANANAS31.xlsx saved OK")
log("\nStep 4 done. All steps complete — ready to commit.")
