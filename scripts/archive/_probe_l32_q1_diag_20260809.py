"""Diagnose L-32 extra row + Q1 A.5 text search."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT     = Path(r"D:\Data back test\whale_data")
CASE_DIR = ROOT / "docs" / "case_files"
DOCS     = ROOT / "docs"

def log(s=""): print(s)

# ── DAILY: find the extra row(s) ─────────────────────────────────────────────
log("--- L-32 DIAG: CASE_BANANAS31 DAILY ---")
wb = openpyxl.load_workbook(str(CASE_DIR / "CASE_BANANAS31.xlsx"), data_only=True)
ws = wb["DAILY"]

dates_seen = {}
extra_rows = []
empty_rows  = []
for r in range(2, ws.max_row + 1):
    d = ws.cell(r, 1).value
    if d in (None, ""):
        empty_rows.append(r)
    else:
        ds = str(d)
        if ds in dates_seen:
            extra_rows.append((r, ds, "DUPLICATE of row", dates_seen[ds]))
        else:
            dates_seen[ds] = r

log(f"  max_row: {ws.max_row}")
log(f"  Unique date rows: {len(dates_seen)}")
log(f"  Empty date rows: {empty_rows[:10]}")
log(f"  Duplicate dates: {extra_rows[:10]}")

# Show first 3 and last 5 rows
log(f"  First 3 rows (date col):")
for r in range(2, 5):
    log(f"    row {r}: {ws.cell(r, 1).value}")
log(f"  Last 5 rows (date col):")
for r in range(ws.max_row - 4, ws.max_row + 1):
    log(f"    row {r}: {ws.cell(r, 1).value}")

wb.close()

# ── A.5 text search ──────────────────────────────────────────────────────────
log("\n--- Q1 DIAG: Phụ lục A.5 content ---")
playbook = (DOCS / "analysis_playbook.md").read_text(encoding="utf-8")
if "A.5" in playbook:
    start = playbook.index("### A.5")
    block = playbook[start:start+1500]
    log(f"  A.5 block found (showing first 1000 chars):")
    log(block[:1000])
    # Check for various spellings of "Quy trình"
    for term in ["Quy trình", "Quy tr", "quy trinh", "Quy trinh", "quy tr"]:
        log(f"  Search '{term}' in block: {term in block}")
else:
    log("  A.5 NOT FOUND in playbook")
