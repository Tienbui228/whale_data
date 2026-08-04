#!/usr/bin/env python3
"""Verify CASE_CLO.xlsx structure and spot-check key values."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl
from pathlib import Path

f = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_CLO.xlsx")
wb = openpyxl.load_workbook(f, data_only=True)
print(f"Sheets: {wb.sheetnames}")

# DAILY spot-checks
ws = wb["DAILY"]
print(f"\nDAILY: {ws.max_row}r x {ws.max_column}c")
print("Header R2 cols 1-12:", [ws.cell(2,j).value for j in range(1,13)])
print("First pre-fw row (R3):", [ws.cell(3,j).value for j in range(1,13)])
# Find a row with full freshwatch data
for i in range(3, ws.max_row+1):
    d = ws.cell(i,1).value
    if str(d) >= "2026-06-05":
        print(f"First fw row ({d}):", [ws.cell(i,j).value for j in range(1,18)])
        break
# Find last row
last = ws.max_row
print(f"Last row ({ws.cell(last,1).value}):", [ws.cell(last,j).value for j in range(1,18)])

# EVENTS
ws3 = wb["EVENTS"]
print(f"\nEVENTS: {ws3.max_row}r x {ws3.max_column}c")
print("Last 3 events:")
for i in range(ws3.max_row-2, ws3.max_row+1):
    print(f"  R{i}:", [ws3.cell(i,j).value for j in range(1,5)])

# OUTCOME
ws5 = wb["OUTCOME"]
print(f"\nOUTCOME rows:")
for i in range(1, ws5.max_row+1):
    row = [ws5.cell(i,j).value for j in range(1,10)]
    if any(v is not None for v in row):
        print(f"  R{i}:", row[:6])
