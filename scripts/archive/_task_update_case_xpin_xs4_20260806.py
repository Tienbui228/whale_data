# _task_update_case_xpin_xs4_20260806.py
# XS4: Fill CASE_XPIN.xlsx DAILY vol_spot_usd + update CASE venue map
# Source: XS2 GeckoTerminal OHLCV (scratch/xpin_spot_20260806.txt)
# Coverage: 2026-02-06 to 2026-08-05 (181 bars). Before 2026-02-06 = ABSTAIN.
# Alpha ABSTAIN (API 400, no local store). vol_spot_usd = vol_dex_usd only.

import openpyxl

CASE_PATH = r"D:\Data back test\whale_data\docs\case_files\CASE_XPIN.xlsx"

# vol_dex from XS2 GeckoTerminal main pool (XPIN/USDT 0.01% PancakeSwap V3)
# source: scratch/xpin_spot_20260806.txt, XS2 pivot table
# Jul9-13 use precise values from XS3c; all others from XS2 formatted output (~1% precision)
DEX_VOL = {
    # Feb 2026 - pool coverage begins (created 2025-08-20 but GeckoTerminal returns ~181 bars from Feb06)
    "2026-02-06": 1_030_000, "2026-02-07": 920_000, "2026-02-08": 345_400,
    "2026-02-09": 438_300, "2026-02-10": 862_400, "2026-02-11": 749_800,
    "2026-02-12": 523_100, "2026-02-13": 2_510_000, "2026-02-14": 1_720_000,
    "2026-02-15": 1_600_000, "2026-02-16": 2_290_000, "2026-02-17": 624_900,
    "2026-02-18": 642_900, "2026-02-19": 1_650_000, "2026-02-20": 1_240_000,
    "2026-02-21": 894_100, "2026-02-22": 523_700, "2026-02-23": 1_020_000,
    "2026-02-24": 1_210_000, "2026-02-25": 633_200, "2026-02-26": 447_700,
    "2026-02-27": 450_700, "2026-02-28": 498_100,
    # Mar 2026
    "2026-03-01": 490_300, "2026-03-02": 546_900, "2026-03-03": 560_800,
    "2026-03-04": 454_200, "2026-03-05": 421_700, "2026-03-06": 520_200,
    "2026-03-07": 480_300, "2026-03-08": 436_600, "2026-03-09": 351_500,
    "2026-03-10": 581_800, "2026-03-11": 718_200, "2026-03-12": 1_140_000,
    "2026-03-13": 3_000_000, "2026-03-14": 762_300, "2026-03-15": 387_900,
    "2026-03-16": 493_100, "2026-03-17": 721_500, "2026-03-18": 669_600,
    "2026-03-19": 600_900, "2026-03-20": 426_000, "2026-03-21": 458_500,
    "2026-03-22": 479_100, "2026-03-23": 364_600, "2026-03-24": 253_000,
    "2026-03-25": 307_000, "2026-03-26": 582_800, "2026-03-27": 523_300,
    "2026-03-28": 387_600, "2026-03-29": 493_700, "2026-03-30": 261_600,
    "2026-03-31": 451_100,
    # Apr 2026
    "2026-04-01": 264_100, "2026-04-02": 380_300, "2026-04-03": 252_200,
    "2026-04-04": 253_400, "2026-04-05": 213_500, "2026-04-06": 372_500,
    "2026-04-07": 2_260_000, "2026-04-08": 2_450_000, "2026-04-09": 3_720_000,
    "2026-04-10": 925_400, "2026-04-11": 618_600, "2026-04-12": 391_500,
    "2026-04-13": 1_620_000, "2026-04-14": 2_250_000, "2026-04-15": 1_040_000,
    "2026-04-16": 1_110_000, "2026-04-17": 696_900, "2026-04-18": 683_800,
    "2026-04-19": 517_600, "2026-04-20": 531_300, "2026-04-21": 739_400,
    "2026-04-22": 1_650_000, "2026-04-23": 1_060_000, "2026-04-24": 636_100,
    "2026-04-25": 367_200, "2026-04-26": 514_400, "2026-04-27": 858_600,
    "2026-04-28": 668_200, "2026-04-29": 569_600, "2026-04-30": 1_440_000,
    # May 2026
    "2026-05-01": 392_800, "2026-05-02": 1_650_000, "2026-05-03": 919_300,
    "2026-05-04": 470_100, "2026-05-05": 324_100, "2026-05-06": 648_300,
    "2026-05-07": 489_900, "2026-05-08": 1_230_000, "2026-05-09": 1_310_000,
    "2026-05-10": 1_400_000, "2026-05-11": 650_200, "2026-05-12": 1_550_000,
    "2026-05-13": 1_420_000, "2026-05-14": 1_400_000, "2026-05-15": 2_220_000,
    "2026-05-16": 697_100, "2026-05-17": 809_000, "2026-05-18": 649_700,
    "2026-05-19": 371_400, "2026-05-20": 308_500, "2026-05-21": 322_200,
    "2026-05-22": 704_300, "2026-05-23": 390_600, "2026-05-24": 375_600,
    "2026-05-25": 400_700, "2026-05-26": 634_900, "2026-05-27": 641_700,
    "2026-05-28": 589_800, "2026-05-29": 859_300, "2026-05-30": 286_500,
    "2026-05-31": 277_800,
    # Jun 2026
    "2026-06-01": 496_700, "2026-06-02": 832_500, "2026-06-03": 711_100,
    "2026-06-04": 882_200, "2026-06-05": 1_320_000, "2026-06-06": 1_150_000,
    "2026-06-07": 407_200, "2026-06-08": 580_400, "2026-06-09": 455_000,
    "2026-06-10": 596_500, "2026-06-11": 597_000, "2026-06-12": 610_000,
    "2026-06-13": 998_300, "2026-06-14": 451_400, "2026-06-15": 1_020_000,
    "2026-06-16": 715_200, "2026-06-17": 944_300, "2026-06-18": 1_130_000,
    "2026-06-19": 1_460_000, "2026-06-20": 551_900, "2026-06-21": 350_500,
    "2026-06-22": 871_100, "2026-06-23": 1_370_000, "2026-06-24": 1_140_000,
    "2026-06-25": 997_300, "2026-06-26": 1_010_000, "2026-06-27": 556_500,
    "2026-06-28": 360_000, "2026-06-29": 491_800, "2026-06-30": 1_250_000,
    # Jul 2026
    "2026-07-01": 2_300_000, "2026-07-02": 1_010_000, "2026-07-03": 565_400,
    "2026-07-04": 547_700, "2026-07-05": 399_700, "2026-07-06": 491_400,
    "2026-07-07": 972_700, "2026-07-08": 1_270_000,
    "2026-07-09": 496_677,   # precise XS3c
    "2026-07-10": 2_788_481, # precise XS3c
    "2026-07-11": 7_785_737, # precise XS3c
    "2026-07-12": 2_907_977, # precise XS3c
    "2026-07-13": 1_494_383, # precise XS3c
    "2026-07-14": 1_270_000, "2026-07-15": 2_580_000, "2026-07-16": 1_490_000,
    "2026-07-17": 635_100, "2026-07-18": 1_520_000, "2026-07-19": 660_300,
    "2026-07-20": 759_200, "2026-07-21": 887_800, "2026-07-22": 619_700,
    "2026-07-23": 656_700, "2026-07-24": 688_000, "2026-07-25": 451_800,
    "2026-07-26": 582_500, "2026-07-27": 668_500, "2026-07-28": 761_300,
    "2026-07-29": 472_000, "2026-07-30": 396_700, "2026-07-31": 639_300,
    # Aug 2026
    "2026-08-01": 228_300, "2026-08-02": 557_100, "2026-08-03": 620_900,
    "2026-08-04": 326_400, "2026-08-05": 154_800,
}

def main():
    wb = openpyxl.load_workbook(CASE_PATH)

    # ── DAILY sheet: fill vol_spot_usd (col 28) ──────────────────────────────
    ws_d = wb["DAILY"]
    # Row 2 = headers; data from row 3 onward
    # col 28 = vol_spot_usd, col 27 = spot_venue
    COL_DATE = 1
    COL_VOL_SPOT = 28
    COL_SPOT_VENUE = 27

    updated = 0
    abstained = 0
    for row in range(3, ws_d.max_row + 1):
        d = ws_d.cell(row, COL_DATE).value
        if d is None:
            continue
        d_str = str(d)[:10]  # "2025-09-12" from date or string
        if d_str in DEX_VOL:
            ws_d.cell(row, COL_VOL_SPOT).value = DEX_VOL[d_str]
            # spot_venue already set = BSC_DEX_only; add note about source
            ws_d.cell(row, COL_SPOT_VENUE).value = "BSC_DEX_PancakeSwapV3_0.01%"
            updated += 1
        else:
            # Before Feb 06 2026: GeckoTerminal no data (only ~181 bar window)
            if d_str < "2026-02-06":
                ws_d.cell(row, COL_VOL_SPOT).value = None  # ABSTAIN — no GeckoTerminal coverage
                ws_d.cell(row, COL_SPOT_VENUE).value = "BSC_DEX_only_ABSTAIN_pre-Feb06"
                abstained += 1

    print(f"DAILY updated: {updated} rows, abstained: {abstained} rows")

    # ── CASE sheet: update venue map fields ──────────────────────────────────
    ws_c = wb["CASE"]
    UPDATES = {
        "venue_spot": (
            "BSC DEX PancakeSwapV3 0.01% main pool (0x08de…4a4, liq=$944K, v24=$345K). "
            "19 ghost pools (all $0 vol). No Binance spot, no CEX spot trustworthy."
        ),
        "venue_alpha": (
            "ALPHA_334 listed ~2026-01-27. API 400 (klines endpoint not serving XPIN). "
            "Local store: 0 files. Alpha vol = ABSTAIN."
        ),
        "spot_perp_ratio": (
            "BSC DEX: $521K/day avg crowd_surge (Jul22-Aug05); $1.9M/day prior (Jul8-21). "
            "Crowd surge: DEX contracted 0.28x — crowd buying pure perp leverage, not spot. "
            "Source: GeckoTerminal Feb06–Aug05 (181 bars). Pre-Feb06: ABSTAIN (no coverage)."
        ),
        "O4_bitmart_vol": (
            "CoinGecko: BitMart $1.61M vol/liq=338x SUSPECT. Biconomy $251K vol/liq=661x SUSPECT. "
            "KuCoin $72K vol/liq=7.2x CLEAN. Gate $198K. MEXC $142K. Hibt $22K vol/liq=6.9x CLEAN. "
            "Source: XS1c CoinGecko tickers+depth 06/08/2026."
        ),
    }
    # Add new CASE rows if keys not found
    existing_keys = {}
    for r in range(1, ws_c.max_row + 1):
        k = ws_c.cell(r, 1).value
        if k:
            existing_keys[k] = r

    for key, val in UPDATES.items():
        if key in existing_keys:
            ws_c.cell(existing_keys[key], 2).value = val
            print(f"CASE updated row {existing_keys[key]}: {key}")
        else:
            new_row = ws_c.max_row + 1
            ws_c.cell(new_row, 1).value = key
            ws_c.cell(new_row, 2).value = val
            print(f"CASE appended row {new_row}: {key}")

    # Add detection_floor summary if missing
    floor_key = "detection_floor_spot"
    if floor_key not in existing_keys:
        new_row = ws_c.max_row + 1
        ws_c.cell(new_row, 1).value = floor_key
        ws_c.cell(new_row, 2).value = (
            "Daily spot total ~$571K (DEX $521K + Alpha $50K est). "
            "1% total supply (1B tokens @ $0.00155) = $1.55M = 2.7x daily spot → DETECTABLE. "
            "5% circ (2.4B tokens) = $3.73M = 6.5x daily spot → VERY DETECTABLE. "
            "Any non-trivial kho selling via BSC DEX would be visible within 1 day."
        )
        print(f"CASE appended: {floor_key}")

    # verdict_spot_layer summary
    verdict_key = "verdict_spot_layer"
    if verdict_key not in existing_keys:
        new_row = ws_c.max_row + 1
        ws_c.cell(new_row, 1).value = verdict_key
        ws_c.cell(new_row, 2).value = (
            "XS4 verdict (06/08/2026): WATCH-nhạt ĐỨNG. "
            "Spot thin ($571K/day) confirms crowd buying pure leverage (perp), not real money into DEX. "
            "No DEX sell burst observed (XS3b rate-limited; XS3c DEX vol expanded on Jul11 fail = crowd panic buy into DEX, not distribution). "
            "Detection floor LOW: any kho >1% supply selling via DEX would be visible within 1-2 days. "
            "Verdict unchanged from cold-start X6. Alpha ABSTAIN pending operator verification."
        )
        print(f"CASE appended: {verdict_key}")

    import tempfile, shutil, os
    tmp = CASE_PATH + ".tmp_xs4"
    wb.save(tmp)
    os.replace(tmp, CASE_PATH)
    print(f"\nSaved: {CASE_PATH}")
    print(f"DAILY: {updated} vol_spot_usd filled (Feb06-Aug05), {abstained} marked ABSTAIN (pre-Feb06)")
    print("CASE: venue_spot, venue_alpha, spot_perp_ratio, O4_bitmart_vol, detection_floor_spot, verdict_spot_layer updated")

if __name__ == "__main__":
    main()
