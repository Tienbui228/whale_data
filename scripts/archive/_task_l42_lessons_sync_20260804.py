#!/usr/bin/env python3
"""Task: append L-42 to LESSONS sheet, add scope to L-05/06/07, regen lessons_learned.md."""
import io, sys
from pathlib import Path
from datetime import date
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
REPO_DIR = Path(r"D:\Data back test\whale_data")
INDEX_PATH = CASE_DIR / "CASE_INDEX.xlsx"
MD_PATH = REPO_DIR / "docs" / "lessons_learned.md"

# ── L-42 content (operator ky 04/08/2026) ──────────────────────────────────
L42 = {
    "id": "L-42",
    "date": "2026-08-04",
    "topic": "IDOL/TAG/BANK — kho spot MM: phan loai theo SO HUU truoc khi doc on-chain",
    "body": (
        "Kho spot cua MM khong tu no la 'hang se xa' — ham muc tieu phu thuoc SO HUU kho. "
        "Phan 3 loai TRUOC khi doc bat ky transfer nao: "
        "LOAI A (kho MUON tu team — khong gia tri voi MM; P&L 100% o perp; moi deposit/withdrawal la THONG DIEP nhung nguoi doc on-chain, khong phai bang chung; "
        "endgame kho round-trip nguyen dai; dong co DAP GIA hau campaign de mua lai tra no kho — canh kho hoi day o day). "
        "LOAI B (kho unlock/dau tu — so huu that, cost thap; doanh thu phai hien thuc hoa; neo lich unlock; deposit-len-san <=48h sau ngay unlock = tell hang nhat — BANK 17/07->18/07). "
        "LOAI C (tu bo tien mua tai cho — conviction that, cost basis = floor co nghia, khong deadline; loai duy nhat 'di cung thuyen' duoc). "
        "Tell phan loai: (1) NGUON kho — deployer/team wallet = nghi A; vesting contract dung lich = B; gom le nhieu nguon dai han = C; "
        "(2) nhip voi lich unlock; (3) hanh vi kho HAU DINH — round-trip nguyen = A, ban that = B; "
        "(4) do 'dien' cua transfer — chia may deu + Sent=0 trung bay = tin hieu re, nghi A. "
        "Bang chung: IDOL 3 man mot dao cu (8% vao 03-05/07 -> rut 273M staging shape 13 vi -> redeposit 102M) trong khi ket that o perp ~$1.2-1.8M; "
        "TAG 90B round-trip May->Jul; BANK chan unlock ban that (lai A+B); AKE 'perp vol 70x spot — spot la can gat'."
    ),
    "type": "rule",
    "status": "OPERATOR KY 04/08",
    "action": (
        "(a) cot mm_archetype{A/B/C/lai} + kho_nguon cham hoi to 12 case + case song; "
        "(b) BR kho 6% trace nguon = viec phan loai dau tien — neu team wallet -> Loai A -> cu flush gLSR 03/08 doc lai duoi anh sang 'dao cu man hai'; "
        "(c) moi dong PnL spot chua co sell-side xac nhan ha nhan 'unverified inference' (IDOL 111M, BANK $12-28M); "
        "(d) K3 sell-side Pancake IDOL van la test phan xu — PENDING."
    ),
    "scope": "Ap dung cho tat ca case co transfer on-chain. Cap nhat cho L-05/L-06/L-07: scope = Loai B/C; Loai A doc theo L-42 truoc.",
}

# ── Scope note text for old lessons ────────────────────────────────────────
SCOPE_L05 = "[Pham vi L-42]: Loai B/C — withdrawal gan dinh = bang chung phan phoi THAT. Loai A — withdrawal la dao cu, khong vao bang chung; chuyen sang doc theo L-42."
SCOPE_L06 = "[Pham vi L-42]: Ap dung khi kho xac dinh duoc Loai B hoac C. Loai A — V_venue phan xu 'mua' hay 'hang co san' khong con nghia (hang KHONG thuoc MM); chuyen sang L-42."
SCOPE_L07 = "[Pham vi L-42]: Ap dung khi xac nhan Loai B/C (kho that). Loai A — kho len san = dao cu chuan bi dan, khong phai leg keo; endgame la round-trip tra no. Doc L-42 truoc L-07."

wb = openpyxl.load_workbook(INDEX_PATH)
ws = wb["LESSONS"]

# ── Verify current max ──────────────────────────────────────────────────────
max_row = ws.max_row
last_id = ws.cell(max_row, 1).value
print(f"Current last row: {max_row}, last lesson: {last_id}")

# ── Find rows for L-05, L-06, L-07 ─────────────────────────────────────────
scope_targets = {"L-05": SCOPE_L05, "L-06": SCOPE_L06, "L-07": SCOPE_L07}
for r in range(1, ws.max_row+1):
    lid = ws.cell(r, 1).value
    if lid in scope_targets:
        existing = ws.cell(r, 8).value
        if existing:
            print(f"  {lid} (row {r}) C8 already has: {str(existing)[:60]}... — skipping")
        else:
            ws.cell(r, 8).value = scope_targets[lid]
            print(f"  {lid} (row {r}): added scope note to C8")

# ── Append L-42 ─────────────────────────────────────────────────────────────
new_row = max_row + 1
ws.cell(new_row, 1).value = L42["id"]
ws.cell(new_row, 2).value = L42["date"]
ws.cell(new_row, 3).value = L42["topic"]
ws.cell(new_row, 4).value = L42["body"]
ws.cell(new_row, 5).value = L42["type"]
ws.cell(new_row, 6).value = L42["status"]
ws.cell(new_row, 7).value = L42["action"]
ws.cell(new_row, 8).value = L42["scope"]
print(f"\nAppended L-42 at row {new_row}")

wb.save(INDEX_PATH)
print("Saved CASE_INDEX.xlsx")

# ── Regen lessons_learned.md ─────────────────────────────────────────────────
print("\nRegenerating lessons_learned.md...")
wb2 = openpyxl.load_workbook(INDEX_PATH)
ws2 = wb2["LESSONS"]

lessons = []
for r in range(3, ws2.max_row+1):  # skip header rows 1-2
    lid = ws2.cell(r, 1).value
    if not lid or not str(lid).startswith("L-"):
        continue
    lessons.append({
        "id": lid,
        "date": ws2.cell(r, 2).value or "",
        "topic": ws2.cell(r, 3).value or "",
        "body": ws2.cell(r, 4).value or "",
        "type": ws2.cell(r, 5).value or "",
        "status": ws2.cell(r, 6).value or "",
        "action": ws2.cell(r, 7).value or "",
        "scope": ws2.cell(r, 8).value or "",
    })

total = len(lessons)
first_id = lessons[0]["id"] if lessons else "L-01"
last_id_regen = lessons[-1]["id"] if lessons else "L-??"

lines = []
lines.append("# Lessons Learned — FreshWatch\n")
lines.append(f"> Auto-generated tu CASE_INDEX.xlsx sheet LESSONS. Khong edit file nay truc tiep.")
lines.append(f"> Lan sinh: 2026-08-04 · Tong: {total} bai hoc ({first_id} → {last_id_regen})\n")

for L in lessons:
    type_str = f"`{L['type']}`" if L["type"] else ""
    status_str = f" · {L['status']}" if L["status"] else ""
    header = f"## {L['id']} [{L['date']}] {L['topic']} — {type_str}{status_str}"
    lines.append(header)
    if L["body"]:
        lines.append(L["body"])
    if L["action"]:
        lines.append(f"*Hanh dong tiep: {L['action']}*")
    if L["scope"]:
        lines.append(f"*Pham vi / Cross-ref: {L['scope']}*")
    lines.append("")

md_content = "\n".join(lines)
MD_PATH.write_text(md_content, encoding="utf-8")
print(f"Written: {MD_PATH}")
print(f"Total lessons: {total}")
print(f"Last: {last_id_regen}")
print("DONE")
