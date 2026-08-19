#!/usr/bin/env python3
"""[A3] Append EVENTS row for Gate-anomaly (án con, chưa điều tra sâu)."""
import sys, io
from pathlib import Path
import openpyxl

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_GUA.xlsx")

wb = openpyxl.load_workbook(str(XLSX), data_only=False)
if "EVENTS" not in wb.sheetnames:
    print("[ERR] no EVENTS sheet"); wb.close(); raise SystemExit(1)
ws = wb["EVENTS"]

# guard L-50: capture max_col before + after
max_col_before = ws.max_column
max_row_before = ws.max_row
print(f"[before] EVENTS max_row={max_row_before}, max_col={max_col_before}")

next_row = max_row_before + 1
row = (
    "2026-08-19",
    "GATE_PRICE_ANOMALY_4.5X",
    "Gate GUA_USDT = $0.19 vs Alpha/perp $0.042 (4.5×) — contract khớp bit-equal 3/3 (A0 SEAL 19/08). Anomaly nội tại Gate pair. Nghi: thanh khoản mỏng / thị trường cô lập / deposit-withdraw. ÁN CON riêng, chưa backfill Gate vào băng-thông-xả.",
    "gua_identity_forensics_20260819",
    "open",
)
for i, v in enumerate(row, 1):
    ws.cell(next_row, i).value = v
print(f"[append] row {next_row}: {row}")

# L-50 write-after-guard
assert ws.max_column >= max_col_before, f"L-50 fail: max_col shrunk {max_col_before} → {ws.max_column}"
assert ws.max_row == max_row_before + 1, f"L-50 fail: row count wrong {max_row_before + 1} vs {ws.max_row}"

wb.save(str(XLSX))
print(f"[after] EVENTS max_row={ws.max_row}, max_col={ws.max_column}  → L-50 PASS")

# verify: reopen data_only + spot-check
wb2 = openpyxl.load_workbook(str(XLSX), data_only=True, read_only=True)
ws2 = wb2["EVENTS"]
print(f"\n[verify] EVENTS row {next_row}:")
for c in range(1, 6):
    v = ws2.cell(next_row, c).value
    v_str = str(v)[:80] if v else "(None)"
    print(f"  col {c}: {v_str}")
wb2.close()
print("[DONE]")
