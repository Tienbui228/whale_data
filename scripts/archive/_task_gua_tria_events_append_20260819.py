#!/usr/bin/env python3
"""Append EVENTS rows to CASE_GUA + CASE_TRIA per operator F.request."""
import sys, io
from pathlib import Path
import openpyxl

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")

def append_events(xlsx_path, rows_to_add):
    if not xlsx_path.exists():
        print(f"  [SKIP] {xlsx_path.name} NOT FOUND"); return
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    if "EVENTS" not in wb.sheetnames:
        print(f"  [SKIP] {xlsx_path.name}: no EVENTS sheet"); wb.close(); return
    ws = wb["EVENTS"]
    next_row = ws.max_row + 1
    print(f"  {xlsx_path.name} EVENTS: max_row={ws.max_row}, appending at row {next_row}")
    for row in rows_to_add:
        for i, val in enumerate(row, 1):
            ws.cell(next_row, i).value = val
        print(f"    row {next_row}: {row}")
        next_row += 1
    wb.save(str(xlsx_path))
    print(f"  [save] {xlsx_path.name}")

# GUA EVENTS
gua_rows = [
    ("2026-08-10", "GATE_SPIKE_5X",
     "Gate spike $128,814 = 5.8× nền T7 Jul ($22K med), CÙNG NGÀY perp $180.5M pump peak. NOT lead — chase.",
     "venue_backfill_19/08", "source=probe"),
    ("2026-07-01→2026-08-18", "GATE_BASELINE",
     "Gate baseline med $22K/day (n=31 Jul). Pump 05-08 = 1.23× nền, post-pump = 1.00× (phẳng). Van nhỏ, không dùng làm xả.",
     "venue_backfill_19/08", "source=probe"),
]
print("=== GUA ===")
append_events(CASE_DIR / "CASE_GUA.xlsx", gua_rows)

# TRIA EVENTS
tria_rows = [
    ("2026-08-14", "SPOT-EXT-LEAD",
     "Bybit turnover $406K spike (khuôn LEAD: Bybit tăng TRƯỚC perp/alpha động)",
     "venue_backfill_19/08", "unresolved"),
    ("2026-07-25", "SPOT-EXT-ANOMALY",
     "Bybit spike $734K dị thường, chưa rõ context",
     "venue_backfill_19/08", "unresolved"),
]
print("\n=== TRIA ===")
append_events(CASE_DIR / "CASE_TRIA.xlsx", tria_rows)

print("\nDONE.")
