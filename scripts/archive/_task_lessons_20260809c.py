"""Steps 5-6 only: CASE_HEI EVENTS + INDEX.md update."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl
from openpyxl.styles import Font

ROOT     = Path(r"D:\Data back test\whale_data")
CASE_DIR = ROOT / "docs" / "case_files"
DOCS     = ROOT / "docs"
def log(s=""): print(s)

# ── 5. CASE_HEI EVENTS ───────────────────────────────────────────────────────
log("=== STEP 5: CASE_HEI.xlsx EVENTS ===")
hei_path = CASE_DIR / "CASE_HEI.xlsx"
wb_hei   = openpyxl.load_workbook(str(hei_path))
ws_ev    = wb_hei["EVENTS"]
last_ev_r = ws_ev.max_row
while last_ev_r > 1 and all(ws_ev.cell(last_ev_r, c).value in (None, "") for c in range(1, 6)):
    last_ev_r -= 1
new_ev_r = last_ev_r + 1
# Check if cross-ref already written
already = any("CROSS_REF" in str(ws_ev.cell(r, 3).value or "") or
              "BANANAS31" in str(ws_ev.cell(r, 2).value or "")
              for r in range(1, ws_ev.max_row + 1))
if already:
    log(f"  EVENTS cross-ref already present — skip")
else:
    ws_ev.cell(new_ev_r, 1).value = "2026-08-09"
    ws_ev.cell(new_ev_r, 1).font  = Font(size=9)
    ws_ev.cell(new_ev_r, 2).value = ("09/08: HEI duoc dung lam doi chung DUONG trong cap L-46 "
                                      "(vs BANANAS31 doi chung AM) — xem Phu luc A.5 analysis_playbook.md")
    ws_ev.cell(new_ev_r, 2).font  = Font(size=9)
    ws_ev.cell(new_ev_r, 3).value = "CROSS_REF"
    ws_ev.cell(new_ev_r, 3).font  = Font(size=9)
    wb_hei.save(str(hei_path))
    log(f"  CASE_HEI.xlsx EVENTS: row {new_ev_r} appended")

# ── 6. INDEX.md ──────────────────────────────────────────────────────────────
log("\n=== STEP 6: docs/INDEX.md ===")
idx_md_path = DOCS / "INDEX.md"
idx_text    = idx_md_path.read_text(encoding="utf-8")

lines_idx = idx_text.splitlines()
new_lines  = []
i = 0
hei_fixed = False
while i < len(lines_idx):
    ln = lines_idx[i]
    # Detect the broken merged form: "| `docs/case_files/CASE_HEI.xlsx"
    # followed by BANANAS31 content on next line OR same line
    if "CASE_HEI.xlsx" in ln and not hei_fixed:
        if "BANANAS31" in ln:
            # single merged line
            new_lines.append("| `docs/case_files/CASE_HEI.xlsx` | HEI | TRACKING (doi chung DUONG cap L-46) |")
            new_lines.append("| `docs/case_files/CASE_BANANAS31.xlsx` | BANANAS31 | COLD_START — PENDING (doi chung AM cap L-46, TRACKING-passive) |")
            log(f"  Fixed merged HEI+BANANAS31 line at idx {i}")
            hei_fixed = True
            i += 1
            continue
        elif i + 1 < len(lines_idx) and "BANANAS31" in lines_idx[i + 1]:
            # two-line broken form
            new_lines.append("| `docs/case_files/CASE_HEI.xlsx` | HEI | TRACKING (doi chung DUONG cap L-46) |")
            new_lines.append("| `docs/case_files/CASE_BANANAS31.xlsx` | BANANAS31 | COLD_START — PENDING (doi chung AM cap L-46, TRACKING-passive) |")
            log(f"  Fixed two-line broken HEI+BANANAS31 at idx {i},{i+1}")
            hei_fixed = True
            i += 2
            continue
        else:
            # HEI row is clean — just add annotation if not present
            if "TRACKING" in ln and "doi chung" not in ln and "đối chứng" not in ln:
                ln = ln.rstrip("|").rstrip() + " (doi chung DUONG cap L-46) |"
                log(f"  Added DUONG annotation to HEI row at idx {i}")
            hei_fixed = True
    # Detect existing BANANAS31 row without annotation
    if "CASE_BANANAS31.xlsx` | BANANAS31" in ln:
        if "doi chung" not in ln and "đối chứng" not in ln:
            ln = ln.replace("COLD_START — PENDING |",
                            "COLD_START — PENDING (doi chung AM cap L-46, TRACKING-passive) |")
            log(f"  Added AM annotation to BANANAS31 row at idx {i}")
    new_lines.append(ln)
    i += 1

idx_text = "\n".join(new_lines)

# Update lessons range
for old, new in [
    ("sheet LESSONS (không edit trực tiếp)",
     "sheet LESSONS (không edit trực tiếp) · L-01→L-48"),
    ("sheet LESSONS (khong edit truc tiep)",
     "sheet LESSONS (khong edit truc tiep) · L-01→L-48"),
]:
    if old in idx_text and "L-01→L-48" not in idx_text:
        idx_text = idx_text.replace(old, new)
        log("  Lesson range updated: L-01→L-48")
        break

# Update version date
for old, new in [
    ("khi thêm/xóa file canonical", "09/08/2026 (L-46/47/48 + Phụ lục A.5)"),
    ("khi them/xoa file canonical", "09/08/2026 (L-46/47/48 + Phu luc A.5)"),
]:
    if old in idx_text:
        idx_text = idx_text.replace(old, new)
        log("  Version date updated")
        break

idx_md_path.write_text(idx_text, encoding="utf-8")
log(f"  docs/INDEX.md saved ({len(idx_text.splitlines())} lines)")

log("\nSteps 5-6 done. Close CASE_BANANAS31.xlsx then run _task_lessons_20260809d.py")
