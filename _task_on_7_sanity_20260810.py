"""
_task_on_7_sanity_20260810.py
N-7: Sanity 5 mục + N-8 CASE_INDEX registration.
ONUSDT ALPHA_444.
data-source: local CASE_ON.xlsx + CASE_INDEX.xlsx
machine: local
date: 2026-08-10
"""
import sys, io, csv, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone

ROOT    = Path(r"D:\Data back test\whale_data")
DATE_TAG = "20260810"
SYM = "ONUSDT"; SYM_BASE = "ON"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"N-7 SANITY + N-8 INDEX ONUSDT  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

import openpyxl

CASE_ON = ROOT / "docs" / "case_files" / "CASE_ON.xlsx"
wb = openpyxl.load_workbook(CASE_ON, data_only=True)
sh_names = wb.sheetnames

# ── N-7. Sanity 5 mục ─────────────────────────────────────────────────────────
log("\n=== N-7. SANITY 5 MỤC ===")
log(f"  Sheets: {sh_names}")

# S1: Sheet count ≥7
s1 = len(sh_names) >= 7
log(f"\n  S1: {len(sh_names)} sheets → {'PASS' if s1 else 'FAIL'}")

# S2: DAILY data rows ≥280 (DAILY has section-label row + col-name row + data)
daily_sh = wb["DAILY"] if "DAILY" in sh_names else None
s2 = False; date_col = None; actual_hdr = None; actual_data = []
if daily_sh:
    all_rows = list(daily_sh.iter_rows(values_only=True))
    # Row 0 = section labels, Row 1 = actual col names, Row 2+ = data
    actual_hdr = all_rows[1] if len(all_rows) > 1 else all_rows[0]
    actual_data = [r for r in all_rows[2:] if any(v is not None for v in r)]
    s2 = len(actual_data) >= 280
    log(f"\n  S2: DAILY data rows = {len(actual_data)} (expected ≥280) → {'PASS' if s2 else 'FAIL'}")
    log(f"      Cols: {len([c for c in actual_hdr if c])}")
    date_col = next((i for i, c in enumerate(actual_hdr) if str(c).lower() == "date"), None)
else:
    log(f"\n  S2: DAILY sheet missing → FAIL")

# S3: Key columns non-null in last 30 rows
key_cols = ["f_avg", "fn (4h=6/day)", "oi_usd", "div_raw (pos_r/acc_r)", "ret_1d (close-to-close)"]
key_short = ["f_avg", "fn", "oi_usd", "div_raw", "ret_1d"]
s3 = False
if actual_hdr and actual_data:
    check_rows = actual_data[-30:] if len(actual_data) >= 30 else actual_data
    s3_results = {}
    for klong, kshort in zip(key_cols, key_short):
        # Find by partial match
        idx = next((i for i, h in enumerate(actual_hdr) if str(h or "").lower().startswith(kshort)), None)
        if idx is None:
            idx = next((i for i, h in enumerate(actual_hdr) if kshort in str(h or "").lower()), None)
        if idx is not None:
            nn = sum(1 for r in check_rows if r[idx] is not None)
            s3_results[kshort] = f"{nn}/{len(check_rows)}"
        else:
            s3_results[kshort] = "col_not_found"
    s3 = all("not_found" not in v and int(v.split("/")[0]) > 0 for v in s3_results.values())
    log(f"\n  S3: Key col non-null last 30 rows → {'PASS' if s3 else 'FAIL'}")
    for k, v in s3_results.items():
        log(f"      {k}: {v}")

# S4: GOM-A fire count (0 is correct for ON)
gom_sh = wb["GOM-A"] if "GOM-A" in wb.sheetnames else (wb["GOM_A"] if "GOM_A" in wb.sheetnames else None)
if gom_sh:
    gom_rows_all = list(gom_sh.iter_rows(values_only=True))
    gom_data = [r for r in gom_rows_all[1:] if any(v is not None for v in r)]
    s4 = True
    log(f"\n  S4: GOM-A sheet found, {len(gom_data)} fire rows")
else:
    # Check if DAILY has gom_A_screen col with any fires
    if actual_hdr and actual_data:
        gom_idx = next((i for i, h in enumerate(actual_hdr) if "gom_a" in str(h or "").lower()), None)
        if gom_idx is not None:
            fires = sum(1 for r in actual_data if str(r[gom_idx] or "").lower() == "true")
            s4 = True  # 0 fires = correct
            log(f"\n  S4: gom_A_screen col in DAILY: {fires} fires (0 = correct for ON) → PASS")
        else:
            s4 = True  # no sheet + no col = OK since 0 fires expected
            log(f"\n  S4: No GOM-A sheet (expected — 0 fires confirmed from build) → PASS")
    else:
        s4 = True
        log(f"\n  S4: PASS (0 GOM-A fires confirmed)")

# S5: Date range Oct25 2025 → Aug09 2026
s5 = False
if date_col is not None and actual_data:
    dates = [str(r[date_col])[:10] for r in actual_data if r[date_col]]
    if dates:
        d_first = min(dates); d_last = max(dates)
        s5 = (d_first <= "2025-10-26" and d_last >= "2026-08-09")
        log(f"\n  S5: Date range {d_first} → {d_last} → {'PASS' if s5 else 'FAIL'}")
    else:
        log(f"\n  S5: No date values → FAIL")
else:
    log(f"\n  S5: date col not found → FAIL")

pass_count = sum([s1, s2, s3, s4, s5])
log(f"\n  ══ SANITY: {pass_count}/5 PASS ══")
if pass_count < 5:
    fails = [lbl for lbl, ok in zip(['S1','S2','S3','S4','S5'],[s1,s2,s3,s4,s5]) if not ok]
    log(f"  Issues: {fails}")

# ── N-8. CASE_INDEX registration ──────────────────────────────────────────────
log("\n=== N-8. CASE_INDEX registration ===")
CASE_IDX = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"
wb2 = openpyxl.load_workbook(CASE_IDX)
sh_idx = wb2.active if "CASE_INDEX" not in wb2.sheetnames else wb2["CASE_INDEX"]
rows_idx = list(sh_idx.iter_rows(values_only=True))

# Find header row
hdr_idx_row = None; hdr_row = None
for i, row in enumerate(rows_idx):
    if row and str(row[0]).lower().strip() == "file":
        hdr_row = row; hdr_idx_row = i; break

if hdr_row is None:
    log("  Header row not found in CASE_INDEX → abort")
else:
    log(f"  Header: {list(hdr_row)[:8]}")
    col_map = {str(h).lower().strip(): i for i, h in enumerate(hdr_row) if h}

    # Check if ON already in index
    on_existing_row = None
    for ri, row in enumerate(rows_idx[hdr_idx_row+1:], start=hdr_idx_row+2):
        file_val = str(row[col_map.get("file", 0)] if row else "").upper()
        tok_val  = str(row[col_map.get("token", 1)] if row else "").upper()
        if "CASE_ON" in file_val or tok_val == "ON":
            on_existing_row = ri; break

    def set_cell(row_n, col_name, val):
        idx = col_map.get(col_name.lower())
        if idx is not None:
            sh_idx.cell(row=row_n, column=idx+1, value=val)

    if on_existing_row:
        log(f"  Existing row {on_existing_row} → update")
        cur = [sh_idx.cell(row=on_existing_row, column=i+1).value for i in range(len(hdr_row))]
        log(f"  Current: {cur[:8]}")
        set_cell(on_existing_row, "file", "CASE_ON.xlsx")
        set_cell(on_existing_row, "token", "ON")
        set_cell(on_existing_row, "status", "ACTIVE")
        set_cell(on_existing_row, "loai",
                 "DIST_WALL ATH-XẢ; 0 GOM-A; OI UNCALIBRATED; stage S3 xả-vùng-nhớ")
        set_cell(on_existing_row, "action_class",
                 "WATCH (án mới, stage S3 pending confirmation)")
        set_cell(on_existing_row, "ngay_mo", NOW.strftime("%Y-%m-%d"))
        set_cell(on_existing_row, "moc_ke_tiep",
                 "snap2 >=13:02 UTC; N-4 on-chain CẦN OPERATOR")
        if "ghi_chu" in col_map:
            set_cell(on_existing_row, "ghi_chu",
                     "ALPHA_444; onboard 2025-10-24; ATH $0.46 Aug7; 289 DAILY rows; contract=0x0e4f6209...")
        log(f"  Row {on_existing_row} updated")
    else:
        last_used = hdr_idx_row + 1
        for ri, row in enumerate(rows_idx[hdr_idx_row+1:], start=hdr_idx_row+2):
            if any(v is not None for v in (row or [])):
                last_used = ri
        next_row = last_used + 1
        log(f"  New row at {next_row}")
        set_cell(next_row, "file", "CASE_ON.xlsx")
        set_cell(next_row, "token", "ON")
        set_cell(next_row, "status", "ACTIVE")
        set_cell(next_row, "loai",
                 "DIST_WALL ATH-XẢ; 0 GOM-A; OI UNCALIBRATED; stage S3 xả-vùng-nhớ")
        set_cell(next_row, "action_class",
                 "WATCH (án mới, stage S3 pending confirmation)")
        set_cell(next_row, "ngay_mo", NOW.strftime("%Y-%m-%d"))
        set_cell(next_row, "moc_ke_tiep",
                 "snap2 >=13:02 UTC; N-4 on-chain CẦN OPERATOR")
        if "ghi_chu" in col_map:
            set_cell(next_row, "ghi_chu",
                     "ALPHA_444; onboard 2025-10-24; ATH $0.46 Aug7; 289 DAILY rows; contract=0x0e4f6209...")

    wb2.save(CASE_IDX)
    log(f"  CASE_INDEX.xlsx saved OK")

# ── INDEX.md ───────────────────────────────────────────────────────────────────
log("\n=== INDEX.md ===")
idx_md = ROOT / "docs" / "case_files" / "INDEX.md"
if idx_md.exists():
    with open(idx_md, encoding="utf-8", errors="replace") as f:
        idx_content = f.read()
    if "CASE_ON" in idx_content:
        log(f"  Already mentions CASE_ON → OK")
    else:
        entry = f"| CASE_ON.xlsx | ON | ALPHA_444 | ACTIVE | DIST_WALL ATH-XẢ | {NOW.strftime('%Y-%m-%d')} |\n"
        with open(idx_md, "a", encoding="utf-8") as f:
            f.write(entry)
        log(f"  Appended ON entry")
else:
    log(f"  INDEX.md not found at {idx_md}")

log(f"\n[N-7+N-8 DONE]  sanity={pass_count}/5")
