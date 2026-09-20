#!/usr/bin/env python3
"""
Generate docs/case_files/case_state_map.md — VIEW SINH read-only từ CASE_INDEX.xlsx.

Format §2: token · state_ký · ngày_ký · verdict_ref (loại).
Không phải nơi ký. Operator ký state trong CASE_INDEX.xlsx như cũ.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

REPO = Path(r"D:\Data back test\whale_data")
XLSX = REPO / "docs" / "case_files" / "CASE_INDEX.xlsx"
OUT  = REPO / "docs" / "case_files" / "case_state_map.md"
OUT_MD5 = REPO / "data" / "state" / "case_index_source.md5"   # L1/19-09 · md5 nguồn cho VPS đối chiếu (32 ký tự, không SCP xlsx)

wb = openpyxl.load_workbook(str(XLSX), data_only=True, read_only=True)
ws = wb["INDEX"]

# Header row = row 4 per inspection
HEADER_ROW = 4
COL = {
    "file":    1,
    "token":   2,
    "status":  3,
    "loai":    4,
    "action":  5,
    "ngay_mo": 6,
    "T0":      7,
    "moc":     8,
    "han":     9,
    "MM_ID":  10,
}

# Special row 3 (TST orphan before header)
orphan_rows = [3]

entries = []
for r in orphan_rows + list(range(HEADER_ROW+1, ws.max_row+1)):
    tok = ws.cell(r, COL["token"]).value
    if not tok or str(tok).strip().lower() == "token": continue
    entries.append({
        "token":  str(tok).strip(),
        "state":  ws.cell(r, COL["status"]).value or "?",
        "ngay":   ws.cell(r, COL["ngay_mo"]).value or "?",
        "loai":   ws.cell(r, COL["loai"]).value or "",
        "action": ws.cell(r, COL["action"]).value or "",
        "file":   ws.cell(r, COL["file"]).value or "",
        "MM_ID":  ws.cell(r, COL["MM_ID"]).value or "",
    })
wb.close()

# Write markdown
now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
# [L-74 ENFORCEMENT (b) KÝ 01/09] header source_md5 + generated_at cho consumer đối chiếu
import hashlib as _hashlib
_source_md5 = _hashlib.md5(XLSX.read_bytes()).hexdigest()
_generated_at_iso = datetime.now(timezone.utc).isoformat()
lines = [
    f"# case_state_map — VIEW SINH read-only từ `CASE_INDEX.xlsx` sheet INDEX",
    f"",
    f"<!-- L-74 VIEW-AUTO-REGEN header · consumer đối chiếu md5 · lệch → in [VIEW-STALE age=Xd] -->",
    f"<!-- source_md5: {_source_md5} -->",
    f"<!-- generated_at: {_generated_at_iso} -->",
    f"<!-- source_path: docs/case_files/CASE_INDEX.xlsx -->",
    f"",
    f"**Generated:** {now}  ·  **Source:** `docs/case_files/CASE_INDEX.xlsx` (md5=`{_source_md5[:8]}`)",
    f"**Nguyên tắc (spec PHASE-TAG v0.3 §2, D1 chốt v0.3 #19):**",
    f"- Operator ký state TRONG `CASE_INDEX.xlsx` (canonical). File này là ống dẫn read-only.",
    f"- Đọc map = đọc verdict người ký, KHÔNG phạm Δ5.",
    f"- `case_tracker.md` giữ nguyên (hành chính án OPEN/RESOLVED/BLOCKED — khác sổ).",
    f"- Agent regenerate khi CASE_INDEX đổi. CẤM edit trực tiếp file này.",
    f"",
    f"---",
    f"",
    f"| token | state_ký | ngày_ký | verdict_ref (loại) | action_class | MM_ID | file |",
    f"|---|---|---|---|---|---|---|",
]
for e in entries:
    ng = e["ngay"]
    if hasattr(ng, "strftime"):
        ng = ng.strftime("%Y-%m-%d")
    verdict = str(e["loai"])[:80]
    lines.append(f"| {e['token']} | {e['state']} | {ng} | {verdict} | {e['action']} | {e['MM_ID']} | {e['file']} |")

lines += [
    f"",
    f"---",
    f"",
    f"**Tổng:** {len(entries)} case",
    f"",
    f"**Machine-readable JSON view (join-key cho phase_tag.py):**",
    f"",
    f"```json",
    "{",
]
for i, e in enumerate(entries):
    ng = e["ngay"]
    if hasattr(ng, "strftime"):
        ng = ng.strftime("%Y-%m-%d")
    comma = "," if i < len(entries)-1 else ""
    lines.append(f'  "{e["token"]}": {{"state": "{e["state"]}", "ngay_ky": "{ng}", "verdict_ref": "{str(e["loai"])[:60]}"}}{comma}')
lines += ["}", "```"]

OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
# L1/19-09 · ghi md5 nguồn ra file 32-byte cho VPS đối chiếu (không SCP xlsx)
OUT_MD5.parent.mkdir(parents=True, exist_ok=True)
OUT_MD5.write_text(_source_md5 + "\n", encoding="utf-8")
print(f"[gen] wrote {len(entries)} entries → {OUT}")
print(f"[gen] wrote source_md5={_source_md5[:12]}… → {OUT_MD5}")
print()
# Show first 5 lines of table
in_table = 0
for line in lines:
    if line.startswith("|") and "---" not in line:
        in_table += 1
        if 1 <= in_table <= 6:
            print(line)
print(f"...\ntotal: {len(entries)} case")
