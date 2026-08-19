"""_task_tria_verify_20260819.py — verify CASE_TRIA.xlsx structure"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_TRIA.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=True)
print(f"Sheets: {wb.sheetnames}")

ws = wb["DAILY"]
print(f"\nDAILY dims: {ws.dimensions}, max_row: {ws.max_row}")
row2 = [c.value for c in ws[2]]
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}

# print first and last rows
print(f"\nFirst DAILY row (3):")
for k, cx in list(ci.items())[:6]:
    v = ws.cell(3, cx).value
    print(f"  {k}: {v}")

print(f"\nLast DAILY row ({ws.max_row}):")
for k in ["date","c","vol_perp_usd","f_avg","fn (4h=6/day)","oi_usd","pos_r","acc_r","gLSR","div_avg3 (avg3p/avg3a)","data_source_lsr"]:
    if k in ci:
        v = ws.cell(ws.max_row, ci[k]).value
        print(f"  {k}: {v}")

# CASE sheet
ws_c = wb["CASE"]
print(f"\nCASE sheet, first 20 rows:")
for r in range(1, min(21, ws_c.max_row+1)):
    v1, v2 = ws_c.cell(r,1).value, ws_c.cell(r,2).value
    if v1 or v2:
        print(f"  r{r}: {v1!r} | {v2!r}")

# EVENTS
ws_e = wb["EVENTS"]
print(f"\nEVENTS ({ws_e.max_row} rows):")
for r in range(1, ws_e.max_row+1):
    row_v = [ws_e.cell(r,c).value for c in range(1,6)]
    if any(v is not None for v in row_v):
        print(f"  r{r}: {row_v}")

print("\n[DONE verify]")
