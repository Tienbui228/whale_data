"""
Register CASE_HEI in CASE_INDEX.xlsx (new row after last row).
L-34: PermissionError → raise.
"""
import sys, io
from pathlib import Path
from datetime import datetime, timezone

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import Font

CASE_INDEX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")

try:
    wb = openpyxl.load_workbook(CASE_INDEX)
except PermissionError as e:
    print(f"PERMISSION ERROR: {e}")
    raise

ws = wb["INDEX"]
print(f"  INDEX max_row before: {ws.max_row}")

# Check if HEI already registered
for ri in range(1, ws.max_row + 1):
    v = ws.cell(ri, 2).value
    if v and "HEI" in str(v).upper():
        print(f"  HEI already at row {ri} — updating")
        ws.cell(ri, 1).value = "CASE_HEI.xlsx"
        ws.cell(ri, 2).value = "HEI"
        ws.cell(ri, 3).value = "COLD_START — PENDING"
        ws.cell(ri, 4).value = "cold-start perp+spot, spike events 05-29/08-05/08-06"
        ws.cell(ri, 5).value = "COLD_START"
        ws.cell(ri, 6).value = "2026-08-08"
        ws.cell(ri, 7).value = "unknown"
        ws.cell(ri, 8).value = "duyệt operator"
        ws.cell(ri, 9).value = ""
        ws.cell(ri, 10).value = "PENDING"
        ws.cell(ri, 11).value = "cold-start 08-08; contract unconfirmed; L-23 2nd depth MISSING"
        break
else:
    # Append new row
    new_row = ws.max_row + 1
    print(f"  Appending HEI at row {new_row}")
    ws.cell(new_row, 1).value = "CASE_HEI.xlsx"
    ws.cell(new_row, 2).value = "HEI"
    ws.cell(new_row, 3).value = "COLD_START — PENDING"
    ws.cell(new_row, 4).value = "cold-start perp+spot, spike events 05-29/08-05/08-06"
    ws.cell(new_row, 5).value = "COLD_START"
    ws.cell(new_row, 6).value = "2026-08-08"
    ws.cell(new_row, 7).value = "unknown"
    ws.cell(new_row, 8).value = "duyệt operator"
    ws.cell(new_row, 9).value = ""
    ws.cell(new_row, 10).value = "PENDING"
    ws.cell(new_row, 11).value = "cold-start 08-08; contract unconfirmed; L-23 2nd depth MISSING"
    for ci in range(1, 12):
        ws.cell(new_row, ci).font = Font(size=9)

print(f"  INDEX max_row after: {ws.max_row}")

try:
    wb.save(CASE_INDEX)
    print(f"  Saved: {CASE_INDEX}")
except PermissionError as e:
    print(f"PERMISSION ERROR: {e}")
    print("  L-34: file mở ở Excel? Đóng lại rồi chạy lại.")
    raise

# Verify
wb2 = openpyxl.load_workbook(CASE_INDEX, data_only=True)
ws2 = wb2["INDEX"]
print(f"  Verify: last row = {ws2.max_row}")
last = [ws2.cell(ws2.max_row, ci).value for ci in range(1, 12)]
print(f"  Last row content: {last[:4]}")
print("[REGISTER HEI DONE]")
