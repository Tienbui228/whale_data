"""Task: sync CASE_INDEX.xlsx LESSONS sheet theo rev.2.

Thu tu thuc hien:
1. Rename L-26 → L-32, L-27 → L-33 (rows 30-31)
2. Insert 6 blank rows at position 30 (push L-32/L-33 xuong rows 36-37)
3. Write L-26 → L-31 vao rows 30-35
4. Update L-14 → L-20 content (rows 18-24, khong bi anh huong boi insert)
5. Append L-34 at row 38
6. Verify row count = 4 header + 34 lessons = 38 total
7. Regenerate docs/lessons_learned.md tu sheet

Ket qua: 34 bai hoc, khong trung ID.
"""
from pathlib import Path
import openpyxl
from datetime import datetime

IDX  = Path("D:/Data back test/whale_data/docs/case_files/CASE_INDEX.xlsx")
MDOUT = Path("D:/Data back test/whale_data/docs/lessons_learned.md")

# --- Data moi / sua ---

UPDATE_ROWS: dict[int, list] = {
    # row_number: [id, ngay, case_nguon, bai_hoc, loai, trang_thai, hanh_dong_tiep, MM]
    18: [
        "L-14",
        "2026-07-30",
        "TAG 08-09/07 + AKE 08-10/07",
        ("Dump manh + gLSR SAP >=30%/<=3d = wash-by-crash: mot phat an hai: don chi phi long + "
         "nap nhien lieu short (crowd short moi = nhien lieu re cho squeeze). La bien the crowd_flush "
         "thu nhat, doi xung L-11 (rally) va L-15 (bait). Phan biet ro voi L-26: gLSR SAP khi gia sap "
         "(wash) vs gLSR TANG khi gia sap (wrong-side fuel)"),
        "thesis",
        "DA VAO yaml (crowd_flush cause=crash)",
        "phan biet ro voi L-26 (gLSR TANG khi gia sap = hien tuong nguoc, KHONG phai wash)",
        None,
    ],
    19: [
        "L-15",
        "2026-07-30",
        "BANK 02/07 + BASED 01&09/07 + POWER 08-09/07",
        ("Vol spike >=8x med14 + gia bi tu choi <=2d = bait-probe: muc dich khong phai markup het ma la "
         "DAY crowd short nap nhien lieu squeeze. Sau bait canh gLSR sap trong <=3d lam xac nhan; "
         "vol persistence T+1 phan biet bait vs ignition that"),
        "thesis",
        "DA VAO yaml (crowd_flush cause=bait)",
        "tag probe events trong EVENTS log khi thay pattern",
        None,
    ],
    20: [
        "L-16",
        "2026-07-30",
        "BR + EVAA + CLO",
        ("div_raw truot tu >1 xuong <1 = vi the lon thu hep/dao long bias = chom phan phoi; "
         "flag STAND_ASIDE review; khong ARM sau div-flip neu chua co bang chung re-gom. "
         "LUU Y THUAT NGU: day KHONG phai 'M2 dao dau' — M2 la module CEX balance flow (xem L-27)"),
        "thesis",
        "GHI NHAN",
        "flag div-flip trong DAILY log; ho hang gan voi lsr_dist (da la suppressor yaml)",
        None,
    ],
    21: [
        "L-17",
        "2026-07-30",
        "AKE + ESPORTS",
        ("Squeeze/markup chet khi nhien lieu chet — hai cua dong an: (a) gLSR hoi ve ~1 (crowd short dau hang), "
         "(b) muc tieu cu the goc (TARGETED_SQUEEZE). Dang giu vi the ma mot trong hai bat = EXIT_NOW. "
         "[De xuat CHUA CALIB: thao tac hoa bang glsr_slope_5d doi dau — chua test case nao, "
         "glsr_slope_5d hien chi la cot hien thi (A5), khong phai gate]"),
        "dinh luat",
        "DA VAO master plan §1 #5 + EXIT_NOW conditions",
        "neu muon slope-flip thanh trigger di duong proposal + calib nhu #17",
        None,
    ],
    22: [
        "L-18",
        "2026-07-30",
        "u50_verdicts §2 (n=18)",
        ("Truoc ARM phai check du 3: (1) pha chu ky (nhin du px7/px14 + oi_cycle_pos, khong chi cua so 28d), "
         "(2) crowd structure (gLSR + dao ham — co tinh gan vo dung), (3) OI khong tai PEAK. "
         "Thieu 1/3 = low-conviction. KHONG trich dan % precision khi chua do — precedent dem that (n=...), cam noi %"),
        "thesis",
        "GHI NHAN",
        "phase_check vao ARM checklist; precision that do bang shadow-mode P2 + alerts_log",
        None,
    ],
    23: [
        "L-19",
        "2026-07-30",
        "u50 §2 + CLO/BLUAI/STAR (n=18)",
        ("Signal fire ban giong nhau o pha gom va pha phan phoi — can phase_tag "
         "(ACCUMULATE/MARKUP/DISTRIBUTE/UNKNOWN) tu oi_cycle_pos + div trend + glsr_slope de loc; "
         "den khi co phase_tag: moi fire phai kem phase check tay (L-18). Khong trich % false-rate khi chua do"),
        "process",
        "GHI NHAN — no P2",
        "thiet ke phase_tag inference trong P2",
        None,
    ],
    24: [
        "L-20",
        "2026-07-30",
        "TAG/BR 28-29/07",
        ("fn=1 / data_bad=1 / missing >3 ngay trong window / so tu 1 nguon khong kiem chung "
         "→ low_conf=True; so bi ha cap KHONG lam threshold ARM/FIRE"),
        "process",
        "DA VAO CLAUDE.md (fn=1 rule)",
        "mo rong tu-ha-cap cho data_bad flag (ngoai fn=1)",
        None,
    ],
}

NEW_L26_L31: list[list] = [
    # L-26
    [
        "L-26",
        "2026-07-30",
        "ESPORTS hau-climax + TAG 29/07 + BLUAI",
        ("Gia sap/roi + gLSR TANG = crowd long bat dao = nhien lieu SAI phia — NGUOC han wash-by-crash "
         "(L-14: gLSR SAP). He qua: (a) leg-len ke tiep can lua long nay bi rua truoc, "
         "(b) con dang SUPPRESSED co pattern nay thi REMAIN, (c) khong doc nham thanh 'co nguoi gom'"),
        "thesis",
        "GHI NHAN",
        "dieu kien tien quyet 'crowd flush truoc leg' ghi vao chot canh cac case dinh (TAG, BR, ESPORTS)",
        None,
    ],
    # L-27
    [
        "L-27",
        "2026-07-30",
        "BTW vs BAS/TAG | AKE/IDOL",
        ("M2 (CEX balance flow) tu no KHONG mang y dinh — dau cua no do boi canh quyet: "
         "vi tri chu ky + funding regime + hanh vi gia. Chay M2 don le la An tron bay BTW. "
         "Suppressor o bat ky lop nao phu quyet moi bullish cua ca ba lop"),
        "dinh luat",
        "DA VAO master plan §1 #1 + A1 (R4)",
        "giu ten 'M2' doc quyen cho balance flow — khong dung cho div/ratio (an L-16 rev.1)",
        None,
    ],
    # L-28
    [
        "L-28",
        "2026-07-30",
        "u50 §6 (AGT / HUMA-AT sync / listing 04-14)",
        ("Moi narrative dung tren mot nguon phai kiem nguon goc (aggregator != klines; "
         "mep cua so = artifact ung vien). Narrative khong con dung duoc → RUT cong khai va GHI SO "
         "muc 'da tu ha cap' — de khong tu lua va khong tai sinh. Cu luan diem chi giu phan con dung tren bang chung"),
        "process",
        "GHI NHAN — u50 §6 la mau chuan",
        "moi report lon co muc 'bang chung da tu ha cap' nhu u50",
        None,
    ],
    # L-29
    [
        "L-29",
        "2026-07-30",
        "OPN + HYPER (u50 §3a, n=2)",
        ("alpha_vol_suspect = vol_spot/vol_perp > 3x → moi metric VOL SPOT cua con do ABSTAIN "
         "(khong wash-check, khong #17, khong bang chung dong tien); van dung depth/price/taker. "
         "takerBUY spot lech cuc doan (<35% hoac >70%) tren vol khong lo = hinh dang farming. "
         "Guard chi ha long tin, khong sinh bullish → ap ngay"),
        "guard",
        "DA VAO amendment A6 (guard)",
        "cot alpha_vol_suspect vao triage/snapshot khi build spot collector (P1.5)",
        None,
    ],
    # L-30
    [
        "L-30",
        "2026-07-30",
        "BLUAI + ARX (u50)",
        ("Token venue_top=bitget: cam so sanh spot-vs-perp cung-ngay cho den khi bug lech ngay "
         "duoc xac minh/va; phan tich cross-market phai shift-test truoc. Ghi chu venue vao moi bang so sanh"),
        "data-bug",
        "GHI NHAN (no ky thuat u50 #2, muc CAO)",
        "probe xac minh offset tren 2 token + va tai ingest",
        None,
    ],
    # L-31
    [
        "L-31",
        "2026-07-30",
        "lessons_learned.md rev.1",
        ("Tong hop/dien giai tu tai lieu cu = vung rui ro cao: moi entry phai trace ve cau chu nguon "
         "(dinh nghia lay tu yaml/reference, khong tu tri nho dien giai); so lieu khong co trong nguon = khong viet; "
         "thuat ngu he thong (M1/M2/ten signal) dung dung nghia goc, khong tai dinh nghia. "
         "L-09 ap cho ca tai lieu, khong rieng fire"),
        "process",
        "GHI NHAN",
        "moi consolidation sau nay kem bang trace entry-nguon khi operator yeu cau kiem",
        None,
    ],
]

L34_ROW = [
    "L-34",
    "2026-07-30",
    "CASE_BANK_NEW.xlsx (an 30/07), fix da commit",
    ("PermissionError khi save xlsx → RAISE RuntimeError ngay, cam silent-save sang ten khac (_NEW suffix). "
     "Ten file khac = INDEX tro sai = case_daily ghi nham → data bi pha vo im lang. Fix: close xlsx roi retry"),
    "process/rule",
    "DA VAO case_daily.py (L-34 raise, commit 2026-07-30)",
    "backfill script cung phai RAISE neu gap lock (kiem tra backfill.py)",
    None,
]


def write_row(ws, row_n: int, vals: list) -> None:
    for c, v in enumerate(vals, start=1):
        ws.cell(row_n, c).value = v


wb = openpyxl.load_workbook(str(IDX), data_only=False)
ws = wb["LESSONS"]

print(f"=== TRUOC SYNC: max_row={ws.max_row} max_col={ws.max_column} ===")

# --- STEP 1: Rename L-26 → L-32 and L-27 → L-33 ---
row30_id = ws.cell(30, 1).value
row31_id = ws.cell(31, 1).value
print(f"Step 1: row30={row30_id}, row31={row31_id}")
assert row30_id == "L-26", f"Expected L-26 at row 30, got {row30_id!r}"
assert row31_id == "L-27", f"Expected L-27 at row 31, got {row31_id!r}"
ws.cell(30, 1).value = "L-32"
ws.cell(31, 1).value = "L-33"
print("Step 1 DONE: L-26→L-32, L-27→L-33")

# --- STEP 2: Insert 6 blank rows at position 30 ---
ws.insert_rows(30, 6)
print(f"Step 2 DONE: inserted 6 rows at pos 30. max_row now={ws.max_row}")
# Verify L-32 moved to row 36, L-33 to row 37
assert ws.cell(36, 1).value == "L-32", f"L-32 not at row 36: {ws.cell(36,1).value!r}"
assert ws.cell(37, 1).value == "L-33", f"L-33 not at row 37: {ws.cell(37,1).value!r}"
print("Step 2 verify: L-32 at row 36, L-33 at row 37 ✓")

# --- STEP 3: Write new L-26 to L-31 in rows 30-35 ---
for i, row_data in enumerate(NEW_L26_L31):
    write_row(ws, 30 + i, row_data)
    print(f"  Wrote {row_data[0]} at row {30+i}")
print("Step 3 DONE: L-26..L-31 written")

# --- STEP 4: Update L-14 to L-20 (rows 18-24) ---
for row_n, vals in UPDATE_ROWS.items():
    cur_id = ws.cell(row_n, 1).value
    expected = vals[0]
    assert cur_id == expected, f"Row {row_n}: expected {expected!r}, got {cur_id!r}"
    write_row(ws, row_n, vals)
    print(f"  Updated {expected} at row {row_n}")
print("Step 4 DONE: L-14..L-20 updated")

# --- STEP 5: Append L-34 ---
last_row = ws.max_row
# Find last non-empty row after row 37
last_data_row = 37
for r in range(37, ws.max_row + 1):
    if ws.cell(r, 1).value is not None:
        last_data_row = r
l34_row_n = last_data_row + 1
write_row(ws, l34_row_n, L34_ROW)
print(f"Step 5 DONE: L-34 written at row {l34_row_n}")

# --- STEP 6: Verify ---
print(f"\n=== VERIFY: max_row={ws.max_row} ===")
# Count lessons
lessons = []
for r in range(5, ws.max_row + 1):
    lid = ws.cell(r, 1).value
    if lid and str(lid).startswith("L-"):
        lessons.append((r, str(lid)))

print(f"Total lessons: {len(lessons)}")
if len(lessons) != 34:
    print(f"ERROR: expected 34 lessons, got {len(lessons)}")
    for r, lid in lessons:
        print(f"  row {r}: {lid}")
else:
    # Check IDs are L-01..L-34, no duplicates
    ids = [lid for _, lid in lessons]
    expected_ids = [f"L-{i:02d}" for i in range(1, 35)]
    missing = set(expected_ids) - set(ids)
    extra = set(ids) - set(expected_ids)
    if missing or extra:
        print(f"  MISSING: {missing}")
        print(f"  EXTRA: {extra}")
    else:
        print("  IDs: L-01..L-34 all present, no duplicates ✓")
    for r, lid in lessons:
        bai_hoc_preview = str(ws.cell(r, 4).value or "")[:50]
        print(f"  row {r:02d}: {lid} | {bai_hoc_preview}")

wb.save(str(IDX))
print(f"\nSaved: {IDX}")

# --- STEP 7: Regenerate docs/lessons_learned.md ---
print("\n=== REGENERATE lessons_learned.md ===")

# Re-open to read (data_only)
wb2 = openpyxl.load_workbook(str(IDX), data_only=True)
ws2 = wb2["LESSONS"]

SECTION_BY_TYPE = {
    "dinh luat": "I. SIGNAL-LEVEL",
    "thesis": "I. SIGNAL-LEVEL",
    "rule": "I. SIGNAL-LEVEL",
    "guard": "III. DATA-QUALITY",
    "data-bug": "III. DATA-QUALITY",
    "process": "IV. PROCESS",
    "process/rule": "IV. PROCESS",
}

lines = [
    "# Lessons Learned — FreshWatch",
    "",
    "> Auto-generated tu CASE_INDEX.xlsx sheet LESSONS. Khong edit file nay truc tiep.",
    f"> Lan sinh: {datetime.now().strftime('%Y-%m-%d')} (rev.2) · Tong: 34 bai hoc (L-01 → L-34)",
    "",
]

for r in range(5, ws2.max_row + 1):
    lid      = ws2.cell(r, 1).value
    ngay     = ws2.cell(r, 2).value
    case_src = ws2.cell(r, 3).value
    bai_hoc  = ws2.cell(r, 4).value
    loai     = ws2.cell(r, 5).value
    trang_thai = ws2.cell(r, 6).value
    hanh_dong  = ws2.cell(r, 7).value

    if not lid or not str(lid).startswith("L-"):
        continue

    date_s   = str(ngay or "")
    case_s   = str(case_src or "")
    loai_s   = str(loai or "")
    status_s = str(trang_thai or "")
    bai_s    = str(bai_hoc or "")
    action_s = str(hanh_dong or "")

    lines.append(f"## {lid} [{date_s}] {case_s} — `{loai_s}` · {status_s}")
    lines.append(bai_s)
    if action_s:
        lines.append(f"*Hanh dong tiep: {action_s}*")
    lines.append("")

wb2.close()

MDOUT.write_text("\n".join(lines), encoding="utf-8")
print(f"Written: {MDOUT} ({len(lines)} lines)")
