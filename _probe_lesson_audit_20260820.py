"""_probe_lesson_audit_20260820.py — Station 2: LESSONS sheet in CASE_INDEX.xlsx."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=True)
print(f"CASE_INDEX sheets: {wb.sheetnames}")

if "LESSONS" not in wb.sheetnames:
    print("!! LESSONS sheet KHÔNG có trong CASE_INDEX.xlsx")
    sys.exit(1)

ws = wb["LESSONS"]
print(f"LESSONS dims: {ws.dimensions}, max_row: {ws.max_row}, max_col: {ws.max_column}")

# Read row 1 header + col map
row1 = [c.value for c in ws[1]]
print(f"Headers: {row1}")

# Find ID column
id_col = None
for i, v in enumerate(row1, start=1):
    if v and str(v).strip().lower() in ("id", "lesson_id", "l_id"):
        id_col = i; break
if id_col is None:
    # try col 1 by default
    id_col = 1

# List all IDs
ids = []
for r in range(2, ws.max_row+1):
    v = ws.cell(r, id_col).value
    if v is not None:
        ids.append((r, str(v).strip()))

print(f"\nTotal LESSONS rows: {len(ids)}")
print(f"IDs full list:")
for row, id_ in ids:
    print(f"  row {row:>3}: {id_}")

# Max ID (parse numeric suffix)
import re
max_num = 0
max_id = None
for _, id_ in ids:
    m = re.match(r"L-?(\d+)", id_)
    if m:
        n = int(m.group(1))
        if n > max_num:
            max_num = n; max_id = id_
print(f"\nMax lesson-ID: {max_id} (numeric={max_num})")
