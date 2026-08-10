"""
_task_br3_2_build_20260810.py
BR3-2b: Rename scratch files to broccolif3b_* convention + run build_case.py
"""
import sys, io, json, shutil, subprocess
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT    = Path(r"D:\Data back test\whale_data")
SCRATCH = ROOT / "scratch"
CASE_DIR = ROOT / "docs" / "case_files"
DATE_TAG = "20260810"
SYM_PERP = "BROCCOLIF3BUSDT"
SYM_BASE = "BROCCOLIF3B"
SYM_LOWER = "broccolif3b"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"BR3-2b BUILD  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

CASE_Q = ("Có ai gom hàng không? Gom từ bao giờ, bao nhiêu "
          "(§3 ba phương pháp), đang ở pha nào vòng đời?")

# ── Step 1: Copy/rename br3_* → broccolif3b_* ─────────────────────────────────
log("\n=== Step 1: Copy scratch files to broccolif3b_* naming ===")
copies = [
    (f"br3_perp_{DATE_TAG}.csv",        f"{SYM_LOWER}_perp_{DATE_TAG}.csv"),
    (f"br3_lsr_raw_{DATE_TAG}.csv",     f"{SYM_LOWER}_lsr_raw_{DATE_TAG}.csv"),
    (f"br3_funding_{DATE_TAG}.csv",     f"{SYM_LOWER}_funding_{DATE_TAG}.csv"),
    (f"br3_oi_hist_{DATE_TAG}.csv",     f"{SYM_LOWER}_oi_hist_{DATE_TAG}.csv"),
    (f"br3_spot_{DATE_TAG}.csv",        f"{SYM_LOWER}_spot_{DATE_TAG}.csv"),
    (f"br3_bapi_{DATE_TAG}.json",       f"{SYM_LOWER}_bapi_{DATE_TAG}.json"),
    (f"br3_perp_info_{DATE_TAG}.json",  f"{SYM_LOWER}_perp_info_{DATE_TAG}.json"),
]
for src_name, dst_name in copies:
    src = SCRATCH / src_name
    dst = SCRATCH / dst_name
    if src.exists():
        shutil.copy2(src, dst)
        log(f"  Copied: {src_name} → {dst_name} ({dst.stat().st_size:,} bytes)")
    else:
        log(f"  MISSING src: {src_name}")

# ── Step 2: L-34 backup ────────────────────────────────────────────────────────
log("\n=== Step 2: L-34 backup ===")
case_xlsx = CASE_DIR / f"CASE_{SYM_BASE}.xlsx"
if case_xlsx.exists():
    bkp = SCRATCH / f"CASE_{SYM_BASE}_bak_{NOW.strftime('%Y%m%d_%H%M%S')}.xlsx"
    shutil.copy2(case_xlsx, bkp)
    log(f"  Backup: {bkp.name}")
else:
    log(f"  No existing CASE_{SYM_BASE}.xlsx — fresh build")

# ── Step 3: build_case.py ─────────────────────────────────────────────────────
log("\n=== Step 3: build_case.py ===")
cmd = [sys.executable, str(ROOT / "build_case.py"), SYM_PERP, DATE_TAG, CASE_Q]
log(f"  CMD: {' '.join(cmd[:3])} ...")
result = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8",
                        errors="replace", cwd=str(ROOT))
print(result.stdout)
if result.stderr:
    print("STDERR:", result.stderr[:500])
log(f"  BUILD exit={result.returncode}")

# ── Step 4: Verify CASE_BROCCOLIF3B.xlsx ──────────────────────────────────────
log("\n=== Step 4: Verify CASE_BROCCOLIF3B.xlsx ===")
if case_xlsx.exists():
    wb = openpyxl.load_workbook(str(case_xlsx), read_only=True)
    log(f"  Sheets: {wb.sheetnames}")
    ws = wb["DAILY"]
    max_row = ws.max_row
    log(f"  DAILY: {max_row} rows × {ws.max_column} cols")
    # Read last 5 data rows (full 50 cols) — BR3-2 verify tay
    log(f"\n  === LAST 5 DAILY ROWS FULL (verify tay) ===")
    # Get headers from row 2
    headers = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    log(f"  Headers (row 2): {headers}")
    # Get last 5 data rows
    data_rows = list(ws.iter_rows(min_row=3, values_only=True))
    log(f"  Total data rows: {len(data_rows)}")
    for row in data_rows[-5:]:
        # Show all 50 cols
        row_dict = {h: v for h, v in zip(headers, row)}
        log(f"\n  ── Row (date={row[0]}) ──")
        for h, v in zip(headers, row):
            if v is not None:
                log(f"    {h}: {v}")
            else:
                log(f"    {h}: None")
    # Read CASE sheet
    log(f"\n  === CASE sheet ===")
    cs = wb["CASE"]
    for r in cs.iter_rows(values_only=True):
        if r[0]:
            log(f"  {r[0]}: {str(r[1])[:100] if r[1] else ''}")
    wb.close()
    log(f"\n  File size: {case_xlsx.stat().st_size:,} bytes")
else:
    log(f"  CASE_{SYM_BASE}.xlsx NOT CREATED")

log(f"\n[BR3-2b BUILD DONE]")
