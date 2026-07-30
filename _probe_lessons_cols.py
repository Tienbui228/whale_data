"""Probe: dump header rows + first 2 data rows of LESSONS sheet."""
import openpyxl
from pathlib import Path

IDX = Path("D:/Data back test/whale_data/docs/case_files/CASE_INDEX.xlsx")
wb = openpyxl.load_workbook(IDX, data_only=True)
ws = wb["LESSONS"]
print(f"max_row={ws.max_row}  max_col={ws.max_column}")
for r in range(1, 8):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    print(f"row {r:02d}: {vals}")
wb.close()
