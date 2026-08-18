"""_task_gua_b2_b3_verify_20260818.py
B1 (append Aug17) + B2 (LSR-repair Jul01→Aug17 window from freshwatch + backup)
+ recompute derived LSR-dependent cols + flags.

Write-after-guard: verify acceptance (Jul29 acc_r≈6.03, gLSR≈5.93 · Jul31 acc_r=6.1124
· Aug07 gLSR=2.8745) BEFORE save. Fail → DUMP diff, do NOT touch xlsx.
"""
import sys, io, csv, json, urllib.request, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT    = Path(r"D:\Data back test\whale_data")
SCRATCH = ROOT / "scratch"
EXPORT  = ROOT / "data" / "export"
CASES   = ROOT / "docs" / "case_files"

SYM     = "GUA"
TICKER  = "GUAUSDT"
XLSX    = CASES / "CASE_GUA.xlsx"
BACKUP  = SCRATCH / "CASE_GUA_BACKUP_20260817_0315.xlsx"

DATE_AUG17 = "2026-08-17"
WINDOW_LO  = "2026-07-01"
WINDOW_HI  = "2026-08-17"

def safe_float(v):
    if v in (None, "", "None"): return None
    try: return float(v)
    except (ValueError, TypeError): return None

# ── 1. LOAD current CASE_GUA.xlsx ─────────────────────────────────────────
print("="*70)
print("STEP 1: Load current CASE_GUA.xlsx DAILY")
print("="*70)
wb = openpyxl.load_workbook(XLSX)
ws = wb["DAILY"]
row1 = [c.value for c in ws[1]]  # section titles
row2 = [c.value for c in ws[2]]  # field names
n_cols = len(row2)
print(f"  cols: {n_cols}, dims: {ws.dimensions}")

# Build col_idx map
col_idx = {}
for i, name in enumerate(row2, start=1):
    if name is not None:
        col_idx[str(name).strip()] = i
print(f"  key cols: date={col_idx.get('date')} pos_r={col_idx.get('pos_r')} "
      f"gLSR={col_idx.get('gLSR')} div_avg3={col_idx.get('div_avg3')}")

# Load data rows (row 3+) into dict {date: {col_name: value}}
data_by_date = {}
data_dates   = []  # ordered
for r in range(3, ws.max_row + 1):
    date_val = ws.cell(r, col_idx["date"]).value
    if not date_val: continue
    d_str = str(date_val)
    if len(d_str) > 10: d_str = d_str[:10]
    row_data = {}
    for name, ci in col_idx.items():
        row_data[name] = ws.cell(r, ci).value
    row_data["__row"] = r
    data_by_date[d_str] = row_data
    data_dates.append(d_str)

print(f"  data rows loaded: {len(data_dates)}, "
      f"range={data_dates[0]}→{data_dates[-1]}")
LAST_ROW_XLSX = max(v["__row"] for v in data_by_date.values())
print(f"  max row in xlsx: {LAST_ROW_XLSX}")

# ── 2. APPEND Aug17 perp/funding/OI ───────────────────────────────────────
print("\n" + "="*70)
print("STEP 2: Build Aug17 row from fetched scratch CSVs")
print("="*70)

def load_csv(path):
    if not path.exists(): return []
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))

perp_new = load_csv(SCRATCH / "gua_perp_aug17_only_20260818.csv")
fund_new = load_csv(SCRATCH / "gua_funding_aug17_only_20260818.csv")
oi_new   = load_csv(SCRATCH / "gua_oi_hist_20260818.csv")

perp17 = perp_new[0] if perp_new else None
print(f"  perp Aug17: c={perp17['close'] if perp17 else 'MISSING'}")

# funding aggregate
fund17 = [r for r in fund_new if r["date"] == DATE_AUG17]
fn = len(fund17)
frs = [float(r["fundingRate"]) for r in fund17]
f_avg = round(sum(frs)/len(frs), 8) if frs else None
f_min = round(min(frs), 8) if frs else None
f_max = round(max(frs), 8) if frs else None
print(f"  funding Aug17: fn={fn} f_avg={f_avg} f_min={f_min} f_max={f_max}")

# OI Aug17 from dedup CSV
oi17 = [r for r in oi_new if r["date"] == DATE_AUG17]
oi_row = oi17[0] if oi17 else {}
oi_usd = safe_float(oi_row.get("oi_usd"))
oi_tok = safe_float(oi_row.get("oi_token"))
print(f"  OI Aug17: oi_usd={oi_usd} oi_token={oi_tok}")

# taker_scale_check: perp CSV writes taker_buy_pct in PERCENT (0-100). Verify.
if perp17:
    tkr = safe_float(perp17.get("taker_buy_pct"))
    if tkr is not None and 0 <= tkr <= 1.5:
        raise SystemExit(f"TAKER SCALE FAIL: got {tkr}, expected percent (>1). "
                         f"Refuse to write — check fetch code.")
    print(f"  taker_scale_check: OK (taker_buy_pct={tkr}% in percent scale)")

# ret_1d Aug17 vs Aug16 close
prev = data_by_date.get("2026-08-16", {})
c_prev = safe_float(prev.get("c"))
c_new  = safe_float(perp17.get("close")) if perp17 else None
ret1d_17 = round((c_new - c_prev) / c_prev * 100, 4) if (c_new and c_prev) else None
print(f"  ret_1d Aug17: c_prev={c_prev} c_new={c_new} ret_1d={ret1d_17}%")

# ── 3. LOAD LSR sources ───────────────────────────────────────────────────
print("\n" + "="*70)
print("STEP 3: Load LSR sources (freshwatch data/export + backup)")
print("="*70)

# 3a. backup CASE_GUA_BACKUP_20260817_0315.xlsx — extract freshwatch rows
print(f"\n  → backup: {BACKUP.name}")
wb_bak = openpyxl.load_workbook(BACKUP, read_only=True, data_only=True)
ws_bak = wb_bak["DAILY"]
bak_row2 = [c.value for c in ws_bak[2]]
bak_ci   = {str(n).strip(): i+1 for i, n in enumerate(bak_row2) if n}
lsr_from_backup = {}  # {date: {pos_r, acc_r, gLSR, taker_lsr}}
for r in range(3, ws_bak.max_row + 1):
    d_val = ws_bak.cell(r, bak_ci["date"]).value
    if not d_val: continue
    d_str = str(d_val)[:10]
    src   = ws_bak.cell(r, bak_ci["data_source_lsr"]).value
    if src != "freshwatch": continue
    lsr_from_backup[d_str] = {
        "pos_r":     safe_float(ws_bak.cell(r, bak_ci["pos_r"]).value),
        "acc_r":     safe_float(ws_bak.cell(r, bak_ci["acc_r"]).value),
        "gLSR":      safe_float(ws_bak.cell(r, bak_ci["gLSR"]).value),
        "taker_lsr": safe_float(ws_bak.cell(r, bak_ci["taker_lsr"]).value),
    }
wb_bak.close()
print(f"  backup freshwatch rows: {len(lsr_from_backup)}, "
      f"range={min(lsr_from_backup)}→{max(lsr_from_backup)}")

# 3b. freshwatch store (data/export/) — canonical
print(f"\n  → freshwatch data/export/")
lsr_from_store = {}
fw_files = sorted(EXPORT.glob("freshwatch_daily_2026-*.csv"))
for fp in fw_files:
    with open(fp, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("symbol", "").upper() != TICKER: continue
            d_str = row.get("date","")[:10]
            if not d_str: continue
            pos_r = safe_float(row.get("lsr_pos_r"))
            if not pos_r: continue  # skip if primary LSR missing
            lsr_from_store[d_str] = {
                "pos_r":     pos_r,
                "acc_r":     safe_float(row.get("lsr_acc_r")),
                "gLSR":      safe_float(row.get("lsr_global_r")),
                "taker_lsr": safe_float(row.get("taker_ratio")),
            }
print(f"  freshwatch_store rows: {len(lsr_from_store)}, "
      f"range={min(lsr_from_store)}→{max(lsr_from_store)}")

# 3c. also load scratch/freshwatch_daily_*.csv as fallback (may have extra)
lsr_scratch = {}
for fp in sorted(SCRATCH.glob("freshwatch_daily_2026-*.csv")):
    with open(fp, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("symbol", "").upper() != TICKER: continue
            d_str = row.get("date","")[:10]
            pos_r = safe_float(row.get("lsr_pos_r"))
            if not d_str or not pos_r: continue
            lsr_scratch[d_str] = {
                "pos_r":     pos_r,
                "acc_r":     safe_float(row.get("lsr_acc_r")),
                "gLSR":      safe_float(row.get("lsr_global_r")),
                "taker_lsr": safe_float(row.get("taker_ratio")),
            }
print(f"  freshwatch_scratch rows: {len(lsr_scratch)}, "
      f"range={min(lsr_scratch)}→{max(lsr_scratch)}")

# ── 4. VERIFY ACCEPTANCE (BEFORE any write) ───────────────────────────────
print("\n" + "="*70)
print("STEP 4: ACCEPTANCE VERIFY (án E exhibit numbers) — DO OR DIE")
print("="*70)

def get_lsr(d):
    """Priority: freshwatch_store > freshwatch_scratch > backup_0817"""
    if d in lsr_from_store: return lsr_from_store[d], "freshwatch_store"
    if d in lsr_scratch:    return lsr_scratch[d],   "freshwatch_scratch"
    if d in lsr_from_backup:return lsr_from_backup[d], "backup_0817"
    return None, None

acceptance = [
    ("2026-07-29", "acc_r", 6.03,   0.05),  # tolerance ±0.05
    ("2026-07-29", "gLSR",  5.93,   0.05),
    ("2026-07-31", "acc_r", 6.1124, 0.005),
    ("2026-08-07", "gLSR",  2.8745, 0.005),
]

verify_pass = True
for date, field, expected, tol in acceptance:
    r, src = get_lsr(date)
    if r is None:
        print(f"  ✗ {date} {field}: MISSING from all sources — FAIL")
        verify_pass = False
        continue
    got = r.get(field)
    if got is None:
        print(f"  ✗ {date} {field}: got None (src={src}) — FAIL")
        verify_pass = False
        continue
    diff = abs(got - expected)
    ok = diff <= tol
    marker = "✓" if ok else "✗"
    print(f"  {marker} {date} {field}: got={got:.4f} expected={expected} "
          f"diff={diff:.4f} tol={tol} src={src} {'PASS' if ok else 'FAIL'}")
    if not ok: verify_pass = False

if not verify_pass:
    print("\n!!!! ACCEPTANCE VERIFY FAILED !!!!")
    print("!!!! DUMP diff above · NOT writing xlsx · exit 1 !!!!")
    sys.exit(1)

print("\n  → ACCEPTANCE PASS ✓ · proceed to build in-memory + write")

# ── 5. BUILD LSR-repair map for full window ────────────────────────────────
print("\n" + "="*70)
print("STEP 5: Build LSR-repair map for Jul01→Aug17 window")
print("="*70)

repair_map = {}  # {date: {pos_r, acc_r, gLSR, taker_lsr, source}}
# window dates
from datetime import date as _date, timedelta
d0 = _date.fromisoformat(WINDOW_LO)
d1 = _date.fromisoformat(WINDOW_HI)
cur = d0
while cur <= d1:
    ds = cur.isoformat()
    r, src = get_lsr(ds)
    if r:
        repair_map[ds] = {**r, "source": src}
    cur += timedelta(days=1)

n_store = sum(1 for v in repair_map.values() if v["source"]=="freshwatch_store")
n_scr   = sum(1 for v in repair_map.values() if v["source"]=="freshwatch_scratch")
n_bak   = sum(1 for v in repair_map.values() if v["source"]=="backup_0817")
print(f"  window days: {(d1-d0).days+1}, filled: {len(repair_map)}")
print(f"    freshwatch_store: {n_store}")
print(f"    freshwatch_scratch: {n_scr}")
print(f"    backup_0817: {n_bak}")
missing = [d for d in (d0+timedelta(n) for n in range((d1-d0).days+1))
           if d.isoformat() not in repair_map]
if missing:
    print(f"  MISSING dates in window ({len(missing)}): {[d.isoformat() for d in missing]}")

# ── 6. IN-MEMORY APPEND Aug17 + WRITE LSR + recompute derived ─────────────
print("\n" + "="*70)
print("STEP 6: Build merged in-memory dataset + recompute derived")
print("="*70)

# Ensure Aug17 exists in data_by_date (create if not)
if DATE_AUG17 not in data_by_date:
    new_row = {name: None for name in col_idx}
    new_row["__row"] = LAST_ROW_XLSX + 1
    data_by_date[DATE_AUG17] = new_row
    data_dates.append(DATE_AUG17)
    print(f"  ADD new row for {DATE_AUG17} at xlsx row {new_row['__row']}")

# Populate Aug17 perp/funding/OI in memory
row17 = data_by_date[DATE_AUG17]
row17["date"] = DATE_AUG17
row17["px_avg (H+L+C)/3"] = round((float(perp17["high"])+float(perp17["low"])+float(perp17["close"]))/3, 8)
row17["o"] = safe_float(perp17["open"])
row17["h"] = safe_float(perp17["high"])
row17["l"] = safe_float(perp17["low"])
row17["c"] = safe_float(perp17["close"])
row17["vol_perp_usd"] = safe_float(perp17["vol_quote_usd"])
row17["taker_buy_pct (k[10]/k[7])"] = safe_float(perp17["taker_buy_pct"])
row17["f_avg"] = f_avg
row17["f_min"] = f_min
row17["f_max"] = f_max
row17["fn (4h=6/day)"] = fn
row17["oi_usd"] = oi_usd
row17["oi_token"] = oi_tok
row17["data_ok (fn>=6)"] = (fn >= 6)
row17["low_conf (fn<6)"]  = (fn < 6)
row17["ret_1d (close-to-close)"] = ret1d_17
row17["data_source_perp"] = "fapi_klines_1d"

# Populate LSR for full window
for d, lsr in repair_map.items():
    if d not in data_by_date: continue
    r = data_by_date[d]
    r["pos_r"]     = lsr["pos_r"]
    r["acc_r"]     = lsr["acc_r"]
    r["gLSR"]      = lsr["gLSR"]
    r["taker_lsr"] = lsr["taker_lsr"]
    r["data_source_lsr"] = lsr["source"]

# Recompute derived cols (window Jul01→Aug17)
# We need FULL arrays to compute glsr_slope_5d (needs t-5), div_avg3 (needs t-2), d_pos_3d (needs t-3)
# So iterate over ALL data_dates (sorted) and rebuild these cols where needed.
sorted_dates = sorted(data_dates)
# Rebuild pos_r_all / acc_r_all / glsr_all arrays FRESH from repair_map for the window,
# keep existing values outside window
pos_arr = []; acc_arr = []; glsr_arr = []
for d in sorted_dates:
    r = data_by_date[d]
    pos_arr.append(safe_float(r.get("pos_r")))
    acc_arr.append(safe_float(r.get("acc_r")))
    glsr_arr.append(safe_float(r.get("gLSR")))

def recompute_row(i, d):
    r = data_by_date[d]
    # div_raw = pos_r/acc_r
    p, a = pos_arr[i], acc_arr[i]
    r["div_raw (pos_r/acc_r)"] = round(p/a, 4) if (p and a and a != 0) else None
    # div_avg3 = avg3(pos)/avg3(acc)
    if i >= 2:
        p3 = [v for v in pos_arr[i-2:i+1] if v is not None]
        a3 = [v for v in acc_arr[i-2:i+1] if v is not None]
        if len(p3)==3 and len(a3)==3:
            aa = sum(a3)/3
            r["div_avg3 (avg3p/avg3a)"] = round((sum(p3)/3)/aa, 4) if aa != 0 else None
        else:
            r["div_avg3 (avg3p/avg3a)"] = None
    # glsr_slope_5d = gLSR(t)/gLSR(t-5)
    if i >= 5:
        g, gp = glsr_arr[i], glsr_arr[i-5]
        r["glsr_slope_5d (ratio t/t-5)"] = round(g/gp, 4) if (g and gp and gp > 0) else None
    # d_pos_3d, d_acc_3d
    if i >= 3:
        pp = pos_arr[i-3]; ap = acc_arr[i-3]
        pn = pos_arr[i];   an = acc_arr[i]
        r["d_pos_3d (%chg)"] = round((pn-pp)/pp*100, 4) if (pn and pp and pp != 0) else None
        r["d_acc_3d (%chg)"] = round((an-ap)/ap*100, 4) if (an and ap and ap != 0) else None

# LSR-dependent flags: flag_hedge_struct (div_avg3<=0.5), flag_knife_crowd (gLSR↑&ret<-3)
def recompute_flags(i, d):
    r = data_by_date[d]
    d3 = r.get("div_avg3 (avg3p/avg3a)")
    r["flag_hedge_struct (div_avg3<=0.5)"] = bool(d3 is not None and d3 <= 0.5)
    if i >= 1:
        g, gp = glsr_arr[i], glsr_arr[i-1]
        ret = safe_float(r.get("ret_1d (close-to-close)"))
        knife = (g is not None and gp is not None and ret is not None
                 and g > gp and ret < -3.0)
        r["flag_knife_crowd (gLSR↑&ret<-3%)"] = bool(knife)

# Apply recompute across ALL rows (LSR affects arrays globally)
for i, d in enumerate(sorted_dates):
    recompute_row(i, d)
    recompute_flags(i, d)

# Sanity: after repair, count non-null LSR in window
in_window = [d for d in sorted_dates if WINDOW_LO <= d <= WINDOW_HI]
n_null = sum(1 for d in in_window if data_by_date[d].get("pos_r") is None)
print(f"  window rows: {len(in_window)}, pos_r null after repair: {n_null}")

# ── 7. GUARD CHECK before write ────────────────────────────────────────────
print("\n" + "="*70)
print("STEP 7: FINAL GUARD before write")
print("="*70)
# Re-verify acceptance from in-memory data_by_date
for date, field, expected, tol in acceptance:
    xlsx_col = {"acc_r": "acc_r", "gLSR": "gLSR"}[field]
    got = safe_float(data_by_date[date].get(xlsx_col))
    ok  = got is not None and abs(got - expected) <= tol
    print(f"  {'✓' if ok else '✗'} in-memory {date} {field}={got} (expected {expected})")
    if not ok:
        sys.exit(f"GUARD FAIL post-repair: {date} {field}")

# Additional: A-0b coverage check (last 30 rows pos_r null <=50%)
last30 = sorted_dates[-30:]
null30 = sum(1 for d in last30 if data_by_date[d].get("pos_r") is None)
null_pct = null30/30*100
print(f"  A-0b coverage last30: pos_r null={null30}/30 ({null_pct:.1f}%)")
if null_pct > 50:
    sys.exit(f"A-0b guard FAIL: {null_pct:.1f}% > 50%")
print("  A-0b: PASS")

# ── 8. WRITE cells ─────────────────────────────────────────────────────────
print("\n" + "="*70)
print("STEP 8: Write cells to xlsx (post-guard)")
print("="*70)

write_count = 0
for d in sorted_dates:
    r = data_by_date[d]
    row_num = r["__row"]
    for field, ci in col_idx.items():
        val = r.get(field)
        # only touch if changed
        cur_val = ws.cell(row_num, ci).value
        if cur_val != val:
            ws.cell(row_num, ci).value = val
            write_count += 1

# For Aug17 new row (if added), make sure date col is set
r17 = data_by_date[DATE_AUG17]
if r17["__row"] > LAST_ROW_XLSX:
    print(f"  new Aug17 row at {r17['__row']}: {sum(1 for v in r17.values() if v is not None and v != '')} non-null cells")

print(f"  cell writes: {write_count}")
wb.save(XLSX)
print(f"  SAVED: {XLSX}")

# Summary of Aug17 row for report
print("\n=== Aug17 row summary ===")
for k in ["date","o","h","l","c","vol_perp_usd","taker_buy_pct (k[10]/k[7])",
          "f_avg","fn (4h=6/day)","oi_usd","oi_token",
          "ret_1d (close-to-close)","pos_r","acc_r","gLSR","taker_lsr",
          "div_raw (pos_r/acc_r)","div_avg3 (avg3p/avg3a)","glsr_slope_5d (ratio t/t-5)",
          "d_pos_3d (%chg)","d_acc_3d (%chg)",
          "flag_hedge_struct (div_avg3<=0.5)","flag_knife_crowd (gLSR↑&ret<-3%)",
          "data_source_perp","data_source_lsr"]:
    print(f"  {k}: {r17.get(k)}")

print("\n[DONE B1+B2 verify+write]")
