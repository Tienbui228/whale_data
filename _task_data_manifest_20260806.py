#!/usr/bin/env python3
"""D1: Generate scratch/data_manifest.json — inventory of all CASE_*.xlsx DAILY state.
   For each case: dates present, last_date, stale rows (data_bad=1), fn_low rows.
"""
import sys, json
from pathlib import Path
from datetime import datetime, timezone

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

import openpyxl

WHALE    = Path(__file__).parent
CASE_DIR = WHALE / "docs" / "case_files"
OUT      = WHALE / "scratch" / "data_manifest.json"

manifest = {
    "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    "source": "docs/case_files/CASE_*.xlsx DAILY sheets",
    "cases": {}
}

for xlsx_f in sorted(CASE_DIR.glob("CASE_*.xlsx")):
    if xlsx_f.name == "CASE_INDEX.xlsx":
        continue
    token = xlsx_f.stem.replace("CASE_", "")

    wb = openpyxl.load_workbook(str(xlsx_f), read_only=True, data_only=True)
    if "DAILY" not in wb.sheetnames:
        wb.close()
        manifest["cases"][token] = {"error": "no_DAILY_sheet"}
        continue

    ws = wb["DAILY"]
    dates = []
    stale_dates = []
    fn_low_dates = []

    for row in ws.iter_rows(min_row=3, values_only=True):
        date_v = row[0] if len(row) > 0 else None
        if date_v is None:
            continue
        date_s = date_v.strftime("%Y-%m-%d") if isinstance(date_v, datetime) else str(date_v)[:10]
        if not date_s or date_s == "None" or len(date_s) != 10:
            continue

        dates.append(date_s)

        col26 = row[25] if len(row) > 25 else None
        if col26 == 1 or col26 is True:
            stale_dates.append(date_s)

        fn_val = row[8] if len(row) > 8 else None  # col 9 = index 8
        if fn_val is None or (isinstance(fn_val, (int, float)) and fn_val < 6):
            fn_low_dates.append(date_s)

    wb.close()

    dates_sorted = sorted(set(dates))
    manifest["cases"][token] = {
        "file": xlsx_f.name,
        "n_dates": len(dates_sorted),
        "first_date": dates_sorted[0] if dates_sorted else None,
        "last_date": dates_sorted[-1] if dates_sorted else None,
        "stale_count": len(stale_dates),
        "stale_dates": sorted(stale_dates),
        "fn_low_count": len(fn_low_dates),
        "fn_low_dates": sorted(fn_low_dates)[:5],  # first 5 only to keep compact
    }
    print(f"  {token}: {len(dates_sorted)} dates, {len(stale_dates)} stale, "
          f"{len(fn_low_dates)} fn_low")

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"\nWrote: {OUT}")
print(f"Cases: {len(manifest['cases'])}")
