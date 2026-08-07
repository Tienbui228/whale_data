"""
Patch CASE_HEI.xlsx — patch ①:
  (a) div_raw = pos_r / acc_r  (không phải pos_r × gLSR)
  (b) div_avg3 = avg3(pos_r) / avg3(acc_r)  (không phải mean(div_raw,3d))
  (c) glsr_slope_5d = gLSR(t)/gLSR(t-5) ratio  (không phải %chg)
  (d) Re-chấm GOM-A A2: slope < 1.0  (ratio format)
  (e) EVENTS: ghi 3 án bug + append note
  (f) Append DAILY row 2026-08-07

Backup: scratch/CASE_HEI_BACKUP_20260808.xlsx trước khi sửa.
L-34: PermissionError → raise.
L-32: verify row count sau save.
"""
import sys, io, json, csv, shutil
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import defaultdict
import urllib.request

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CASE_DIR  = Path(r"D:\Data back test\whale_data\docs\case_files")
SCRATCH   = Path(r"D:\Data back test\whale_data\scratch")
XLSX_PATH = CASE_DIR / "CASE_HEI.xlsx"
BACKUP    = SCRATCH  / "CASE_HEI_BACKUP_20260808.xlsx"

PIN_LO = 4e-5; PIN_HI = 6e-5

def log(s=""): print(s)
def safe_float(v, d=None):
    try: return float(v) if v not in (None,"","None","ABSTAIN(no_alpha_id)","BANNED(k9==k5)","ABSTAIN(Nansen_LOCKED)","MISSING(>30d)","fapi_LSR_30d","fapi_klines_1d") else d
    except: return d
def fetch_json(url, timeout=20):
    try:
        req = urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode())
    except Exception as e:
        return {"_error": str(e)}

log("="*70)
log(f"FIX CASE_HEI div/slope  |  {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
log("="*70)

# ── Backup ────────────────────────────────────────────────────────────────────
log(f"\n  Backup: {XLSX_PATH.name} → {BACKUP.name}")
shutil.copy2(XLSX_PATH, BACKUP)
log(f"  Backup done: {BACKUP}")

# ── Load workbook ─────────────────────────────────────────────────────────────
try:
    wb = openpyxl.load_workbook(XLSX_PATH)
except PermissionError as e:
    log(f"  PERMISSION ERROR: {e} — L-34: đóng Excel, chạy lại")
    raise

ws = wb["DAILY"]
ws_ev = wb["EVENTS"]

# ── Read header map ───────────────────────────────────────────────────────────
# Row 2 = column headers
hdr_map = {}
for ci in range(1, ws.max_column + 1):
    v = ws.cell(2, ci).value
    if v:
        hdr_map[v] = ci

log(f"  DAILY columns: {ws.max_column}")
log(f"  Data rows before fix: {ws.max_row - 2}")

# Use the column names from DAILY_HEADERS (from build script, col index matches)
COL = {
    "date":           1,
    "px_avg":         2,
    "o":              3, "h": 4, "l": 5, "c": 6,
    "vol_perp_usd":   7,
    "taker_buy_pct":  8,
    "f_avg":          9, "f_min": 10, "f_max": 11, "fn": 12,
    "oi_usd":         13, "oi_token": 14,
    "data_ok":        15, "low_conf": 16,
    "ret_1d":         17,
    "pin_streak":     18, "pin_day": 19,
    "oi_vs_med14":    20, "vol_vs_med14": 21,
    "div_raw":        22,
    "div_avg3":       23,
    "glsr_slope_5d":  24,
    "d_pos_3d":       25, "d_acc_3d": 26,
    "pos_r":          27,
    "acc_r":          28,
    "gLSR":           29,
    "taker_lsr":      30,
    "vol_spot_usd":   31, "taker_spot_pct": 32, "spot_perp_ratio": 33, "l29_flag": 34,
    "vol_alpha":      35, "taker_alpha": 36,
    "on_chain_source":37,
    "gom_A_screen":   38, "gom_B_screen": 39, "gom_note": 40,
    "data_source_perp": 41, "data_source_lsr": 42,
}

DATA_ROW_START = 3
n_data = ws.max_row - 2  # 540

# ── Extract full DAILY arrays ─────────────────────────────────────────────────
dates    = []
pos_rs   = []
acc_rs   = []
glsrs    = []
vol_vals = []
taker_vals = []
f_avgs   = []
ret_vals = []
pin_days_raw = []

for ri in range(DATA_ROW_START, ws.max_row + 1):
    dates.append(ws.cell(ri, COL["date"]).value)
    pos_rs.append(safe_float(ws.cell(ri, COL["pos_r"]).value))
    acc_rs.append(safe_float(ws.cell(ri, COL["acc_r"]).value))
    glsrs.append(safe_float(ws.cell(ri, COL["gLSR"]).value))
    vol_vals.append(safe_float(ws.cell(ri, COL["vol_perp_usd"]).value))
    taker_vals.append(safe_float(ws.cell(ri, COL["taker_buy_pct"]).value))
    f_avgs.append(safe_float(ws.cell(ri, COL["f_avg"]).value))
    ret_vals.append(safe_float(ws.cell(ri, COL["ret_1d"]).value))
    pin_days_raw.append(ws.cell(ri, COL["pin_day"]).value)

log(f"  Loaded {len(dates)} rows: {dates[0]} → {dates[-1]}")
log(f"  Rows with pos_r: {sum(1 for v in pos_rs if v is not None)}")
log(f"  Rows with gLSR: {sum(1 for v in glsrs if v is not None)}")

# ── vol_med14 (for GOM-A A1) ─────────────────────────────────────────────────
vol_med14 = [None] * len(vol_vals)
for i in range(14, len(vol_vals)):
    w = [v for v in vol_vals[i-14:i] if v is not None]
    if len(w) >= 7:
        ws2 = sorted(w); n = len(ws2)
        vol_med14[i] = (ws2[n//2-1]+ws2[n//2])/2 if n%2==0 else ws2[n//2]

# pos_r_persist_cnt (for GOM-B B1)
pos_r_persist = [0]*len(pos_rs)
cnt_ = 0
for i in range(len(pos_rs)):
    if pos_rs[i] and pos_rs[i] > 1.5: cnt_ += 1
    else: cnt_ = 0
    pos_r_persist[i] = cnt_

# oi_med14 (for GOM-B B4)
oi_vals = [safe_float(ws.cell(r, COL["oi_usd"]).value) for r in range(DATA_ROW_START, ws.max_row+1)]
oi_med14 = [None]*len(oi_vals)
for i in range(14, len(oi_vals)):
    w = [v for v in oi_vals[i-14:i] if v is not None]
    if len(w) >= 7:
        ws2 = sorted(w); n = len(ws2)
        oi_med14[i] = (ws2[n//2-1]+ws2[n//2])/2 if n%2==0 else ws2[n//2]

# ── Compute corrected derived vars ────────────────────────────────────────────

# (a) div_raw = pos_r / acc_r
new_div_raw = []
for i in range(len(pos_rs)):
    p = pos_rs[i]; a = acc_rs[i]
    new_div_raw.append(round(p/a, 4) if (p and a and a != 0) else None)

# (b) div_avg3 = avg3(pos_r) / avg3(acc_r)
new_div_avg3 = [None]*len(pos_rs)
for i in range(2, len(pos_rs)):
    p3 = [v for v in pos_rs[i-2:i+1] if v is not None]
    a3 = [v for v in acc_rs[i-2:i+1] if v is not None]
    if len(p3) == 3 and len(a3) == 3:
        ap = sum(p3)/3; aa = sum(a3)/3
        new_div_avg3[i] = round(ap/aa, 4) if aa != 0 else None

# (c) glsr_slope_5d = gLSR(t) / gLSR(t-5) — ratio
new_glsr_slope = [None]*len(glsrs)
for i in range(5, len(glsrs)):
    g = glsrs[i]; gp = glsrs[i-5]
    if g is not None and gp is not None and gp > 0:
        new_glsr_slope[i] = round(g/gp, 4)

log(f"\n  div_raw non-null: {sum(1 for v in new_div_raw if v is not None)}")
log(f"  div_avg3 non-null: {sum(1 for v in new_div_avg3 if v is not None)}")
log(f"  glsr_slope non-null: {sum(1 for v in new_glsr_slope if v is not None)}")

# ── Re-chấm GOM-A (A2 now slope < 1.0) ───────────────────────────────────────
gom_a_new = [False]*len(dates)
gom_b_new = [False]*len(dates)
gom_note_new = [""]*len(dates)

for i in range(len(dates)):
    vol_p = vol_vals[i]; vm = vol_med14[i]
    vol_3x = (vol_p/vm >= 3) if (vol_p and vm) else False
    A1 = vol_3x
    A2 = (new_glsr_slope[i] is not None and new_glsr_slope[i] < 1.0)  # FIXED: ratio < 1.0
    A3 = (pos_rs[i] is not None and pos_rs[i] > 1.5)
    A4 = (f_avgs[i] is not None and f_avgs[i] >= PIN_LO)
    A5 = (taker_vals[i] is not None and taker_vals[i] >= 45)
    gom_a_new[i] = A1 and A2 and A3 and A4 and A5

    oi_usd = safe_float(ws.cell(i+DATA_ROW_START, COL["oi_usd"]).value)
    oi_m   = oi_med14[i]
    oi_vs_m = round(oi_usd/oi_m, 3) if (oi_usd and oi_m) else None
    pin_day = pin_days_raw[i]
    ret_d   = ret_vals[i]
    B1 = (pos_r_persist[i] >= 7)
    B2 = (pin_day is True)
    B3 = (vol_3x and ret_d is not None and ret_d < 0 and taker_vals[i] is not None and taker_vals[i] >= 50)
    B4 = (oi_vs_m is not None and oi_vs_m >= 1.5)
    gom_b_new[i] = B1 and (B2 or B3) and B4

    if gom_a_new[i] or gom_b_new[i]:
        parts = []
        if gom_a_new[i]: parts.append(f"A(A1={A1},A2-ratio={A2},A3={A3},A4={A4},A5={A5})")
        if gom_b_new[i]: parts.append(f"B(B1={B1},B2={B2},B3={B3},B4={B4})")
        gom_note_new[i] = " | ".join(parts) + " — gom_detect_spec v1.1 CHUA KY"

log(f"  GOM-A fires (new): {sum(gom_a_new)} dates: {[dates[i] for i,v in enumerate(gom_a_new) if v]}")
log(f"  GOM-B fires (new): {sum(gom_b_new)}")

# ── Write corrected cells ─────────────────────────────────────────────────────
log("\n  Writing corrected cells to DAILY...")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=9)

for i in range(len(dates)):
    ri = i + DATA_ROW_START

    # (a) div_raw
    c = ws.cell(ri, COL["div_raw"])
    c.value = new_div_raw[i]
    c.font  = Font(size=8)

    # (b) div_avg3
    c = ws.cell(ri, COL["div_avg3"])
    c.value = new_div_avg3[i]
    c.font  = Font(size=8)

    # (c) glsr_slope_5d
    c = ws.cell(ri, COL["glsr_slope_5d"])
    c.value = new_glsr_slope[i]
    c.font  = Font(size=8)

    # (d) GOM re-screen
    c_a = ws.cell(ri, COL["gom_A_screen"])
    c_b = ws.cell(ri, COL["gom_B_screen"])
    c_n = ws.cell(ri, COL["gom_note"])
    c_a.value = "TRUE" if gom_a_new[i] else "false"
    c_b.value = "TRUE" if gom_b_new[i] else "false"
    c_n.value = gom_note_new[i]
    c_a.font  = Font(size=8)
    c_b.font  = Font(size=8)
    c_n.font  = Font(size=8)
    # clear highlight first, then re-apply if needed
    c_a.fill  = PatternFill(fill_type=None)
    c_b.fill  = PatternFill(fill_type=None)
    if gom_a_new[i]:
        c_a.fill = PatternFill("solid", fgColor="FFA500")
    if gom_b_new[i]:
        c_b.fill = PatternFill("solid", fgColor="FF6B6B")

# Fix column header labels (row 2) — patch wrong descriptions
ws.cell(2, COL["div_raw"]).value    = "div_raw (pos_r/acc_r)"
ws.cell(2, COL["div_avg3"]).value   = "div_avg3 (avg3(pos_r)/avg3(acc_r))"
ws.cell(2, COL["glsr_slope_5d"]).value = "glsr_slope_5d (gLSR[t]/gLSR[t-5] ratio)"
for ci in [COL["div_raw"], COL["div_avg3"], COL["glsr_slope_5d"]]:
    ws.cell(2, ci).fill = HDR_FILL
    ws.cell(2, ci).font = HDR_FONT

log(f"  Cell update done.")

# ── Fetch 2026-08-07 data ─────────────────────────────────────────────────────
log("\n  Fetching 2026-08-07 data...")
APD_DATE = "2026-08-07"
ts_start = int(datetime(2026, 8, 7, 0, 0, 0, tzinfo=timezone.utc).timestamp() * 1000)
ts_end   = int(datetime(2026, 8, 8, 0, 0, 0, tzinfo=timezone.utc).timestamp() * 1000)

# Perp kline 2026-08-07
perp_url = f"https://fapi.binance.com/fapi/v1/klines?symbol=HEIUSDT&interval=1d&startTime={ts_start}&endTime={ts_end}&limit=2"
perp_k = fetch_json(perp_url)
perp_row = {}
if isinstance(perp_k, list) and perp_k:
    k = perp_k[0]
    perp_row = {
        "o": float(k[1]), "h": float(k[2]), "l": float(k[3]), "c": float(k[4]),
        "vol_quote_usd": float(k[7]),
        "taker_pct": round(float(k[10])/float(k[7])*100, 2) if float(k[7]) else None,
    }
    log(f"  Perp 07/08: o={perp_row['o']} h={perp_row['h']} l={perp_row['l']} c={perp_row['c']} vol=${perp_row['vol_quote_usd']:.0f}")
else:
    log(f"  Perp 07/08: ERROR {perp_k}")

# Spot kline 2026-08-07
spot_url = f"https://api.binance.com/api/v3/klines?symbol=HEIUSDT&interval=1d&startTime={ts_start}&endTime={ts_end}&limit=2"
spot_k = fetch_json(spot_url)
spot_row = {}
if isinstance(spot_k, list) and spot_k:
    k = spot_k[0]
    spot_vol = float(k[7])
    spot_row = {
        "vol_spot_usd": spot_vol,
        "taker_spot_pct": round(float(k[10])/spot_vol*100, 2) if spot_vol else None,
    }
    perp_vol = perp_row.get("vol_quote_usd")
    spot_row["spot_perp_ratio"] = round(spot_vol/perp_vol, 3) if (perp_vol and perp_vol > 0) else None
    spot_row["l29_flag"] = "CLEAR" if (spot_row["spot_perp_ratio"] and spot_row["spot_perp_ratio"] < 3) else "L29_ABSTAIN" if spot_row.get("spot_perp_ratio") else ""
    log(f"  Spot 07/08: vol_spot=${spot_vol:.0f} taker={spot_row['taker_spot_pct']}% ratio={spot_row['spot_perp_ratio']}")
else:
    log(f"  Spot 07/08: ERROR {spot_k}")

# Funding 2026-08-07
fund_url = "https://fapi.binance.com/fapi/v1/fundingRate?symbol=HEIUSDT&limit=1000"
fund_raw = fetch_json(fund_url)
fund_0807 = []
if isinstance(fund_raw, list):
    for fr in fund_raw:
        ts = fr.get("fundingTime")
        if ts:
            dt = datetime.fromtimestamp(int(ts)/1000, tz=timezone.utc).strftime("%Y-%m-%d")
            if dt == APD_DATE:
                fund_0807.append(float(fr.get("fundingRate", 0)))
log(f"  Funding 07/08: fn={len(fund_0807)}, rates={fund_0807}")

# OI hist 2026-08-07
oih_url = "https://fapi.binance.com/futures/data/openInterestHist?symbol=HEIUSDT&period=1d&limit=30"
oih_raw = fetch_json(oih_url)
oi_0807 = {}
if isinstance(oih_raw, list):
    for o in oih_raw:
        ts = o.get("timestamp")
        if ts:
            dt = datetime.fromtimestamp(int(ts)/1000, tz=timezone.utc).strftime("%Y-%m-%d")
            if dt == APD_DATE:
                oi_0807 = {"oi_usd": float(o.get("sumOpenInterestValue",0)), "oi_token": float(o.get("sumOpenInterest",0))}
log(f"  OI 07/08: {oi_0807}")

# LSR 2026-08-07 — from CSV
lsr_csv = SCRATCH / "hei_lsr_raw_20260808.csv"
lsr_0807 = {}
if lsr_csv.exists():
    with open(lsr_csv, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("date") == APD_DATE:
                lsr_0807 = {k: safe_float(v) for k, v in row.items() if k != "date"}
log(f"  LSR 07/08 from CSV: pos_r={lsr_0807.get('pos_r_raw')} gLSR={lsr_0807.get('glsr_raw')}")

# ── Build new DAILY row for 2026-08-07 ───────────────────────────────────────
o  = perp_row.get("o"); h = perp_row.get("h"); l = perp_row.get("l"); c = perp_row.get("c")
px_avg = round((h+l+c)/3, 8) if (h and l and c) else None
vol_p  = perp_row.get("vol_quote_usd")
taker  = perp_row.get("taker_pct")
ret_1d = round((c - safe_float(ws.cell(ws.max_row, COL["c"]).value)) / safe_float(ws.cell(ws.max_row, COL["c"]).value) * 100, 2) if (c and safe_float(ws.cell(ws.max_row, COL["c"]).value)) else None

fn    = len(fund_0807)
f_avg = round(sum(fund_0807)/fn, 8) if fn > 0 else None
f_min = round(min(fund_0807), 8) if fn > 0 else None
f_max = round(max(fund_0807), 8) if fn > 0 else None
data_ok  = fn >= 6 if fn > 0 else None
low_conf = fn < 6  if fn > 0 else True

oi_usd   = oi_0807.get("oi_usd")
oi_token = oi_0807.get("oi_token")

pin_day = all(PIN_LO <= f <= PIN_HI for f in fund_0807) and fn >= 6 if fn > 0 else None

# pin_streak: read last from sheet
last_streak = safe_float(ws.cell(ws.max_row, COL["pin_streak"]).value, 0)
pin_streak_new = (int(last_streak) + 1) if pin_day is True else 0

# vol_med14, oi_med14 for this row (extend arrays)
vol_vals_ext = vol_vals + [vol_p]
vm_ext = None
if len(vol_vals_ext) >= 14:
    w = [v for v in vol_vals_ext[-14:-1] if v is not None]
    if len(w) >= 7:
        ws2 = sorted(w); n = len(ws2)
        vm_ext = (ws2[n//2-1]+ws2[n//2])/2 if n%2==0 else ws2[n//2]
vol_vs_m = round(vol_p/vm_ext, 2) if (vol_p and vm_ext) else None
oi_vals_ext = oi_vals + [oi_usd]
om_ext = None
if len(oi_vals_ext) >= 14:
    w = [v for v in oi_vals_ext[-14:-1] if v is not None]
    if len(w) >= 7:
        ws2 = sorted(w); n = len(ws2)
        om_ext = (ws2[n//2-1]+ws2[n//2])/2 if n%2==0 else ws2[n//2]
oi_vs_m = round(oi_usd/om_ext, 3) if (oi_usd and om_ext) else None

pos_r = lsr_0807.get("pos_r_raw")
acc_r = lsr_0807.get("acc_r_raw")
glsr  = lsr_0807.get("glsr_raw")
tkls  = lsr_0807.get("taker_lsr_raw")

# Corrected formulas for new row
div_raw_new = round(pos_r/acc_r, 4) if (pos_r and acc_r and acc_r != 0) else None
# avg3 for pos_r / acc_r: use last 2 from sheet + new
last_pos = [safe_float(ws.cell(ws.max_row-1, COL["pos_r"]).value), safe_float(ws.cell(ws.max_row, COL["pos_r"]).value)]
last_acc = [safe_float(ws.cell(ws.max_row-1, COL["acc_r"]).value), safe_float(ws.cell(ws.max_row, COL["acc_r"]).value)]
p3 = [v for v in last_pos + [pos_r] if v is not None]
a3 = [v for v in last_acc + [acc_r] if v is not None]
div_avg3_new = round(sum(p3)/len(p3) / (sum(a3)/len(a3)), 4) if (len(p3)==3 and len(a3)==3 and sum(a3)/len(a3) != 0) else None

# glsr_slope: gLSR(t)/gLSR(t-5) — t-5 from 5 rows back in sheet
glsr_5back = safe_float(ws.cell(ws.max_row - 4, COL["gLSR"]).value)
glsr_slope_new = round(glsr/glsr_5back, 4) if (glsr and glsr_5back and glsr_5back > 0) else None

# GOM screen for new row
vol_3x = (vol_p/vm_ext >= 3) if (vol_p and vm_ext) else False
A1 = vol_3x
A2 = (glsr_slope_new is not None and glsr_slope_new < 1.0)
A3 = (pos_r is not None and pos_r > 1.5)
A4 = (f_avg is not None and f_avg >= PIN_LO)
A5 = (taker is not None and taker >= 45)
gom_a_apd = A1 and A2 and A3 and A4 and A5

pos_persist_last = pos_r_persist[-1] if pos_rs[-1] and pos_rs[-1] > 1.5 else 0
if pos_r and pos_r > 1.5: pos_persist_new = pos_persist_last + 1
else: pos_persist_new = 0
B1 = (pos_persist_new >= 7)
B2 = (pin_day is True)
B3 = (vol_3x and ret_1d is not None and ret_1d < 0 and taker is not None and taker >= 50)
B4 = (oi_vs_m is not None and oi_vs_m >= 1.5)
gom_b_apd = B1 and (B2 or B3) and B4

gom_note_apd = ""
if gom_a_apd or gom_b_apd:
    parts = []
    if gom_a_apd: parts.append(f"A(A1={A1},A2={A2},A3={A3},A4={A4},A5={A5})")
    if gom_b_apd: parts.append(f"B(B1={B1},B2={B2},B3={B3},B4={B4})")
    gom_note_apd = " | ".join(parts) + " — gom_detect_spec v1.1 CHUA KY"

vol_s   = spot_row.get("vol_spot_usd")
tk_s    = spot_row.get("taker_spot_pct")
spr     = spot_row.get("spot_perp_ratio")
l29     = spot_row.get("l29_flag","")

new_row_vals = [
    APD_DATE, px_avg, o, h, l, c, vol_p, taker,
    f_avg, f_min, f_max, fn if fn>0 else None, oi_usd, oi_token, data_ok, low_conf,
    ret_1d, pin_streak_new, pin_day, oi_vs_m, vol_vs_m,
    div_raw_new, div_avg3_new, glsr_slope_new, None, None,
    pos_r, acc_r, glsr, tkls,
    vol_s, tk_s, spr, l29,
    "ABSTAIN(no_alpha_id)", "BANNED(k9==k5)", "ABSTAIN(Nansen_LOCKED)",
    "TRUE" if gom_a_apd else "false", "TRUE" if gom_b_apd else "false", gom_note_apd,
    "fapi_klines_1d", "fapi_LSR_30d" if lsr_0807 else "MISSING(>30d)",
]

# Append row
new_ri = ws.max_row + 1
log(f"\n  Appending 2026-08-07 at row {new_ri}")
log(f"  o={o} h={h} l={l} c={c} ret_1d={ret_1d} fn={fn} pos_r={pos_r} gLSR={glsr}")
log(f"  div_raw={div_raw_new} div_avg3={div_avg3_new} glsr_slope={glsr_slope_new}")
log(f"  GOM-A={gom_a_apd} GOM-B={gom_b_apd}")

for ci, val in enumerate(new_row_vals, 1):
    c = ws.cell(new_ri, ci)
    c.value = val
    c.font  = Font(size=8)
    c.alignment = Alignment(vertical="center")
if gom_a_apd:
    ws.cell(new_ri, COL["gom_A_screen"]).fill = PatternFill("solid", fgColor="FFA500")
if gom_b_apd:
    ws.cell(new_ri, COL["gom_B_screen"]).fill = PatternFill("solid", fgColor="FF6B6B")
if not data_ok:
    ws.cell(new_ri, COL["data_ok"]).fill = PatternFill("solid", fgColor="FFFF00")

# ── EVENTS: append bug-fix + append entries ───────────────────────────────────
log("\n  Appending EVENTS entries...")

# Find next EVENTS row
ev_next = ws_ev.max_row + 1
ev_now = datetime.now(timezone.utc).strftime("%Y-%m-%d")

new_events = [
    (ev_now, "BUG_FIX_A", "Patch ①(a): div_raw đổi công thức pos_r×gLSR → pos_r/acc_r (Phụ lục A.1).", "L-31: import từ canonical, không viết lại. Script build tự tái định nghĩa thuật ngữ.", "agent patch 08/08"),
    (ev_now, "BUG_FIX_B", "Patch ①(b): div_avg3 đổi mean(div_raw,3d) → avg3(pos_r)/avg3(acc_r) (Phụ lục A.2).", "Bẫy trung-bình-của-tỷ-số: trung bình tỷ số ≠ tỷ số trung bình.", "agent patch 08/08"),
    (ev_now, "BUG_FIX_C", "Patch ①(c): glsr_slope_5d đổi %chg×100 → ratio gLSR[t]/gLSR[t-5] (Phụ lục A.3). GOM-A A2 re-screen.", "Ngưỡng gắn với format; ratio <1.0 = crowd xả fuel (thay vì %chg<0).", "agent patch 08/08"),
    (APD_DATE, "DAILY_APPEND", f"Append DAILY 2026-08-07: c={c}, ret_1d={ret_1d}%, fn={fn}, pos_r={pos_r}, gLSR={glsr}.", "L-21: 08-08 chưa vào. Thêm sau khi nến 07-08 đóng.", "fapi klines+LSR csv"),
]

for ev in new_events:
    for ci, v in enumerate(ev, 1):
        ws_ev.cell(ev_next, ci).value = v
        ws_ev.cell(ev_next, ci).font  = Font(size=8)
    ev_next += 1

log(f"  EVENTS rows after: {ws_ev.max_row}")

# ── Save ─────────────────────────────────────────────────────────────────────
log(f"\n  Saving {XLSX_PATH.name}...")
try:
    wb.save(XLSX_PATH)
    log(f"  Saved.")
except PermissionError as e:
    log(f"  PERMISSION ERROR: {e} — L-34: đóng Excel"); raise

# ── L-32: Verify ──────────────────────────────────────────────────────────────
wb2 = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws2 = wb2["DAILY"]
actual = ws2.max_row - 2
expected = n_data + 1  # original 540 + appended 1
log(f"\n  L-32: DAILY expected {expected}, actual={actual}")
assert actual == expected, f"ROW COUNT MISMATCH: expected {expected}, got {actual}"
log(f"  L-32: PASS")

# ── Print sample 5 rows (last 5, including new) ───────────────────────────────
log(f"\n  === SAMPLE: last 5 DAILY rows (div_raw, div_avg3, glsr_slope) ===")
log(f"  {'date':12} {'pos_r':8} {'acc_r':8} {'gLSR':8} {'div_raw':9} {'div_avg3':9} {'slope':8} {'gom_A':7}")
for ri in range(ws2.max_row - 4, ws2.max_row + 1):
    d   = ws2.cell(ri, COL["date"]).value
    pr  = ws2.cell(ri, COL["pos_r"]).value
    ar  = ws2.cell(ri, COL["acc_r"]).value
    gl  = ws2.cell(ri, COL["gLSR"]).value
    dr  = ws2.cell(ri, COL["div_raw"]).value
    da  = ws2.cell(ri, COL["div_avg3"]).value
    sl  = ws2.cell(ri, COL["glsr_slope_5d"]).value
    ga  = ws2.cell(ri, COL["gom_A_screen"]).value
    log(f"  {str(d):12} {str(pr)[:8]:8} {str(ar)[:8]:8} {str(gl)[:8]:8} {str(dr)[:9]:9} {str(da)[:9]:9} {str(sl)[:8]:8} {str(ga)[:7]:7}")

log(f"\n[FIX CASE_HEI DONE] — {XLSX_PATH}")
