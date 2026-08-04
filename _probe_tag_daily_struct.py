#!/usr/bin/env python3
"""Probe CASE_TAG.xlsx DAILY structure around row 32."""
import io, sys
from pathlib import Path
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_TAG.xlsx")
wb = openpyxl.load_workbook(XLSX)
ws = wb["DAILY"]

print(f"max_row={ws.max_row}, max_col={ws.max_column}")
print(f"Merged ranges: {ws.merged_cells.ranges}")
print()

# Show rows 30-35
for r in range(30, min(36, ws.max_row+1)):
    row_data = []
    for c in range(1, 15):
        cell = ws.cell(r, c)
        row_data.append(f"C{c}={repr(cell.value)[:20]}")
    print(f"Row {r}: {', '.join(row_data[:8])}")

print()
# Show column headers (row 1-2)
for r in range(1, 3):
    for c in range(1, 25):
        v = ws.cell(r, c).value
        if v:
            print(f"  R{r}C{c}: {v}")
