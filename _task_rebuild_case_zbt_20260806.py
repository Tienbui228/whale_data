"""Rebuild CASE_ZBT.xlsx — 30 days full (07/07→05/08).
Sources:
 - Export Jul+Aug CSV: px_avg, vol_perp, taker, f_avg/min/max/fn, oi, LSR (canonical where available)
 - fapi klines: OHLCV + taker fallback for gaps (24-26/07)
 - fapi LSR endpoints: pos_r, acc_r, gLSR (30d coverage)
 - fapi OI hist: oi_usd (30d)
 - api/v3 spot: vol_spot
"""
import sys, io, json, urllib.request, csv, statistics
from pathlib import Path
from datetime import datetime, timezone
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

SYM = "ZBTUSDT"
case_dir = Path(r"D:\Data back test\whale_data\docs\case_files")
export_dir = Path(r"D:\Check onchain fresh wallets\data\export")

def get(url):
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as r:
            return json.loads(r.read())
    except Exception as e:
        return {"ERROR": str(e)}

def ts_date(ms):
    return datetime.fromtimestamp(ms/1000, tz=timezone.utc).strftime("%Y-%m-%d")

BASE_F = "https://fapi.binance.com"
BASE_S = "https://api.binance.com"

# ── 1. Export CSV (Jul + Aug) — canonical source ──────────────────────────────
export = {}
for fname in ["freshwatch_daily_2026-07.csv", "freshwatch_daily_2026-08.csv"]:
    fp = export_dir / fname
    if not fp.exists():
        print(f"  WARN: {fname} not found")
        continue
    with open(fp, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["symbol"] != SYM:
                continue
            d = row["date"]
            taker = float(row["taker_buy_pct"])
            if taker <= 1:
                taker *= 100  # normalize if fraction
            export[d] = {
                "px_avg":   float(row["px_avg"]),
                "vol_perp": float(row["vol_quote"]),
                "taker":    taker,
                "f_avg":    float(row["funding_avg"]) if row.get("funding_avg") else None,
                "f_min":    float(row["funding_min"]) if row.get("funding_min") else None,
                "f_max":    float(row["funding_max"]) if row.get("funding_max") else None,
                "fn":       int(row["funding_n"]) if row.get("funding_n") else None,
                "oi":       float(row["oi_usd"]) if row.get("oi_usd") else None,
                "pos_r":    float(row["lsr_pos_r"]) if row.get("lsr_pos_r") else None,
                "acc_r":    float(row["lsr_acc_r"]) if row.get("lsr_acc_r") else None,
                "gLSR":     float(row["lsr_global_r"]) if row.get("lsr_global_r") else None,
                "data_ok":  row.get("data_ok", "True") == "True",
            }
print(f"[Export] {len(export)} dates: {min(export)} → {max(export)}")

# ── 2. fapi perp klines 1d (fallback OHLCV + taker) ─────────────────────────
kl = get(f"{BASE_F}/fapi/v1/klines?symbol={SYM}&interval=1d&limit=32")
klines = {}
if isinstance(kl, list):
    for k in kl:
        d = ts_date(k[0])
        klines[d] = {
            "open": float(k[1]), "high": float(k[2]),
            "low":  float(k[3]), "close": float(k[4]),
            "vol_perp": float(k[7]),
            "taker": round(float(k[10])/float(k[7])*100, 4) if float(k[7]) > 0 else None,
        }
print(f"[Klines] {len(klines)} days: {min(klines)} → {max(klines)}")

# ── 3. fapi LSR endpoints (30d) ───────────────────────────────────────────────
def fetch_lsr(endpoint, label):
    d = get(f"{BASE_F}/futures/data/{endpoint}?symbol={SYM}&period=1d&limit=32")
    if isinstance(d, list):
        out = {}
        for rec in d:
            dt = ts_date(rec["timestamp"])
            out[dt] = float(rec["longShortRatio"])
        print(f"[{label}] {len(out)} days: {min(out)} → {max(out)}")
        return out
    print(f"[{label}] ERROR: {d}")
    return {}

pos_r  = fetch_lsr("topLongShortPositionRatio", "pos_r")
acc_r  = fetch_lsr("topLongShortAccountRatio",  "acc_r")
glsr_d = fetch_lsr("globalLongShortAccountRatio","gLSR")

# ── 4. fapi OI hist ───────────────────────────────────────────────────────────
oi_d = get(f"{BASE_F}/futures/data/openInterestHist?symbol={SYM}&period=1d&limit=32")
oi_hist = {}
if isinstance(oi_d, list):
    for rec in oi_d:
        d = ts_date(rec["timestamp"])
        oi_hist[d] = float(rec["sumOpenInterestValue"])
print(f"[OI hist] {len(oi_hist)} days: {min(oi_hist)} → {max(oi_hist)}")

# ── 5. Spot klines ─────────────────────────────────────────────────────────────
sk = get(f"{BASE_S}/api/v3/klines?symbol={SYM}&interval=1d&limit=32")
spot = {}
if isinstance(sk, list):
    for k in sk:
        d = ts_date(k[0])
        spot[d] = float(k[7])
print(f"[Spot] {len(spot)} days: {min(spot)} → {max(spot)}")

# ── 6. BUILD ROW LIST ──────────────────────────────────────────────────────────
today = datetime.now(tz=timezone.utc).strftime("%Y-%m-%d")
all_dates = sorted(set(
    [d for d in klines if d < today] +
    [d for d in export if d < today] +
    [d for d in oi_hist if d < today]
))

ROWS = []
for d in all_dates:
    ex = export.get(d, {})
    kl = klines.get(d, {})

    # px_avg: export canonical; fallback perp close
    px_avg = ex.get("px_avg") or kl.get("close")

    # vol_perp: export canonical
    vol_perp = ex.get("vol_perp") or kl.get("vol_perp")

    # taker: export canonical (reliable); fallback klines
    taker = ex.get("taker") or kl.get("taker")

    # funding: export canonical
    f_avg = ex.get("f_avg")
    f_min = ex.get("f_min")
    f_max = ex.get("f_max")
    fn    = ex.get("fn")

    # OI: export canonical; fallback OI hist
    oi_v = ex.get("oi") or oi_hist.get(d)

    # LSR: export canonical; fallback fapi endpoints
    pr = ex.get("pos_r") or pos_r.get(d)
    ar = ex.get("acc_r") or acc_r.get(d)
    gl = ex.get("gLSR")  or glsr_d.get(d)

    # spot vol
    vs = spot.get(d)

    # data_bad: 1 if fn missing or fn<6 or no LSR
    data_ok = ex.get("data_ok", True)
    db = 0
    if not data_ok:
        db = 1
    elif fn is None or fn < 6:
        db = 1
    elif pr is None:
        db = 1

    # OHLC from klines (for CASE context — not in DAILY cols but useful for context)
    o, h, lo, c = (kl.get("open"), kl.get("high"), kl.get("low"), kl.get("close"))

    ROWS.append({
        "date": d, "px_avg": px_avg, "vol_perp": vol_perp,
        "taker": taker, "f_avg": f_avg, "f_min": f_min, "f_max": f_max, "fn": fn,
        "oi": oi_v, "pos_r": pr, "acc_r": ar, "gLSR": gl,
        "vol_spot": vs, "data_bad": db,
        "_open": o, "_high": h, "_low": lo, "_close": c,
    })

print(f"\n[Compiled] {len(ROWS)} days: {ROWS[0]['date']} → {ROWS[-1]['date']}")
bad = [r["date"] for r in ROWS if r["data_bad"]]
print(f"  data_bad=1: {len(bad)} days: {bad}")

# ── 7. DERIVED METRICS ─────────────────────────────────────────────────────────
# OI series for med14 / cycle_pos (use full oi_hist for wider window)
oi_series = {}
for d in oi_hist:
    oi_series[d] = oi_hist[d]
for r in ROWS:
    if r["oi"]:
        oi_series[r["date"]] = r["oi"]
oi_sorted = sorted(oi_series.keys())

def med14(d):
    idx = oi_sorted.index(d) if d in oi_sorted else -1
    if idx < 0:
        return None
    window = [oi_series[oi_sorted[i]] for i in range(max(0, idx-13), idx+1)]
    return statistics.median(window) if len(window) == 14 else None

def oi_cycle(d, oi_v):
    idx = oi_sorted.index(d) if d in oi_sorted else -1
    if idx < 0 or oi_v is None:
        return None
    window = [oi_series[oi_sorted[i]] for i in range(max(0, idx-13), idx+1)]
    if len(window) < 14:
        return None
    mn, mx = min(window), max(window)
    return (oi_v - mn) / (mx - mn) if mx != mn else 0.0

computed = []
for i, r in enumerate(ROWS):
    d     = r["date"]
    pr    = r["pos_r"]
    ar    = r["acc_r"]
    gl    = r["gLSR"]
    oi_v  = r["oi"]
    fa    = r["f_avg"]
    px    = r["px_avg"]

    # ret_1d (px_avg sequential, both export)
    ret_1d = None
    if i > 0 and px and ROWS[i-1]["px_avg"]:
        ret_1d = round((px / ROWS[i-1]["px_avg"] - 1) * 100, 4)

    # pin_streak
    pin = None
    if fa is not None and 4e-5 <= fa <= 6e-5:
        streak = 1
        j = i - 1
        while j >= 0 and ROWS[j]["f_avg"] is not None and 4e-5 <= ROWS[j]["f_avg"] <= 6e-5:
            streak += 1
            j -= 1
        pin = streak

    # oi_vs_med14, oi_cycle_pos
    m14 = med14(d)
    oi_med = round(oi_v / m14, 4) if (oi_v and m14) else None
    ocp = oi_cycle(d, oi_v)
    ocp = round(ocp, 4) if ocp is not None else None

    # div_raw
    div_raw = round(pr / ar, 4) if (pr and ar) else None

    # div_avg3
    div_avg3 = None
    if i >= 2:
        prs = [ROWS[k]["pos_r"] for k in range(i-2, i+1)]
        ars = [ROWS[k]["acc_r"] for k in range(i-2, i+1)]
        if all(v is not None for v in prs + ars):
            div_avg3 = round(sum(prs)/3 / (sum(ars)/3), 4)

    # glsr_slope_5d
    slope5 = None
    if i >= 5 and gl is not None and ROWS[i-5]["gLSR"] is not None:
        slope5 = round(gl / ROWS[i-5]["gLSR"], 4)

    # d_pos_3d, d_acc_3d
    d_pos3 = None
    d_acc3 = None
    if i >= 3:
        if pr is not None and ROWS[i-3]["pos_r"] is not None:
            d_pos3 = round(pr - ROWS[i-3]["pos_r"], 4)
        if ar is not None and ROWS[i-3]["acc_r"] is not None:
            d_acc3 = round(ar - ROWS[i-3]["acc_r"], 4)

    computed.append({
        "date": d,
        "px_avg": px, "ret_1d": ret_1d,
        "vol_perp_usd": r["vol_perp"],
        "taker_buy_pct": r["taker"],
        "f_avg": fa, "f_min": r["f_min"], "f_max": r["f_max"], "f_n": r["fn"],
        "funding_regime": None, "pin_streak": pin,
        "oi_usd": oi_v, "oi_vs_med14": oi_med, "oi_cycle_pos": ocp,
        "lsr_pos_r": pr, "lsr_acc_r": ar, "lsr_global_r": gl,
        "div_raw": div_raw, "div_avg3": div_avg3, "glsr_slope_5d": slope5,
        "d_pos_3d": d_pos3, "d_acc_3d": d_acc3,
        "signals_bullish": None, "signals_suppress": None, "screen_rank": None,
        "data_bad": r["data_bad"],
        "spot_venue": "binance_spot", "vol_spot_usd": r["vol_spot"],
        "taker_spot_pct": None,  # ABSTAIN — units invalid
        "depth_bid2_usd": None, "depth_ask2_usd": None,
        "bid_wall_px": None, "bid_wall_usd": None,
        "ask_wall_px": None, "ask_wall_usd": None,
        "alpha_suspect": "DELISTED",
        "exch_balance_delta": None, "wallet_activity": None, "dominance_pct": None,
        "note_operator": None, "mm_note": None,
    })

# ── 8. PREVIEW TABLE ──────────────────────────────────────────────────────────
print("\n=== PREVIEW (all rows) ===")
print(f"{'date':<12} {'px':>8} {'ret%':>6} {'vol_perp':>12} {'taker':>6} {'f_avg':>9} {'fn':>3} {'oi':>11} {'pos_r':>7} {'acc_r':>7} {'gLSR':>7} {'div_avg3':>8} {'slope5':>7} {'pin':>4} db")
print("-"*145)
for rec in computed:
    def s(v, w, d=4): return f"{v:{w}.{d}f}" if v is not None else "—".rjust(w)
    def si(v, w): return f"{v:{w},.0f}" if v is not None else "—".rjust(w)
    def ss(v, w): return str(v).rjust(w) if v is not None else "—".rjust(w)
    print(f"  {rec['date']}  {s(rec['px_avg'],8,5)}  {s(rec['ret_1d'],6,2)}  {si(rec['vol_perp_usd'],12)}  "
          f"{s(rec['taker_buy_pct'],6,1)}  {s(rec['f_avg'],9,6)}  {ss(rec['f_n'],3)}  "
          f"{si(rec['oi_usd'],11)}  {s(rec['lsr_pos_r'],7,4)}  {s(rec['lsr_acc_r'],7,4)}  "
          f"{s(rec['lsr_global_r'],7,4)}  {s(rec['div_avg3'],8,4)}  {s(rec['glsr_slope_5d'],7,4)}  "
          f"{ss(rec['pin_streak'],4)}  {rec['data_bad']}")

# ── 9. WRITE XLSX ──────────────────────────────────────────────────────────────
COLS = [
    "date","px_avg","ret_1d","vol_perp_usd","taker_buy_pct",
    "f_avg","f_min","f_max","f_n","funding_regime","pin_streak",
    "oi_usd","oi_vs_med14","oi_cycle_pos",
    "lsr_pos_r","lsr_acc_r","lsr_global_r",
    "div_raw","div_avg3","glsr_slope_5d","d_pos_3d","d_acc_3d",
    "signals_bullish","signals_suppress","screen_rank","data_bad",
    "spot_venue","vol_spot_usd","taker_spot_pct",
    "depth_bid2_usd","depth_ask2_usd",
    "bid_wall_px","bid_wall_usd","ask_wall_px","ask_wall_usd",
    "alpha_suspect",
    "exch_balance_delta","wallet_activity","dominance_pct",
    "note_operator","mm_note",
]
assert len(COLS) == 41

wb = openpyxl.Workbook()
wb.remove(wb.active)
for sh in ["CASE","DAILY","EVENTS","TRIGGERS","OUTCOME","MM_PROFILE","UNLOCK"]:
    wb.create_sheet(sh)

# CASE sheet
ws_c = wb["CASE"]
rows_case = [
    ("CASE FILE — ZBT", None),
    ("Nguon: case_tracker v1.1 · verdict chi operator ghi · so kem nguon · PENDING khong doan", None),
    ("Status",        "COLD_REFRESH — PENDING (verdict trong tay operator)"),
    ("Loai",          "cold-refresh (diem mu sau U50, chua phan loai)"),
    ("Ngay phat hien","06/08/2026 — cold-refresh lan dau, window 07/07→05/08"),
    ("T0",            "unknown — chua xac dinh"),
    ("action_class",  "PENDING — verdict trong tay operator"),
    ("Vi the",        "none"),
    ("Tom tat input", ("ZBTUSDT cold-refresh 30d (07/07→05/08). "
                       "Export: Jul+Aug CSV. LSR+OI: fapi 30d. "
                       "Alpha DELISTED (ALPHA_435 fullyDelisted, HTTP 400). "
                       "07-16/07: fn tu export, LSR tu fapi. "
                       "Pin streak 10d (27/07→05/08), f_avg=5e-05 tu 17/07. "
                       "pos_r>=3.0 tu 01/08 (5 ngay). div_avg3 4.1-4.2 tu 02/08. "
                       "glsr_slope_5d >1.0 toan window. "
                       "Vol spike: 13/07=21M, 14/07=84M (flush event). "
                       "OI spike 14/07: 11.2M (max 30d). "
                       "Xem: docs/reports/_collect_zbt_20260806.md")),
    ("Verdict (operator)", None),
    ("Dieu gi chung minh SAI", None),
    ("LEGEND", ("O vang = operator dien/quyet · Header nau vang (DAILY) = cot manual · "
                "Dong to cam (DAILY) = data_bad=1 · "
                "Cot cong thuc da tinh san (ret_1d, div, slope, oi_vs_med14, oi_cycle_pos).")),
]
for row in rows_case:
    ws_c.append(list(row))

# DAILY sheet
ws = wb["DAILY"]
SECTION_ROW = [None]*41
SECTION_ROW[0]  = "PERP (may)"
SECTION_ROW[26] = "SPOT/ALPHA (may)"
SECTION_ROW[36] = "ON-CHAIN (ban tu dong)"
SECTION_ROW[39] = "MANUAL (operator)"

ws.append(SECTION_ROW)
ws.append(COLS)

fill_bad  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
fill_hdr  = PatternFill(start_color="BF8F00", end_color="BF8F00", fill_type="solid")
bold      = Font(bold=True)

for i, v in enumerate(SECTION_ROW, 1):
    if v:
        ws.cell(1, i).font = bold
for i in range(1, 42):
    ws.cell(2, i).font = bold
for c_idx in [40, 41]:
    ws.cell(2, c_idx).fill = fill_hdr
    ws.cell(2, c_idx).font = Font(bold=True, color="FFFFFF")

for rec in computed:
    ws.append([rec[c] for c in COLS])

for r_idx, rec in enumerate(computed, start=3):
    if rec["data_bad"] == 1:
        for c_idx in range(1, 42):
            ws.cell(r_idx, c_idx).fill = fill_bad

ws.column_dimensions["A"].width = 12
for c in range(2, 42):
    ws.column_dimensions[get_column_letter(c)].width = 13

# Other sheets
wb["EVENTS"].append(["D. NHAT KY DIEN BIEN (append-only, moi nhat duoi cung)"])
wb["EVENTS"].append(["date","event","source","note"])
wb["TRIGGERS"].append(["C. CHOT THEO DOI — doc moi ngay sau nen dong"])
wb["TRIGGERS"].append(["#","trigger","met?","date_met","note"])
wb["OUTCOME"].append(["E. OUTCOME — dien tai moc; verdict_check"])
wb["MM_PROFILE"].append(["I. MM / OPERATOR PROFILE — van tay hanh vi"])
ws_ul = wb["UNLOCK"]
ws_ul.append(["J. TOKENOMICS / UNLOCK — dong cung co lich"])
ws_ul.append(["ZBT unlock schedule: UNVERIFIED (L-28 — can operator xac minh nguon)"])

out_path = case_dir / "CASE_ZBT.xlsx"
wb.save(out_path)
print(f"\n[SAVED] {out_path}")
print(f"  DAILY: {len(computed)} rows (rows 3-{2+len(computed)})")
bad_rows = [(r["date"], r["data_bad"]) for r in computed if r["data_bad"]]
print(f"  data_bad=1 rows ({len(bad_rows)}): {[d for d,_ in bad_rows]}")
print("\n[DONE]")
