"""Fix: overwrite row 17 in CASE_INDEX with correct COOKIE data, regenerate INDEX.md"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
CASE_DIR = ROOT / "docs" / "case_files"
NOW = datetime.now(timezone.utc)

idx_path = CASE_DIR / "CASE_INDEX.xlsx"
wb = openpyxl.load_workbook(str(idx_path))
ws = wb["INDEX"]

# Columns (row 4 = headers):
# col1=file, col2=token, col3=status, col4=loai, col5=action_class,
# col6=ngay_mo, col7=T0, col8=moc_ke_tiep, col9=han, col10=MM_ID, col11=ghi_chu
cookie_data = [
    "CASE_COOKIE.xlsx",                                          # file
    "COOKIE",                                                    # token
    "OPEN — audit-collect complete, awaiting operator verdict",  # status
    "GOM-C fire 08-03; pump pattern; ATH-97.9%; OI $1.1M",    # loai
    "COLD_START",                                                # action_class
    "2026-08-09",                                                # ngay_mo
    "2026-08-03",                                                # T0 (fire date)
    "depth snap #2 ≥05:50 UTC",                                  # moc_ke_tiep
    None,                                                        # han
    None,                                                        # MM_ID
    "fullyDelisted=True; Q1=0(chain); oi_cycle=0.289; VWAP=$0.00774",  # ghi_chu
]

# Overwrite row 17
for col_i, val in enumerate(cookie_data, start=1):
    ws.cell(row=17, column=col_i, value=val)

wb.save(str(idx_path))
print(f"Saved: {idx_path.name}")

# Verify row 17
wb2 = openpyxl.load_workbook(str(idx_path), read_only=True)
ws2 = wb2["INDEX"]
print(f"\nRow 17 after fix:")
for row in ws2.iter_rows(min_row=17, max_row=17, values_only=True):
    for i, v in enumerate(row, 1):
        if v is not None:
            print(f"  col{i}: {v}")

# Regenerate INDEX.md
all_rows = list(ws2.iter_rows(values_only=True))
wb2.close()

index_md = CASE_DIR / "INDEX.md"
# Get actual data rows (row 5+ where row 4 = headers)
headers = all_rows[3]  # row 4 (0-indexed: 3)
data_rows = all_rows[4:]  # row 5+ (0-indexed: 4+)

lines = [
    "# CASE INDEX\n",
    f"Generated: {NOW.strftime('%Y-%m-%d %H:%M UTC')}\n\n",
    "| " + " | ".join(str(h) if h else "" for h in headers) + " |\n",
    "|" + "|".join([" --- " for _ in headers]) + "|\n",
]
for row in all_rows[2:]:  # from row 3 (TST before headers) through end
    if any(v is not None for v in row):
        cells = [str(v) if v is not None else "" for v in row]
        lines.append("| " + " | ".join(cells) + " |\n")

index_md.write_text("".join(lines), encoding="utf-8")
print(f"\nINDEX.md written: {index_md.stat().st_size:,} bytes")
print("Done")
