"""
Rebuild DAILY sheet trong CASE_HEI.xlsx — full 541 rows (2025-02-13 → 2026-08-07).
Backup đã có tại scratch/CASE_HEI_BACKUP_20260808.xlsx (từ fix script trước).
Công thức canonical:
  div_raw    = pos_r / acc_r
  div_avg3   = avg3(pos_r) / avg3(acc_r)
  glsr_slope = gLSR(t) / gLSR(t-5)  [ratio]
L-32: verify row count sau save.
"""
import sys, io, json, csv
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import defaultdict
import urllib.request

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SCRATCH  = Path(r"D:\Data back test\whale_data\scratch")
CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
XLSX     = CASE_DIR / "CASE_HEI.xlsx"
PIN_LO, PIN_HI = 4e-5, 6e-5

def log(s=""): print(s)
def safe_float(v, d=None):
    try:
        return float(v) if v not in (None,"","None") else d
    except: return d
def fetch_json(url, timeout=20):
    try:
        req = urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode())
    except Exception as e:
        return {"_error": str(e)}

log("="*70)
log(f"REBUILD CASE_HEI DAILY  |  {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
log("="*70)

# ── Load CSVs ─────────────────────────────────────────────────────────────────
def load_csv(path):
    if not Path(path).exists(): return []
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))

perp_rows = load_csv(SCRATCH / "hei_perp_full_20260808.csv")
spot_rows = load_csv(SCRATCH / "hei_spot_full_20260808.csv")
lsr_rows  = load_csv(SCRATCH / "hei_lsr_raw_20260808.csv")
log(f"  CSVs: perp={len(perp_rows)}, spot={len(spot_rows)}, lsr={len(lsr_rows)}")

perp_by_d = {r["date"]: r for r in perp_rows}
spot_by_d = {r["date"]: r for r in spot_rows}
lsr_by_d  = {r["date"]: r for r in lsr_rows}

# ── Re-fetch funding + OI + 07/08 live data ───────────────────────────────────
log("  Fetching funding rates...")
fund_raw = fetch_json("https://fapi.binance.com/fapi/v1/fundingRate?symbol=HEIUSDT&limit=1000")
fund_by_day = defaultdict(list)
if isinstance(fund_raw, list):
    for fr in fund_raw:
        ts = fr.get("fundingTime")
        if ts:
            dt = datetime.fromtimestamp(int(ts)/1000, tz=timezone.utc).strftime("%Y-%m-%d")
            fund_by_day[dt].append(float(fr.get("fundingRate", 0)))
log(f"  Funding days: {len(fund_by_day)}")

log("  Fetching OI hist...")
oih_raw = fetch_json("https://fapi.binance.com/futures/data/openInterestHist?symbol=HEIUSDT&period=1d&limit=30")
oi_by_d = {}
if isinstance(oih_raw, list):
    for o in oih_raw:
        ts = o.get("timestamp")
        if ts:
            dt = datetime.fromtimestamp(int(ts)/1000, tz=timezone.utc).strftime("%Y-%m-%d")
            oi_by_d[dt] = {"oi_usd": float(o.get("sumOpenInterestValue",0)),
                           "oi_token": float(o.get("sumOpenInterest",0))}
log(f"  OI days: {len(oi_by_d)}")

# Fetch 2026-08-07 live klines (not in CSVs)
log("  Fetching 2026-08-07 perp+spot klines...")
ts_0807s = int(datetime(2026,8,7,0,0,0,tzinfo=timezone.utc).timestamp()*1000)
ts_0807e = int(datetime(2026,8,8,0,0,0,tzinfo=timezone.utc).timestamp()*1000)

pk = fetch_json(f"https://fapi.binance.com/fapi/v1/klines?symbol=HEIUSDT&interval=1d&startTime={ts_0807s}&endTime={ts_0807e}&limit=2")
perp_0807 = {}
if isinstance(pk, list) and pk:
    k = pk[0]
    vol_q = float(k[7])
    taker_q = float(k[10])
    perp_0807 = {"o":float(k[1]),"h":float(k[2]),"l":float(k[3]),"c":float(k[4]),
                 "vol_quote_usd":vol_q, "taker_pct":round(taker_q/vol_q*100,2) if vol_q else None,
                 "ret_1d":None}
    log(f"  Perp 07/08: o={perp_0807['o']} c={perp_0807['c']} vol=${vol_q:.0f}")

sk = fetch_json(f"https://api.binance.com/api/v3/klines?symbol=HEIUSDT&interval=1d&startTime={ts_0807s}&endTime={ts_0807e}&limit=2")
spot_0807 = {}
if isinstance(sk, list) and sk:
    k = sk[0]
    vol_s = float(k[7]); taker_s = float(k[10])
    vol_p = perp_0807.get("vol_quote_usd") or 1
    spot_0807 = {"vol_spot_usd":vol_s,"taker_spot_pct":round(taker_s/vol_s*100,2) if vol_s else None,
                 "spot_perp_ratio":round(vol_s/vol_p,3),"l29_flag":"CLEAR" if vol_s/vol_p < 3 else "L29_ABSTAIN"}
    log(f"  Spot 07/08: vol_spot=${vol_s:.0f} ratio={spot_0807['spot_perp_ratio']}")

# Inject 07/08 into lookup dicts
perp_by_d["2026-08-07"] = perp_0807
spot_by_d["2026-08-07"] = spot_0807

# ── Build all_dates (2025-02-13 → 2026-08-07) ────────────────────────────────
all_dates = sorted(perp_by_d.keys())
log(f"  Date range: {all_dates[0]} → {all_dates[-1]} ({len(all_dates)} dates)")

# ── Compute global rolling arrays ─────────────────────────────────────────────
vol_vals = [safe_float(perp_by_d.get(d,{}).get("vol_quote_usd")) for d in all_dates]
oi_vals  = [oi_by_d.get(d,{}).get("oi_usd") for d in all_dates]
pos_r_all= [safe_float(lsr_by_d.get(d,{}).get("pos_r_raw")) for d in all_dates]
acc_r_all= [safe_float(lsr_by_d.get(d,{}).get("acc_r_raw")) for d in all_dates]
glsr_all = [safe_float(lsr_by_d.get(d,{}).get("glsr_raw")) for d in all_dates]

def rolling_med14(vals):
    out = [None]*len(vals)
    for i in range(14, len(vals)):
        w = [v for v in vals[i-14:i] if v is not None]
        if len(w) >= 7:
            s = sorted(w); n = len(s)
            out[i] = (s[n//2-1]+s[n//2])/2 if n%2==0 else s[n//2]
    return out

vol_med14 = rolling_med14(vol_vals)
oi_med14  = rolling_med14(oi_vals)

# (a) div_raw = pos_r / acc_r
div_raw_all = [round(pos_r_all[i]/acc_r_all[i],4) if (pos_r_all[i] and acc_r_all[i] and acc_r_all[i]!=0) else None for i in range(len(all_dates))]

# (b) div_avg3 = avg3(pos_r) / avg3(acc_r)
div_avg3_all = [None]*len(all_dates)
for i in range(2, len(all_dates)):
    p3 = [v for v in pos_r_all[i-2:i+1] if v is not None]
    a3 = [v for v in acc_r_all[i-2:i+1] if v is not None]
    if len(p3)==3 and len(a3)==3:
        ap = sum(p3)/3; aa = sum(a3)/3
        div_avg3_all[i] = round(ap/aa,4) if aa != 0 else None

# (c) glsr_slope_5d = gLSR(t)/gLSR(t-5) ratio
slope_all = [None]*len(all_dates)
for i in range(5, len(all_dates)):
    g = glsr_all[i]; gp = glsr_all[i-5]
    if g is not None and gp is not None and gp > 0:
        slope_all[i] = round(g/gp, 4)

# pos_r persist count (for GOM-B B1)
pos_r_persist = [0]*len(all_dates)
cnt_ = 0
for i in range(len(all_dates)):
    if pos_r_all[i] and pos_r_all[i] > 1.5: cnt_ += 1
    else: cnt_ = 0
    pos_r_persist[i] = cnt_

log(f"  div_raw non-null: {sum(1 for v in div_raw_all if v is not None)}")
log(f"  div_avg3 non-null: {sum(1 for v in div_avg3_all if v is not None)}")
log(f"  slope non-null: {sum(1 for v in slope_all if v is not None)}")

# ── Build DAILY rows ──────────────────────────────────────────────────────────
DAILY_HEADERS = [
    "date","px_avg","o","h","l","c","vol_perp_usd","taker_buy_pct",
    "f_avg","f_min","f_max","fn","oi_usd","oi_token","data_ok","low_conf",
    "ret_1d","pin_streak","pin_day","oi_vs_med14","vol_vs_med14",
    "div_raw","div_avg3","glsr_slope_5d","d_pos_3d","d_acc_3d",
    "pos_r","acc_r","gLSR","taker_lsr",
    "vol_spot_usd","taker_spot_pct","spot_perp_ratio","l29_flag",
    "vol_alpha","taker_alpha","on_chain_source",
    "gom_A_screen","gom_B_screen","gom_note",
    "data_source_perp","data_source_lsr",
]
assert len(DAILY_HEADERS) == 42

daily_data = []
pin_streak = 0

for i, dt in enumerate(all_dates):
    pr = perp_by_d.get(dt, {})
    sr = spot_by_d.get(dt, {})
    ls = lsr_by_d.get(dt, {})
    fd = fund_by_day.get(dt, [])
    oi = oi_by_d.get(dt, {})

    o  = safe_float(pr.get("o"))
    h  = safe_float(pr.get("h"))
    l  = safe_float(pr.get("l"))
    c  = safe_float(pr.get("c"))
    px_avg = round((h+l+c)/3, 8) if (h and l and c) else None
    vol_p  = safe_float(pr.get("vol_quote_usd"))
    taker  = safe_float(pr.get("taker_pct"))
    ret_1d = safe_float(pr.get("ret_1d"))

    fn    = len(fd)
    f_avg = round(sum(fd)/fn, 8) if fn > 0 else None
    f_min = round(min(fd), 8)    if fn > 0 else None
    f_max = round(max(fd), 8)    if fn > 0 else None
    data_ok  = fn >= 6 if fn > 0 else None
    low_conf = fn < 6  if fn > 0 else True

    oi_usd   = oi.get("oi_usd")
    oi_token = oi.get("oi_token")

    pin_day = all(PIN_LO <= f <= PIN_HI for f in fd) and fn >= 6 if fn > 0 else None
    if pin_day is True: pin_streak += 1
    elif pin_day is False: pin_streak = 0
    else: pin_streak = 0

    oi_vs_m  = round(oi_usd/oi_med14[i], 3)   if (oi_usd  and oi_med14[i])  else None
    vol_vs_m = round(vol_p/vol_med14[i], 2)    if (vol_p   and vol_med14[i]) else None

    pos_r = pos_r_all[i]
    acc_r = acc_r_all[i]
    glsr  = glsr_all[i]
    tkls  = safe_float(ls.get("taker_lsr_raw"))

    d_pos = None; d_acc = None
    if i >= 3:
        pp = pos_r_all[i-3]; ap = acc_r_all[i-3]
        if pos_r and pp and pp > 0: d_pos = round((pos_r-pp)/pp*100, 2)
        if acc_r and ap and ap > 0: d_acc = round((acc_r-ap)/ap*100, 2)

    vol_s = safe_float(sr.get("vol_spot_usd"))
    tk_s  = safe_float(sr.get("taker_spot_pct"))
    spr   = safe_float(sr.get("spot_perp_ratio"))
    l29   = sr.get("l29_flag","") or ""

    # GOM-A (A2: slope < 1.0 canonical)
    vol_3x = (vol_p/vol_med14[i] >= 3) if (vol_p and vol_med14[i]) else False
    A1 = vol_3x
    A2 = (slope_all[i] is not None and slope_all[i] < 1.0)
    A3 = (pos_r is not None and pos_r > 1.5)
    A4 = (f_avg is not None and f_avg >= PIN_LO)
    A5 = (taker is not None and taker >= 45)
    gom_a = A1 and A2 and A3 and A4 and A5

    B1 = (pos_r_persist[i] >= 7)
    B2 = (pin_day is True)
    B3 = (vol_3x and ret_1d is not None and ret_1d < 0 and taker is not None and taker >= 50)
    B4 = (oi_vs_m is not None and oi_vs_m >= 1.5)
    gom_b = B1 and (B2 or B3) and B4

    gom_note = ""
    if gom_a or gom_b:
        parts = []
        if gom_a: parts.append(f"A(A1={A1},A2-ratio={A2},A3={A3},A4={A4},A5={A5})")
        if gom_b: parts.append(f"B(B1={B1},B2={B2},B3={B3},B4={B4})")
        gom_note = " | ".join(parts) + " — gom_detect_spec v1.1 CHUA KY"

    row = {
        "date":dt, "px_avg":px_avg,
        "o":o,"h":h,"l":l,"c":c,
        "vol_perp_usd":vol_p, "taker_buy_pct":taker,
        "f_avg":f_avg,"f_min":f_min,"f_max":f_max,
        "fn":fn if fn>0 else None,
        "oi_usd":oi_usd,"oi_token":oi_token,
        "data_ok":data_ok,"low_conf":low_conf,
        "ret_1d":ret_1d,
        "pin_streak":pin_streak,"pin_day":pin_day,
        "oi_vs_med14":oi_vs_m,"vol_vs_med14":vol_vs_m,
        "div_raw":div_raw_all[i],"div_avg3":div_avg3_all[i],
        "glsr_slope_5d":slope_all[i],
        "d_pos_3d":d_pos,"d_acc_3d":d_acc,
        "pos_r":pos_r,"acc_r":acc_r,"gLSR":glsr,"taker_lsr":tkls,
        "vol_spot_usd":vol_s,"taker_spot_pct":tk_s,"spot_perp_ratio":spr,"l29_flag":l29,
        "vol_alpha":"ABSTAIN(no_alpha_id)","taker_alpha":"BANNED(k9==k5)",
        "on_chain_source":"ABSTAIN(Nansen_LOCKED)",
        "gom_A_screen":"TRUE" if gom_a else "false",
        "gom_B_screen":"TRUE" if gom_b else "false",
        "gom_note":gom_note,
        "data_source_perp":"fapi_klines_1d",
        "data_source_lsr":"fapi_LSR_30d" if ls else "MISSING(>30d)",
    }
    daily_data.append(row)

log(f"  DAILY rows built: {len(daily_data)}")
log(f"  GOM-A fires: {sum(1 for r in daily_data if r['gom_A_screen']=='TRUE')}")
log(f"  GOM-A dates: {[r['date'] for r in daily_data if r['gom_A_screen']=='TRUE']}")
log(f"  GOM-B fires: {sum(1 for r in daily_data if r['gom_B_screen']=='TRUE')}")

# ── Load existing workbook (keep all other sheets) ────────────────────────────
try:
    wb = openpyxl.load_workbook(XLSX)
except PermissionError as e:
    log(f"  PERMISSION ERROR: {e}"); raise

# ── Rebuild DAILY sheet ───────────────────────────────────────────────────────
# Remove old DAILY and create fresh
if "DAILY" in wb.sheetnames:
    idx = wb.sheetnames.index("DAILY")
    del wb["DAILY"]
    ws_daily = wb.create_sheet("DAILY", idx)
else:
    ws_daily = wb.create_sheet("DAILY", 1)

HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
HDR_FONT  = Font(bold=True, color="FFFFFF", size=9)
GRP_FILLS = {
    "PERP":  PatternFill("solid", fgColor="BDD7EE"),
    "DERIV": PatternFill("solid", fgColor="E2EFDA"),
    "CROWD": PatternFill("solid", fgColor="FCE4D6"),
    "SPOT":  PatternFill("solid", fgColor="FFF2CC"),
    "ALPHA": PatternFill("solid", fgColor="F2F2F2"),
    "CHAIN": PatternFill("solid", fgColor="EAD1DC"),
    "GOM":   PatternFill("solid", fgColor="D9D9D9"),
    "SRC":   PatternFill("solid", fgColor="EDEDED"),
}

COL_GROUPS = [
    ("PERP MACHINE (fapi klines 1d + funding)", 1, 16, "PERP"),
    ("DERIVED (cong thuc canonical — Phu luc A)", 17, 26, "DERIV"),
    ("CROWD — LSR (fapi, 30d window — L-38)", 27, 30, "CROWD"),
    ("SPOT (binance api/v3 k[10]/k[7])", 31, 34, "SPOT"),
    ("ALPHA (ABSTAIN — no alpha_id)", 35, 36, "ALPHA"),
    ("ON-CHAIN (ABSTAIN — Nansen LOCKED)", 37, 37, "CHAIN"),
    ("GOM SCREEN (gom_detect_spec v1.1 CHUA KY)", 38, 40, "GOM"),
    ("NGUON", 41, 42, "SRC"),
]

for label, cs, ce, grp in COL_GROUPS:
    ws_daily.merge_cells(start_row=1, start_column=cs, end_row=1, end_column=ce)
    c = ws_daily.cell(1, cs)
    c.value = label
    c.fill  = GRP_FILLS.get(grp, HDR_FILL)
    c.font  = Font(bold=True, size=8)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

HDR_LABELS = [
    "date","px_avg (H+L+C)/3","o","h","l","c",
    "vol_perp_usd","taker_buy_pct (k[10]/k[7])","f_avg","f_min","f_max","fn (4h=6/day)",
    "oi_usd","oi_token","data_ok (fn>=6)","low_conf (fn<6)",
    "ret_1d (close-to-close)","pin_streak","pin_day ([4e-5,6e-5] all fn)","oi_vs_med14","vol_vs_med14",
    "div_raw (pos_r/acc_r)","div_avg3 (avg3(pos_r)/avg3(acc_r))","glsr_slope_5d (ratio t/t-5)","d_pos_3d (%chg)","d_acc_3d (%chg)",
    "pos_r","acc_r","gLSR","taker_lsr (buy/sell)",
    "vol_spot_usd","taker_spot_pct (k[10]/k[7])","spot_perp_ratio","l29_flag",
    "vol_alpha (ABSTAIN)","taker_alpha (BANNED)",
    "on_chain_source",
    "gom_A_screen","gom_B_screen","gom_note (CHUA KY)",
    "data_source_perp","data_source_lsr",
]
assert len(HDR_LABELS) == 42

for ci, lbl in enumerate(HDR_LABELS, 1):
    c = ws_daily.cell(2, ci)
    c.value = lbl
    c.fill  = HDR_FILL
    c.font  = HDR_FONT
    c.alignment = Alignment(wrap_text=True, vertical="center")

ws_daily.row_dimensions[1].height = 28
ws_daily.row_dimensions[2].height = 36

# Write data rows
for ri_off, row in enumerate(daily_data):
    ri = ri_off + 3
    for ci, key in enumerate(DAILY_HEADERS, 1):
        val = row.get(key)
        # Always write a string for None to ensure row exists in XML
        c = ws_daily.cell(ri, ci)
        c.value = val if val is not None else ""
        c.font  = Font(size=8)
        c.alignment = Alignment(vertical="center")
    # Highlights
    if row.get("gom_A_screen") == "TRUE":
        ws_daily.cell(ri, 38).fill = PatternFill("solid", fgColor="FFA500")
    if row.get("gom_B_screen") == "TRUE":
        ws_daily.cell(ri, 39).fill = PatternFill("solid", fgColor="FF6B6B")
    if not row.get("data_ok"):
        ws_daily.cell(ri, 15).fill = PatternFill("solid", fgColor="FFFF00")

ws_daily.freeze_panes = "B3"
ws_daily.column_dimensions["A"].width = 12
for ci in range(2, 43):
    ws_daily.column_dimensions[get_column_letter(ci)].width = 12

log(f"\n  DAILY sheet rebuilt. Saving...")
try:
    wb.save(XLSX)
    log(f"  Saved: {XLSX}")
except PermissionError as e:
    log(f"  PERMISSION ERROR: {e}"); raise

# ── L-32: Verify ──────────────────────────────────────────────────────────────
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2["DAILY"]
actual = ws2.max_row - 2
log(f"\n  L-32: expected {len(daily_data)} rows, max_row-2={actual}")

# More robust count: count non-empty date cells
cnt = sum(1 for ri in range(3, ws2.max_row+100) if ws2.cell(ri,1).value not in (None,""))
log(f"  L-32 (count non-empty date): {cnt}")
assert cnt == len(daily_data), f"ROW COUNT MISMATCH: expected {len(daily_data)}, got {cnt}"
log(f"  L-32: PASS")
log(f"  Sheets: {wb2.sheetnames}")

# ── Sample 5 rows ─────────────────────────────────────────────────────────────
log(f"\n  === SAMPLE: last 5 DAILY rows ===")
log(f"  {'date':12} {'pos_r':7} {'acc_r':7} {'gLSR':7} {'div_raw':8} {'div_avg3':9} {'slope':7} {'gom_A':6}")
for ri in range(ws2.max_row-4, ws2.max_row+1):
    vals = [ws2.cell(ri, ci).value for ci in [1,27,28,29,22,23,24,38]]
    log(f"  {str(vals[0]):12} {str(vals[1])[:7]:7} {str(vals[2])[:7]:7} {str(vals[3])[:7]:7} {str(vals[4])[:8]:8} {str(vals[5])[:9]:9} {str(vals[6])[:7]:7} {str(vals[7])[:6]:6}")

log(f"\n[REBUILD DONE]")
