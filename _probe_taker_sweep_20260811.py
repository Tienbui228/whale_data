"""B-sweep: scan all CASE_*.xlsx for taker_buy_pct fraction scale.

Identifies which packs were built with taker as fraction (0-1) vs correct pct (20-80).
Output: rebuild list + ANALYSIS warning.

data-source: docs/case_files/CASE_*.xlsx (local)
machine: local
date-range: snapshot 2026-08-11
code-commit: see git log
"""
from __future__ import annotations
import sys, io, math
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.cell.cell import MergedCell

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")
WINDOW = 30

def _normalize_col(s: str) -> str:
    return s.split("(")[0].strip()

def _find_header_row(ws):
    for r in range(1, 6):
        row = ws[r]
        names = {}
        for i, cell in enumerate(row):
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            raw = str(cell.value)
            norm = _normalize_col(raw)
            names[raw] = i + 1
            if norm != raw:
                names[norm] = i + 1
        aliases = {"f_n": "fn", "lsr_pos_r": "pos_r", "lsr_acc_r": "acc_r",
                   "lsr_global_r": "gLSR"}
        for v1, canon in aliases.items():
            if v1 in names and canon not in names:
                names[canon] = names[v1]
        if "date" in names or "px_avg" in names:
            return r, names
    return 0, {}

results = []

for xlsx_path in sorted(CASE_DIR.glob("CASE_*.xlsx")):
    if xlsx_path.name.startswith("~$"):
        continue
    token = xlsx_path.stem.replace("CASE_", "")
    try:
        wb = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
        if "DAILY" not in wb.sheetnames:
            wb.close()
            results.append((token, "NO_DAILY_SHEET", None, None, "?"))
            continue
        ws = wb["DAILY"]
        _, headers = _find_header_row(ws)
        col_idx = headers.get("taker_buy_pct")
        if col_idx is None:
            wb.close()
            results.append((token, "COL_MISSING", None, None, "REBUILD?"))
            continue

        data_rows = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row[0] not in (None, ""):
                data_rows.append(row)

        wb.close()

        window_rows = data_rows[-WINDOW:] if len(data_rows) >= WINDOW else data_rows
        n = len(window_rows)
        col_0 = col_idx - 1

        vals = []
        for row in window_rows:
            v = row[col_0] if col_0 < len(row) else None
            if v is not None and str(v).strip() not in ("", "ABSTAIN", "MISSING", "N/A"):
                try:
                    fv = float(v)
                    if not math.isnan(fv):
                        vals.append(fv)
                except (TypeError, ValueError):
                    pass

        if not vals:
            results.append((token, "ALL_NULL", None, None, "CHECK"))
        elif max(vals) <= 1.0:
            results.append((token, "FRACTION", min(vals), max(vals), "REBUILD"))
        elif max(vals) == min(vals):
            results.append((token, "CONSTANT", min(vals), max(vals), "CHECK"))
        else:
            results.append((token, "OK", min(vals), max(vals), "OK"))

    except Exception as exc:
        results.append((token, f"ERROR: {exc}", None, None, "CHECK"))

print(f"\n{'='*70}")
print(f"TAKER SWEEP — CASE_*.xlsx ({len(results)} files)")
print(f"{'='*70}")
print(f"  {'Token':<20} {'Status':<12} {'Min':>8} {'Max':>8}  Action")
print(f"  {'-'*60}")

rebuild_list = []
for tok, status, mn, mx, action in results:
    mn_s = f"{mn:.4f}" if mn is not None else "N/A"
    mx_s = f"{mx:.4f}" if mx is not None else "N/A"
    flag = " ◄" if action == "REBUILD" else ""
    print(f"  {tok:<20} {status:<12} {mn_s:>8} {mx_s:>8}  {action}{flag}")
    if action == "REBUILD":
        rebuild_list.append(tok)

print(f"\n{'='*70}")
print(f"REBUILD LIST ({len(rebuild_list)} packs):")
for t in rebuild_list:
    print(f"  {t}")
print()
print("ANALYSIS WARNING (cột taker các pack trên = zero rác, đừng chấm):")
print("Các pack có taker FRACTION được ghi 0.52 thay vì 52%.")
print("B3-streak, flag_dist_wall, flag_decay trong các pack này đều SAI.")
print("Rebuild sau khi fix build_case.py (taker scale normalized)")
print("Đã SCP code mới lên VPS chưa → confirm trước khi rebuild bất kỳ pack.")
