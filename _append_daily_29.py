#!/usr/bin/env python3
"""Append 2026-07-29 row to DAILY sheet in CASE_TAG.xlsx + CASE_BR.xlsx.
  TAG: fill row 31 (empty). BR: fill missing cells in existing row 67.
"""
import re, openpyxl
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"D:\Data back test\whale_data")
TAG_FILE = ROOT / "docs" / "case_files" / "CASE_TAG.xlsx"
BR_FILE  = ROOT / "docs" / "case_files" / "CASE_BR.xlsx"

def inc_formula(formula, by=1):
    """Increment all row numbers in an Excel formula string by `by`."""
    return re.sub(r'([A-Z]+)(\d+)', lambda m: m.group(1) + str(int(m.group(2)) + by), formula)

def copy_formula_row(ws_src_row, ws, new_row_num):
    """Copy formula cells from source row, incrementing row numbers by 1."""
    for cell in ws_src_row:
        if cell.value and str(cell.value).startswith("="):
            target = ws.cell(row=new_row_num, column=cell.column)
            if target.value is None:  # only fill empty cells
                target.value = inc_formula(str(cell.value))

# ─────────────────────────────────────────────────────────────
# CASE_TAG: append row 31 = 2026-07-29
# ─────────────────────────────────────────────────────────────
print("Opening CASE_TAG.xlsx...")
wb_tag = load_workbook(str(TAG_FILE))
ws_tag = wb_tag["DAILY"]

ROW = 31  # target new row

# Raw data — col letter → value
raw_tag = {
    "A": "2026-07-29",
    "B": None,             # px_avg empty (fn=1, no OHLC)
    "D": None,             # vol_perp_usd empty
    "E": None,             # taker_buy_pct empty
    "F": 0.00014019,       # funding_avg
    "G": 0.00014019,       # funding_min (fn=1 → same as avg)
    "H": 0.00014019,       # funding_max (fn=1 → same as avg)
    "I": 1,                # funding_n
    "J": "",               # funding_regime — low_conf, leave blank
    "L": 11172261,         # oi_usd (from export)
    "N": "PEAK",           # oi_cycle_pos (above 75th pctile Jul)
    "O": 1.6204,           # lsr_pos_r
    "P": 1.3719,           # lsr_acc_r
    "Q": 1.8129,           # lsr_global_r (fn=1 SUSPECT — see note)
    "Z": "",               # data_bad = blank (data_ok=True, just low_conf)
    "AA": "gate",          # spot_venue
    "AN": "gLSR 1.8129 fn=1 low_conf — single snapshot, suspect spike; acc_r 0.999->1.372 keo div_raw 1.597->1.181",
}

# Fill raw values
for col_letter, val in raw_tag.items():
    if val is not None:
        ws_tag[f"{col_letter}{ROW}"] = val

# Copy formulas from row 30, incrementing row references
src_row = ws_tag[30]
copy_formula_row(src_row, ws_tag, ROW)

# Verify what we wrote
print(f"  TAG row {ROW} written:")
for col_letter in ["A","B","C","F","G","H","I","L","M","N","O","P","Q","R","S","T","U","V","AN"]:
    cell = ws_tag[f"{col_letter}{ROW}"]
    if cell.value is not None:
        print(f"    {col_letter}{ROW} = {cell.value}")

wb_tag.save(str(TAG_FILE))
print(f"  Saved {TAG_FILE.name}")

# ─────────────────────────────────────────────────────────────
# CASE_BR: row 67 already exists — fill missing cells only
# ─────────────────────────────────────────────────────────────
print("\nOpening CASE_BR.xlsx...")
wb_br = load_workbook(str(BR_FILE))
ws_br = wb_br["DAILY"]

BR_ROW = 67

# Only write to cells that are currently empty
missing_br = {
    "G": 0.00017978,   # f_min (fn=1 → same as avg)
    "H": 0.00017978,   # f_max (fn=1 → same as avg)
    "I": 1,            # f_n
    # do NOT overwrite O, P, Q, L, F, B, C, AK, AL, AN, AO — operator data
}

for col_letter, val in missing_br.items():
    cell = ws_br[f"{col_letter}{BR_ROW}"]
    if cell.value is None:
        cell.value = val
        print(f"  BR row {BR_ROW} filled: {col_letter}{BR_ROW} = {val}")
    else:
        print(f"  BR row {BR_ROW} SKIP (already has value): {col_letter}{BR_ROW} = {cell.value}")

# Show full row 67 state
print(f"  BR row {BR_ROW} final state:")
for col_letter in ["A","B","F","G","H","I","J","L","N","O","P","Q","R","S","T","AN","AO"]:
    cell = ws_br[f"{col_letter}{BR_ROW}"]
    if cell.value is not None:
        print(f"    {col_letter}{BR_ROW} = {cell.value}")

wb_br.save(str(BR_FILE))
print(f"  Saved {BR_FILE.name}")

# ─────────────────────────────────────────────────────────────
# Verify
# ─────────────────────────────────────────────────────────────
print("\nVerify TAG row 31:")
wb2 = load_workbook(str(TAG_FILE), data_only=True)
ws2 = wb2["DAILY"]
row31 = {ws2.cell(row=31, column=c).column_letter: ws2.cell(row=31, column=c).value
         for c in range(1, 42) if ws2.cell(row=31, column=c).value is not None}
print(" ", row31)
wb2.close()

print("\nVerify BR row 67:")
wb3 = load_workbook(str(BR_FILE), data_only=True)
ws3 = wb3["DAILY"]
row67 = {ws3.cell(row=67, column=c).column_letter: ws3.cell(row=67, column=c).value
         for c in range(1, 42) if ws3.cell(row=67, column=c).value is not None}
print(" ", row67)
wb3.close()

print("\nDone.")
