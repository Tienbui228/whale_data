"""Verify CASE_STBL.xlsx last 3 rows + CASE meta head"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_STBL.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["DAILY"]
row2 = [c.value for c in ws[2]]
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
print(f"DAILY dims: {ws.dimensions}, max_row: {ws.max_row}")

print(f"\nLast 3 rows key fields:")
for r in range(ws.max_row-2, ws.max_row+1):
    d = ws.cell(r, ci["date"]).value
    print(f"\n  row {r} ({d}):")
    for k in ["c","vol_perp_usd","oi_usd","f_avg","fn (4h=6/day)","pos_r","gLSR",
             "div_avg3 (avg3p/avg3a)","vol_spot_usd","vol_alpha","vol_spot_ext_usd","spot_ext_venue","data_source_lsr"]:
        if k in ci:
            print(f"    {k}: {ws.cell(r, ci[k]).value}")

ws_c = wb["CASE"]
print(f"\nCASE first 12 rows:")
for r in range(1, 13):
    v1, v2 = ws_c.cell(r,1).value, ws_c.cell(r,2).value
    if v1 or v2:
        print(f"  r{r}: {v1!r} | {v2!r}")
