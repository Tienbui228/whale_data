"""V1 + V2: kiểm vol ×N.
V1: 5 CASE file vol_perp_usd vs fresh fapi k[7] toàn range → paste bảng ratio ≠ 1.00 ±2%.
V2: freshwatch store vs fapi k[7] cho 5 case symbols + 5 random khác fw-139.
"""
import sys, io, json, urllib.request, time, csv, random
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone, date as _date, timedelta
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
CASES = ROOT / "docs" / "case_files"
EXPORT = ROOT / "data" / "export"
FW139 = ROOT / "scratch" / "freshwatch_symbols.txt"

def fetch_json(url, params=None, timeout=15):
    if params:
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        url = f"{url}?{qs}"
    req = urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode())
    except Exception as e:
        return {"_error": str(e)}
def to_ms(d): return int(datetime.strptime(d,"%Y-%m-%d").replace(tzinfo=timezone.utc).timestamp()*1000)
def ms_to_date(ms): return datetime.fromtimestamp(int(ms)/1000, tz=timezone.utc).strftime("%Y-%m-%d")

def fetch_klines_full(ticker, start, end):
    s_ms = to_ms(start); e_ms = to_ms(end) + 86400000 - 1
    all_k = []; cur = s_ms
    while cur <= e_ms:
        d = fetch_json("https://fapi.binance.com/fapi/v1/klines",
                       {"symbol":ticker,"interval":"1d","startTime":cur,"endTime":e_ms,"limit":1000})
        time.sleep(0.3)
        if not isinstance(d, list) or not d: break
        all_k.extend(d)
        last = int(d[-1][0])
        if last <= cur or last >= e_ms: break
        cur = last + 86400000
    by_d = {}
    for k in all_k:
        d_str = ms_to_date(int(k[0]))
        if start <= d_str <= end:
            by_d[d_str] = {"k7": round(float(k[7]), 2)}  # quote asset volume USD
    return by_d

def safe_float(v):
    if v in (None, "", "None"): return None
    try: return float(v)
    except: return None

# ─── V1: CASE files vs fresh k[7] ───
print("="*70)
print("V1 — 5 CASE file vol_perp_usd vs fresh fapi k[7]")
print("="*70)

CASE_CONFIG = [
    ("BANK", "BANKUSDT"),
    ("TUT",  "TUTUSDT"),
    ("AKE",  "AKEUSDT"),
    ("BTW",  "BTWUSDT"),
    ("CYS",  "CYSUSDT"),
]

v1_results = {}

for sym, ticker in CASE_CONFIG:
    fp = CASES / f"CASE_{sym}.xlsx"
    if not fp.exists():
        print(f"\n[{sym}] file missing"); continue
    wb = openpyxl.load_workbook(fp, data_only=True)
    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
    vp_c = ci.get("vol_perp_usd")

    # Collect stored per date
    stored = {}
    for r in range(3, ws.max_row+1):
        d = ws.cell(r, 1).value
        if not d: continue
        d_str = str(d)[:10]
        v = safe_float(ws.cell(r, vp_c).value) if vp_c else None
        stored[d_str] = v

    if not stored:
        print(f"\n[{sym}] no stored data"); continue
    dates = sorted(stored.keys())
    start = dates[0]; end = dates[-1]

    print(f"\n[{sym}] range {start}→{end} · fetching fresh...")
    fresh = fetch_klines_full(ticker, start, end)

    mismatches = []
    for d_str in dates:
        s = stored[d_str]
        f = fresh.get(d_str, {}).get("k7")
        if s is None or f is None or f == 0:
            continue
        ratio = s / f
        if abs(ratio - 1.0) > 0.02:  # >2% diff
            mismatches.append((d_str, s, f, ratio))

    v1_results[sym] = mismatches
    print(f"  total dates: {len(dates)} · mismatches (>2%): {len(mismatches)}")
    if mismatches:
        print(f"  {'date':<12} {'stored':>18} {'fresh':>18} {'ratio':>8}")
        for d, s, f, r in mismatches[:30]:
            print(f"  {d:<12} {s:>18,.2f} {f:>18,.2f} {r:>8.3f}")
        if len(mismatches) > 30:
            print(f"  ... +{len(mismatches)-30} more")

# ─── V2: freshwatch store vs fapi k[7] ───
print("\n" + "="*70)
print("V2 — freshwatch store vol vs fapi k[7] cho 5 case + 5 random fw-139")
print("="*70)

with open(FW139, encoding="utf-8", errors="ignore") as f:
    fw139_symbols = [l.strip().upper() for l in f if l.strip()]
random.seed(42)
random_5 = random.sample([s for s in fw139_symbols if s not in ["BANKUSDT","TUTUSDT","AKEUSDT","BTWUSDT","CYSUSDT"]], 5)
CHECK_SYMBOLS = ["BANKUSDT","TUTUSDT","AKEUSDT","BTWUSDT","CYSUSDT"] + random_5
print(f"Random 5 fw-139: {random_5}")

# Load freshwatch vol per symbol per date
fw_data = {}  # sym → date → vol
for fp in sorted(EXPORT.glob("freshwatch_daily_2026-*.csv")):
    with open(fp, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            sym = row.get("symbol","").upper()
            if sym not in CHECK_SYMBOLS: continue
            d_str = row.get("date","")[:10]
            v = safe_float(row.get("vol_quote"))
            if d_str and v:
                fw_data.setdefault(sym, {})[d_str] = v

v2_results = {}
for sym in CHECK_SYMBOLS:
    if sym not in fw_data or not fw_data[sym]:
        print(f"\n[{sym}] no freshwatch data"); continue
    dates = sorted(fw_data[sym].keys())
    start = dates[0]; end = dates[-1]
    print(f"\n[{sym}] freshwatch range {start}→{end} · fetching fresh...")
    fresh = fetch_klines_full(sym, start, end)

    mismatches = []
    for d_str in dates:
        s = fw_data[sym][d_str]
        f = fresh.get(d_str, {}).get("k7")
        if s is None or f is None or f == 0:
            continue
        ratio = s / f
        if abs(ratio - 1.0) > 0.02:
            mismatches.append((d_str, s, f, ratio))

    v2_results[sym] = mismatches
    print(f"  freshwatch dates: {len(dates)} · mismatches (>2%): {len(mismatches)}")
    if mismatches:
        print(f"  {'date':<12} {'store':>18} {'fresh':>18} {'ratio':>8}")
        for d, s, f, r in mismatches[:20]:
            print(f"  {d:<12} {s:>18,.2f} {f:>18,.2f} {r:>8.3f}")
        if len(mismatches) > 20:
            print(f"  ... +{len(mismatches)-20} more")

# Summary
print("\n" + "="*70)
print("SUMMARY V1 + V2")
print("="*70)
print(f"\nV1 CASE file mismatches:")
for sym in ["BANK","TUT","AKE","BTW","CYS"]:
    m = v1_results.get(sym, [])
    print(f"  {sym:<6}: {len(m)} ngày ≠1 (>2%)" if m else f"  {sym:<6}: 0 ngày ≠1")

print(f"\nV2 freshwatch store mismatches:")
store_contaminated = []
for sym in CHECK_SYMBOLS:
    m = v2_results.get(sym, [])
    if m:
        print(f"  {sym:<15}: {len(m)} ngày ≠1 → STORE-NHIỄM")
        store_contaminated.append(sym)
    else:
        print(f"  {sym:<15}: 0 ngày ≠1")

if store_contaminated:
    print(f"\n⚠ STORE-NHIỄM confirmed: {store_contaminated}")
    print(f"P0 protocol: KHÔNG sửa store (vai DEV), báo và DỪNG")
else:
    print(f"\n✓ freshwatch store CLEAN — không nhiễm")

# Save for V3
import pickle
with open(ROOT/"scratch"/"v1_v2_results_20260821.pkl", "wb") as f:
    pickle.dump({"v1": v1_results, "v2": v2_results, "store_contaminated": store_contaminated}, f)
print(f"\nSaved: scratch/v1_v2_results_20260821.pkl")
