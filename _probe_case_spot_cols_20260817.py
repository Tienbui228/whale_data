"""
_probe_case_spot_cols_20260817.py
Check DAILY sheet structure for AIN, GUA, APR — what columns exist, how many rows,
whether col 31/32/33 (vol_spot_usd / taker_spot_pct / spot_perp_ratio) are present.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

CASES = Path(r"D:\Data back test\whale_data\docs\case_files")

for sym in ["AIN", "GUA", "APR"]:
    path = CASES / f"CASE_{sym}.xlsx"
    if not path.exists():
        print(f"{sym}: CASE file MISSING")
        continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "DAILY" not in wb.sheetnames:
        print(f"{sym}: no DAILY sheet — sheets={wb.sheetnames}")
        wb.close()
        continue
    ws = wb["DAILY"]
    # Read header row
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    print(f"\n{sym} DAILY:")
    print(f"  Rows: {ws.max_row-1} data rows (max_row={ws.max_row})")
    print(f"  Total cols: {len(headers)}")
    # Print all headers with index
    for i, h in enumerate(headers, 1):
        print(f"    col{i:2d}: {h}")
    # Check col 31-33 specifically
    cols_of_interest = {
        "vol_spot_usd": None,
        "taker_spot_pct": None,
        "spot_perp_ratio": None,
    }
    for i, h in enumerate(headers, 1):
        if h in cols_of_interest:
            cols_of_interest[h] = i
            print(f"  *** {h} found at col {i}")
    for name, idx in cols_of_interest.items():
        if idx is None:
            print(f"  --- {name}: NOT FOUND (will need to add at col {len(headers)+1}+)")
    # Print last 2 data rows
    rows = list(ws.iter_rows(min_row=ws.max_row-1, values_only=True))
    print(f"  Last rows (sample):")
    for r in rows:
        print(f"    {list(r[:5])} ... {list(r[-5:])}")
    wb.close()
