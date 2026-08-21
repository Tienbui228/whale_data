"""Evidence paste R nhịp 1: sample 3 rows quanh key dates + R1+ BANK verify."""
import sys, io, pickle
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

CASES = Path(r"D:\Data back test\whale_data\docs\case_files")
SCRATCH = Path(r"D:\Data back test\whale_data\scratch")

def paste_around(sym, target_dates):
    fp = CASES / f"CASE_{sym}.xlsx"
    wb = openpyxl.load_workbook(fp, data_only=True)
    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}

    print(f"\n[{sym}] sample rows quanh {target_dates}:")
    for target in target_dates:
        # Find target row
        for r in range(3, ws.max_row + 1):
            d = ws.cell(r, 1).value
            if d and str(d)[:10] == target:
                # Print target ± 1
                for rr in range(max(3, r-1), min(ws.max_row+1, r+2)):
                    dv = ws.cell(rr, 1).value
                    # Grab common fields
                    fields = {}
                    for k in ["c","c (v2)","px_avg","px_avg (H+L+C)/3","h","l",
                              "vol_perp_usd","oi_usd","pos_r","lsr_pos_r",
                              "acc_r","lsr_acc_r","gLSR","lsr_global_r","f_avg","taker_buy_pct","taker_buy_pct (k[10]/k[7])"]:
                        if k in ci:
                            v = ws.cell(rr, ci[k]).value
                            if v is not None and k not in fields:
                                fields[k] = v
                    # Also col 45 for v1 (BANK/TUT c after R4)
                    if "c" not in fields and len(row2) >= 45 and row2[44] == "c":
                        v = ws.cell(rr, 45).value
                        if v is not None: fields["c(col45)"] = v
                    s = " · ".join(f"{k}={v}" for k, v in list(fields.items())[:8])
                    print(f"    row{rr} ({dv}): {s}")
                break

paste_around("BANK", ["2026-07-10"])
paste_around("TUT", ["2026-07-12"])
paste_around("AKE", ["2026-07-14"])
paste_around("BTW", ["2026-06-20"])
paste_around("CYS", ["2026-08-07"])

# R1+ BANK verify — fresh fetch vs stored
print("\n" + "="*70)
print("R1+ BANK 17-21/07 fresh k[7] vs stored vol_perp_usd")
print("="*70)
with open(SCRATCH / "r_nhip1_fetch_20260821.pkl", "rb") as f:
    ALL = pickle.load(f)
bank_klines = ALL["BANK"]["klines"]

wb = openpyxl.load_workbook(CASES / "CASE_BANK.xlsx", data_only=True)
ws = wb["DAILY"]
row2 = [c.value for c in ws[2]]
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
vp_c = ci.get("vol_perp_usd")

print(f"  {'date':<12} {'stored vol_perp':>18} {'fresh k[7] fapi':>18} {'match?':>7}")
for target in ["2026-07-17","2026-07-18","2026-07-19","2026-07-20","2026-07-21"]:
    stored = None
    for r in range(3, ws.max_row+1):
        d = ws.cell(r, 1).value
        if d and str(d)[:10] == target:
            stored = ws.cell(r, vp_c).value
            break
    fresh = bank_klines.get(target, {}).get("vol_quote_usd")
    match = "SAME" if (stored is not None and fresh is not None and abs(stored-fresh) < 1) else "LỆCH"
    print(f"  {target:<12} {str(stored):>18} {str(fresh):>18} {match:>7}")
