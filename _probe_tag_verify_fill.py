#!/usr/bin/env python3
"""Verify CASE_TAG DAILY spot/alpha fill."""
import io, sys
from pathlib import Path
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_TAG.xlsx")
wb = openpyxl.load_workbook(XLSX)
ws = wb["DAILY"]

print("AA=spot_venue  AB=vol_spot_usd  AC=taker_spot_pct  AJ=alpha_suspect")
print(f"{'row':>4} {'date':12} {'AA':15} {'AB(USD)':>12} {'AC':8} {'AJ_formula':30} {'D_perp_vol':>14}")
print("-"*100)
for r in range(3, ws.max_row + 1):
    date_v = ws.cell(r, 1).value
    if not date_v:
        continue
    aa = ws.cell(r, 27).value
    ab = ws.cell(r, 28).value
    ac = ws.cell(r, 29).value
    aj = ws.cell(r, 36).value
    d  = ws.cell(r, 4).value
    ab_s = f"{ab:,.0f}" if isinstance(ab, (int, float)) else str(ab or "—")
    d_s  = f"{d:,.0f}" if isinstance(d, (int, float)) else str(d or "—")
    print(f"{r:>4} {str(date_v):12} {str(aa or '—'):15} {ab_s:>12} {str(ac or '—'):8} {str(aj or '—')[:30]:30} {d_s:>14}")
