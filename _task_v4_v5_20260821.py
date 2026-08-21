"""V4 + V5: R7 chi tiết BANK null LSR + first_nn OI thực + AKE OI 12-21/07 · schema_version meta."""
import sys, io, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
CASES = ROOT / "docs" / "case_files"
SCRATCH = ROOT / "scratch"

def safe_float(v):
    if v in (None, "", "None"): return None
    try: return float(v)
    except: return None

# Backup pre-V4
for sym in ["BANK","TUT","BTW","CYS","AKE"]:
    src = CASES / f"CASE_{sym}.xlsx"
    dst = SCRATCH / f"CASE_{sym}_bak_pre_v4_20260821.xlsx"
    if src.exists(): shutil.copy2(src, dst)
print("Backups pre-V4 done")

# ─── V4: BANK null LSR list + first_nn OI ───
print("\n" + "="*70)
print("V4 — R7 detail")
print("="*70)

results = {}
for sym in ["BANK","TUT","BTW","CYS","AKE"]:
    fp = CASES / f"CASE_{sym}.xlsx"
    if not fp.exists():
        results[sym] = None; continue
    wb = openpyxl.load_workbook(fp, data_only=True)
    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
    is_v2 = "c" in ci
    pos_col = ci.get("lsr_pos_r") or ci.get("pos_r")
    oi_col = ci.get("oi_usd")

    # Collect all rows
    lsr_null_dates = []
    oi_dates_nonnull = []
    for r in range(3, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not d: continue
        d_str = str(d)[:10]
        pos = safe_float(ws.cell(r, pos_col).value) if pos_col else None
        oi = safe_float(ws.cell(r, oi_col).value) if oi_col else None
        if pos is None:
            lsr_null_dates.append(d_str)
        if oi is not None:
            oi_dates_nonnull.append(d_str)

    results[sym] = {
        "lsr_null_dates": sorted(lsr_null_dates),
        "oi_first_nn": min(oi_dates_nonnull) if oi_dates_nonnull else None,
        "oi_last_nn": max(oi_dates_nonnull) if oi_dates_nonnull else None,
        "oi_count": len(oi_dates_nonnull),
    }
    print(f"\n[{sym}]")
    print(f"  lsr_null_dates ({len(lsr_null_dates)}): {sorted(lsr_null_dates)[:20]}" +
          (" ..." if len(lsr_null_dates) > 20 else ""))
    print(f"  OI first_nn: {results[sym]['oi_first_nn']}, last_nn: {results[sym]['oi_last_nn']}, count: {results[sym]['oi_count']}")

# AKE store OI 12-21/07 kiểm — hiện tại DAILY của AKE có OI cho những ngày này không?
print(f"\n[AKE OI check 12-21/07]:")
fp = CASES / "CASE_AKE.xlsx"
wb = openpyxl.load_workbook(fp, data_only=True)
ws = wb["DAILY"]
row2 = [c.value for c in ws[2]]
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
oi_col = ci.get("oi_usd")
for r in range(3, ws.max_row+1):
    d = ws.cell(r, 1).value
    if not d: continue
    d_str = str(d)[:10]
    if "2026-07-12" <= d_str <= "2026-07-21":
        v = ws.cell(r, oi_col).value
        print(f"  {d_str}: OI = {v}")

# ─── V5: schema_version meta ───
print("\n" + "="*70)
print("V5 — schema_version + consumer header note vào CASE meta")
print("="*70)

for sym in ["BANK","TUT","BTW","CYS","AKE"]:
    fp = CASES / f"CASE_{sym}.xlsx"
    if not fp.exists(): continue
    wb = openpyxl.load_workbook(fp)
    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    n_cols = len(row2)
    is_v2 = "c" in {str(n).strip() for n in row2 if n}

    if sym in ("BANK","TUT"):
        version = f"v1+ext-20260821 (v1 cols 1-41 + R4 append cols 42-52 · o/h/l/c + vol_alpha CHƯA-TRA + cols 51-52 CHƯA-TRA)"
    elif sym in ("BTW","CYS"):
        version = f"v2+ext-20260821 (v2 cols 1-50 + R4 append cols 51-52 CHƯA-TRA)"
    elif sym == "AKE":
        version = f"v2.1 (52 cols, cold-start from CASE_PUMPBTC template — cols 51-52 CHƯA-TRA nhịp 2)"

    ws_case = wb["CASE"]
    # Append meta row for schema_version
    r_next = ws_case.max_row + 2
    ws_case.cell(r_next, 1).value = "META (V5 21/08)"
    ws_case.cell(r_next+1, 1).value = "schema_version"
    ws_case.cell(r_next+1, 2).value = version
    ws_case.cell(r_next+2, 1).value = "consumer_rule"
    ws_case.cell(r_next+2, 2).value = "Consumer đọc theo HEADER row 2 field names (không hardcode col index). Cols 42+ append-only per L-57."

    # V4 R7 update: append LSR null dates list
    r_next2 = r_next + 4
    lsr_nulls = results[sym]["lsr_null_dates"] if results[sym] else []
    ws_case.cell(r_next2, 1).value = "LAYER COVERAGE UPDATE (V4 21/08)"
    ws_case.cell(r_next2+1, 1).value = f"lsr_null_dates ({len(lsr_nulls)})"
    ws_case.cell(r_next2+1, 2).value = ", ".join(lsr_nulls) if lsr_nulls else "(0)"
    ws_case.cell(r_next2+2, 1).value = "OI first_nn (real)"
    ws_case.cell(r_next2+2, 2).value = str(results[sym]["oi_first_nn"]) if results[sym] else "N/A"
    ws_case.cell(r_next2+3, 1).value = "OI last_nn"
    ws_case.cell(r_next2+3, 2).value = str(results[sym]["oi_last_nn"]) if results[sym] else "N/A"

    wb.save(fp)
    print(f"  [{sym}] schema_version + V4 R7 update appended · saved")

print("\n[DONE V4+V5]")
