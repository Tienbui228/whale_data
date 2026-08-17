"""
_probe_spot_write_20260817.py
S-5: Write vol_spot_usd (col31), taker_spot_pct (col32), spot_perp_ratio (col33)
into CASE_AIN, CASE_GUA, CASE_APR DAILY sheets.

Sources:
  AIN: scratch/spot_klines_AIN_20260817.csv (binance_alpha, taker=ABSTAIN)
  GUA: scratch/spot_klines_GUA_20260817.csv (binance_alpha, taker=ABSTAIN)
  APR: scratch/spot_klines_APR_20260817.csv (kucoin, taker=ABSTAIN, wash_flag=True)

DAILY schema (row 2):
  col 1: date, col 13: oi_usd, col 31: vol_spot_usd, col 32: taker_spot_pct, col 33: spot_perp_ratio
Data starts row 3.
"""
import sys, io, csv
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT    = Path(r"D:\Data back test\whale_data")
CASES   = ROOT / "docs" / "case_files"
SCRATCH = ROOT / "scratch"

# Column positions (1-based)
COL_DATE             = 1
COL_OI_USD           = 13
COL_VOL_SPOT_USD     = 31
COL_TAKER_SPOT_PCT   = 32
COL_SPOT_PERP_RATIO  = 33
DATA_START_ROW       = 3   # row 1=section, row 2=field names, row 3+=data

TOKENS = {
    "AIN": {
        "csv":         SCRATCH / "spot_klines_AIN_20260817.csv",
        "venue":       "binance_alpha",
        "wash_flag":   False,
        "wash_ratio":  31.3,
        "taker_note":  "ABSTAIN (Alpha endpoint klines[9]==klines[5])",
    },
    "GUA": {
        "csv":         SCRATCH / "spot_klines_GUA_20260817.csv",
        "venue":       "binance_alpha",
        "wash_flag":   False,
        "wash_ratio":  196.0,
        "taker_note":  "ABSTAIN (Alpha endpoint klines[9]==klines[5])",
    },
    "APR": {
        "csv":         SCRATCH / "spot_klines_APR_20260817.csv",
        "venue":       "kucoin",
        "wash_flag":   True,
        "wash_ratio":  98.4,
        "taker_note":  "ABSTAIN (KuCoin klines no taker field)",
    },
}

def safe_float(v):
    try:
        f = float(v)
        return f if f == f else None  # NaN check
    except (TypeError, ValueError):
        return None

def load_spot_csv(path):
    """Returns dict: date_str → {quote_vol: float, ...}"""
    if not path.exists():
        return {}
    data = {}
    for row in csv.DictReader(open(path, encoding="utf-8")):
        dt = row.get("date", "")
        if not dt:
            continue
        qv = safe_float(row.get("quote_vol"))
        data[dt] = {"quote_vol": qv, "venue": row.get("venue", ""), "close": row.get("close")}
    return data

results = {}

for sym, meta in TOKENS.items():
    case_path = CASES / f"CASE_{sym}.xlsx"
    csv_path  = meta["csv"]

    print(f"\n{'='*60}")
    print(f"{sym} — venue={meta['venue']} wash_flag={meta['wash_flag']}")
    print(f"{'='*60}")

    if not case_path.exists():
        print(f"  CASE file MISSING: {case_path}")
        continue
    if not csv_path.exists():
        print(f"  Spot CSV MISSING: {csv_path}")
        continue

    spot_data = load_spot_csv(csv_path)
    print(f"  Spot CSV: {len(spot_data)} dates, "
          f"{min(spot_data) if spot_data else 'N/A'} → {max(spot_data) if spot_data else 'N/A'}")

    wb = openpyxl.load_workbook(case_path)
    ws = wb["DAILY"]

    # Verify header structure
    h_date  = ws.cell(2, COL_DATE).value
    h_oi    = ws.cell(2, COL_OI_USD).value
    h_vol   = ws.cell(2, COL_VOL_SPOT_USD).value
    h_taker = ws.cell(2, COL_TAKER_SPOT_PCT).value
    h_ratio = ws.cell(2, COL_SPOT_PERP_RATIO).value
    print(f"  Header check: col1={h_date}, col13={h_oi}, col31={h_vol}, col32={h_taker}, col33={h_ratio}")

    written = 0; skipped_no_spot = 0; skipped_no_oi = 0
    oi_null_dates = []
    ratio_computed = 0

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        dt = ws.cell(row_idx, COL_DATE).value
        if dt is None:
            continue
        dt_str = str(dt)[:10]  # handle date objects

        if dt_str not in spot_data:
            skipped_no_spot += 1
            continue

        spot = spot_data[dt_str]
        qv   = spot["quote_vol"]

        if qv is None:
            skipped_no_spot += 1
            continue

        # vol_spot_usd
        ws.cell(row_idx, COL_VOL_SPOT_USD).value = round(qv, 2)

        # taker_spot_pct — always ABSTAIN (NaN) for Alpha and KuCoin
        ws.cell(row_idx, COL_TAKER_SPOT_PCT).value = None

        # spot_perp_ratio = vol_spot_usd / oi_usd
        oi_usd = safe_float(ws.cell(row_idx, COL_OI_USD).value)
        if oi_usd and oi_usd > 0:
            ratio = round(qv / oi_usd, 6)
            ws.cell(row_idx, COL_SPOT_PERP_RATIO).value = ratio
            ratio_computed += 1
        else:
            ws.cell(row_idx, COL_SPOT_PERP_RATIO).value = None
            oi_null_dates.append(dt_str)
            skipped_no_oi += 1

        written += 1

    print(f"  Rows written: {written}")
    print(f"  Rows skipped (no spot date match): {skipped_no_spot}")
    print(f"  ratio computed: {ratio_computed}, oi_null: {skipped_no_oi}")
    if oi_null_dates[:5]:
        print(f"  oi_null sample: {oi_null_dates[:5]}")

    # Sample: print last 5 written rows
    print(f"  Sample last 5 rows:")
    count = 0
    for row_idx in range(ws.max_row, DATA_START_ROW - 1, -1):
        dt = ws.cell(row_idx, COL_DATE).value
        if dt is None:
            continue
        dt_str = str(dt)[:10]
        if dt_str not in spot_data:
            continue
        vol   = ws.cell(row_idx, COL_VOL_SPOT_USD).value
        taker = ws.cell(row_idx, COL_TAKER_SPOT_PCT).value
        ratio = ws.cell(row_idx, COL_SPOT_PERP_RATIO).value
        oi    = ws.cell(row_idx, COL_OI_USD).value
        vol_s   = f"{vol:,.0f}"   if vol   is not None else "N/A"
        ratio_s = f"{ratio:.6f}" if ratio is not None else "N/A"
        print(f"    {dt_str}: vol_spot={vol_s} taker={taker} ratio={ratio_s} oi_usd={oi}")
        count += 1
        if count >= 5:
            break

    # Try saving; fall back to scratch if file locked by Excel
    try:
        wb.save(case_path)
        print(f"  SAVED: {case_path.name}")
    except PermissionError:
        tmp = SCRATCH / f"CASE_{sym}_spot_patch_20260817.xlsx"
        wb.save(tmp)
        print(f"  PermissionError — saved to TEMP: {tmp.name}")
        print(f"  Action needed: close CASE_{sym}.xlsx in Excel, then copy {tmp.name} → {case_path.name}")

    results[sym] = {
        "written": written,
        "ratio_computed": ratio_computed,
        "oi_null": skipped_no_oi,
        "venue": meta["venue"],
        "wash_flag": meta["wash_flag"],
    }

print("\n" + "="*60)
print("SUMMARY")
print("="*60)
for sym, r in results.items():
    print(f"  {sym}: {r['written']} rows written, ratio_ok={r['ratio_computed']}, "
          f"oi_null={r['oi_null']}, venue={r['venue']}, wash_flag={r['wash_flag']}")
print("\n[DONE S-5]")
