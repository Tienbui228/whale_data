"""_task_stbl_create_case_20260819.py
Create docs/case_files/CASE_STBL.xlsx từ template CASE_TRIA.xlsx (v2.1 schema — đã có cols 51/52).
Populate: CASE meta + DAILY (Jun19-Aug18, 61 rows) + EVENTS + UNLOCK.
- vol_alpha, vol_spot_ext_usd (=Gate turnover), spot_ext_venue=gate primary
- vol_spot_usd = N/A-STRUCTURAL (no api/v3 pair)
- taker_alpha = ABSTAIN-STRUCTURAL (Alpha MIRROR L-48 confirmed 31/31)
- Note: contract UNVERIFIED (Alpha map symbol-string only, bài F1 GUA)
"""
import sys, io, shutil, csv, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone, date as _date, timedelta
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
SCRATCH = ROOT / "scratch"
CASES = ROOT / "docs" / "case_files"

TEMPLATE = CASES / "CASE_TRIA.xlsx"  # v2.1 schema (52 cols, already has spot-ext amendment)
TARGET   = CASES / "CASE_STBL.xlsx"

SYM = "STBL"
TICKER = "STBLUSDT"
ALPHA_ID = "ALPHA_376"  # UNVERIFIED symbol-match
CONTRACT_BAPI = "0x8dedf84656fa932157e27c060d8613824e7979e3"  # BSC, UNVERIFIED per L-51
START = "2026-06-19"
END = "2026-08-18"

def safe_float(v):
    if v in (None, "", "None"): return None
    try: return float(v)
    except: return None

# ── 1. Copy template ──
assert not TARGET.exists(), f"CASE_STBL.xlsx already exists — abort"
assert TEMPLATE.exists(), f"template missing: {TEMPLATE}"
shutil.copy2(TEMPLATE, TARGET)
print(f"COPIED template: {TEMPLATE.name} → {TARGET.name}")

wb = openpyxl.load_workbook(TARGET)
print(f"Sheets: {wb.sheetnames}")

# ── 2. Clear DAILY data rows ──
ws = wb["DAILY"]
row2 = [c.value for c in ws[2]]
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
print(f"DAILY cols: {len(row2)}")
print(f"  vol_alpha col: {ci.get('vol_alpha')}")
print(f"  vol_spot_usd col: {ci.get('vol_spot_usd')}")
print(f"  vol_spot_ext_usd col: {ci.get('vol_spot_ext_usd')}")
print(f"  spot_ext_venue col: {ci.get('spot_ext_venue')}")

old_last = ws.max_row
if old_last >= 3:
    ws.delete_rows(3, old_last - 2)
print(f"Cleared rows 3-{old_last}")

# ── 3. Load scratch data ──
def load_csv(p):
    if not p.exists(): return []
    with open(p, encoding="utf-8") as f: return list(csv.DictReader(f))

perp_rows = load_csv(SCRATCH / "stbl_perp_60d_20260819.csv")
perp_by_d = {r["date"]: r for r in perp_rows}
print(f"Perp scratch: {len(perp_by_d)} rows")

alpha_rows = load_csv(SCRATCH / "stbl_alpha_bapi_20260819.csv")
alpha_by_d = {r["date"]: r for r in alpha_rows}
print(f"Alpha scratch: {len(alpha_by_d)} rows")

gate_rows = load_csv(SCRATCH / "stbl_gate_spot_20260819.csv")
gate_by_d = {r["date"]: r for r in gate_rows}
print(f"Gate scratch: {len(gate_by_d)} rows")

kucoin_rows = load_csv(SCRATCH / "stbl_kucoin_20260819.csv")
kucoin_by_d = {r["date"]: r for r in kucoin_rows}
print(f"KuCoin scratch: {len(kucoin_by_d)} rows")

mexc_rows = load_csv(SCRATCH / "stbl_mexc_20260819.csv")
mexc_by_d = {r["date"]: r for r in mexc_rows}
print(f"MEXC scratch: {len(mexc_by_d)} rows")

# ── 4. Re-fetch funding+OI+LSR (didn't save these) ──
import urllib.request, time as _time
def fetch_json(url, params=None, timeout=15):
    if params:
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        url = f"{url}?{qs}"
    req = urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode())
    except Exception as e:
        return {"_error": str(e)}
def to_ms(d): return int(datetime.strptime(d,"%Y-%m-%d").replace(tzinfo=timezone.utc).timestamp()*1000)
def ms_to_date(ms): return datetime.fromtimestamp(int(ms)/1000, tz=timezone.utc).strftime("%Y-%m-%d")

print("\n=== Re-fetch funding + OI + LSR ===")
s_ms = to_ms(START); e_ms = to_ms(END) + 86400000 - 1
# funding (paginate)
fund_all = []; cur = s_ms
while cur <= e_ms:
    d = fetch_json("https://fapi.binance.com/fapi/v1/fundingRate",
                   {"symbol": TICKER, "startTime": cur, "endTime": e_ms, "limit": 1000})
    _time.sleep(0.3)
    if not isinstance(d, list) or not d: break
    fund_all.extend(d)
    last = int(d[-1]["fundingTime"])
    if last <= cur or last >= e_ms: break
    cur = last + 1
fund_by_d = {}
for it in fund_all:
    ts = int(it["fundingTime"]); d_str = ms_to_date(ts)
    if START <= d_str <= END:
        fund_by_d.setdefault(d_str, []).append(float(it["fundingRate"]))
print(f"  funding: {len(fund_by_d)} days")

# OI 30d
d = fetch_json("https://fapi.binance.com/futures/data/openInterestHist",
               {"symbol": TICKER, "period":"1d", "limit": 30})
_time.sleep(0.3)
oi_by_d = {}
for it in d if isinstance(d, list) else []:
    ts = int(it["timestamp"]); d_str = ms_to_date(ts)
    oi_by_d[d_str] = {"oi_usd": float(it["sumOpenInterestValue"] or 0),
                       "oi_token": float(it["sumOpenInterest"] or 0)}
print(f"  OI: {len(oi_by_d)}")

# LSR freshwatch
EXPORT = ROOT / "data" / "export"
lsr_fw = {}
for fp in sorted(EXPORT.glob("freshwatch_daily_2026-*.csv")):
    with open(fp, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("symbol","").upper() != TICKER: continue
            d_str = row.get("date","")[:10]
            pr = safe_float(row.get("lsr_pos_r"))
            if not d_str or not pr: continue
            lsr_fw[d_str] = {
                "pos_r": pr,
                "acc_r": safe_float(row.get("lsr_acc_r")),
                "gLSR":  safe_float(row.get("lsr_global_r")),
                "taker_lsr": safe_float(row.get("taker_ratio")),
                "src": "freshwatch_store",
            }
print(f"  freshwatch LSR: {len(lsr_fw)}")

# fapi LSR fallback
def fetch_ep(ep):
    d = fetch_json(f"https://fapi.binance.com/futures/data/{ep}",
                   {"symbol": TICKER, "period":"1d", "limit": 30})
    _time.sleep(0.3)
    by_d = {}
    for it in d if isinstance(d, list) else []:
        ts = int(it["timestamp"]); by_d[ms_to_date(ts)] = float(it.get("longShortRatio",0) or 0)
    return by_d
lsr_top_pos = fetch_ep("topLongShortPositionRatio")
lsr_top_acc = fetch_ep("topLongShortAccountRatio")
lsr_glo     = fetch_ep("globalLongShortAccountRatio")

# ── 5. Build 61 rows Jun19-Aug18 ──
d0 = _date.fromisoformat(START); d1 = _date.fromisoformat(END)
days = []
cur = d0
while cur <= d1:
    days.append(cur.isoformat()); cur += timedelta(days=1)

rows_data = []
prev_c = None
for d_str in days:
    p = perp_by_d.get(d_str)
    if not p: continue
    o = safe_float(p["o"]); h = safe_float(p["h"]); l = safe_float(p["l"]); c = safe_float(p["c"])
    frs = fund_by_d.get(d_str, [])
    fn = len(frs)
    f_avg = round(sum(frs)/fn, 8) if fn else None
    oi = oi_by_d.get(d_str, {})
    fw = lsr_fw.get(d_str)
    if fw:
        pos_r, acc_r, gl, tk_lsr, src = fw["pos_r"], fw["acc_r"], fw["gLSR"], fw["taker_lsr"], fw["src"]
    else:
        pos_r = lsr_top_pos.get(d_str); acc_r = lsr_top_acc.get(d_str); gl = lsr_glo.get(d_str)
        tk_lsr = None
        src = "fapi_LSR_30d" if pos_r else None

    ret_1d = round((c - prev_c)/prev_c*100, 4) if (c and prev_c) else None
    prev_c = c if c else prev_c

    # Spot venues
    a = alpha_by_d.get(d_str)
    g = gate_by_d.get(d_str)
    row = {
        "date": d_str,
        "px_avg (H+L+C)/3": round((h+l+c)/3, 8),
        "o": o, "h": h, "l": l, "c": c,
        "vol_perp_usd": safe_float(p["vol_quote_usd"]),
        "taker_buy_pct (k[10]/k[7])": safe_float(p["taker_buy_pct"]),
        "f_avg": f_avg, "f_min": min(frs) if frs else None, "f_max": max(frs) if frs else None,
        "fn (4h=6/day)": fn,
        "oi_usd": oi.get("oi_usd"), "oi_token": oi.get("oi_token"),
        "data_ok (fn>=6)": (fn >= 6), "low_conf (fn<6)": (fn < 6),
        "ret_1d (close-to-close)": ret_1d,
        "pos_r": pos_r, "acc_r": acc_r, "gLSR": gl, "taker_lsr": tk_lsr,
        # Structural tags (spot L-51 v2.1)
        "vol_spot_usd": "N/A-STRUCTURAL",
        "taker_spot_pct": "ABSTAIN-STRUCTURAL",
        "vol_alpha": safe_float(a["quote_vol_usd"]) if a else None,
        "taker_alpha": "ABSTAIN-STRUCTURAL",
        "vol_spot_ext_usd": safe_float(g.get("vol_quote_usd")) if g else None,
        "spot_ext_venue": "gate",
        "data_source_perp": "fapi_klines_1d",
        "data_source_lsr":  src or "MISSING",
    }
    rows_data.append(row)

print(f"\n  built {len(rows_data)} rows")

# derived cols (need arrays)
pos_arr = [r["pos_r"] for r in rows_data]
acc_arr = [r["acc_r"] for r in rows_data]
gl_arr  = [r["gLSR"]  for r in rows_data]
for i, r in enumerate(rows_data):
    p, a, g = pos_arr[i], acc_arr[i], gl_arr[i]
    r["div_raw (pos_r/acc_r)"] = round(p/a, 4) if (p and a) else None
    if i >= 2:
        p3 = [v for v in pos_arr[i-2:i+1] if v]
        a3 = [v for v in acc_arr[i-2:i+1] if v]
        if len(p3)==3 and len(a3)==3:
            aa = sum(a3)/3
            r["div_avg3 (avg3p/avg3a)"] = round((sum(p3)/3)/aa, 4) if aa else None
    if i >= 5 and gl_arr[i-5]:
        r["glsr_slope_5d (ratio t/t-5)"] = round(g/gl_arr[i-5], 4) if g else None
    if i >= 3:
        pp, ap = pos_arr[i-3], acc_arr[i-3]
        r["d_pos_3d (%chg)"] = round((p-pp)/pp*100,4) if (p and pp) else None
        r["d_acc_3d (%chg)"] = round((a-ap)/ap*100,4) if (a and ap) else None

# ── 6. Write DAILY ──
for ridx, r in enumerate(rows_data, start=3):
    for name, ciix in ci.items():
        v = r.get(name)
        if v is not None:
            ws.cell(ridx, ciix).value = v
print(f"\n  Wrote {len(rows_data)} rows (row 3-{2+len(rows_data)})")

# ── 7. CASE meta ──
ws_case = wb["CASE"]
for r in range(1, ws_case.max_row+1):
    for c in range(1, 10):
        ws_case.cell(r, c).value = None

case_meta = [
    ["CASE FILE — STBL", None],
    ["Nguồn: probe _stbl_probe_20260819 · verdict chỉ operator ghi · số kèm nguồn · PENDING không đoán", None],
    [None, None],
    ["Status", "TRACKING (cold-start 19/08 — chờ operator §S4 relay contract + on-chain)"],
    ["Loại", "Alpha + Perp + multi-venue spot (Gate, KuCoin, MEXC) · hedge episode ongoing (start ~2026-06-27 per watchlist)"],
    ["Perp fapi", "STBLUSDT TRADING · onboardDate=2025-09-17 (11 tháng history)"],
    ["Contract (BAPI symbol-match, ⚠ UNVERIFIED)", f"{CONTRACT_BAPI} on BSC (Alpha bapi symbol-match, chờ operator confirm — bài F1 GUA)"],
    ["Alpha ID", f"{ALPHA_ID} · name=STBL · chain=BSC · fullyDelisted=False · offline=False · MIRROR TAKER L-48 (k[9]≡k[5], k[10]≡k[7] 31/31)"],
    ["Spot venues", "Gate STBL_USDT (BSC dep+wd OPEN ✓) · KuCoin STBL-USDT (enabled) · MEXC STBLUSDT (Trading) · Alpha ALPHA_376 (taker broken)"],
    ["Spot api/v3 Binance", "KHÔNG có (Alpha + Gate/KuCoin/MEXC only) → vol_spot_usd = N/A-STRUCTURAL, taker_spot_pct = ABSTAIN-STRUCTURAL"],
    ["Ngày phát hiện", "2026-08-19 (probe ANALYSIS approved)"],
    ["Watchlist v1.3", "TRIA-style episode: hedge_start=2026-06-27, struct_n=49 @ Aug17 (dài, ổn định)"],
    [None, None],
    ["A. HEDGE EPISODE (đọc watchlist_log Aug17)", None],
    ["hedge_start", "2026-06-27"],
    ["struct_n Aug17", "49 days (episode dài, div_avg3 ≤ 0.5 liên tục)"],
    ["div_avg3 window 20/07→18/08", "min 0.236 (Aug11) · max 0.336 (Jul25) · Aug18 = 0.306 (đang leo từ min)"],
    ["hedge_start_px", "~$0.02"],
    ["px_dist% Aug18", "≈+30% từ ~$0.02 hedge start (0.02603 Aug18 vs ~$0.02)"],
    [None, None],
    ["B. FUNDING PROFILE (30d cửa sổ)", None],
    ["Pin streak 20/07→18/08", "30 days (f_avg = +0.00005 mọi ngày, band [4e-5, 6e-5])"],
    ["Un-pin event", "CHƯA có (khuôn TRIA — captive committed cực chặt)"],
    [None, None],
    ["C. UN-HEDGE PROGRESSION 15-19/08", None],
    ["pos_r trend Aug14-18", "0.917 → 0.933 → 0.933 → 0.978 → 0.973 (cross-1 imminent, chưa vượt)"],
    ["acc_r trend Aug16-18", "3.888 → 2.997 → 2.535 (giảm -35% trong 2d — top-trader hedge unwind)"],
    ["div_raw trend Aug16-18", "0.240 → 0.326 → 0.384 (jump lên đỉnh window)"],
    ["Aug18 verdict", "UN-HEDGE-EARLY signal (acc_r rơi + div jump), OI vẫn peak $2.06M — chưa cover thật, có thể re-hedge hoặc S4"],
    [None, None],
    ["D. OI CYCLE + TEXTURE", None],
    ["OI 30d min", "$1,271,717 (2026-07-21)"],
    ["OI 30d max", "$2,062,661 (2026-08-18)"],
    ["oi_cycle_pos hiện tại", "1.0000 (AT PEAK BUILD, sensitive đến cover)"],
    ["Texture Aug16-18", "S3-BIDWALK peak (Aug17 perp vol × 6, spot tổng × 2.9 nền) → S4-TRANSITION ứng viên Aug18"],
    [None, None],
    ["E. CHÂN HÀNG (multi-venue, identity-check PASS)", None],
    ["Nền T7 perp", "$1,134,194/day (Jul20-Aug01)"],
    ["Nền T7 spot tổng", "$190,545/day (Gate $94K + MEXC $71K + KuCoin $19K + Alpha $6K)"],
    ["Peak day Aug17 spot tổng", "$553,712 (× 2.9 nền T7)"],
    ["Deposit/withdraw flags", "Gate BSC: OPEN ✓ (van THÔNG). KuCoin/MEXC: chưa fetch flags"],
    [None, None],
    ["F. IDENTITY-CHECK L-51 (3 mẫu/venue)", None],
    ["Gate", "bit-eq=False (3/3) · implied avg ≈ close ✓ PASS"],
    ["KuCoin", "bit-eq=False (3/3) · implied avg ≈ close ✓ PASS"],
    ["MEXC", "bit-eq=False (3/3) · implied avg ≈ close ✓ PASS"],
    ["Alpha ALPHA_376", "MIRROR 31/31 (L-48 k[9]≡k[5], k[10]≡k[7]) → taker=ABSTAIN-STRUCTURAL"],
    [None, None],
    ["G. FALSIFIERS / TRIGGERS TO WATCH", None],
    ["Un-hedge THẬT", "pos_r vượt 1.0 + OI giảm + div_raw > 0.5 + acc_r tiếp rơi"],
    ["Re-hedge (không có xả)", "acc_r bò lại + div_raw quay về 0.25 + OI giữ peak"],
    ["S4 xác nhận", "OI rơi >10% từ peak $2.06M + gLSR về ~1 + funding un-pin (âm)"],
    [None, None],
    ["H. §S4 CẦN OPERATOR (chưa có, mọi nhãn treo [OWNERSHIP-UNVERIFIED])", None],
    ["Contract 0x + chain", "PENDING (Alpha bapi trả 0x8ded...BSC — cần confirm khớp, bài F1 GUA)"],
    ["Exchange balance trend 30d", "PENDING"],
    ["Whale movements top100", "PENDING"],
    ["Circ supply + FDV + unlock", "PENDING"],
    ["Top10/50/100 concentration", "PENDING"],
    [None, None],
    ["I. NGUỒN + REPORT", None],
    ["Probe main", "_probe_stbl_20260819.py"],
    ["Probe supp", "_probe_stbl_venues_supp_20260819.py"],
    ["Report", "docs/reports/_stbl_probe_20260819.md"],
    ["Data provenance", "fapi/v1 (perp/funding/OI/LSR) + freshwatch_daily (fw-139 in-universe) + bapi/defi/v1 (Alpha, MIRROR) + api.gateio.ws + api.kucoin.com + api.mexc.com"],
]
for r_idx, (a, b) in enumerate(case_meta, start=1):
    if a is not None: ws_case.cell(r_idx, 1).value = a
    if b is not None: ws_case.cell(r_idx, 2).value = b
print(f"  CASE meta: {len(case_meta)} rows")

# ── 8. EVENTS ──
ws_ev = wb["EVENTS"]
for r in range(1, ws_ev.max_row+1):
    for c in range(1, 8):
        ws_ev.cell(r, c).value = None

ev_headers = ["date", "type", "description", "ref", "resolved"]
for c, h in enumerate(ev_headers, start=1):
    ws_ev.cell(1, c).value = h

events = [
    ["2025-09-17", "LISTING_PERP", "STBLUSDT onboard fapi (perp)", "fapi exchangeInfo", "N/A"],
    ["2026-06-27", "HEDGE_EPISODE_START", "Episode #1 start per watchlist_log (div_avg3 ≤ 0.5 first day, hedge_start_px~$0.02)", "watchlist_log_2026-08.csv", "N/A"],
    ["2026-07-22", "PUMP_D1 + VOL_2X", "+11.76% ret1d, vol 2.56M (2× nền)", "fapi klines", "N/A"],
    ["2026-07-27", "FLUSH_D1", "-10.23% single day", "fapi klines", "N/A"],
    ["2026-08-10", "FLUSH_D1", "-6.61% flush, OI leo 1.63→1.80M (short bổ sung + long hedge giữ)", "fapi klines+OI", "N/A"],
    ["2026-08-16", "PUMP_D1 + SPOT_IGNITE", "+10.63% ret1d + Alpha × 5 nền + Gate × 1.6 + MEXC × 1.1 (spot cùng ignite)", "fapi klines + multi-venue", "N/A"],
    ["2026-08-17", "VOL_5X_SPIKE + OI_BUILD", "Perp vol × 6 ($6.9M) + Alpha × 13.6 + Gate × 3.2 + OI +22% (1.60→1.96M) + pos_r 0.978 near-cross-1 + acc_r rơi 3.89→3.00 (top-trader hedge unwind start)", "fapi + venues + LSR", "N/A"],
    ["2026-08-18", "UN-HEDGE-EARLY / ACC_R_UNWIND", "acc_r 3.00→2.54 (-15%), div_raw jump 0.326→0.384 (peak window), OI vẫn ở peak $2.06M (cycle_pos=1.0), price -3.66% — un-hedge sơ khai (chưa cover thật)", "fapi + LSR", "unresolved"],
]
for r, ev in enumerate(events, start=2):
    for c, val in enumerate(ev, start=1):
        ws_ev.cell(r, c).value = val
print(f"  EVENTS: {len(events)} rows")

# ── 9. UNLOCK skeleton ──
ws_ul = wb["UNLOCK"]
for r in range(1, ws_ul.max_row+1):
    for c in range(1, 8):
        ws_ul.cell(r, c).value = None
ws_ul.cell(1,1).value = "date"
ws_ul.cell(1,2).value = "event_type"
ws_ul.cell(1,3).value = "amount_tokens"
ws_ul.cell(1,4).value = "notes"
ws_ul.cell(1,5).value = "source"
ws_ul.cell(2,1).value = "2026-08-19"
ws_ul.cell(2,2).value = "DATA-GAP"
ws_ul.cell(2,4).value = ("[DATA-GAP] chưa xác định (§S4 pending operator): "
                          "contract 0x confirmed (Alpha bapi trả 0x8dedf84656fa932157e27c060d8613824e7979e3 on BSC — UNVERIFIED symbol match), "
                          "total/circ supply, unlock schedule, exchange balance trend, whale top100, "
                          "MC tracker source. Không có contract → không tính Doctrine #10 đầy đủ.")
ws_ul.cell(2,5).value = "task create 19/08"

# ── 10. Clear TRIGGERS / OUTCOME / MM_PROFILE ──
for sname in ["TRIGGERS", "OUTCOME", "MM_PROFILE"]:
    if sname not in wb.sheetnames: continue
    wsx = wb[sname]
    for r in range(1, wsx.max_row+1):
        for c in range(1, 10):
            wsx.cell(r, c).value = None
    wsx.cell(1,1).value = f"{sname} (STBL) — PENDING operator to define"

# ── 11. Save ──
wb.save(TARGET)
print(f"\nSAVED: {TARGET} ({TARGET.stat().st_size} B)")

# ── 12. Update CASE_INDEX.xlsx ──
INDEX = CASES / "CASE_INDEX.xlsx"
wb_idx = openpyxl.load_workbook(INDEX)
ws_idx = wb_idx["INDEX"]
header_row = None
token_col = state_col = None
for r in range(1, 6):
    for c in range(1, 20):
        v = ws_idx.cell(r, c).value
        if v and str(v).strip().lower() in ("token","symbol"):
            header_row = r; token_col = c
        if v and str(v).strip().lower() in ("state","status","trạng thái"):
            state_col = c

existing = []
for r in range(header_row+1, ws_idx.max_row+1):
    t = ws_idx.cell(r, token_col).value
    if t: existing.append((r, str(t).strip().upper()))

if "STBL" not in [t for _,t in existing]:
    next_r = ws_idx.max_row + 1
    ws_idx.cell(next_r, token_col).value = "STBL"
    if state_col:
        ws_idx.cell(next_r, state_col).value = "TRACKING (cold-start Aug19, contract UNVERIFIED)"
    print(f"CASE_INDEX: ADDED STBL at row {next_r}")
    wb_idx.save(INDEX)
else:
    print("STBL already in CASE_INDEX")

print("\n[DONE CASE_STBL.xlsx created]")
