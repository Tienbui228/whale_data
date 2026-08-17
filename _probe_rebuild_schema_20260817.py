#!/usr/bin/env python3
"""
Rebuild COOKIE / BROCCOLIF3B / KOMA DAILY sheets.

Bad rows were written by case_daily with old EXPORT_COL_MAP (col4=vol, col5=taker*100,
col6=f_avg, ...) into new-schema files where those positions are h/l/c/vol_perp_usd.
Result: vol at col7 = garbage (funding_min ≈ -5e-4), taker at col8 = garbage (funding_max).

This probe:
  1. Detects bad rows: col7 (vol_perp_usd) < 1000 USD — clearly not real vol
  2. Clears those rows
  3. Re-fetches from freshwatch_daily export CSV
  4. Writes using NEW_SCHEMA_COL_MAP (correct positions)
  5. Saves, validates with validate_written_pack
  6. Prints taker values after rebuild for operator verification

Run: python _probe_rebuild_schema_20260817.py > scratch/rebuild_schema.log 2>&1
"""
import csv
import sys
import io
import math
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

sys.path.insert(0, str(Path(__file__).parent))
from scripts.case_guard import validate_written_pack

WHALE      = Path(r"D:\Data back test\whale_data")
CASE_DIR   = WHALE / "docs" / "case_files"
EXPORT_DIR = Path(r"D:\Check onchain fresh wallets\data\export")

FIRST_DATA_ROW = 3

# New-schema column map (correct positions after o/h/l/c insertion)
NEW_COL_MAP = {
    1:  ("date",          "str"),
    2:  ("px_avg",        "float"),
    # cols 3-6 = o/h/l/c — not in freshwatch_daily export, skip
    7:  ("vol_quote",     "float"),      # vol_perp_usd
    8:  ("taker_buy_pct", "float"),      # ×100 (fraction → percent)
    9:  ("funding_avg",   "float"),
    10: ("funding_min",   "float"),
    11: ("funding_max",   "float"),
    12: ("funding_n",     "int"),
    13: ("oi_usd",        "float"),
    15: ("data_ok",       "invert_bool"),
    27: ("lsr_pos_r",     "float"),
    28: ("lsr_acc_r",     "float"),
    29: ("lsr_global_r",  "float"),
}
TAKER_COL  = 8   # apply ×100 here
FORMULA_COLS_NEW = {17, 22, 23, 24, 25, 26, 30}

CASE_TARGETS = {
    "CASE_COOKIE.xlsx":      "COOKIEUSDT",
    "CASE_BROCCOLIF3B.xlsx": "BROCCOLIF3BUSDT",
    "CASE_KOMA.xlsx":        "KOMAUSDT",
}


def _parse_val(raw, parse_type):
    if raw is None or str(raw).strip() == "":
        return None
    raw = str(raw).strip()
    if parse_type == "str":
        return raw
    if parse_type == "float":
        try:
            v = float(raw); return None if math.isnan(v) else v
        except ValueError:
            return None
    if parse_type == "int":
        try:
            return int(float(raw))
        except ValueError:
            return None
    if parse_type == "invert_bool":
        return 1 if raw == "False" else None
    return None


def _new_formulas(n):
    return {
        17: f'=IFERROR(F{n}/F{n-1}-1,"")',
        22: f'=IFERROR(AA{n}/AB{n},"")',
        23: f'=IFERROR(AVERAGE(V{n-2}:V{n}),"")',
        24: f'=IFERROR(AC{n}/AC{n-5},"")',
        25: f'=IFERROR(AA{n}/AA{n-3}-1,"")',
        26: f'=IFERROR(AB{n}/AB{n-3}-1,"")',
        30: f'=IF(OR(AC{n}="",H{n}="",H{n}=0),"",AC{n}/H{n})',
    }


def _clear_row(ws, row_n):
    for c in range(1, (ws.max_column or 44) + 1):
        cell = ws.cell(row_n, c)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None
        cell.comment = None


def _load_export(months):
    """Load export rows indexed by (symbol, date)."""
    result = {}
    for ym in months:
        csv_path = EXPORT_DIR / f"freshwatch_daily_{ym}.csv"
        if not csv_path.exists():
            print(f"  [WARN] export CSV missing: {csv_path}")
            continue
        with open(csv_path, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                key = (row["symbol"], row["date"])
                result[key] = row
    return result


def _write_row(ws, row_n, csv_row):
    formulas = _new_formulas(row_n)
    max_col = ws.max_column or 44
    for c in range(1, max_col + 1):
        cell = ws.cell(row_n, c)
        if isinstance(cell, MergedCell):
            continue
        if c in FORMULA_COLS_NEW:
            cell.value = formulas.get(c)
            continue
        if c not in NEW_COL_MAP:
            continue
        field, parse_type = NEW_COL_MAP[c]
        raw = csv_row.get(field, "")
        cell.value = _parse_val(raw, parse_type)
        if c == TAKER_COL and isinstance(cell.value, (int, float)):
            cell.value = round(cell.value * 100, 6)


DIV = "=" * 70
print(__doc__)

# Load export data for Aug 2026
export = _load_export(["2026-08", "2026-07"])
print(f"[export] loaded {len(export)} (symbol, date) keys")

for fname, sym_usdt in CASE_TARGETS.items():
    print(f"\n{DIV}")
    print(f"FILE: {fname}  symbol={sym_usdt}")
    fpath = CASE_DIR / fname
    if not fpath.exists():
        print(f"  [ERROR] file not found: {fpath}")
        continue

    wb = openpyxl.load_workbook(str(fpath), data_only=False)
    ws = wb["DAILY"]

    # [1] Detect bad rows: col7 (vol_perp_usd) < 1000 USD
    bad_rows = {}   # {date_str: row_n}
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        date_v = ws.cell(r, 1).value
        if not date_v or not isinstance(date_v, str):
            continue
        vol_v = ws.cell(r, 7).value
        # Real vol for any active perp is >> $1000; garbage from old map is < $1
        if vol_v is None or (isinstance(vol_v, (int, float)) and abs(vol_v) < 1000):
            bad_rows[str(date_v)] = r

    print(f"  [detect] {len(bad_rows)} bad rows: {sorted(bad_rows.keys())}")

    # [2] Clear bad rows
    for date_str, row_n in bad_rows.items():
        _clear_row(ws, row_n)
    print(f"  [clear] {len(bad_rows)} rows cleared")

    # [3] Rewrite from export
    written = []
    missing_export = []
    for date_str, row_n in sorted(bad_rows.items()):
        key = (sym_usdt, date_str)
        csv_row = export.get(key)
        if csv_row is None:
            print(f"  [WARN] no export data for {sym_usdt} {date_str} — row stays blank")
            missing_export.append(date_str)
            continue
        _write_row(ws, row_n, csv_row)
        taker_v = ws.cell(row_n, TAKER_COL).value
        vol_v   = ws.cell(row_n, 7).value
        pos_v   = ws.cell(row_n, 27).value
        written.append((date_str, row_n, taker_v, vol_v, pos_v))

    print(f"  [write] {len(written)} rows written, {len(missing_export)} skipped (no export)")

    # [4] Save + validate
    wb.save(str(fpath))
    print(f"  [save] {fpath.name}")

    try:
        validate_written_pack(fpath)
        print(f"  [validate] PASS ✓")
    except RuntimeError as e:
        print(f"  [validate] FAIL:\n    {e}")

    # [5] Report rebuilt rows
    print(f"\n  Rebuilt rows (taker should be ~40-55% = percent scale):")
    print(f"  {'date':>12}  {'row':>5}  {'taker_col8_pct':>14}  {'vol_col7_usd':>14}  {'pos_r_col27':>12}")
    for date_str, row_n, taker_v, vol_v, pos_v in written:
        tv = f"{taker_v:.2f}%" if isinstance(taker_v, (int, float)) else str(taker_v)
        vv = f"{vol_v:,.0f}" if isinstance(vol_v, (int, float)) else str(vol_v)
        pv = f"{pos_v:.4f}" if isinstance(pos_v, (int, float)) else str(pos_v)
        print(f"  {date_str:>12}  {row_n:>5}  {tv:>14}  {vv:>14}  {pv:>12}")

    if missing_export:
        print(f"\n  Missing export dates (rows stay blank): {missing_export}")

print(f"\n{DIV}")
print("DUNG.")
