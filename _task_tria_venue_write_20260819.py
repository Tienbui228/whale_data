"""_task_tria_venue_write_20260819.py
Write vol_alpha + vol_spot_ext_usd/spot_ext_venue into CASE_TRIA.xlsx DAILY.
- Add 2 new cols at end (col 51, 52) with row1 section header 'SPOT-EXT (v2.1)'
- Fill vol_alpha (col 35) from tria_alpha_vol_20260819.csv
- Fill vol_spot_ext_usd (col 51) + spot_ext_venue (col 52) from tria_bybit_spot_20260819.csv
- Explicit tag: vol_spot_usd (col 31) = 'N/A-STRUCTURAL'
                taker_spot_pct (col 32) = 'ABSTAIN-STRUCTURAL'
                taker_alpha (col 36) = 'ABSTAIN-STRUCTURAL'
- Write-after-guard (L-50): verify then save
"""
import sys, io, csv, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
SCRATCH = ROOT / "scratch"
XLSX = ROOT / "docs" / "case_files" / "CASE_TRIA.xlsx"
BAK  = SCRATCH / "CASE_TRIA_bak_pre_venue_20260819.xlsx"

# Backup guard
assert BAK.exists(), f"BACKUP MISSING — refuse write. Run backup first."
print(f"Backup OK: {BAK} ({BAK.stat().st_size} B)")

# Load sources
def load_csv(p):
    with open(p, encoding="utf-8") as f: return list(csv.DictReader(f))

alpha_rows = load_csv(SCRATCH / "tria_alpha_vol_20260819.csv")
bybit_rows = load_csv(SCRATCH / "tria_bybit_spot_20260819.csv")

alpha_by_d = {r["date"]: r for r in alpha_rows}
bybit_by_d = {r["date"]: r for r in bybit_rows}
print(f"Alpha rows: {len(alpha_by_d)}, Bybit rows: {len(bybit_by_d)}")

# Load xlsx
wb = openpyxl.load_workbook(XLSX)
ws = wb["DAILY"]
row1 = [c.value for c in ws[1]]  # section titles
row2 = [c.value for c in ws[2]]  # field names
ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
n_cols_before = len(row2)
print(f"Cols before: {n_cols_before}")
print(f"  vol_alpha col: {ci.get('vol_alpha')}")
print(f"  vol_spot_usd col: {ci.get('vol_spot_usd')}")
print(f"  taker_spot_pct col: {ci.get('taker_spot_pct')}")
print(f"  taker_alpha col: {ci.get('taker_alpha')}")
print(f"  spot_perp_ratio col: {ci.get('spot_perp_ratio')}")

# Add 2 new cols at end (col 51, 52)
NEW_EXT_USD_COL = n_cols_before + 1  # 51
NEW_EXT_VENUE_COL = n_cols_before + 2  # 52

# Row 1 header (SPOT-EXT section)
ws.cell(1, NEW_EXT_USD_COL).value = "SPOT-EXT (v2.1 amendment 19/08)"
# Row 2 field names
ws.cell(2, NEW_EXT_USD_COL).value = "vol_spot_ext_usd"
ws.cell(2, NEW_EXT_VENUE_COL).value = "spot_ext_venue"

print(f"\nAdded new cols: {NEW_EXT_USD_COL}={'vol_spot_ext_usd'}, {NEW_EXT_VENUE_COL}={'spot_ext_venue'}")

# Walk rows 3-32, fill:
# - vol_alpha (col 35)
# - vol_spot_usd = 'N/A-STRUCTURAL'  (col 31)
# - taker_spot_pct = 'ABSTAIN-STRUCTURAL' (col 32)
# - taker_alpha = 'ABSTAIN-STRUCTURAL' (col 36)
# - vol_spot_ext_usd (col 51)
# - spot_ext_venue (col 52)
# - spot_perp_ratio (col 33) — computed as (vol_alpha + bybit_turnover) / oi_usd
#   for TRIA doctrine #10 use aggregate spot; also keep as canonical spot_perp
#   but since vol_spot_usd = N/A-STRUCTURAL, we log a compound ratio elsewhere
#   → leave col 33 = None (since vol_spot_usd is N/A, formal ratio can't compute)

vol_alpha_c = ci["vol_alpha"]
vol_spot_c = ci["vol_spot_usd"]
taker_spot_c = ci["taker_spot_pct"]
taker_alpha_c = ci["taker_alpha"]
spot_perp_ratio_c = ci["spot_perp_ratio"]

updated = 0
missing_dates = []
for r in range(3, ws.max_row + 1):
    d = ws.cell(r, ci["date"]).value
    if not d: continue
    d_str = str(d)[:10]

    # vol_alpha
    a = alpha_by_d.get(d_str)
    if a:
        ws.cell(r, vol_alpha_c).value = round(float(a["quote_vol_usd"]), 2)
    else:
        missing_dates.append(("alpha", d_str))

    # Structural tags
    ws.cell(r, vol_spot_c).value = "N/A-STRUCTURAL"
    ws.cell(r, taker_spot_c).value = "ABSTAIN-STRUCTURAL"
    ws.cell(r, taker_alpha_c).value = "ABSTAIN-STRUCTURAL"
    # spot_perp_ratio: N/A (vol_spot_usd is N/A). Report will use aggregate ratio elsewhere.
    ws.cell(r, spot_perp_ratio_c).value = None

    # spot-ext (Bybit)
    b = bybit_by_d.get(d_str)
    if b:
        ws.cell(r, NEW_EXT_USD_COL).value = round(float(b["turnover_usd"]), 2)
        ws.cell(r, NEW_EXT_VENUE_COL).value = "bybit"
    else:
        missing_dates.append(("bybit", d_str))
        ws.cell(r, NEW_EXT_USD_COL).value = None
        ws.cell(r, NEW_EXT_VENUE_COL).value = "bybit"  # venue known even if that day missing

    updated += 1

print(f"\nRows updated: {updated}, missing: {missing_dates if missing_dates else 'none'}")

# GUARD before save
non_null_alpha = sum(1 for r in range(3, ws.max_row+1)
                    if ws.cell(r, vol_alpha_c).value is not None)
non_null_bybit = sum(1 for r in range(3, ws.max_row+1)
                    if ws.cell(r, NEW_EXT_USD_COL).value is not None)
print(f"Guard: non_null vol_alpha={non_null_alpha}/{updated}, non_null vol_spot_ext_usd={non_null_bybit}/{updated}")

if non_null_alpha < updated * 0.9:
    sys.exit(f"GUARD FAIL: alpha coverage < 90% ({non_null_alpha}/{updated})")
if non_null_bybit < updated * 0.9:
    sys.exit(f"GUARD FAIL: bybit coverage < 90% ({non_null_bybit}/{updated})")

# Save
wb.save(XLSX)
print(f"\nSAVED: {XLSX}")

# Print sample of updated last-5 rows for verification
print(f"\n=== Verify last 5 rows ===")
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2["DAILY"]
row2_new = [c.value for c in ws2[2]]
ci2 = {str(n).strip(): i+1 for i, n in enumerate(row2_new) if n}
for r in range(ws2.max_row-4, ws2.max_row+1):
    d = ws2.cell(r, ci2["date"]).value
    print(f"\n  row {r} ({d}):")
    for key in ["c","vol_perp_usd","oi_usd","vol_alpha","vol_spot_usd","taker_spot_pct",
                "taker_alpha","spot_perp_ratio","vol_spot_ext_usd","spot_ext_venue"]:
        if key in ci2:
            v = ws2.cell(r, ci2[key]).value
            print(f"    {key}: {v}")

print("\n[DONE write]")
