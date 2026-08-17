#!/usr/bin/env python3
"""Check taker scale in col 8 for COOKIE/BROCCOLIF3B/KOMA across date ranges."""
from pathlib import Path
import openpyxl

CASE_DIR = Path(r"D:\Data back test\whale_data\docs\case_files")

BAD_DATES = {
    "CASE_COOKIE.xlsx":      ["2026-08-14", "2026-08-15", "2026-08-16"],
    "CASE_BROCCOLIF3B.xlsx": [f"2026-08-{d:02d}" for d in range(10, 17)],
    "CASE_KOMA.xlsx":        [f"2026-08-{d:02d}" for d in range(10, 17)],
}

for fname, bad_dates in BAD_DATES.items():
    fpath = CASE_DIR / fname
    print(f"\n{'='*70}")
    print(f"FILE: {fname}  (bad dates: {bad_dates[0]} … {bad_dates[-1]})")
    wb = openpyxl.load_workbook(str(fpath), data_only=True, read_only=True)
    ws = wb["DAILY"]

    # Print last 20 rows: date, col2(px_avg), col7(vol), col8(taker), col9(f_avg),
    #                     col12(fn), col13(oi), col27(pos_r), col28(acc_r), col29(gLSR)
    max_row = ws.max_row
    print(f"  {'date':>12}  {'col2_px':>10}  {'col7_vol':>14}  {'col8_taker':>12}  "
          f"{'col9_favg':>12}  {'col12_fn':>8}  {'col13_oi':>14}  "
          f"{'col27_pos':>10}  {'col28_acc':>10}  {'col29_gLSR':>10}  STATUS")
    for r in range(max(3, max_row - 19), max_row + 1):
        d   = ws.cell(r, 1).value
        p   = ws.cell(r, 2).value
        vol = ws.cell(r, 7).value
        tak = ws.cell(r, 8).value
        fa  = ws.cell(r, 9).value
        fn  = ws.cell(r, 12).value
        oi  = ws.cell(r, 13).value
        pos = ws.cell(r, 27).value
        acc = ws.cell(r, 28).value
        gls = ws.cell(r, 29).value
        bad = "BAD" if str(d) in bad_dates else ""
        def f(v, w=12): return f"{v:{w}.4g}" if isinstance(v, (int, float)) else f"{str(v)[:w]:>{w}}"
        print(f"  {str(d):>12}  {f(p,10)}  {f(vol,14)}  {f(tak,12)}  "
              f"{f(fa,12)}  {f(fn,8)}  {f(oi,14)}  "
              f"{f(pos,10)}  {f(acc,10)}  {f(gls,10)}  {bad}")
    wb.close()

print("\nDUNG.")
