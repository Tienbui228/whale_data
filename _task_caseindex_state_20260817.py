"""_task_caseindex_state_20260817.py
Sync state column in CASE_INDEX.xlsx per §1 of analysis_handover_20260817.md.
AIN→TRACKING-ACTIVE, APR→DISTRIBUTION-WATCH, GUA→S2-WATCH·DISTRIBUTION,
BTW→DISTRIBUTION-WATCH, ON→DISTRIBUTION-WATCH·S2-leg3.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
xlsx = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"

STATE_MAP = {
    "AIN":  "TRACKING-ACTIVE",
    "APR":  "DISTRIBUTION-WATCH",
    "GUA":  "S2-WATCH · DISTRIBUTION",
    "BTW":  "DISTRIBUTION-WATCH",
    "ON":   "DISTRIBUTION-WATCH · S2-leg3",
}

wb = openpyxl.load_workbook(xlsx)
print(f"Sheets: {wb.sheetnames}")

# Try to find the main/index sheet
# Common names: "INDEX", "CASE_INDEX", first sheet
target_sheet = None
for name in wb.sheetnames:
    if "index" in name.lower() or "case" in name.lower():
        target_sheet = wb[name]
        break
if target_sheet is None:
    target_sheet = wb.active

ws = target_sheet
print(f"Using sheet: {ws.title}")
print(f"Dimensions: {ws.dimensions}")

# Find header row — look for "token" and "state" columns
header_row = None
token_col  = None
state_col  = None

for row_idx in range(1, 6):
    for col_idx in range(1, 20):
        val = ws.cell(row=row_idx, column=col_idx).value
        if val is None:
            continue
        val_str = str(val).strip().lower()
        if val_str in ("token", "symbol", "ticker"):
            header_row = row_idx
            token_col  = col_idx
        if val_str in ("state", "status", "trạng thái"):
            state_col  = col_idx

print(f"Header row: {header_row}, token col: {token_col}, state col: {state_col}")

if token_col is None or state_col is None:
    # Print first 5 rows to inspect
    print("Could not find header. First 5 rows:")
    for r in range(1, 6):
        row_data = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        print(f"  Row {r}: {row_data}")
    sys.exit(1)

# Walk rows and update
updated = []
for row_idx in range(header_row + 1, ws.max_row + 1):
    token_val = ws.cell(row=row_idx, column=token_col).value
    if token_val is None:
        continue
    token_str = str(token_val).strip().upper()
    if token_str in STATE_MAP:
        old_state = ws.cell(row=row_idx, column=state_col).value
        new_state = STATE_MAP[token_str]
        ws.cell(row=row_idx, column=state_col).value = new_state
        updated.append((token_str, old_state, new_state))

if not updated:
    print("No tokens matched — check sheet structure manually")
    # Print first 10 data rows
    for r in range(header_row, min(header_row+10, ws.max_row+1)):
        tok = ws.cell(row=r, column=token_col).value
        st  = ws.cell(row=r, column=state_col).value
        print(f"  Row {r}: token={tok!r} state={st!r}")
    sys.exit(1)

for tok, old, new in updated:
    print(f"  {tok}: {old!r} → {new!r}")

wb.save(xlsx)
print(f"\nSaved: {xlsx.name} ({len(updated)} rows updated)")
print("[DONE]")
