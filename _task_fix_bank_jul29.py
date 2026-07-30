"""Task: retroactively fix CASE_BANK.xlsx row 37 (Jul29).
- col 5 (taker_buy_pct): restore native 0.49892... (was 0.5046 from Zone-A)
- cols 15/16/17 (pos_r/acc_r/gLSR): add SUSPECT fn<6 comment (fn=1, src=native_1d)
Native values from _probe_bank_jul29_taker.py:
  taker_buy_pct = 0.49892268377010573
  lsr_pos_r     = 1.3307
  lsr_acc_r     = 0.8567
  lsr_global_r  = 0.9384
"""
import openpyxl
from openpyxl.comments import Comment
from pathlib import Path

CASE_BANK = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_BANK.xlsx")

TARGET_DATE   = "2026-07-29"
NATIVE_TAKER  = 0.49892268377010573
FN_N          = "1"
SRC_RATIO     = "native_1d"
SUSPECT_NOTE  = f"SUSPECT fn<6 (fn={FN_N}, src={SRC_RATIO})"

wb = openpyxl.load_workbook(str(CASE_BANK), data_only=False)
ws = wb["DAILY"]

target_row = None
for r in range(3, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v and str(v).strip() == TARGET_DATE:
        target_row = r
        break

if target_row is None:
    print(f"ERROR: row for {TARGET_DATE} not found in DAILY sheet")
else:
    print(f"Found {TARGET_DATE} at row {target_row}")

    # Before
    print(f"  col5 (taker) BEFORE: {ws.cell(target_row, 5).value}")
    print(f"  col15 (pos_r) BEFORE: {ws.cell(target_row, 15).value}")
    print(f"  col16 (acc_r) BEFORE: {ws.cell(target_row, 16).value}")
    print(f"  col17 (gLSR) BEFORE:  {ws.cell(target_row, 17).value}")

    # Fix col 5: restore native taker
    c5 = ws.cell(target_row, 5)
    c5.value = NATIVE_TAKER
    c5.comment = None  # clear old Zone-A comment

    # Add SUSPECT comment to ratio cols
    for col in [15, 16, 17]:
        cell = ws.cell(target_row, col)
        cell.comment = Comment(SUSPECT_NOTE, "case_daily.py")

    print(f"\n  col5 (taker) AFTER:  {ws.cell(target_row, 5).value}")
    print(f"  SUSPECT comment added to cols 15/16/17: {SUSPECT_NOTE}")

    wb.save(str(CASE_BANK))
    print(f"\nSaved: {CASE_BANK}")
