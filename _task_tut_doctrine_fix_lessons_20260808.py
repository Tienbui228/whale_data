"""
Fix: L-43/L-44 vua append truoc do phai doi thanh L-44/L-45
de giu L-43 cho operator pending lesson ("khong suy dien da-ky tu im lang").
"""
import sys, io
from pathlib import Path
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

INDEX_PATH = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")

wb = openpyxl.load_workbook(INDEX_PATH)
ws = wb["LESSONS"]

hdrs = {c.value: c.column for c in ws[1] if c.value}
col_id = 1  # lesson_id is col 1 based on previous run

# Find and update the two rows we just added
changes = 0
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, col_id).value
    if v == "L-43":
        ws.cell(r, col_id).value = "L-44"
        # Update body text too
        body_cell = ws.cell(r, 2)
        if body_cell.value:
            body_cell.value = body_cell.value.replace("L-43(a)", "L-44(a)").replace("L-43", "L-44")
        print(f"  row {r}: L-43 → L-44")
        changes += 1
    elif v == "L-44":
        ws.cell(r, col_id).value = "L-45"
        body_cell = ws.cell(r, 2)
        if body_cell.value:
            body_cell.value = body_cell.value.replace("L-44(b)", "L-45(b)").replace("L-44", "L-45")
        print(f"  row {r}: L-44 → L-45")
        changes += 1

if changes == 0:
    print("  WARN: no L-43/L-44 found to rename — check sheet")
else:
    wb.save(INDEX_PATH)
    print(f"\nSaved. {changes} rows updated.")
    print("L-43 slot preserved for operator pending lesson (khong suy dien da-ky tu im lang).")
    print("Our drafts now at L-44 (campaign termination) and L-45 (duration scaling).")

print("\n[DONE]")
