"""
_task_koma_8_register_20260810.py
K-8: Register CASE_KOMA in CASE_INDEX.xlsx + regenerate INDEX.md + L-37 backup
"""
import sys, io, json, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT     = Path(r"D:\Data back test\whale_data")
DATE_TAG = "20260810"
SYM_BASE = "KOMA"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"K-8 REGISTER  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

IDX_PATH  = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"
IMDX_PATH = ROOT / "docs" / "case_files" / "INDEX.md"
CASE_PATH = ROOT / "docs" / "case_files" / f"CASE_{SYM_BASE}.xlsx"
SCRATCH   = ROOT / "scratch"

# ── Step 1: Inspect + find header row ─────────────────────────────────────
log("\n=== Step 1: Inspect CASE_INDEX ===")
wb = openpyxl.load_workbook(str(IDX_PATH))
ws = wb["INDEX"]
log(f"  max_row={ws.max_row}")

# Find header row (contains 'file' and 'token')
header_row = None; headers = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    row_str = " ".join(str(c) for c in row if c).lower()
    if "file" in row_str and "token" in row_str:
        header_row = i; headers = list(row); break

log(f"  Header row: {header_row} → {headers}")

# ── Step 2: Check existing + target row ───────────────────────────────────
log("\n=== Step 2: Check existing KOMA registration ===")
existing_row = None
if header_row:
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= header_row: continue
        row_str = " ".join(str(c) for c in row if c).upper()
        if "KOMA" in row_str and "CASE_KOMA" in row_str:
            existing_row = i
            log(f"  Existing at row {i}: {[str(c)[:30] if c else None for c in row[:6]]}")
            break
    if not existing_row:
        log(f"  Not registered — will append")

# ── Step 3: Column map ────────────────────────────────────────────────────
col_map = {}
if headers:
    for ci, h in enumerate(headers, start=1):
        if h:
            col_map[str(h).lower().strip()] = ci

log(f"  Column map: {col_map}")

# ── Step 4: Build row data ────────────────────────────────────────────────
new_row = {
    "file":         f"CASE_{SYM_BASE}.xlsx",
    "token":        SYM_BASE,
    "status":       "OPEN",
    "loai":         "GOM-C PROVISIONAL",
    "action_class": "COLD_START",
    "ngay_mo":      "2026-08-10",
    "t0":           "",
    "moc_ke_tiep":  "depth snap#2 >=05:51 UTC 2026-08-10; GOM-C gate 15/08",
    "han":          "",
    "mm_id":        "ALPHA_1",
    "ghi_chu":      "K audit-collect 10/08; pump 07-31 +100%/$604M; OI peak $19.4M UNCALIB; GOM-C PROV 07-23+08-02",
}

# ── Step 5: Write row ──────────────────────────────────────────────────────
log("\n=== Step 5: Write row ===")
target_row = existing_row if existing_row else (ws.max_row + 1)
log(f"  Writing to row {target_row}")

for col_key, value in new_row.items():
    if col_key in col_map:
        ci = col_map[col_key]
        ws.cell(row=target_row, column=ci, value=value)
        log(f"  col {ci} ({col_key}): {value}")
    else:
        # Try partial match
        matched = next((hk for hk in col_map if col_key in hk or hk in col_key), None)
        if matched:
            ci = col_map[matched]
            ws.cell(row=target_row, column=ci, value=value)
            log(f"  col {ci} ({matched}≈{col_key}): {value}")
        else:
            log(f"  SKIP {col_key} — no column match")

wb.save(str(IDX_PATH)); wb.close()
log(f"  Saved CASE_INDEX.xlsx")

# ── Step 6: Verify ────────────────────────────────────────────────────────
log("\n=== Step 6: Verify ===")
wb2 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws2 = wb2["INDEX"]
for i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
    row_str = " ".join(str(c) for c in row if c).upper()
    if "KOMA" in row_str and "CASE_KOMA" in row_str:
        log(f"  Row {i}: {[str(c)[:40] if c else None for c in row]}")
wb2.close()

# ── Step 7: Regenerate INDEX.md ────────────────────────────────────────────
log("\n=== Step 7: Regenerate INDEX.md ===")
wb3 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws3 = wb3["INDEX"]
all_rows = list(ws3.iter_rows(values_only=True))
wb3.close()

hdr_i = None; hdr = None
for i, row in enumerate(all_rows):
    row_str = " ".join(str(c) for c in row if c).lower()
    if "file" in row_str and "token" in row_str:
        hdr_i = i; hdr = row; break

lines = [f"# FRESHWATCH CASE INDEX", "", f"Generated: {NOW.strftime('%Y-%m-%d %H:%M UTC')}", ""]
if hdr_i is not None:
    data_rows = [r for r in all_rows[hdr_i+1:] if any(c for c in r)]
    hdr_c = [str(c) if c else "" for c in hdr]
    lines.append("| " + " | ".join(hdr_c) + " |")
    lines.append("| " + " | ".join("---" for _ in hdr_c) + " |")
    for row in data_rows:
        cells = [str(c)[:60] if c is not None else "" for c in row]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")

IMDX_PATH.write_text("\n".join(lines), encoding="utf-8")
log(f"  INDEX.md: {IMDX_PATH.stat().st_size:,} bytes, {len(lines)} lines")

# ── Step 8: L-37 backup ────────────────────────────────────────────────────
log("\n=== Step 8: L-37 backup ===")
if CASE_PATH.exists():
    bkp = SCRATCH / f"CASE_{SYM_BASE}_bak_reg_{NOW.strftime('%Y%m%d_%H%M%S')}.xlsx"
    shutil.copy2(CASE_PATH, bkp)
    log(f"  Backup: {bkp.name} ({bkp.stat().st_size:,} bytes)")

log(f"\n[K-8 REGISTER DONE]")
