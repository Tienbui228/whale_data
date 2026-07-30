#!/usr/bin/env python3
"""Append L-35, L-36, L-37 to CASE_INDEX.xlsx LESSONS sheet and regenerate lessons_learned.md."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from pathlib import Path
from datetime import datetime

IDX   = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
MDOUT = Path(r"D:\Data back test\whale_data\docs\lessons_learned.md")

NEW_LESSONS = [
    [
        "L-35",
        "2026-07-31",
        "binance-bot-signals P0 regression suite",
        ("Bo calib (12 file _result.txt + GATE dict + replay.py) phai song trong git cung voi code — "
         "da dung, ghi nhan de giu. Quy tac: moi repo moi can regression suite commit ngay tu dau, "
         "khong de calib file ngoai version control. Loi that nhat la khi calib tach khoi code: "
         "code duoc update ma calib cu → false pass / false fail."),
        "process/rule",
        "GHI NHAN",
        "moi repo moi commit regression suite cung voi code trong 24h dau — khong phan tach calib khoi git",
        None,
    ],
    [
        "L-36",
        "2026-07-31",
        "P0 gate session 31/07 (TAG/BTW reconstruct)",
        ("Tai tao data file phai test bang suite nguyen ban (replay.py goc tu git), KHONG tu viet harness moi. "
         "Harness tu viet dan den 3 assertion sai: (1) AKE lsr_accum: viet fires_in thay vi never_in; "
         "(2) BASED crowd_flush: viet fires_in thay vi never_in; "
         "(3) BULLA pos_rerate: dung file tai tao (pos 07-20=1.75) thay vi file goc (1.99). "
         "An: bao 40/43 gia, mat thoi gian debug. "
         "Fix: dung replay.py goc → 43/43 PASS."),
        "process/rule",
        "GHI NHAN",
        "quy trinh tai tao data: viet file → chay replay.py goc tu git → PASS truoc khi commit bat ky thu gi",
        None,
    ],
    [
        "L-37",
        "2026-07-31",
        "B0 bundle backup 4 repo (31/07)",
        ("Repo chi local (khong remote git server) = single point of failure — mot loi o dia hay "
         "OneDrive deauth mat het. Giai phap toi thieu: git bundle → SCP len VPS /root/archive/ "
         "trong vong 24h tao repo, sau do bundle lai tai moi milestone lon. "
         "Da ap 4 repo ngay 31/07: binance-bot-signals, whale_data, Check onchain, Whales_scan. "
         "Repo OneDrive them an: file-lock + single-copy = combo sinh an (CASE_BANK_NEW, L-34)."),
        "rule",
        "DA VAO — bundle 4 repo len VPS /root/archive/ ngay 31/07",
        "moi repo moi phai co bundle tren VPS trong 24h; nhac lai tai moi milestone lon",
        None,
    ],
]

wb = openpyxl.load_workbook(str(IDX), data_only=False)
ws = wb["LESSONS"]

print(f"TRUOC: max_row={ws.max_row}")

# Find last L-xx row
last_lesson_row = 4
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v and str(v).startswith("L-"):
        last_lesson_row = r

last_id_cell = ws.cell(last_lesson_row, 1).value
print(f"Last lesson: {last_id_cell} at row {last_lesson_row}")
assert last_id_cell == "L-34", f"Expected L-34, got {last_id_cell!r}"

# Append L-35, L-36, L-37
for i, row_data in enumerate(NEW_LESSONS):
    row_n = last_lesson_row + 1 + i
    for c, v in enumerate(row_data, start=1):
        ws.cell(row_n, c).value = v
    print(f"  Wrote {row_data[0]} at row {row_n}")

wb.save(str(IDX))
print(f"Saved: {IDX}")

# Verify
wb2 = openpyxl.load_workbook(str(IDX), data_only=True)
ws2 = wb2["LESSONS"]
lessons = [(r, ws2.cell(r, 1).value) for r in range(5, ws2.max_row + 1)
           if ws2.cell(r, 1).value and str(ws2.cell(r, 1).value).startswith("L-")]
print(f"\nVerify: {len(lessons)} lessons total")
for r, lid in lessons[-5:]:
    print(f"  row {r:02d}: {lid}")

# Regenerate lessons_learned.md
print("\n=== Regenerate lessons_learned.md ===")
lines = [
    "# Lessons Learned — FreshWatch",
    "",
    "> Auto-generated tu CASE_INDEX.xlsx sheet LESSONS. Khong edit file nay truc tiep.",
    f"> Lan sinh: {datetime.now().strftime('%Y-%m-%d')} (rev.3) · Tong: 37 bai hoc (L-01 → L-37)",
    "",
]

for r in range(5, ws2.max_row + 1):
    lid        = ws2.cell(r, 1).value
    ngay       = ws2.cell(r, 2).value
    case_src   = ws2.cell(r, 3).value
    bai_hoc    = ws2.cell(r, 4).value
    loai       = ws2.cell(r, 5).value
    trang_thai = ws2.cell(r, 6).value
    hanh_dong  = ws2.cell(r, 7).value

    if not lid or not str(lid).startswith("L-"):
        continue

    date_s   = str(ngay or "")
    case_s   = str(case_src or "")
    loai_s   = str(loai or "")
    status_s = str(trang_thai or "")
    bai_s    = str(bai_hoc or "")
    action_s = str(hanh_dong or "")

    lines.append(f"## {lid} [{date_s}] {case_s} — `{loai_s}` · {status_s}")
    lines.append(bai_s)
    if action_s:
        lines.append(f"*Hanh dong tiep: {action_s}*")
    lines.append("")

wb2.close()

MDOUT.write_text("\n".join(lines), encoding="utf-8")
print(f"Written: {MDOUT} ({len(lines)} lines, {len([l for l in lines if l.startswith('##')])} lessons)")
