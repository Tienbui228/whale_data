"""Audit v2: đọc taker_spot_pct không dùng data_only=True.
So sánh PRE (HEAD~1 từ git) vs POST (file hiện tại).
"""
import sys, io
from pathlib import Path
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

PRE  = Path(r"D:\Data back test\whale_data\scratch\CASE_ZBT_pre_patch.xlsx")
POST = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_ZBT.xlsx")

def read_col(path, sheet="DAILY", header_row=2, col_name="taker_spot_pct", date_col="date"):
    wb = openpyxl.load_workbook(path)  # NO data_only
    ws = wb[sheet]
    hdrs = {c.value: c.column for c in ws[header_row] if c.value}
    ci_tsp  = hdrs.get(col_name)
    ci_date = hdrs.get(date_col)
    if not ci_tsp:
        return {}, f"col '{col_name}' not found"
    result = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        d = row[ci_date - 1]
        v = row[ci_tsp  - 1]
        if d:
            result[str(d)[:10]] = v
    return result, None

pre_vals,  err1 = read_col(PRE)
post_vals, err2 = read_col(POST)

if err1: print(f"PRE ERROR: {err1}")
if err2: print(f"POST ERROR: {err2}")

print(f"{'date':<12} {'PRE':>12} {'POST':>12} {'status':>12}")
print("-" * 52)
changed = 0
for d in sorted(set(list(pre_vals) + list(post_vals))):
    p = pre_vals.get(d); q = post_vals.get(d)
    diff = (p != q)
    if diff: changed += 1
    flag = "CHANGED" if diff else "same"
    ps = f"{p:.2f}%" if isinstance(p, (int, float)) else str(p)
    qs = f"{q:.2f}%" if isinstance(q, (int, float)) else str(q)
    print(f"  {d:<12} {ps:>12} {qs:>12} {flag:>12}")

print(f"\nTotal changed: {changed} / {len(pre_vals)} rows")

# ── Chỉ đọc trực tiếp POST taker_spot_pct để xác nhận giá trị ──────────────
print()
print("POST file — taker_spot_pct values (direct read):")
for d, v in sorted(post_vals.items()):
    vstr = f"{v:.2f}%" if isinstance(v, (int, float)) else str(v)
    print(f"  {d}: {vstr}")
