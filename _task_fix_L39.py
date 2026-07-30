#!/usr/bin/env python3
"""Fix bad L-39 row (text in col0) + append correct row + regenerate lessons_learned.md."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from pathlib import Path

CASE_INDEX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
LESSONS_MD = Path(r"D:\Data back test\whale_data\docs\lessons_learned.md")

# LESSONS sheet structure (confirmed from probe):
# col0=id, col1=date, col2=context, col3=text, col4=category, col5=status, col6=action, col7=None

wb = openpyxl.load_workbook(str(CASE_INDEX))
ws = wb["LESSONS"]

# Find and delete the bad last row (row 43: text in col0, no id)
last_row_idx = ws.max_row
last_row_vals = [ws.cell(last_row_idx, c).value for c in range(1, 9)]
print(f"Last row ({last_row_idx}): {last_row_vals[:4]}")

# Verify it's the bad L-39 (col0 is full text, col1 is date, col2 is None)
if last_row_vals[0] and last_row_vals[0].startswith("Quy") and last_row_vals[2] is None:
    ws.delete_rows(last_row_idx)
    print(f"Deleted bad row {last_row_idx}")
else:
    print(f"WARNING: last row does not look like bad L-39 — aborting delete")
    wb.close()
    sys.exit(1)

# Append correct L-39 row
L39 = [
    "L-39",
    "2026-07-31",
    "proposal v2 (31/07) — boundary canonical",
    (
        "Quy ước boundary phải TEST bằng đối chiếu nguồn (21/21 diff 0%) trước khi văn bản hóa "
        "— proposal v1 hợp lý trên giấy nhưng ngược 180° với thực nghiệm; "
        "n=1 suy luận đẹp vẫn phải qua n=21."
    ),
    "rule",
    "GHI NHAN — proposal v2 ký 31/07, CLAUDE.md updated",
    "trước khi văn bản hóa quy ước metric nào: test >= 21 điểm thực nghiệm đối chiếu nguồn",
    None,
]
ws.append(L39)
print(f"Appended L-39: {L39[:4]}")

wb.save(str(CASE_INDEX))
print(f"Saved: {CASE_INDEX}")

# Regenerate lessons_learned.md
wb2 = openpyxl.load_workbook(str(CASE_INDEX), read_only=True, data_only=True)
ws2 = wb2["LESSONS"]
all_rows = list(ws2.iter_rows(values_only=True))
wb2.close()

data_rows = [r for r in all_rows[1:] if r[0] is not None and str(r[0]).startswith("L-")]
print(f"\nData rows with L-XX id: {len(data_rows)}")

lines = [
    "# Lessons Learned",
    f"_Generated from CASE_INDEX.xlsx sheet LESSONS — {len(data_rows)} entries_",
    "",
]
for r in data_rows:
    lid = r[0] or ""       # col0: ID
    dt = r[1] or ""        # col1: date
    ctx = r[2] or ""       # col2: context
    txt = r[3] or ""       # col3: lesson text
    cat = r[4] or ""       # col4: category
    action = r[6] or ""    # col6: action
    lines.append(f"**{lid}** ({dt}): {txt}")
    lines.append("")

content = "\n".join(lines)
LESSONS_MD.write_text(content, encoding="utf-8")
print(f"Regenerated: {LESSONS_MD} (rev.5, {len(data_rows)} lessons)")

print("\nLast 3 lessons:")
for r in data_rows[-3:]:
    print(f"  {r[0]} | {r[1]} | {str(r[3])[:80]}...")
