"""
_task_br3_8_register_20260810.py
BR3-8: Register CASE_BROCCOLIF3B in CASE_INDEX.xlsx + regenerate INDEX.md
"""
import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT     = Path(r"D:\Data back test\whale_data")
DATE_TAG = "20260810"
SYM_BASE = "BROCCOLIF3B"

def log(s=""): print(s)
NOW = datetime.now(timezone.utc)
log(f"BR3-8 REGISTER  |  {NOW.strftime('%Y-%m-%d %H:%M UTC')}")

IDX_PATH  = ROOT / "docs" / "case_files" / "CASE_INDEX.xlsx"
IMDX_PATH = ROOT / "docs" / "case_files" / "INDEX.md"
CASE_PATH = ROOT / "docs" / "case_files" / f"CASE_{SYM_BASE}.xlsx"

# ── Step 1: Inspect CASE_INDEX structure ──────────────────────────────────
log("\n=== Step 1: Inspect CASE_INDEX structure ===")
wb = openpyxl.load_workbook(str(IDX_PATH))
ws = wb["INDEX"]
log(f"  max_row={ws.max_row} max_col={ws.max_column}")
# Find header row
header_row = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if row[0] and str(row[0]).lower() in ("file","case_file","filename"):
        header_row = i
        headers = list(row)
        log(f"  Header row: {i} → {headers}")
        break
    if row[0]:
        log(f"  Row {i}: {[str(c)[:40] if c else None for c in row[:6]]}")

if not header_row:
    # fallback: scan for row with 'file' or 'token' in cols
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        row_str = " ".join(str(c) for c in row if c).lower()
        if "file" in row_str and "token" in row_str:
            header_row = i
            headers = list(row)
            log(f"  Header row (fallback): {i} → {headers}")
            break

log(f"  Header row found: {header_row}")

# ── Step 2: Check if BROCCOLIF3B already registered ──────────────────────
log("\n=== Step 2: Check existing registration ===")
existing_row = None
if header_row:
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= header_row:
            continue
        row_str = " ".join(str(c) for c in row if c).upper()
        if "BROCCOLIF3B" in row_str:
            existing_row = i
            log(f"  Found existing at row {i}: {[str(c)[:30] if c else None for c in row[:6]]}")
            break
    if not existing_row:
        log(f"  Not yet registered — will add")

# ── Step 3: Map column positions ─────────────────────────────────────────
log("\n=== Step 3: Column map ===")
col_map = {}
if header_row and headers:
    for ci, h in enumerate(headers, start=1):
        if h:
            col_map[str(h).lower().strip()] = ci
            log(f"  col {ci}: '{h}'")

# Build row data — match columns from CASE_INDEX structure
# Standard columns: file, token, status, loai, action_class, ngay_mo, T0, moc_ke_tiep, han, MM_ID, ghi_chu
new_row_data = {
    "file":         f"CASE_{SYM_BASE}.xlsx",
    "token":        SYM_BASE,
    "status":       "OPEN",
    "loai":         "COLD_START",
    "action_class": "COLD_START",
    "ngay_mo":      "2026-08-10",
    "t0":           "",
    "moc_ke_tiep":  f"depth snap#2 >=04:45 UTC 2026-08-10",
    "han":          "",
    "mm_id":        "",
    "ghi_chu":      "BR3 audit-collect 10/08; two-broccoli confirmed distinct; DUNG cho doc",
}

# ── Step 4: Write row ─────────────────────────────────────────────────────
log("\n=== Step 4: Write to INDEX ===")
if existing_row:
    target_row = existing_row
    log(f"  Overwriting existing row {target_row}")
else:
    target_row = ws.max_row + 1
    log(f"  Appending at row {target_row}")

# Write by column position using col_map
for col_key, value in new_row_data.items():
    if col_key in col_map:
        ci = col_map[col_key]
        ws.cell(row=target_row, column=ci, value=value)
        log(f"  col {ci} ({col_key}): {value}")
    else:
        # Try partial match
        matched = None
        for hk in col_map:
            if col_key in hk or hk in col_key:
                matched = hk; break
        if matched:
            ci = col_map[matched]
            ws.cell(row=target_row, column=ci, value=value)
            log(f"  col {ci} ({matched}≈{col_key}): {value}")
        else:
            log(f"  SKIP {col_key} — no matching column")

wb.save(str(IDX_PATH))
wb.close()
log(f"  Saved CASE_INDEX.xlsx")

# ── Step 5: Verify written row ─────────────────────────────────────────────
log("\n=== Step 5: Verify ===")
wb2 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws2 = wb2["INDEX"]
for i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
    row_str = " ".join(str(c) for c in row if c).upper()
    if "BROCCOLIF3B" in row_str:
        log(f"  Row {i}: {[str(c)[:40] if c else None for c in row]}")
wb2.close()

# ── Step 6: Regenerate INDEX.md ───────────────────────────────────────────
log("\n=== Step 6: Regenerate INDEX.md ===")
wb3 = openpyxl.load_workbook(str(IDX_PATH), read_only=True)
ws3 = wb3["INDEX"]
all_rows = list(ws3.iter_rows(values_only=True))
wb3.close()

# Find header row again
hdr_i = None; hdr = None
for i, row in enumerate(all_rows):
    row_str = " ".join(str(c) for c in row if c).lower()
    if "file" in row_str and "token" in row_str:
        hdr_i = i; hdr = row; break

lines = []
lines.append(f"# FRESHWATCH CASE INDEX")
lines.append(f"")
lines.append(f"Generated: {NOW.strftime('%Y-%m-%d %H:%M UTC')}")
lines.append(f"")

if hdr_i is not None:
    # Build markdown table
    data_rows = [r for r in all_rows[hdr_i+1:] if any(c for c in r)]
    # Header
    hdr_clean = [str(c) if c else "" for c in hdr]
    lines.append("| " + " | ".join(hdr_clean) + " |")
    lines.append("| " + " | ".join("---" for _ in hdr_clean) + " |")
    for row in data_rows:
        cells = [str(c)[:60] if c is not None else "" for c in row]
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
else:
    for row in all_rows:
        if any(c for c in row):
            lines.append(" | ".join(str(c)[:40] if c else "" for c in row))

IMDX_PATH.write_text("\n".join(lines), encoding="utf-8")
log(f"  INDEX.md: {IMDX_PATH.stat().st_size:,} bytes, {len(lines)} lines")

# ── Step 7: L-37 backup ────────────────────────────────────────────────────
log("\n=== Step 7: L-37 backup of CASE_BROCCOLIF3B ===")
import shutil
SCRATCH = ROOT / "scratch"
if CASE_PATH.exists():
    bkp = SCRATCH / f"CASE_{SYM_BASE}_bak_reg_{NOW.strftime('%Y%m%d_%H%M%S')}.xlsx"
    shutil.copy2(CASE_PATH, bkp)
    log(f"  Backup: {bkp.name} ({bkp.stat().st_size:,} bytes)")

log(f"\n[BR3-8 REGISTER DONE]")
