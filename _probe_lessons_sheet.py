#!/usr/bin/env python3
"""Probe CASE_INDEX.xlsx LESSONS sheet — columns and last rows."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from pathlib import Path

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
print(f"Sheets: {wb.sheetnames}")

ws = wb["LESSONS"]
rows = list(ws.iter_rows(values_only=True))
print(f"\nLESSONS rows total: {len(rows)}")
print(f"Header (row 0): {rows[0]}")
print(f"\nLast 3 data rows:")
for r in rows[-3:]:
    print(r)
wb.close()
