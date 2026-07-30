#!/usr/bin/env python3
"""Append L-39 to CASE_INDEX.xlsx LESSONS sheet + regenerate lessons_learned.md."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
from pathlib import Path
from datetime import date

CASE_INDEX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
LESSONS_MD = Path(r"D:\Data back test\whale_data\docs\lessons_learned.md")

# --- 1. Append L-39 to LESSONS sheet ---
wb = openpyxl.load_workbook(str(CASE_INDEX))
ws = wb["LESSONS"]

rows = list(ws.iter_rows(values_only=True))
header = rows[0]
print(f"Header: {header}")
print(f"Current rows (data): {len(rows) - 1}")

# Find last lesson number
last_num = 0
for r in rows[1:]:
    if r[0] and str(r[0]).startswith("L-"):
        try:
            n = int(str(r[0])[2:])
            if n > last_num:
                last_num = n
        except ValueError:
            pass
print(f"Last lesson: L-{last_num:02d}")

# L-39 data — map to header columns
# Header likely: (id, date, text, ...) — check actual header
print(f"Header cols: {list(enumerate(header))}")

# Construct L-39 row matching header structure
L39_ID = "L-39"
L39_DATE = "2026-07-31"
L39_TEXT = (
    "Quy ước boundary phải TEST bằng đối chiếu nguồn (21/21 diff 0%) trước khi văn bản hóa "
    "— proposal v1 hợp lý trên giấy nhưng ngược 180° với thực nghiệm; "
    "n=1 suy luận đẹp vẫn phải qua n=21."
)

# Build row matching header length
col_map = {str(c).lower(): i for i, c in enumerate(header) if c}
print(f"col_map: {col_map}")

new_row = [None] * len(header)
# Try common column names
for key, idx in col_map.items():
    if "id" in key or key == "lesson" or key == "#":
        new_row[idx] = L39_ID
    elif "date" in key:
        new_row[idx] = L39_DATE
    elif "text" in key or "lesson" in key or "note" in key or "description" in key:
        new_row[idx] = L39_TEXT

# Fallback: if col_map matching didn't land id+text, set positionally
if new_row[0] is None:
    new_row[0] = L39_ID
if len(new_row) > 1 and new_row[1] is None:
    # check if col 1 is date-like
    sample = rows[2][1] if len(rows) > 2 else None
    if sample and (isinstance(sample, str) and "-" in str(sample)):
        new_row[1] = L39_DATE
    elif len(new_row) > 2:
        new_row[1] = L39_DATE
if not any(v == L39_TEXT for v in new_row):
    # find first non-id non-date text col
    for i in range(2, len(new_row)):
        if new_row[i] is None:
            new_row[i] = L39_TEXT
            break

ws.append(new_row)
print(f"Appended: {new_row[:4]}")

wb.save(str(CASE_INDEX))
print(f"Saved: {CASE_INDEX}")

# --- 2. Regenerate lessons_learned.md ---
# Re-read to verify
wb2 = openpyxl.load_workbook(str(CASE_INDEX), read_only=True, data_only=True)
ws2 = wb2["LESSONS"]
all_rows = list(ws2.iter_rows(values_only=True))
header2 = all_rows[0]
data_rows = [r for r in all_rows[1:] if any(c is not None for c in r)]
wb2.close()

print(f"\nTotal lesson rows after append: {len(data_rows)}")

# Find id/text col indices
h_lower = [str(c).lower() if c else "" for c in header2]
id_col = next((i for i, h in enumerate(h_lower) if "id" in h or h == "#" or h == "lesson"), 0)
date_col = next((i for i, h in enumerate(h_lower) if "date" in h), 1)
text_col = next((i for i, h in enumerate(h_lower) if "text" in h or "lesson" in h or "note" in h or "description" in h), 2)
print(f"id_col={id_col}, date_col={date_col}, text_col={text_col}")

lines = ["# Lessons Learned", f"_Generated from CASE_INDEX.xlsx sheet LESSONS — {len(data_rows)} entries_", ""]
for r in data_rows:
    lid = r[id_col] if r[id_col] is not None else ""
    dt = r[date_col] if r[date_col] is not None else ""
    txt = r[text_col] if r[text_col] is not None else ""
    lines.append(f"**{lid}** ({dt}): {txt}")
    lines.append("")

content = "\n".join(lines)
LESSONS_MD.write_text(content, encoding="utf-8")
print(f"Regenerated: {LESSONS_MD} ({len(data_rows)} lessons)")
print("\nLast 3 lessons:")
for r in data_rows[-3:]:
    print(f"  {r[id_col]} | {r[date_col]} | {str(r[text_col])[:80]}")
