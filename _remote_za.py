#!/usr/bin/env python3
"""Probe Zone A sample rows in CAMPAIGNS after R1 fix."""
import openpyxl
wb = openpyxl.load_workbook('/root/Data_back_test/whale_data/data/ledger/ledger_v2.xlsx', read_only=True, data_only=True)
ws = wb['CAMPAIGNS']
header = [str(c) if c else '' for c in next(ws.iter_rows(max_row=1, values_only=True))]
print("CAMPAIGNS header:", header[:25])
print()
targets = {'TAG', 'CLO', 'BANK', 'BR'}
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row:
        continue
    sym = str(row[1]) if len(row) > 1 else ''
    if sym in targets:
        d = {header[i]: row[i] for i in range(min(len(header), len(row))) if row[i] is not None and header[i]}
        print(f"  {sym}:")
        for k, v in d.items():
            print(f"    {k}: {v}")
        print()
