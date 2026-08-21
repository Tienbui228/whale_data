"""_probe_case_inventory_step_a_20260821.py — STEP A read-only inventory 5 cases.
Chưa write gì. Chỉ paste bảng 5 dòng.
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

CASES = Path(r"D:\Data back test\whale_data\docs\case_files")
TOKENS = ["BANK", "TUT", "AKE", "BTW", "CYS"]
REQUIRED_SHEETS = ["CASE", "DAILY", "EVENTS", "TRIGGERS", "OUTCOME", "MM_PROFILE", "UNLOCK"]

def safe_float(v):
    if v in (None, "", "None"): return None
    try: return float(v)
    except: return None

def inventory(sym):
    fp = CASES / f"CASE_{sym}.xlsx"
    print(f"\n{'='*70}")
    print(f"[{sym}] {fp}")
    print(f"{'='*70}")
    if not fp.exists():
        print(f"  FILE KHÔNG TỒN TẠI — CHƯA CHẠY (bỏ inventory)")
        return None

    wb = openpyxl.load_workbook(fp, data_only=True)
    sheets = wb.sheetnames
    print(f"  Sheets: {sheets}")
    missing = [s for s in REQUIRED_SHEETS if s not in sheets]
    extra = [s for s in sheets if s not in REQUIRED_SHEETS]
    print(f"  Missing: {missing} · Extra: {extra}")

    if "DAILY" not in sheets:
        print(f"  DAILY sheet MISSING — dừng inventory")
        return None

    ws = wb["DAILY"]
    row2 = [c.value for c in ws[2]]
    n_cols = len(row2)
    print(f"  DAILY cols: {n_cols} · dims: {ws.dimensions}")

    # Schema check: cols 51-52 tồn tại? cols 53-54 KHÔNG được?
    col_51 = row2[50] if len(row2) > 50 else None
    col_52 = row2[51] if len(row2) > 51 else None
    col_53 = row2[52] if len(row2) > 52 else None
    col_54 = row2[53] if len(row2) > 53 else None
    print(f"  Schema: col51={col_51!r} · col52={col_52!r}")
    print(f"          col53={col_53!r} · col54={col_54!r} (không được tồn tại)")

    # DAILY date range + rows
    dates = []
    for r in range(3, ws.max_row+1):
        d = ws.cell(r, 1).value
        if d:
            dates.append(str(d)[:10])
    if not dates:
        print(f"  DAILY: 0 data rows")
        return None
    dates.sort()
    print(f"  DAILY: {len(dates)} rows · min={dates[0]} · max={dates[-1]}")

    # missing dates
    from datetime import date as _date, timedelta
    d0 = _date.fromisoformat(dates[0])
    d1 = _date.fromisoformat(dates[-1])
    expected_all = set()
    cur = d0
    while cur <= d1:
        expected_all.add(cur.isoformat()); cur += timedelta(days=1)
    actual = set(dates)
    missing_dates = sorted(expected_all - actual)
    dup_dates = sorted({d for d in dates if dates.count(d) > 1})
    print(f"  Missing dates: {len(missing_dates)}" + (f" (first 5: {missing_dates[:5]})" if missing_dates else ""))
    print(f"  Duplicate dates: {len(dup_dates)}" + (f" {dup_dates[:5]}" if dup_dates else ""))

    # Null% per key columns
    ci = {str(n).strip(): i+1 for i, n in enumerate(row2) if n}
    key_fields = {
        "pos_r": ["pos_r", "lsr_pos_r"],
        "acc_r": ["acc_r", "lsr_acc_r"],
        "gLSR": ["gLSR", "lsr_global_r"],
        "funding": ["f_avg"],
        "OI": ["oi_usd"],
        "vol_perp": ["vol_perp_usd"],
        "taker_buy_pct": ["taker_buy_pct (k[10]/k[7])", "taker_buy_pct"],
        "vol_alpha (col35)": ["vol_alpha"],
        "vol_spot_ext (col51)": ["vol_spot_ext_usd"],
    }
    n = len(dates)
    print(f"  Null% (n={n}):")
    for label, candidates in key_fields.items():
        col_idx = None; found_name = None
        for c in candidates:
            if c in ci:
                col_idx = ci[c]; found_name = c; break
        if col_idx is None:
            print(f"    {label:<25}: COL MISSING (candidates: {candidates})")
            continue
        nulls = 0
        for r in range(3, ws.max_row+1):
            d = ws.cell(r, 1).value
            if not d: continue
            v = ws.cell(r, col_idx).value
            if v in (None, "", "None"):
                nulls += 1
        pct = nulls/n*100 if n else 0
        flag = " ⚠HALT" if (label == "pos_r" and pct > 50) else ""
        print(f"    {label:<25}: {nulls}/{n} = {pct:.1f}% (col={col_idx} '{found_name}'){flag}")

    # CASE meta row 7 contract_canonical
    if "CASE" in sheets:
        ws_case = wb["CASE"]
        row7 = ws_case.cell(7, 1).value
        row7_b = ws_case.cell(7, 2).value
        print(f"  CASE row 7 (contract_canonical check): {row7!r} | {row7_b!r}")
    return {"cols": n_cols, "rows": n, "min": dates[0], "max": dates[-1],
            "missing": len(missing_dates), "dup": len(dup_dates)}

results = {}
for sym in TOKENS:
    r = inventory(sym)
    results[sym] = r

# Summary table
print("\n" + "="*70)
print("SUMMARY 5 CASES")
print("="*70)
print(f"{'TOKEN':<8} {'cols':>5} {'rows':>6} {'min':>12} {'max':>12} {'missing':>8} {'dup':>4}")
for sym in TOKENS:
    r = results[sym]
    if r is None:
        print(f"{sym:<8}  (CHƯA CHẠY — file not found)")
    else:
        print(f"{sym:<8} {r['cols']:>5} {r['rows']:>6} {r['min']:>12} {r['max']:>12} {r['missing']:>8} {r['dup']:>4}")

print("\n[DONE STEP A read-only inventory]")
