"""_task_lesson_p3_20260821.py — Part A: L-60/L-61 insert + L-51 archive-ref
+ Part B: V2 batch full insert (L-55/L-56/L-57/L-58/L-59)
"""
import sys, io, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_INDEX.xlsx")
BAK = Path(r"D:\Data back test\whale_data\scratch\CASE_INDEX_bak_pre_p3_20260821.xlsx")
shutil.copy2(XLSX, BAK)
print(f"Backup P3: {BAK}")

wb = openpyxl.load_workbook(XLSX)
ws = wb["LESSONS"]
print(f"LESSONS before: max_row={ws.max_row}")

# ── PART A ──
# A.1 L-51 archive-ref note (extend existing action_next)
old_action = ws.cell(54, 7).value
new_action = (str(old_action) if old_action else "") + " · ARCHIVE-REF: old L-51 content (GOM-C C1 spec drift 11/08, commit 49f1cd6 và trước) archived in git history — canonical L-51 từ nay = Step-0 identity discipline."
ws.cell(54, 7).value = new_action
print(f"\nL-51 archive-ref note added to row 54")

# A.2 Insert L-60 (schema-drift-câm) - renumbered from V1 draft-49
# Find next empty row
next_r = ws.max_row + 1
for r in range(ws.max_row, 4, -1):
    if ws.cell(r, 1).value:
        next_r = r + 1
        break

r_60 = next_r
ws.cell(r_60, 1).value = "L-60"
ws.cell(r_60, 2).value = "2026-08-17"
ws.cell(r_60, 3).value = "án E LSR-merge (GUA rebuild 17/08) — schema mismatch silent"
ws.cell(r_60, 4).value = (
    "SCHEMA-DRIFT-CÂM: schema mismatch giữa producer + consumer = corrupt SILENT (không alert, hệ tin data đúng nhưng đọc sai). "
    "Evidence án E LSR-merge autopsy 17/08 (docs/reports/_probe_lsr_merge_20260817.md): "
    "fetch script `_probe_gua_apr_rebuild_20260817.py` fetch_lsr() viết CSV với schema Binance raw "
    "(`longShortRatio/longAccount/shortAccount`); build_case.py A-0a chỉ nhận `pos_r_raw|pos_r`, `acc_r_raw|acc_r`, "
    "`glsr_raw|gLSR|glsr` — schema mismatch → tất cả 30 dates Jul19-Aug17 vào `lsr_by_d` với all-None values, "
    "freshwatch loop skip vì dates đã occupied → CASE_GUA.xlsx corrupted (pos_r null 100%). "
    "Rule chốt: (1) mọi CSV producer khai schema explicitly ở header + provenance metadata; "
    "(2) mọi consumer assert schema-check TRƯỚC compute — mismatch = RAISE + block, không silent None; "
    "(3) fetch scripts + consumers cùng amendment cùng chuyến khi schema đổi. "
    "Ghi chú: renumbered from V1 draft-49 — va số với chủ cũ L-49 (11/08 GUARD-CÂM state.json), phát hiện 21/08 audit. "
    "Nguồn text: proposals/lessons_pending_batch_20260818.md §C.* context + án E report + chat-ký 21/08 operator."
)
ws.cell(r_60, 5).value = "SCHEMA"
ws.cell(r_60, 6).value = "DA VAO — 21/08 (renumbered from V1 draft-49)"
ws.cell(r_60, 7).value = (
    "Proposal Fix áp cho build_case.py A-0a: schema-check assert + Fix B nguồn (sửa fetch_lsr() output canonical schema). "
    "Ref: án E report §4-§5 fix proposals A/B/C."
)
ws.cell(r_60, 8).value = "V1-renum-20260821"
print(f"L-60 inserted at row {r_60}")

# A.3 Insert L-61 (write-before-guard) - renumbered from V1 draft-50
r_61 = r_60 + 1
ws.cell(r_61, 1).value = "L-61"
ws.cell(r_61, 2).value = "2026-08-17"
ws.cell(r_61, 3).value = "án E LSR-merge (GUA rebuild 17/08) — XLSX ghi trước guard check"
ws.cell(r_61, 4).value = (
    "WRITE-BEFORE-GUARD anti-pattern (fix = WRITE-AFTER-GUARD): mọi write persistent (xlsx save, csv commit, DB write) "
    "PHẢI SAU khi ALL guards pass — write TRƯỚC guard = corrupt state on-disk khi guard fail. "
    "Evidence án E: build_case.py XLSX write ~line 800-856 (wb.save()) TRƯỚC A-0b guard check line 857-868. "
    "Khi A-0b raise ValueError (pos_r null > 50%) → exit≠0 nhưng XLSX đã corrupted trên disk (GUA 100% null, APR 75%). "
    "Guard fires DID fire correctly (đúng logic), nhưng vô nghĩa vì corrupted state đã persist. "
    "Rule chốt: (1) build dataframe fully in-memory; (2) run ALL guards on in-memory df; "
    "(3) write to disk CHỈ SAU guards pass; (4) guard fail = discard in-memory + exit không đụng disk. "
    "Áp dụng: đã tuân thủ trong GUA refresh 18/08 + TRIA venue backfill 19/08 + STBL/HEMI/PUMPBTC create 19-20/08. "
    "Ghi chú: renumbered from V1 draft-50 — va số với chủ cũ L-50 (11/08 FFUSDT P0), phát hiện 21/08 audit. "
    "Nguồn text: proposals/lessons_pending_batch_20260818.md §C.* + án E fix proposals + chat-ký 21/08 operator."
)
ws.cell(r_61, 5).value = "WRITE-GUARD"
ws.cell(r_61, 6).value = "DA VAO — 21/08 (renumbered from V1 draft-50)"
ws.cell(r_61, 7).value = (
    "Refactor build_case.py: move all guards trước wb.save(). Same rule for any script producing xlsx/csv/db writes. "
    "Verification: mọi task từ 18/08+ đã tuân thủ write-after-guard."
)
ws.cell(r_61, 8).value = "V1-renum-20260821"
print(f"L-61 inserted at row {r_61}")

# ── PART B: V2 batch full (L-55/L-56/L-57/L-58 already-signed/L-59) ──
# L-58 already exists in V2 file - insert to sheet too

r_55 = r_61 + 1
ws.cell(r_55, 1).value = "L-55"
ws.cell(r_55, 2).value = "2026-08-21"
ws.cell(r_55, 3).value = "PHASE-TAG go-live 20/08 · save_snapshot fail-câm · 🎓 badge"
ws.cell(r_55, 4).value = (
    "GO-LIVE ≠ VERIFY: deploy xong chưa phải xong — go-live chỉ ĐÓNG khi verify OUTPUT THẬT trên production "
    "(file sinh · đúng schema · đúng giờ cron). "
    "Evidence: phase_tag CSV F1 code-đúng-file-không-sinh 3 ngày · save_snapshot fail-câm · "
    "🎓 (badge graduation) sống lần đầu chỉ sau verify 20/08. "
    "Rule: mọi go-live commit phải kèm verify checklist (output artifact + schema + timing + integration). "
    "Signed batch V2 · operator ký chat 21/08."
)
ws.cell(r_55, 5).value = "GO-LIVE"
ws.cell(r_55, 6).value = "DA VAO — batch V2 21/08"
ws.cell(r_55, 7).value = "Go-live commit template: paste verify output artifact + schema + timing + integration test result."
ws.cell(r_55, 8).value = "V2-20260821"

r_56 = r_55 + 1
ws.cell(r_56, 1).value = "L-56"
ws.cell(r_56, 2).value = "2026-08-21"
ws.cell(r_56, 3).value = "6 relay-loss/tuần (L-48 số hiệu · GUA contract · E-R4b ≥-drift · handover · lib · va số 49/50)"
ws.cell(r_56, 4).value = (
    "MERGE-BEFORE-GATE: amendment ký chat PHẢI merge canonical + paste diff TRƯỚC gate/backfill run. "
    "Mọi run gate in `spec_version` + `amendments_applied`. "
    "Dòng 'ĐÃ KÝ' trong sổ PHẢI trỏ được về chat-ký commit/URL — không trỏ được = CHỜ-KÝ. "
    "Evidence 6 relay-loss/tuần 15-21/08: (1) L-48 số hiệu drift · (2) GUA contract chưa relay verified · "
    "(3) E-R4b ≥1.15 vs <1.15 drift · (4) handover ký-không-tồn-tại · (5) lib đổi-nghĩa silent · "
    "(6) va số L-49/L-50 chủ cũ vs V1 batch (P3 audit 21/08). "
    "Signed batch V2 · operator ký chat 21/08."
)
ws.cell(r_56, 5).value = "GATE-DISCIPLINE"
ws.cell(r_56, 6).value = "DA VAO — batch V2 21/08"
ws.cell(r_56, 7).value = "Gate/backfill runner: mandatory print(spec_version, amendments_applied). Commit template chat-ký link required."
ws.cell(r_56, 8).value = "V2-20260821"

r_57 = r_56 + 1
ws.cell(r_57, 1).value = "L-57"
ws.cell(r_57, 2).value = "2026-08-21"
ws.cell(r_57, 3).value = "Schema evolution (cols 51-52 v2.1 · ep_len/struct_n · vol_spot/vol_alpha)"
ws.cell(r_57, 4).value = (
    "SCHEMA-CHANGE-APPEND: thay đổi schema = APPEND-ONLY + changelog + KHÔNG đè cột cũ. "
    "Cột đổi NGHĨA phải ĐỔI TÊN (không giữ tên cũ mang nghĩa mới — bài F1 'một-thứ-hai-tên'). "
    "Evidence: (1) cols 51-52 v2.1 spot-ext amendment (append canonical, không đè col 31); "
    "(2) ep_len vs struct_n (2 tên cho 2 nghĩa khác nhau — không alias); "
    "(3) vol_spot vs vol_alpha (col 31 vs col 35, không mix). "
    "Signed batch V2 · operator ký chat 21/08."
)
ws.cell(r_57, 5).value = "SCHEMA-EVOLUTION"
ws.cell(r_57, 6).value = "DA VAO — batch V2 21/08"
ws.cell(r_57, 7).value = "PR review checklist: schema change = append-only? changelog? tên cột đúng nghĩa mới? Nếu có drift → block merge."
ws.cell(r_57, 8).value = "V2-20260821"

r_58 = r_57 + 1
ws.cell(r_58, 1).value = "L-58"
ws.cell(r_58, 2).value = "2026-08-21"
ws.cell(r_58, 3).value = "E-R7d/e drift incidents + P0 CHECK watchlist inline vs lib (2ba669f)"
ws.cell(r_58, 4).value = (
    "CONFORMANCE FIXTURE: pin output MỌI consumer, không-pass-không-merge, đổi ngữ nghĩa = amendment ký riêng, "
    "fixture chạy trong mọi gate run. "
    "Mỗi consumer (watchlist · registry · phase_tag · gom_scan · downstream analyzers) có fixture file "
    "(`tests/fixtures/<consumer>_conformance_<YYYYMMDD>.json`) pin OUTPUT actual (events/episodes/labels/metric values) "
    "trên data snapshot cụ thể (freshwatch csv với md5 lock). "
    "CI/pre-commit hook check fixture — commit đụng consumer path mà fixture FAIL → block merge. Không exception. "
    "Đổi ngữ nghĩa = ký amendment spec + regen fixture + update cùng commit + bump version. "
    "Fixture chạy TRƯỚC compute trong mọi gate/backfill — mismatch = STOP. "
    "Signed batch V2 · operator ký chat 21/08. "
    "Ref: proposals/lessons_pending_batch_V2.md §A.1."
)
ws.cell(r_58, 5).value = "CONFORMANCE"
ws.cell(r_58, 6).value = "DA VAO — batch V2 21/08 (fixture chưa build)"
ws.cell(r_58, 7).value = (
    "Kickoff phase-2: build fixture files 5 tokens (HEMI/ON/GUA/CARV/ARIA) cho registry+watchlist+phase_tag. "
    "Add pre-commit hook + CI script trước phase-2 cron 08:25."
)
ws.cell(r_58, 8).value = "V2-20260821"

r_59 = r_58 + 1
ws.cell(r_59, 1).value = "L-59"
ws.cell(r_59, 2).value = "2026-08-21"
ws.cell(r_59, 3).value = "Batch V1 20/08 · P0 audit 21/08 (root-cause 'hiểu-nhầm-nhà')"
ws.cell(r_59, 4).value = (
    "KÝ-KHÔNG-BẰNG-CHỨNG = CHƯA-XONG: lệnh ghi-sổ đã ký ĐÓNG bằng paste bằng-chứng-sau-ghi CÙNG chuyến. "
    "Không bằng chứng = việc CHƯA TỒN TẠI, CHẶN bước phụ thuộc. "
    "Sổ hai nhà (nguồn + view) khai rõ nguồn/view; view CHẾT >3d khi nguồn SỐNG = ALERT. "
    "Evidence #1: batch V1 ký 20/08 — 3 lần đòi không có bằng chứng (L-52/53/54 vào pending_batch nhưng chưa sync sheet), "
    "operator tự phát hiện md lessons_learned CHẾT từ 07/08 (14 ngày stale) trong khi sheet đã append đợt sau. "
    "Evidence #2: root-cause 'hiểu-nhầm-nhà' 21/08 P3 audit (agent commit pending_batch chờ operator paste tay, "
    "operator ký = agent nhập ngay + regen view). "
    "Signed batch V2 · operator ký chat 21/08."
)
ws.cell(r_59, 5).value = "GOVERNANCE"
ws.cell(r_59, 6).value = "DA VAO — batch V2 21/08"
ws.cell(r_59, 7).value = (
    "Mọi commit ký lesson/spec/amendment: mandatory paste bằng-chứng-sau-ghi (mtime · sheet row · md section · commit sha). "
    "Cron alert: staleness check view vs nguồn — view/md mtime lag > 3d = Telegram alert."
)
ws.cell(r_59, 8).value = "V2-20260821"

wb.save(XLSX)
print(f"\nSAVED CASE_INDEX.xlsx — max_row={ws.max_row}")
print(f"Rows inserted: L-60 (r{r_60}), L-61 (r{r_61}), L-55 (r{r_55}), L-56 (r{r_56}), L-57 (r{r_57}), L-58 (r{r_58}), L-59 (r{r_59})")

# Verify
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws2 = wb2["LESSONS"]
print(f"\n=== VERIFY — 7 new rows ===")
for r in range(r_60, r_59+1):
    print(f"row {r}: id={ws2.cell(r,1).value} · loai={ws2.cell(r,5).value} · trang_thai[:40]={str(ws2.cell(r,6).value)[:40]}")

print("\n[DONE P3]")
