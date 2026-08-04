#!/usr/bin/env python3
"""Probe CASE_TAG.xlsx: full column map + check spot/alpha store availability."""
import io, sys, json
from pathlib import Path
import openpyxl
import pandas as pd

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

XLSX = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_TAG.xlsx")
wb = openpyxl.load_workbook(XLSX)

print("=== SHEETS ===")
print(wb.sheetnames)

# ── DAILY sheet full column map ──
ws = wb["DAILY"]
print(f"\n=== DAILY sheet (max_row={ws.max_row}, max_col={ws.max_column}) ===")
print("Header row 2:")
for c in range(1, ws.max_column + 1):
    v = ws.cell(2, c).value
    if v:
        print(f"  col {c} ({chr(64+c) if c<=26 else 'A'+chr(64+c-26)}): {v}")

print("\nSample data rows 3-5 (first 10 cols):")
for r in range(3, 6):
    vals = [ws.cell(r, c).value for c in range(1, 11)]
    print(f"  row {r}: {vals}")

print("\nLast 4 rows (all cols with values):")
for r in range(ws.max_row - 3, ws.max_row + 1):
    row_vals = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None:
            row_vals[c] = v
    print(f"  row {r}: {row_vals}")

# ── Check alpha_id_map for TAG ──
print("\n=== Alpha ID map — TAG ===")
alpha_map = Path(r"D:\Data back test\whale_data\data\config\alpha_id_map.json")
if alpha_map.exists():
    m = json.loads(alpha_map.read_text())
    tag_entry = m.get("TAG", m.get("TAGUSDT", "NOT FOUND"))
    print(f"  TAG → {tag_entry}")
    # Show all entries with TAG in key
    for k, v in m.items():
        if "TAG" in k.upper():
            print(f"  {k} → {v}")
else:
    print("  alpha_id_map.json NOT FOUND")

# ── Check spot store for TAG ──
print("\n=== Spot store — TAGUSDT ===")
store = Path(r"D:\Data back test\whale_data\data\raw\spot_klines_1d")
tag_files = sorted(store.glob("date=*/*.parquet"))
tag_dates = []
for f in tag_files:
    try:
        df = pd.read_parquet(f)
        if "TAGUSDT" in df.get("symbol", pd.Series()).values or \
           "TAG" in str(df.columns.tolist()):
            tag_dates.append(f.parent.name)
    except:
        pass

if not tag_dates:
    # Try reading one file to see schema
    sample = list(store.glob("date=2026-07-25/*.parquet"))
    if sample:
        df = pd.read_parquet(sample[0])
        print(f"  Schema: {list(df.columns)}")
        print(f"  Symbols sample: {df['symbol'].unique()[:10].tolist() if 'symbol' in df.columns else 'no symbol col'}")
        # Filter TAG
        if 'symbol' in df.columns:
            tag_df = df[df['symbol'].str.contains("TAG", na=False)]
            print(f"  TAG rows in {sample[0].parent.name}: {len(tag_df)}")
            if len(tag_df):
                print(f"  TAG row: {tag_df.iloc[0].to_dict()}")

# ── Check alpha store for TAG ──
print("\n=== Alpha store — TAG (venue=binance_alpha) ===")
alpha_store = store  # same store, different venue column
sample_alpha = list(store.glob("date=2026-07-25/*.parquet"))
for f in sample_alpha:
    try:
        df = pd.read_parquet(f)
        if 'venue' in df.columns:
            alpha_df = df[df['venue'] == 'binance_alpha']
            tag_alpha = alpha_df[alpha_df['symbol'].str.contains("TAG", na=False)] if 'symbol' in alpha_df.columns else pd.DataFrame()
            if len(tag_alpha):
                print(f"  TAG alpha row {f.parent.name}: {tag_alpha.iloc[0].to_dict()}")
                break
            elif len(alpha_df):
                print(f"  Alpha rows in {f.parent.name}: {len(alpha_df)}, symbols: {alpha_alpha['symbol'].unique()[:5].tolist()}")
    except Exception as e:
        print(f"  Error reading {f}: {e}")
