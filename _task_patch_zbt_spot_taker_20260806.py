"""Patch CASE_ZBT.xlsx DAILY: điền cột taker_spot_pct từ k[10]/k[7] spot klines.
Root cause fix: probe cũ dùng k[9] (base tokens) ÷ k[7] (USDT) → đơn vị lệch.
Đúng: k[10] (takerBuyQuoteAssetVolume, USDT) ÷ k[7] (quoteVol, USDT).
Nguồn: api.binance.com spot klines limit=60 (đủ phủ 01/07→05/08).
"""
import sys, io, json, urllib.request
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import PatternFill

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

CASE_PATH = Path(r"D:\Data back test\whale_data\docs\case_files\CASE_ZBT.xlsx")

# ── 1. Fetch spot klines ──────────────────────────────────────────────────────
url = "https://api.binance.com/api/v3/klines?symbol=ZBTUSDT&interval=1d&limit=60"
req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
with urllib.request.urlopen(req, timeout=15) as r:
    data = json.loads(r.read())

today_str = datetime.now(tz=timezone.utc).strftime("%Y-%m-%d")
spot_taker = {}
for k in data:
    d   = datetime.fromtimestamp(int(k[0])//1000, tz=timezone.utc).strftime("%Y-%m-%d")
    if d >= today_str:   # bỏ nến chưa đóng
        continue
    qv  = float(k[7])   # quoteAssetVolume (USDT)
    tbq = float(k[10])  # takerBuyQuoteAssetVolume (USDT)
    spot_taker[d] = round(tbq / qv * 100, 2) if qv > 0 else None

print(f"Spot taker fetched: {len(spot_taker)} days ({min(spot_taker)} → {max(spot_taker)})")
for d in sorted(spot_taker)[-5:]:
    print(f"  {d}: {spot_taker[d]:.1f}%")

# ── 2. Open workbook và tìm cột taker_spot_pct ───────────────────────────────
wb  = openpyxl.load_workbook(CASE_PATH)
ws  = wb["DAILY"]

header_row = 2   # row 2 = column names
cols = {c.value: c.column for c in ws[header_row] if c.value}
assert "taker_spot_pct" in cols, f"Column 'taker_spot_pct' not found in header. Got: {list(cols.keys())}"
col_tsp = cols["taker_spot_pct"]
col_date = cols["date"]

print(f"\ntaker_spot_pct at col {col_tsp}, date at col {col_date}")

# ── 3. Patch từng dòng data ───────────────────────────────────────────────────
patched = 0
missing = 0
for row in ws.iter_rows(min_row=3):
    date_cell = row[col_date - 1]
    if not date_cell.value:
        continue
    d_val = date_cell.value
    if hasattr(d_val, "strftime"):
        d_str = d_val.strftime("%Y-%m-%d")
    else:
        d_str = str(d_val)[:10]

    tsp_cell = row[col_tsp - 1]
    if d_str in spot_taker:
        new_val = spot_taker[d_str]
        old_val = tsp_cell.value
        tsp_cell.value = new_val
        patched += 1
        if old_val is not None and abs((old_val or 0) - (new_val or 0)) > 0.1:
            print(f"  {d_str}: {old_val} → {new_val:.2f}%  (changed)")
        else:
            print(f"  {d_str}: → {new_val:.2f}%")
    else:
        missing += 1
        print(f"  {d_str}: NO DATA (keeping existing: {tsp_cell.value})")

print(f"\nPatched: {patched}  Missing: {missing}")

# ── 4. Save ───────────────────────────────────────────────────────────────────
wb.save(CASE_PATH)
print(f"Saved: {CASE_PATH}")
print("[DONE]")
