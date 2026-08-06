#!/usr/bin/env python3
"""C1: Scan all CASE_*.xlsx DAILY sheets for data_bad=1 (col 26) rows.
   Focus: any row with col 26 = 1 (data_ok==False → invert_bool).
   Also check 2026-08-05 specifically.
"""
import sys
from pathlib import Path
from datetime import datetime

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

import openpyxl

CASE_DIR = Path(__file__).parent / "docs" / "case_files"
DATE_FOCUS = "2026-08-05"

print(f"C1: data_bad=1 scan — {datetime.now().strftime('%Y-%m-%dT%H:%M:%S')}")
print(f"Focus date: {DATE_FOCUS}")
print("=" * 70)

total_bad = 0

for xlsx_f in sorted(CASE_DIR.glob("CASE_*.xlsx")):
    if xlsx_f.name == "CASE_INDEX.xlsx":
        continue

    wb = openpyxl.load_workbook(str(xlsx_f), read_only=True, data_only=True)
    if "DAILY" not in wb.sheetnames:
        wb.close()
        continue

    ws = wb["DAILY"]
    file_bad = []

    for row in ws.iter_rows(min_row=3, values_only=True):
        date_v = row[0] if len(row) > 0 else None
        if date_v is None:
            continue
        date_s = date_v.strftime("%Y-%m-%d") if isinstance(date_v, datetime) else str(date_v)[:10]
        if not date_s or date_s == "None":
            continue

        col26_v = row[25] if len(row) > 25 else None  # col 26 = index 25

        if col26_v == 1 or col26_v is True:
            file_bad.append((date_s, col26_v))

    wb.close()

    if file_bad:
        print(f"\n  {xlsx_f.name}: {len(file_bad)} data_bad=1 row(s)")
        for d, v in file_bad:
            focus = " ← FOCUS DATE" if d == DATE_FOCUS else ""
            print(f"    {d}  col26={v}{focus}")
        total_bad += len(file_bad)
    else:
        focus_note = ""
        # Re-open to check focus date specifically
        wb2 = openpyxl.load_workbook(str(xlsx_f), read_only=True, data_only=True)
        ws2 = wb2["DAILY"]
        focus_row = None
        for row in ws2.iter_rows(min_row=3, values_only=True):
            date_v = row[0] if len(row) > 0 else None
            if date_v is None:
                continue
            date_s = date_v.strftime("%Y-%m-%d") if isinstance(date_v, datetime) else str(date_v)[:10]
            if date_s == DATE_FOCUS:
                col26_v = row[25] if len(row) > 25 else "NO_COL26"
                focus_row = (date_s, col26_v)
                break
        wb2.close()
        if focus_row:
            print(f"  {xlsx_f.name}: 0 bad rows  [{DATE_FOCUS}: col26={focus_row[1]}]")
        else:
            print(f"  {xlsx_f.name}: 0 bad rows  [{DATE_FOCUS}: row NOT FOUND]")

print(f"\n{'='*70}")
print(f"TOTAL data_bad=1 rows: {total_bad}")
