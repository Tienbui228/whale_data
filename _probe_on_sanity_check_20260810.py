"""_probe_on_sanity_check_20260810.py — check CASE_ON.xlsx DAILY columns."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pathlib import Path
import openpyxl

ROOT = Path(r"D:\Data back test\whale_data")
CASE_ON = ROOT / "docs" / "case_files" / "CASE_ON.xlsx"
wb = openpyxl.load_workbook(CASE_ON, data_only=True)
print(f"Sheets: {wb.sheetnames}")

for sh_name in wb.sheetnames:
    sh = wb[sh_name]
    rows = list(sh.iter_rows(values_only=True))
    hdr = rows[0] if rows else ()
    data = [r for r in rows[1:] if any(v is not None for v in r)]
    print(f"\n=== {sh_name} ===")
    print(f"  Cols ({len(hdr)}): {list(hdr)}")
    print(f"  Data rows: {len(data)}")
    if data:
        print(f"  First row: {data[0]}")
        print(f"  Last row:  {data[-1]}")
